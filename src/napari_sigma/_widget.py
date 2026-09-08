from __future__ import annotations

import colorsys
import csv
import gc
import importlib
import os
from concurrent.futures import ThreadPoolExecutor, as_completed
from contextlib import suppress
from dataclasses import dataclass
from time import sleep
from types import MethodType, SimpleNamespace
from typing import Any
from uuid import uuid4

import dask.array as da
import numpy as np
from napari.layers import Labels as NapariLabels
from qtpy.QtCore import (
    QEvent,
    QItemSelectionModel,
    QLocale,
    QObject,
    QPoint,
    Qt,
    QThread,
    QTimer,
    Signal,
)
from qtpy.QtGui import QFont, QKeySequence, QShortcut
from qtpy.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QCheckBox,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QDockWidget,
    QDoubleSpinBox,
    QFileDialog,
    QGridLayout,
    QGroupBox,
    QHBoxLayout,
    QHeaderView,
    QLabel,
    QLineEdit,
    QMessageBox,
    QProgressBar,
    QPushButton,
    QScrollArea,
    QSizePolicy,
    QSpinBox,
    QStyle,
    QTableWidget,
    QTableWidgetItem,
    QTabWidget,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)
from scipy import ndimage as ndi
from skimage.exposure import rescale_intensity

from ._analysis import (
    AnalysisCancelledError,
    SegmentAnalysisMixin,
    _analysis_branch_rows_from_rows,
    _component_mask_from_layer,
    _filter_analysis_rows_by_min_size,
    _get_group_id,
    _get_pixel_size_tuple,
    _get_units_from_layer,
    _is_analysis_aux_layer,
    _spatial_units_for_ndim,
    _squeeze_leading_singletons,
    analyze_binary_components,
    write_analysis_export_file,
)
from ._proximity import (
    ProximityCancelledError,
    ProximityResult,
    ProximityRoiSpec,
    _proximity_roi_label_mask,
    _proximity_roi_specs_from_shapes,
    _normalized_layer_array,
    compute_proximity_result,
    is_proximity_raw_candidate_layer,
    is_proximity_segmentation_candidate_layer,
)
from ._image_io import (
    load_image_tc_zyx,
)
from ._layer_candidates import is_binary_mask_data
from ._tracking import (
    MatchSummary,
    TrackingCancelledError,
    TrackingConfig,
    TrackingResult,
    add_manual_tracking_link,
    compute_match_summary,
    ensure_manual_tracking_link_candidate,
    rebuild_tracking_lineages,
    restore_tracking_link_selection,
    set_tracking_link_selected,
    track_segmentations_ilp,
)
from ._writer import write_single_image
from ._reader import napari_get_reader as napari_get_reader, tczyx_to_layer_data  # legacy import path
from ._metadata import layer_dims_tag as _layer_dims_tag, unit_from_metadata


def _prefer_numba_workqueue() -> None:
    """Avoid OpenMP initialization for napari label operations when possible."""
    try:
        import numba
    except (ImportError, OSError):
        return

    try:
        numba.threading_layer()
    except ValueError:
        configured = str(getattr(numba.config, "THREADING_LAYER", "default"))
        if configured == "default":
            os.environ.setdefault("NUMBA_THREADING_LAYER", "workqueue")
            numba.config.THREADING_LAYER = "workqueue"


def _metadata_units_for_ndim(metadata: dict[str, Any] | None, ndim: int):
    if not hasattr(NapariLabels, "units"):
        return None
    md = metadata or {}
    unit = md.get("PhysicalSizeXUnit") or md.get("unit") or md.get("Units")
    if not unit:
        return None
    return (unit,) * max(int(ndim), 1)


def _spatial_unit_kwargs(layer, ndim: int) -> dict[str, Any]:
    units = _spatial_units_for_ndim(layer, ndim)
    return {"units": units} if units is not None else {}


def _limit_display_points(points: np.ndarray, max_points: int) -> np.ndarray:
    """Limit viewer-only markers without changing tracking sample data."""
    points = np.asarray(points)
    max_points = max(int(max_points), 1)
    if points.ndim != 2 or points.shape[0] <= max_points:
        return points
    indices = np.linspace(0, points.shape[0] - 1, max_points, dtype=int)
    return points[indices]

_TORCH_UNAVAILABLE = object()
_torch_module = None
_segmentation_fn = None
_frangi_filter_cls = None
_DEFAULT_EM_FOREGROUND_SAMPLE_POINTS = None
_DEFAULT_EM_BACKGROUND_SAMPLE_POINTS = 1_000_000
_GAUSSIAN_BACKGROUND_SIGMA_XY = 10.0
_GAUSSIAN_BACKGROUND_ALPHA = 0.5
_EM_SAMPLE_PRESETS = (
    100_000,
    250_000,
    500_000,
    1_000_000,
    2_000_000,
    5_000_000,
    10_000_000,
)
_STRUCTURAL_RESPONSE_LAYER_ROLE = "structural_response"
_LAYER_COMBO_METADATA_KEYS = (
    "is_analysis_labels",
    "is_analysis_highlight",
    "is_analysis_topology",
    "is_proximity_result",
    "is_proximity_overlap_display",
    "is_proximity_component_highlight",
    "is_proximity_source_object_highlight",
    "is_proximity_roi",
    "is_proximity_roi_label",
    "is_proximity_roi_mask",
    "is_upsample",
    "is_denoise",
    "is_frangi",
    "sigma_layer_role",
    "sigma_saved_structure",
    "is_segmentation",
    "is_tracking",
)

def _get_torch_module():
    global _torch_module
    if _torch_module is None:
        try:
            _torch_module = importlib.import_module("torch")
        except (ImportError, OSError):  # pragma: no cover
            _torch_module = _TORCH_UNAVAILABLE
    return None if _torch_module is _TORCH_UNAVAILABLE else _torch_module


def _parse_em_sample_points_text(text: str) -> int | None:
    value = str(text).strip()
    if value.casefold() == "all":
        return None
    digits = value.replace(",", "").replace("_", "").replace(" ", "")
    if not digits.isdigit() or int(digits) <= 0:
        raise ValueError("EM sample points must be a positive integer or All.")
    return int(digits)


def _format_em_sample_points(value: int | None) -> str:
    return "All" if value is None else f"{int(value):,}"


def _get_segmentation_fn():
    global _segmentation_fn
    if _segmentation_fn is None:
        module = importlib.import_module("napari_sigma.segmentation")
        _segmentation_fn = module.segmentation
    return _segmentation_fn


def _spatial_scale_for_ndim(layer, ndim: int) -> tuple[float, ...] | None:
    scale = getattr(layer, "scale", None)
    if scale is None:
        return None
    scale_tuple = tuple(float(v) for v in scale)
    if len(scale_tuple) < ndim:
        return None
    return scale_tuple[-ndim:]


def _spatial_translate_for_ndim(layer, ndim: int) -> tuple[float, ...] | None:
    translate = getattr(layer, "translate", None)
    if translate is None:
        return None
    translate_tuple = tuple(float(v) for v in translate)
    if len(translate_tuple) < ndim:
        return None
    return translate_tuple[-ndim:]


def _get_frangi_filter_cls():
    global _frangi_filter_cls
    if _frangi_filter_cls is None:
        module = importlib.import_module("frangi_filter.frangi_filter")
        _frangi_filter_cls = module.FrangiFilter
    return _frangi_filter_cls


def _normalize_0_255(arr) -> np.ndarray:
    arr = np.asarray(arr, dtype=np.float32)
    vmin = float(np.min(arr))
    vmax = float(np.max(arr))
    rng = vmax - vmin
    if rng <= 0:
        return np.zeros_like(arr, dtype=np.float32)
    return (arr - vmin) / rng * 255.0


def _gaussian_background_subtract_0_255(
    arr,
    sigma_xy: float = _GAUSSIAN_BACKGROUND_SIGMA_XY,
    alpha: float = _GAUSSIAN_BACKGROUND_ALPHA,
) -> np.ndarray:
    """Subtract a broad XY Gaussian background and return a 0..255 image."""
    image = np.asarray(arr, dtype=np.float32)
    if image.ndim not in {2, 3}:
        raise ValueError("Gaussian background subtraction expects a 2D or 3D image.")
    sigma_xy = float(sigma_xy)
    if not np.isfinite(sigma_xy) or sigma_xy <= 0:
        raise ValueError("Gaussian background sigma must be finite and positive.")
    alpha = float(alpha)
    if not np.isfinite(alpha) or not 0.0 <= alpha <= 1.0:
        raise ValueError("Gaussian background alpha must be finite and within [0, 1].")
    sigma = (0.0,) * (image.ndim - 2) + (sigma_xy, sigma_xy)
    background = ndi.gaussian_filter(image, sigma=sigma, mode="reflect")
    corrected = np.clip(image - alpha * background, 0.0, None)
    return _normalize_0_255(corrected)


def _frangi_result_name_for_mode(mode: str) -> str:
    return {
        "sheetness": "sheetness",
        "combined": "combined max",
    }.get(str(mode).lower(), "vesselness")


def _rescale_frangi_per_frame(raw, low: float, temporal_dims: str | None) -> np.ndarray:
    raw = np.asarray(raw)

    def _rescale_one(frame):
        lo = float(np.percentile(frame, low))
        hi = float(np.percentile(frame, 100.0 - low))
        if hi <= lo:
            hi = lo + 1e-6
        return rescale_intensity(frame, in_range=(lo, hi))

    if temporal_dims in {"TYX", "TZYX"} and raw.ndim >= 3:
        return np.stack([_rescale_one(raw[idx]) for idx in range(raw.shape[0])], axis=0)
    return np.asarray(_rescale_one(raw))


def _upsampled_xy_scale(scale, factor: int) -> tuple[float, ...]:
    scale = tuple(float(value) for value in scale)
    factor = int(factor)
    if len(scale) < 2 or factor < 2:
        raise ValueError("XY upsampling requires at least two axes and factor >= 2.")
    return (*scale[:-2], scale[-2] / factor, scale[-1] / factor)


def _upsample_xy_bilinear(
    image,
    factor: int,
    *,
    progress=None,
    cancel_check=None,
) -> np.ndarray:
    import cv2

    shape = tuple(int(value) for value in getattr(image, "shape", ()))
    factor = int(factor)
    if len(shape) < 2 or factor < 2:
        raise ValueError("XY upsampling requires image ndim >= 2 and factor >= 2.")
    height, width = shape[-2:]
    leading_shape = shape[:-2]
    dtype = getattr(image, "dtype", None)
    if dtype is None:
        dtype = np.asarray(image).dtype
    output = np.empty(
        (*leading_shape, height * factor, width * factor),
        dtype=np.dtype(dtype),
    )
    total = int(np.prod(leading_shape, dtype=np.int64)) if leading_shape else 1
    for flat_index in range(total):
        if cancel_check is not None and cancel_check():
            raise InterruptedError("XY upsampling cancelled.")
        index = np.unravel_index(flat_index, leading_shape) if leading_shape else ()
        frame = np.asarray(image[index] if leading_shape else image)
        if not frame.dtype.isnative:
            frame = frame.astype(frame.dtype.newbyteorder("="), copy=False)
        output[index] = cv2.resize(
            frame,
            (width * factor, height * factor),
            interpolation=cv2.INTER_LINEAR,
        )
        if progress is not None:
            progress(flat_index + 1, total)
    return output


def _median_filter_chunk_bounds(
    axis_length: int,
    kernel_radius: int,
    n_workers: int,
) -> list[tuple[int, int]]:
    axis_length = int(axis_length)
    kernel_radius = max(int(kernel_radius), 0)
    n_workers = max(int(n_workers), 1)
    chunk_length = max(
        2 * kernel_radius + 1,
        (axis_length + n_workers * 2 - 1) // (n_workers * 2),
    )
    return [
        (start, min(start + chunk_length, axis_length))
        for start in range(0, axis_length, chunk_length)
    ]


def _median_filter_chunked_exact(
    image,
    kernel_size: int | tuple[int, ...],
    *,
    n_workers: int | None = None,
    progress=None,
    cancel_check=None,
) -> np.ndarray:
    """Exact zero-padded median filter using parallel halo chunks."""
    array = np.asarray(image)
    if array.ndim not in {2, 3}:
        raise ValueError("Chunked median filtering expects a 2D or 3D array.")
    if np.isscalar(kernel_size):
        kernel = (int(kernel_size),) * array.ndim
    else:
        kernel = tuple(int(value) for value in kernel_size)
    if len(kernel) != array.ndim or any(value < 1 or value % 2 == 0 for value in kernel):
        raise ValueError(f"Median kernel must contain odd positive sizes, got {kernel!r}.")

    worker_count = min(max(int(n_workers or min(os.cpu_count() or 1, 4)), 1), 4)
    radius = kernel[0] // 2
    bounds = _median_filter_chunk_bounds(array.shape[0], radius, worker_count)
    total = max(len(bounds), 1)
    output = np.empty_like(array)

    def _filter_chunk(start: int, stop: int):
        halo_start = max(0, start - radius)
        halo_stop = min(array.shape[0], stop + radius)
        filtered = ndi.median_filter(
            array[halo_start:halo_stop],
            size=kernel,
            mode="constant",
            cval=0,
        )
        crop_start = start - halo_start
        return start, stop, filtered[crop_start : crop_start + (stop - start)]

    if cancel_check is not None and cancel_check():
        raise InterruptedError("Median filtering cancelled.")
    if len(bounds) == 1 or worker_count == 1:
        for done, (start, stop) in enumerate(bounds, start=1):
            if cancel_check is not None and cancel_check():
                raise InterruptedError("Median filtering cancelled.")
            _, _, chunk = _filter_chunk(start, stop)
            output[start:stop] = chunk
            if progress is not None:
                progress(done, total)
        return output

    with ThreadPoolExecutor(max_workers=min(worker_count, len(bounds))) as executor:
        futures = [executor.submit(_filter_chunk, start, stop) for start, stop in bounds]
        for done, future in enumerate(as_completed(futures), start=1):
            if cancel_check is not None and cancel_check():
                for pending in futures:
                    pending.cancel()
                raise InterruptedError("Median filtering cancelled.")
            start, stop, chunk = future.result()
            output[start:stop] = chunk
            if progress is not None:
                progress(done, total)
    return output


def _downsample_binary_xy_preserve_components(
    mask,
    factor: int,
    *,
    target_xy: tuple[int, int] | None = None,
) -> tuple[np.ndarray, dict[str, int]]:
    """Downsample XY while retaining one disconnected region per component.

    Any foreground voxel in a source XY block is retained. A protected seed
    is still retained for every component so small objects cannot disappear.
    """
    binary = np.asarray(mask) > 0
    factor = int(factor)
    if binary.ndim not in {2, 3}:
        raise ValueError("Component-preserving XY downsampling expects a 2D or 3D mask.")
    if factor < 2:
        raise ValueError("Component-preserving XY downsampling requires factor >= 2.")

    if target_xy is None:
        target_xy = (
            max(int(binary.shape[-2] // factor), 1),
            max(int(binary.shape[-1] // factor), 1),
        )
    target_y, target_x = (int(target_xy[0]), int(target_xy[1]))
    if target_y < 1 or target_x < 1:
        raise ValueError(f"Invalid downsample target XY shape: {target_xy!r}")
    target_shape = (*binary.shape[:-2], target_y, target_x)
    structure = ndi.generate_binary_structure(binary.ndim, 1)
    labels, component_count = ndi.label(binary, structure=structure)
    component_count = int(component_count)
    if component_count == 0:
        return np.zeros(target_shape, dtype=np.uint8), {
            "components_before": 0,
            "components_after": 0,
            "relocated_seeds": 0,
        }

    records = []
    for component_id, component_slice in enumerate(ndi.find_objects(labels), start=1):
        if component_slice is None:
            continue
        local_coords = np.argwhere(labels[component_slice] == component_id)
        starts = np.asarray([axis_slice.start for axis_slice in component_slice], dtype=np.int64)
        high_coords = local_coords + starts
        low_coords = high_coords.copy()
        low_coords[:, -2] //= factor
        low_coords[:, -1] //= factor
        low_coords[:, -2] = np.clip(low_coords[:, -2], 0, target_y - 1)
        low_coords[:, -1] = np.clip(low_coords[:, -1], 0, target_x - 1)
        flat = np.ravel_multi_index(tuple(low_coords.T), target_shape)
        unique_flat, block_counts = np.unique(flat, return_counts=True)
        projected_coords = np.column_stack(np.unravel_index(unique_flat, target_shape)).astype(np.int64)
        center = high_coords.mean(axis=0, dtype=np.float64)
        center[-2:] /= factor
        center = np.clip(center, 0, np.asarray(target_shape, dtype=np.float64) - 1)
        records.append(
            {
                "id": int(component_id),
                "coords": projected_coords,
                "counts": block_counts.astype(np.int64, copy=False),
                "center": center,
                "high_size": int(high_coords.shape[0]),
            }
        )

    owner = np.zeros(target_shape, dtype=np.int32)

    def _is_available(coord: np.ndarray, own_id: int = 0) -> bool:
        coord_tuple = tuple(int(value) for value in coord)
        value = int(owner[coord_tuple])
        if value not in {0, own_id}:
            return False
        for axis in range(owner.ndim):
            for delta in (-1, 1):
                neighbor = coord.copy()
                neighbor[axis] += delta
                if neighbor[axis] < 0 or neighbor[axis] >= owner.shape[axis]:
                    continue
                neighbor_value = int(owner[tuple(int(value) for value in neighbor)])
                if neighbor_value not in {0, own_id}:
                    return False
        return True

    def _nearest_available_seed(center: np.ndarray) -> np.ndarray | None:
        center_int = np.rint(center).astype(np.int64)
        center_int = np.clip(center_int, 0, np.asarray(target_shape, dtype=np.int64) - 1)
        max_radius = max(target_shape)
        for radius in range(max_radius + 1):
            lower = np.maximum(center_int - radius, 0)
            upper = np.minimum(center_int + radius + 1, np.asarray(target_shape, dtype=np.int64))
            region_shape = tuple(int(value) for value in (upper - lower))
            for local_coord in np.ndindex(region_shape):
                coord = lower + np.asarray(local_coord, dtype=np.int64)
                if radius and int(np.max(np.abs(coord - center_int))) != radius:
                    continue
                if _is_available(coord):
                    return coord
        return None

    relocated_seeds = 0
    # Protect constrained/small objects first; larger projected regions have
    # more alternative seed locations.
    for record in sorted(records, key=lambda item: (len(item["coords"]), item["high_size"])):
        coords = record["coords"]
        distances = np.sum((coords - record["center"]) ** 2, axis=1)
        order = np.lexsort((distances, -record["counts"]))
        seed = None
        for index in order:
            candidate = coords[int(index)]
            if _is_available(candidate):
                seed = candidate.copy()
                break
        if seed is None:
            seed = _nearest_available_seed(record["center"])
            relocated_seeds += 1
        if seed is None:
            raise RuntimeError(
                "The original grid has insufficient separated pixels to preserve "
                f"all {component_count} connected components after {factor}x XY downsampling."
            )
        record["seed"] = seed
        owner[tuple(int(value) for value in seed)] = int(record["id"])

    # Grow projected shapes around the protected seeds. Other seeds are already
    # present, so no component can grow into or directly touch another one.
    for record in sorted(records, key=lambda item: item["high_size"], reverse=True):
        component_id = int(record["id"])
        coords = record["coords"]
        allowed = np.ones(len(coords), dtype=bool)
        current_values = owner[tuple(coords.T)]
        allowed &= (current_values == 0) | (current_values == component_id)
        for axis in range(owner.ndim):
            for delta in (-1, 1):
                valid = (coords[:, axis] + delta >= 0) & (coords[:, axis] + delta < owner.shape[axis])
                if not np.any(valid):
                    continue
                shifted = coords[valid].copy()
                shifted[:, axis] += delta
                neighbor_values = owner[tuple(shifted.T)]
                allowed[valid] &= (neighbor_values == 0) | (neighbor_values == component_id)

        allowed_coords = coords[allowed]
        seed = np.asarray(record["seed"], dtype=np.int64)
        all_coords = np.vstack([allowed_coords, seed.reshape(1, -1)])
        lower = all_coords.min(axis=0)
        upper = all_coords.max(axis=0) + 1
        local_shape = tuple(int(value) for value in (upper - lower))
        local_mask = np.zeros(local_shape, dtype=bool)
        local_mask[tuple((all_coords - lower).T)] = True
        local_labels, _ = ndi.label(local_mask, structure=structure)
        seed_local = tuple(int(value) for value in (seed - lower))
        seed_label = int(local_labels[seed_local])
        selected = np.argwhere(local_labels == seed_label) + lower
        owner[tuple(selected.T)] = component_id

    output = owner > 0
    _labels_after, components_after = ndi.label(output, structure=structure)
    components_after = int(components_after)
    if components_after != component_count:
        seed_only = np.zeros_like(owner, dtype=bool)
        for record in records:
            seed_only[tuple(int(value) for value in record["seed"])] = True
        _seed_labels, seed_count = ndi.label(seed_only, structure=structure)
        if int(seed_count) != component_count:
            raise RuntimeError(
                f"Connected-component preservation failed: before={component_count}, after={components_after}."
            )
        output = seed_only
        components_after = int(seed_count)

    return np.ascontiguousarray(output.astype(np.uint8) * 255), {
        "components_before": component_count,
        "components_after": components_after,
        "relocated_seeds": int(relocated_seeds),
    }


# ---------------------------------------------------------------------
# Small utilities
# ---------------------------------------------------------------------

def _as_list(x):
    return list(x) if isinstance(x, list | tuple) else ([x] if x is not None else [])


def _get_source_from_layer(layer) -> str | None:
    md = getattr(layer, "metadata", {}) or {}
    for key in ("source", "original_path", "path", "file"):
        if key in md:
            return md[key]
    return getattr(getattr(layer, "source", None), "path", None)


def _ensure_unit_in_metadata(md: dict[str, Any]) -> str:
    unit = unit_from_metadata(md)
    md["unit"] = unit
    return unit


def _ensure_group_id(md: dict[str, Any], fallback: str | None = None) -> str:
    gid = md.get("group_id") or fallback or uuid4().hex
    md["group_id"] = gid
    return gid


def _mark_plugin_managed(md: dict[str, Any]) -> dict[str, Any]:
    md["managed_by_plugin"] = True
    return md


def _layer_display_ndim(layer) -> int:
    """Return the number of axes that napari displays for a layer."""
    with suppress(AttributeError, TypeError, ValueError):
        return max(int(layer.ndim), 0)
    data = getattr(layer, "data", None)
    with suppress(AttributeError, TypeError, ValueError):
        return max(int(data.ndim), 0)
    return 0


def _viewer_axis_for_layer_axis(viewer, layer, layer_axis: int) -> int | None:
    """Map a layer-local axis to napari's right-aligned viewer axes."""
    layer_ndim = _layer_display_ndim(layer)
    if layer_ndim <= 0:
        return None
    axis = int(layer_axis)
    if axis < 0:
        axis += layer_ndim
    if axis < 0 or axis >= layer_ndim:
        return None
    with suppress(AttributeError, TypeError, ValueError):
        viewer_ndim = int(viewer.dims.ndim)
        if viewer_ndim >= layer_ndim:
            return viewer_ndim - layer_ndim + axis
    return axis


def _lineage_rgba(lineage_id: int) -> tuple[float, float, float, float]:
    rng = np.random.default_rng(int(lineage_id) * 9973 + 17)
    rgb = rng.uniform(0.2, 0.95, size=3)
    return (float(rgb[0]), float(rgb[1]), float(rgb[2]), 1.0)


def _time_size_from_layer(layer) -> int:
    data = getattr(layer, "data", None)
    if data is None:
        return 1
    dims_tag = _layer_dims_tag(layer)
    if dims_tag in {"TYX", "TZYX"} and data.ndim >= 3:
        return int(data.shape[0])
    if dims_tag == "TCZYX" and data.ndim >= 5:
        return int(data.shape[0])
    if data.ndim >= 4:
        return int(data.shape[0])
    return 1


def _format_time_series_info(layer) -> str | None:
    data = getattr(layer, "data", None)
    if data is None:
        return None
    dims_tag = _layer_dims_tag(layer)
    ndim = getattr(data, "ndim", 0)
    has_time_axis = "T" in dims_tag if dims_tag else ndim >= 4
    if not has_time_axis:
        return None

    frame_count = _time_size_from_layer(layer)
    frame_word = "frame" if frame_count == 1 else "frames"
    text = f"T={frame_count} {frame_word}"

    metadata = getattr(layer, "metadata", {}) or {}
    interval = None
    for key in ("time_interval", "finterval"):
        try:
            value = float(metadata.get(key))
        except (TypeError, ValueError):
            continue
        if np.isfinite(value) and value > 0:
            interval = value
            break
    if interval is None:
        try:
            fps = float(metadata.get("fps"))
        except (TypeError, ValueError):
            fps = 0.0
        if np.isfinite(fps) and fps > 0:
            interval = 1.0 / fps

    if interval is not None:
        time_span = max(frame_count - 1, 0) * interval
        text += f" | interval={interval:.6g} s | time span={time_span:.6g} s"
    return text


def _time_axis_index(viewer, layer) -> int | None:
    dims_tag = _layer_dims_tag(layer)
    if dims_tag not in {"TYX", "TZYX", "TCZYX"}:
        return None
    layer_ndim = _layer_display_ndim(layer)
    if len(dims_tag) == layer_ndim:
        mapped_axis = _viewer_axis_for_layer_axis(
            viewer,
            layer,
            dims_tag.index("T"),
        )
        if mapped_axis is not None:
            return mapped_axis
    with suppress(Exception):
        labels = tuple(str(v).upper() for v in viewer.dims.axis_labels)
        for i, label in enumerate(labels):
            if label == "T":
                return i
    return 0


def _apply_viewer_time_step(viewer, layer, frame: int) -> None:
    time_axis = _time_axis_index(viewer, layer)
    if time_axis is None:
        return
    with suppress(AttributeError, IndexError, RuntimeError, TypeError, ValueError):
        viewer.dims.set_current_step(time_axis, int(frame))


def _is_tracking_candidate_layer(layer) -> bool:
    if _is_analysis_aux_layer(layer):
        return False
    md = getattr(layer, "metadata", {}) or {}
    dims = str(md.get("dims") or md.get("dims_out") or "").upper()
    if dims not in {"TYX", "TZYX"}:
        return False
    data = getattr(layer, "data", None)
    if getattr(data, "ndim", None) not in {3, 4} or int(getattr(data, "size", 0)) == 0:
        return False
    if md.get("is_segmentation"):
        return True
    with suppress(TypeError, ValueError, AttributeError):
        return is_binary_mask_data(data)
    return False


def _is_tracking_raw_candidate_layer(layer, reference_layer=None) -> bool:
    if not _is_tracking_overlay_layer(layer, reference_layer):
        return False
    layer_type = str(getattr(layer, "_type_string", "") or "").lower()
    layer_class = layer.__class__.__name__.lower()
    if layer_type != "image" and layer_class != "image":
        return False
    md = getattr(layer, "metadata", {}) or {}
    return not (
        md.get("is_frangi")
        or md.get("is_segmentation")
        or md.get("is_tracking")
    )


def _is_frangi_raw_candidate_layer(layer) -> bool:
    if _is_analysis_aux_layer(layer):
        return False
    if layer is None or not hasattr(layer, "data"):
        return False
    layer_type = str(getattr(layer, "_type_string", "") or "").lower()
    layer_class = layer.__class__.__name__.lower()
    if layer_type not in {"image", "labels"} and layer_class not in {"image", "labels"}:
        return False
    md = getattr(layer, "metadata", {}) or {}
    if md.get("is_frangi") or md.get("is_segmentation") or md.get("is_tracking"):
        return False
    data = getattr(layer, "data", None)
    return data is not None and int(getattr(data, "size", 0)) > 0


def _is_tracking_overlay_layer(layer, reference_layer=None) -> bool:
    if layer is None or not hasattr(layer, "data") or _is_analysis_aux_layer(layer):
        return False
    layer_type = str(getattr(layer, "_type_string", "") or "").lower()
    layer_class = layer.__class__.__name__.lower()
    if layer_type not in {"image", "labels"} and layer_class not in {"image", "labels"}:
        return False
    data = getattr(layer, "data", None)
    if getattr(data, "ndim", None) not in {3, 4} or int(getattr(data, "size", 0)) == 0:
        return False
    dims = str(
        (getattr(layer, "metadata", {}) or {}).get("dims")
        or (getattr(layer, "metadata", {}) or {}).get("dims_out")
        or ""
    ).upper()
    if dims and dims not in {"TYX", "TZYX"}:
        return False
    if reference_layer is None or not hasattr(reference_layer, "data"):
        return data.ndim in {3, 4}
    ref = getattr(reference_layer, "data", None)
    if getattr(ref, "ndim", None) not in {3, 4} or int(getattr(ref, "size", 0)) == 0:
        return False
    return tuple(data.shape) == tuple(ref.shape)


def _apply_tracking_overlay_rendering(target_layer, source_layer) -> None:
    """Copy overlay rendering settings without changing opacity."""
    source_type = str(
        getattr(source_layer, "_type_string", "")
        or source_layer.__class__.__name__
    ).lower()
    source_is_image = source_type == "image"
    defaults = {
        "gamma": 0.7,
        "rendering": "attenuated_mip",
        "projection_mode": "none",
        "attenuation": 0.25,
    }
    for attribute, default in defaults.items():
        value = (
            getattr(source_layer, attribute, default)
            if source_is_image
            else default
        )
        with suppress(Exception):
            setattr(target_layer, attribute, value)
    if source_is_image and hasattr(source_layer, "iso_threshold"):
        with suppress(Exception):
            target_layer.iso_threshold = float(source_layer.iso_threshold)


def _safe_axis_labels(viewer, layer):
    # Layer axes are right-aligned in napari. Preserve leading viewer labels
    # when a lower-dimensional layer becomes active in a higher-dimensional
    # viewer (for example, a ZYX match preview inside a TZYX dataset).
    dims_tag = _layer_dims_tag(layer)
    nd = _layer_display_ndim(layer)
    if dims_tag in {"TYX", "TZYX", "TCZYX"} and len(dims_tag) == nd:
        layer_labels = tuple(dims_tag)
    elif nd >= 4:
        layer_labels = ("T", "Z", "Y", "X")[-nd:]
    elif nd == 3:
        layer_labels = ("Z", "Y", "X")
    elif nd == 2:
        layer_labels = ("Y", "X")
    else:
        layer_labels = tuple(str(i) for i in range(nd))
    with suppress(Exception):
        viewer_ndim = int(viewer.dims.ndim)
        current_labels = list(viewer.dims.axis_labels)
        if len(current_labels) != viewer_ndim:
            current_labels = [str(i) for i in range(viewer_ndim)]
        if len(layer_labels) > viewer_ndim:
            layer_labels = layer_labels[-viewer_ndim:]
        offset = viewer_ndim - len(layer_labels)
        current_labels[offset:] = layer_labels
        viewer.dims.axis_labels = tuple(current_labels)


def _describe_size_and_resolution(layer) -> dict[str, Any]:
    data = layer.data
    result: dict[str, Any] = {}
    sc = tuple(getattr(layer, "scale", (1,) * data.ndim))
    unit = _get_units_from_layer(layer)
    dims_tag = _layer_dims_tag(layer)

    if dims_tag == "TYX":
        ny, nx = data.shape[-2], data.shape[-1]
        vy, vx = float(sc[-2]), float(sc[-1])
        phys = (ny * vy, nx * vx)
        result.update(
            {
                "ndim": 2,
                "shape": (ny, nx),
                "voxel": (vy, vx),
                "phys": (float(phys[0]), float(phys[1])),
                "unit": unit,
            }
        )
        return result

    if data.ndim >= 3 and int(data.shape[-3]) > 1:
        nz, ny, nx = data.shape[-3], data.shape[-2], data.shape[-1]
        vz, vy, vx = float(sc[-3]), float(sc[-2]), float(sc[-1])
        phys = (nz * vz, ny * vy, nx * vx)
        result.update(
            {
                "ndim": 3,
                "shape": (nz, ny, nx),
                "voxel": (vz, vy, vx),
                "phys": (float(phys[0]), float(phys[1]), float(phys[2])),
                "unit": unit,
            }
        )
    elif data.ndim == 2:
        ny, nx = data.shape[-2], data.shape[-1]
        vy, vx = float(sc[-2]), float(sc[-1])
        phys = (ny * vy, nx * vx)
        result.update(
            {
                "ndim": 2,
                "shape": (ny, nx),
                "voxel": (vy, vx),
                "phys": (float(phys[0]), float(phys[1])),
                "unit": unit,
            }
        )
    else:
        result.update({"ndim": data.ndim, "shape": tuple(data.shape)})
    return result


def _load_image_tc_zyx(path: str) -> tuple[np.ndarray, dict[str, Any]]:
    data, meta = load_image_tc_zyx(path)
    _ensure_group_id(meta)
    return data, meta


@dataclass
class ProcessingContext:
    image: Any
    dim: int
    pixel_size: tuple[float, ...]
    unit: str
    source: str | None
    group_id: str | None
    channel_index: int | None = None
    temporal_dims: str | None = None
    time_range: tuple[int, int] | None = None


@dataclass
class FrangiContext(ProcessingContext):
    frangi: Any | None = None
    frangi_raw: Any | None = None
    device: str = "cpu"
    temporal_dims: str | None = None
    time_range: tuple[int, int] | None = None


# ---------------------------------------------------------------------
# Background worker
# ---------------------------------------------------------------------

class UpsampleWorker(QObject):
    finished = Signal(object, object)
    progress = Signal(int, int)

    def __init__(self, image, factor: int):
        super().__init__()
        self.image = image
        self.factor = int(factor)
        self._cancelled = False

    def cancel(self) -> None:
        self._cancelled = True

    def _is_cancelled(self) -> bool:
        thread = self.thread()
        return bool(self._cancelled or (thread is not None and thread.isInterruptionRequested()))

    def run(self) -> None:
        try:
            output = _upsample_xy_bilinear(
                self.image,
                self.factor,
                progress=lambda done, total: self.progress.emit(int(done), int(total)),
                cancel_check=self._is_cancelled,
            )
            self.finished.emit(output, None)
        except Exception as error:  # noqa: BLE001 - worker boundary must always notify the UI
            self.finished.emit(None, error)


class DenoiseWorker(QObject):
    finished = Signal(object, object)
    progress = Signal(int, int)

    def __init__(
        self,
        image,
        kernel_size: int,
        *,
        temporal: bool,
        z_range: tuple[int, int] | None = None,
        n_workers: int | None = None,
    ):
        super().__init__()
        self.image = image
        self.kernel_size = int(kernel_size)
        self.temporal = bool(temporal)
        self.z_range = z_range
        self.n_workers = min(max(int(n_workers or min(os.cpu_count() or 1, 4)), 1), 4)
        self._cancelled = False

    def cancel(self) -> None:
        self._cancelled = True

    def _is_cancelled(self) -> bool:
        thread = self.thread()
        return bool(self._cancelled or (thread is not None and thread.isInterruptionRequested()))

    def _target_from_frame(self, frame: np.ndarray) -> np.ndarray:
        if self.z_range is None:
            return frame
        z_start, z_end = self.z_range
        if z_start == z_end:
            return frame[z_start]
        return frame[z_start : z_end + 1]

    def _operation_steps(self, target: np.ndarray) -> int:
        radius = self.kernel_size // 2
        return len(_median_filter_chunk_bounds(target.shape[0], radius, self.n_workers))

    def run(self) -> None:
        try:
            source = np.asarray(self.image)
            frames = [source[index] for index in range(source.shape[0])] if self.temporal else [source]
            total_steps = sum(self._operation_steps(self._target_from_frame(frame)) for frame in frames)
            total_steps = max(int(total_steps), 1)
            output = np.empty(source.shape, dtype=np.float32)
            completed_steps = 0

            for frame_index, frame in enumerate(frames):
                if self._is_cancelled():
                    raise InterruptedError("Median filtering cancelled.")
                target = self._target_from_frame(frame)
                operation_steps = self._operation_steps(target)

                def _progress(done: int, _total: int) -> None:
                    self.progress.emit(completed_steps + int(done), total_steps)

                filtered_target = _median_filter_chunked_exact(
                    target,
                    self.kernel_size,
                    n_workers=self.n_workers,
                    progress=_progress,
                    cancel_check=self._is_cancelled,
                )
                if self.z_range is None:
                    filtered_frame = filtered_target
                else:
                    filtered_frame = np.array(frame, copy=True)
                    z_start, z_end = self.z_range
                    if z_start == z_end:
                        filtered_frame[z_start] = filtered_target
                    else:
                        filtered_frame[z_start : z_end + 1] = filtered_target
                normalized = _normalize_0_255(filtered_frame)
                if self.temporal:
                    output[frame_index] = normalized
                else:
                    output[...] = normalized
                completed_steps += operation_steps

            self.finished.emit(output, None)
        except Exception as error:  # noqa: BLE001 - worker boundary must always notify the UI
            self.finished.emit(None, error)


class GaussianBackgroundWorker(QObject):
    finished = Signal(object, object)
    progress = Signal(int, int)

    def __init__(self, image, sigma_xy: float, alpha: float, *, temporal: bool):
        super().__init__()
        self.image = image
        self.sigma_xy = float(sigma_xy)
        self.alpha = float(alpha)
        self.temporal = bool(temporal)
        self._cancelled = False

    def cancel(self) -> None:
        self._cancelled = True

    def _is_cancelled(self) -> bool:
        thread = self.thread()
        return bool(self._cancelled or (thread is not None and thread.isInterruptionRequested()))

    def run(self) -> None:
        try:
            source = np.asarray(self.image)
            frames = [source[index] for index in range(source.shape[0])] if self.temporal else [source]
            output = np.empty(source.shape, dtype=np.float32)
            total = max(len(frames), 1)
            for frame_index, frame in enumerate(frames):
                if self._is_cancelled():
                    raise InterruptedError("Gaussian background subtraction cancelled.")
                corrected = _gaussian_background_subtract_0_255(
                    frame,
                    self.sigma_xy,
                    self.alpha,
                )
                if self.temporal:
                    output[frame_index] = corrected
                else:
                    output[...] = corrected
                self.progress.emit(frame_index + 1, total)
            self.finished.emit(output, None)
        except Exception as error:  # noqa: BLE001 - worker boundary must always notify the UI
            self.finished.emit(None, error)


class SegWorker(QObject):
    finished = Signal(object, object)  # (payload, error)
    started = Signal()
    progress = Signal(int, float)  # (iteration, delta)
    frame_progress = Signal(int, int)  # (done, total)

    def __init__(self, kwargs: dict[str, Any]):
        super().__init__()
        self.kwargs = kwargs
        self._cancelled = False

    def cancel(self) -> None:
        self._cancelled = True

    def _is_cancelled(self) -> bool:
        thread = self.thread()
        return bool(self._cancelled or (thread is not None and thread.isInterruptionRequested()))

    @staticmethod
    def _cleanup_torch_device(device_name: str | None) -> None:
        torch = _get_torch_module()
        if torch is None or not device_name:
            return
        if device_name == "cuda":
            with suppress(Exception):
                torch.cuda.synchronize()
            with suppress(Exception):
                torch.cuda.empty_cache()
        elif device_name == "mps":
            with suppress(Exception):
                torch.mps.synchronize()
            with suppress(Exception):
                torch.mps.empty_cache()

    def run(self):
        self.started.emit()
        try:
            segmentation = _get_segmentation_fn()
            run_framewise = bool(self.kwargs.get("run_framewise", False))
            spatial_dim = int(self.kwargs.get("spatial_dim", 2))
            device_name = str(self.kwargs.get("device", "cpu"))

            def _progress_cb(it, delta):
                if self._is_cancelled():
                    raise TrackingCancelledError("Segmentation cancelled.")
                it_i = int(it)
                with suppress(RuntimeError, TypeError, ValueError):
                    self.progress.emit(it_i, float(delta))
                if device_name == "mps":
                    # Napari and MPS share the main GPU. Briefly yielding the
                    # worker GIL keeps Qt input and canvas updates responsive.
                    sleep(0.002)

            kwargs = dict(self.kwargs)
            kwargs.pop("run_framewise", None)
            kwargs.pop("spatial_dim", None)
            downsample_xy_factor = int(kwargs.pop("postprocess_downsample_xy_factor", 1))
            downsample_target_xy_raw = kwargs.pop("postprocess_downsample_target_xy", None)
            downsample_target_xy = (
                tuple(int(value) for value in downsample_target_xy_raw)
                if downsample_target_xy_raw is not None
                else None
            )
            downsample_stats: list[dict[str, int]] = []

            def _postprocess_segmentation(segmentation_result):
                result = _squeeze_leading_singletons(np.asarray(segmentation_result), spatial_dim)
                if downsample_xy_factor > 1:
                    result, stats = _downsample_binary_xy_preserve_components(
                        result,
                        downsample_xy_factor,
                        target_xy=downsample_target_xy,
                    )
                    downsample_stats.append(stats)
                return result

            if run_framewise:
                image_source = kwargs.pop("image")
                frangi_source = kwargs.pop("frangi")
                image_shape = tuple(int(value) for value in image_source.shape)
                frangi_shape = tuple(int(value) for value in frangi_source.shape)
                if image_shape != frangi_shape:
                    raise ValueError(
                        f"image and frangi shapes differ: {image_shape} and {frangi_shape}."
                    )
                total = int(image_shape[0])
                seg = None
                info = None
                for idx in range(total):
                    if self._is_cancelled():
                        raise TrackingCancelledError("Segmentation cancelled.")
                    frame_kwargs = dict(kwargs)
                    frame_kwargs["image"] = np.asarray(
                        image_source[idx],
                        dtype=np.float32,
                    )
                    frame_kwargs["frangi"] = np.asarray(
                        frangi_source[idx],
                        dtype=np.float32,
                    )
                    frame_kwargs["progress"] = _progress_cb
                    seg_frame, frame_info = segmentation(**frame_kwargs)
                    seg_frame = _postprocess_segmentation(seg_frame)
                    if seg is None:
                        seg = np.empty((total, *seg_frame.shape), dtype=seg_frame.dtype)
                    seg[idx] = seg_frame
                    info = frame_info
                    del frame_kwargs, seg_frame, frame_info
                    with suppress(RuntimeError, TypeError, ValueError):
                        self.frame_progress.emit(idx + 1, total)
                with suppress(Exception):
                    gc.collect()
                if seg is None:
                    raise RuntimeError("No frames were segmented.")
                # Let the allocator reuse cached blocks between frames. Clearing
                # MPS/CUDA after every frame creates a visible stop-start cadence.
                self._cleanup_torch_device(device_name)
            else:
                kwargs["image"] = np.asarray(kwargs["image"], dtype=np.float32)
                kwargs["frangi"] = np.asarray(kwargs["frangi"], dtype=np.float32)
                kwargs["progress"] = _progress_cb
                seg, info = segmentation(**kwargs)
                self._cleanup_torch_device(device_name)
                seg = _postprocess_segmentation(seg)
            if info is not None and downsample_stats:
                info.downsample_xy_factor = downsample_xy_factor
                info.downsample_target_xy = downsample_target_xy
                info.components_before_downsample = sum(item["components_before"] for item in downsample_stats)
                info.components_after_downsample = sum(item["components_after"] for item in downsample_stats)
                info.downsample_relocated_seeds = sum(item["relocated_seeds"] for item in downsample_stats)
                info.downsample_frame_stats = tuple(dict(item) for item in downsample_stats)
            self.finished.emit((seg, info), None)
        except TrackingCancelledError as e:
            self.finished.emit(None, e)
        except Exception as e:  # noqa: BLE001 - worker boundary must always notify the UI
            self.finished.emit(None, e)


class ProximityWorker(QObject):
    finished = Signal(object, object)

    def __init__(self, layers, kwargs):
        super().__init__()
        # Snapshot layer properties on the GUI thread, not inside run().
        self.layers = [SimpleNamespace(data=layer.data, metadata=dict(layer.metadata), name=layer.name)
                       for layer in layers]
        self.kwargs = kwargs
        self._cancelled = False

    def cancel(self):
        self._cancelled = True

    def run(self):
        try:
            result = compute_proximity_result(*self.layers, **self.kwargs,
                                              cancel_check=lambda: self._cancelled)
            self.finished.emit(result, None)
        except Exception as exc:
            self.finished.emit(None, exc)


class FrangiWorker(QObject):
    finished = Signal(object, object)  # (payload, error)
    started = Signal()
    progress = Signal(int, int)  # (done, total)
    sigma_progress = Signal(int, int, int, int)  # (frame_done, frame_total, sigma_done, sigma_total)

    def __init__(
        self,
        image: np.ndarray,
        *,
        spatial_dim: int,
        sigmas: list[float],
        kernel_size: int,
        device: str,
        zx_ratio: float,
        temporal_dims: str | None,
        psf_ratio: float = 3.0,
        alpha: float = 0.5,
        beta: float = 0.5,
        gamma: float = 2.0,
        response_mode: str = "vesselness",
        response_specs: list[dict[str, Any]] | None = None,
    ):
        super().__init__()
        # Keep lazy/memmap-like image objects lazy until the worker thread
        # actually needs each frame. Eager np.asarray() here blocks the UI on
        # large TZYX stacks before progress can start.
        self.image = image
        self.spatial_dim = int(spatial_dim)
        self.sigmas = list(sigmas)
        self.kernel_size = int(kernel_size)
        self.device = str(device)
        self.zx_ratio = float(zx_ratio)
        self.psf_ratio = float(psf_ratio)
        self.alpha = float(alpha)
        self.beta = float(beta)
        self.gamma = float(gamma)
        self.response_mode = str(response_mode or "vesselness")
        if response_specs is None:
            response_specs = [
                {
                    "name": self.response_mode,
                    "response_mode": self.response_mode,
                    "sigmas": list(self.sigmas),
                }
            ]
        self.response_specs = [dict(spec) for spec in response_specs]
        self.temporal_dims = temporal_dims
        self._cancelled = False

    def cancel(self) -> None:
        self._cancelled = True

    def _is_cancelled(self) -> bool:
        thread = self.thread()
        return bool(self._cancelled or (thread is not None and thread.isInterruptionRequested()))

    def run(self):
        self.started.emit()
        try:
            FrangiFilter = _get_frangi_filter_cls()
            total = int(self.image.shape[0]) if self.temporal_dims else 1
            shape_obj = getattr(self.image, "shape", None)
            shape = tuple(shape_obj) if shape_obj is not None else tuple(np.asarray(self.image).shape)
            z_depth = int(shape[-3]) if self.spatial_dim == 3 and len(shape) >= 3 else 1
            z_pad = int(self.kernel_size // 2)
            # 3D reflect padding requires pad < input size on every spatial
            # axis. Thin Z subranges such as Z=2 with kernel radius 4 should
            # still be usable, so process each Z plane with the 2D Frangi filter
            # and stack the responses back into ZYX.
            process_z_as_2d = self.spatial_dim == 3 and z_depth <= z_pad
            logical_units = (
                total * z_depth
                if self.temporal_dims and process_z_as_2d
                else total
                if self.temporal_dims
                else z_depth
                if process_z_as_2d
                else 1
            )
            response_specs: list[dict[str, Any]] = []
            for spec in self.response_specs:
                sigmas = list(spec.get("sigmas") or self.sigmas)
                if not sigmas:
                    continue
                response_specs.append(
                    {
                        "name": str(spec.get("name") or spec.get("response_mode") or "frangi"),
                        "response_mode": str(spec.get("response_mode") or "vesselness"),
                        "sigmas": sigmas,
                    }
                )
            if not response_specs:
                raise ValueError("No structural response specifications were provided.")
            total_steps = max(
                int(sum(logical_units * len(spec["sigmas"]) for spec in response_specs)),
                1,
            )
            completed_offset = 0
            results: dict[str, np.ndarray] = {}
            completed_specs: list[dict[str, Any]] = []

            for spec in response_specs:
                sigmas = list(spec["sigmas"])
                response_mode = str(spec["response_mode"])
                if self.spatial_dim == 3 and not process_z_as_2d:
                    frangi_filter = FrangiFilter(
                        channels=1,
                        kernel_size=self.kernel_size,
                        sigmas=sigmas,
                        dim=3,
                        zx_ratio=self.zx_ratio,
                        psf_ratio=self.psf_ratio,
                        alpha=self.alpha,
                        beta=self.beta,
                        gamma=self.gamma,
                        response_mode=response_mode,
                        device=self.device,
                    )
                else:
                    frangi_filter = FrangiFilter(
                        channels=1,
                        kernel_size=self.kernel_size,
                        sigmas=sigmas,
                        dim=2,
                        alpha=self.alpha,
                        beta=self.beta,
                        gamma=self.gamma,
                        response_mode=response_mode,
                        device=self.device,
                    )
                _current_unit_idx = [0]

                def _sigma_cb(done, sigma_total):
                    if self._is_cancelled():
                        raise TrackingCancelledError("Structural extraction cancelled.")
                    done_steps = completed_offset + _current_unit_idx[0] * len(sigmas) + int(done)
                    with suppress(RuntimeError, TypeError, ValueError):
                        self.sigma_progress.emit(1, 1, int(done_steps), int(total_steps))

                frames = []
                if self.temporal_dims:
                    for idx in range(total):
                        if self._is_cancelled():
                            raise TrackingCancelledError("Structural extraction cancelled.")
                        frame = _normalize_0_255(self.image[idx])
                        if process_z_as_2d:
                            z_frames = []
                            for z_idx in range(int(frame.shape[0])):
                                if self._is_cancelled():
                                    raise TrackingCancelledError("Structural extraction cancelled.")
                                _current_unit_idx[0] = idx * z_depth + z_idx
                                out_z = frangi_filter(
                                    -np.expand_dims(frame[z_idx], 0),
                                    progress_callback=_sigma_cb,
                                )[0].detach().cpu().numpy()
                                z_frames.append(out_z)
                                with suppress(RuntimeError, TypeError, ValueError):
                                    self.progress.emit(
                                        completed_offset + (_current_unit_idx[0] + 1) * len(sigmas),
                                        total_steps,
                                    )
                            out = np.stack(z_frames, axis=0)
                        else:
                            _current_unit_idx[0] = idx
                            out = frangi_filter(
                                -np.expand_dims(frame, 0),
                                progress_callback=_sigma_cb,
                            )[0].detach().cpu().numpy()
                            with suppress(RuntimeError, TypeError, ValueError):
                                self.progress.emit(
                                    completed_offset + (idx + 1) * len(sigmas),
                                    total_steps,
                                )
                        frames.append(out)
                    result = np.stack(frames, axis=0)
                else:
                    frame = _normalize_0_255(self.image)
                    if process_z_as_2d:
                        z_frames = []
                        for z_idx in range(int(frame.shape[0])):
                            if self._is_cancelled():
                                raise TrackingCancelledError("Structural extraction cancelled.")
                            _current_unit_idx[0] = z_idx
                            out_z = frangi_filter(
                                -np.expand_dims(frame[z_idx], 0),
                                progress_callback=_sigma_cb,
                            )[0].detach().cpu().numpy()
                            z_frames.append(out_z)
                            with suppress(RuntimeError, TypeError, ValueError):
                                self.progress.emit(
                                    completed_offset + (z_idx + 1) * len(sigmas),
                                    total_steps,
                                )
                        result = np.stack(z_frames, axis=0)
                    else:
                        _current_unit_idx[0] = 0
                        result = frangi_filter(
                            -np.expand_dims(frame, 0),
                            progress_callback=_sigma_cb,
                        )[0].detach().cpu().numpy()
                        with suppress(RuntimeError, TypeError, ValueError):
                            self.progress.emit(completed_offset + len(sigmas), total_steps)

                results[str(spec["name"])] = result
                completed_specs.append(dict(spec))
                completed_offset += logical_units * len(sigmas)

            self.finished.emit({"responses": results, "specs": completed_specs}, None)
        except TrackingCancelledError as e:
            self.finished.emit(None, e)
        except Exception as e:  # noqa: BLE001 - worker boundary must always notify the UI
            self.finished.emit(None, e)


class _TrackingOperationWorker(QObject):
    finished = Signal(object, object)

    def __init__(
        self,
        segmentation_data: np.ndarray,
        intensity_data: np.ndarray,
        *,
        spacing: tuple[float, ...],
        config: TrackingConfig,
    ):
        super().__init__()
        self.segmentation_data = segmentation_data
        self.intensity_data = intensity_data
        self.spacing = tuple(float(v) for v in spacing)
        self.config = config
        self._cancelled = False

    def cancel(self) -> None:
        self._cancelled = True

    def _is_cancelled(self) -> bool:
        thread = self.thread()
        return bool(self._cancelled or (thread is not None and thread.isInterruptionRequested()))

    def _compute(self):
        raise NotImplementedError

    def run(self):
        try:
            result = self._compute()
        except TrackingCancelledError as e:
            self.finished.emit(None, e)
            return
        except Exception as e:  # noqa: BLE001 - worker boundary must always notify the UI
            self.finished.emit(None, e)
            return
        self.finished.emit(result, None)


class TrackingWorker(_TrackingOperationWorker):
    def _compute(self):
        return track_segmentations_ilp(
            self.segmentation_data,
            spacing=self.spacing,
            config=self.config,
            intensity_image=self.intensity_data,
            cancel_check=self._is_cancelled,
        )


class MatchWorker(_TrackingOperationWorker):
    def _compute(self):
        return compute_match_summary(
            self.segmentation_data,
            spacing=self.spacing,
            config=self.config,
            intensity_image=self.intensity_data,
            cancel_check=self._is_cancelled,
        )


class AnalysisWorker(QObject):
    finished = Signal(object, object)

    def __init__(
        self,
        binary: np.ndarray,
        *,
        unit: str,
        voxel_size: tuple[float, ...],
        frame_index: int | None,
        frame_key: tuple[int, int | None],
        request_id: int,
    ):
        super().__init__()
        self.binary = np.asarray(binary, dtype=bool)
        self.unit = str(unit)
        self.voxel_size = tuple(float(v) for v in voxel_size)
        self.frame_index = None if frame_index is None else int(frame_index)
        self.frame_key = frame_key
        self.request_id = int(request_id)
        self._cancelled = False

    def cancel(self) -> None:
        self._cancelled = True

    def _is_cancelled(self) -> bool:
        thread = self.thread()
        return bool(self._cancelled or (thread is not None and thread.isInterruptionRequested()))

    def run(self):
        try:
            labels, rows, plot_info, topology_cache = analyze_binary_components(
                self.binary,
                unit=self.unit,
                voxel_size=self.voxel_size,
                frame_index=self.frame_index,
                cancel_check=self._is_cancelled,
            )
        except AnalysisCancelledError as e:
            self.finished.emit(None, e)
            return
        except Exception as e:  # noqa: BLE001 - worker boundary must always notify the UI
            self.finished.emit(None, e)
            return
        payload = {
            "labels": labels,
            "rows": rows,
            "plot_info": plot_info,
            "topology_cache": topology_cache,
            "frame_key": self.frame_key,
            "request_id": self.request_id,
        }
        self.finished.emit(payload, None)


class AnalysisExportWorker(QObject):
    progress = Signal(int, int)
    finished = Signal(object, object)

    def __init__(
        self,
        segmentation_data: np.ndarray,
        *,
        dims_tag: str,
        unit: str,
        voxel_size_2d: tuple[float, float] | None,
        voxel_size_3d: tuple[float, float, float] | None,
        path: str,
        export_mode: str,
        min_object_size: int,
    ):
        super().__init__()
        self.segmentation_data = segmentation_data
        self.dims_tag = str(dims_tag)
        self.unit = str(unit)
        self.voxel_size_2d = voxel_size_2d
        self.voxel_size_3d = voxel_size_3d
        self.path = str(path)
        self.export_mode = str(export_mode)
        self.min_object_size = max(0, int(min_object_size))
        self._cancelled = False

    def cancel(self) -> None:
        self._cancelled = True

    def _is_cancelled(self) -> bool:
        thread = self.thread()
        return bool(self._cancelled or (thread is not None and thread.isInterruptionRequested()))

    @staticmethod
    def _columns_from_rows(rows: list[dict[str, float | int]]) -> list[str]:
        if not rows:
            return []
        preferred = [
            "frame",
            "pixels",
            "voxels",
            "area",
            "volume",
            "branch_number",
            "junction_number",
            "endpoint_number",
            "total_branch_length",
            "mean_branch_length",
            "length source",
        ]
        columns = [column for column in preferred if column in rows[0] and column != "label"]
        columns.extend(column for column in rows[0] if column not in columns and column != "label")
        return columns

    def _column_labels(self, rows: list[dict[str, float | int]]) -> list[str]:
        labels: list[str] = []
        for column in self._columns_from_rows(rows):
            if column == "frame":
                labels.append("frame")
            elif column == "pixels":
                labels.append("pixels (count)")
            elif column == "voxels":
                labels.append("voxels (count)")
            elif column == "area":
                labels.append(f"area ({self.unit}^2)")
            elif column == "volume":
                labels.append(f"volume ({self.unit}^3)")
            elif column == "total_branch_length":
                labels.append(f"total branch length ({self.unit})")
            elif column == "branch_number":
                labels.append("branch number")
            elif column == "junction_number":
                labels.append("junction number")
            elif column == "mean_branch_length":
                labels.append(f"mean branch length ({self.unit})")
            elif column == "endpoint_number":
                labels.append("endpoint number")
            elif column == "length source":
                labels.append("length source")
            else:
                labels.append(column)
        return labels

    def _branch_headers_and_rows(self, rows: list[dict[str, float | int]]) -> tuple[list[str], list[list[str]]]:
        if not rows:
            return [], []
        sample = rows[0]
        preferred = ["frame", "object_label", "branch_index", "branch_length"]
        columns = [column for column in preferred if column in sample and not str(column).startswith("_")]
        headers: list[str] = []
        for column in columns:
            if column == "frame":
                headers.append("frame")
            elif column == "object_label":
                headers.append("object label")
            elif column == "branch_index":
                headers.append("branch")
            elif column == "branch_length":
                headers.append(f"branch length ({self.unit})")
            else:
                headers.append(column)
        out_rows: list[list[str]] = []
        for row in rows:
            out_rows.append([f"{row[column]:.6g}" if isinstance(row[column], float) else str(row[column]) for column in columns])
        return headers, out_rows

    def run(self):
        try:
            if self.dims_tag not in {"TYX", "TZYX"}:
                raise ValueError("Selected layer is not temporal.")
            total_frames = int(self.segmentation_data.shape[0])
            self.progress.emit(0, total_frames)
            all_rows: list[dict[str, float | int]] = []
            all_branch_rows: list[dict[str, float | int]] = []
            for frame_index in range(total_frames):
                if self._is_cancelled():
                    raise AnalysisCancelledError("Analysis export cancelled.")
                binary = np.asarray(self.segmentation_data[frame_index]) > 0
                voxel_size = self.voxel_size_3d if binary.ndim == 3 else self.voxel_size_2d
                if voxel_size is None:
                    raise ValueError("Missing voxel size for analysis export.")
                _labels, rows, _plot_info, _topology_cache = analyze_binary_components(
                    binary,
                    unit=self.unit,
                    voxel_size=voxel_size,
                    frame_index=frame_index,
                    cancel_check=self._is_cancelled,
                )
                filtered_rows = _filter_analysis_rows_by_min_size(rows, self.min_object_size)
                if self.export_mode == "branches":
                    all_branch_rows.extend(
                        _analysis_branch_rows_from_rows(filtered_rows, _topology_cache, frame_index)
                    )
                else:
                    all_rows.extend(filtered_rows)
                self.progress.emit(frame_index + 1, total_frames)
            if self.export_mode == "branches":
                if not all_branch_rows:
                    raise ValueError("No branch lengths were found across frames.")
                headers, rows_out = self._branch_headers_and_rows(all_branch_rows)
            else:
                if not all_rows:
                    raise ValueError("No measurements were found across frames.")
                columns = self._columns_from_rows(all_rows)
                headers = self._column_labels(all_rows)
                rows_out = []
                for row in all_rows:
                    rows_out.append([f"{row[column]:.6g}" if isinstance(row[column], float) else str(row[column]) for column in columns])
            write_analysis_export_file(self.path, headers, rows_out)
        except AnalysisCancelledError as e:
            self.finished.emit(None, e)
            return
        except Exception as e:  # noqa: BLE001 - worker boundary must always notify the UI
            self.finished.emit(None, e)
            return
        self.finished.emit({"path": self.path, "export_mode": self.export_mode}, None)


# ---------------------------------------------------------------------
# Main widget
# ---------------------------------------------------------------------

class SIGMAWidget(SegmentAnalysisMixin, QWidget):
    _main_thread_dispatch = Signal(object, object)

    def __init__(self, napari_viewer):
        super().__init__()
        self._disposed = False
        self._hooks_installed = False
        self._viewer_connections = []
        self._menu_patch = None
        self._zoom_callback = None
        self._dock_widget = None
        self._ui_ready = False
        _prefer_numba_workqueue()
        self._main_thread_dispatch.connect(
            self._execute_main_thread_callback,
            Qt.QueuedConnection,
        )
        self.viewer = napari_viewer
        self._group_layers: list = []
        self._frangi_ctx: FrangiContext | None = None
        self._frangi_layer = None
        self._frangi_component_raw: dict[str, np.ndarray] = {}
        self._frangi_component_layers: dict[str, Any] = {}
        self._frangi_component_specs: dict[str, dict[str, Any]] = {}
        self._frangi_component_temporal_dims: str | None = None
        self._upsample_layer = None
        self._denoise_layer = None
        self._median_layer = None
        self._gaussian_layer = None
        self._segmentation_layer = None
        self._denoise_ctx: ProcessingContext | None = None
        self._view_initialized = False  # only on first load
        self._scale_conn = None
        self._outer_layout = None
        self._content_layout = None
        self._scroll = None
        self._top_grid = None
        self._info_grid = None
        self._denoise_grid = None
        self._frangi_grid = None
        self._seg_grid = None
        self._row_upsample_layout = None
        self._row_med_layout = None
        self._row_gaussian_layout = None
        self._analysis_summary = None
        self._analysis_table = None
        self._analysis_branch_table = None
        self._analysis_selected_table = None
        self._analysis_export_btn = None
        self._analysis_branch_export_btn = None
        self._analysis_export_progress = None
        self._analysis_layer_combo = None
        self._analysis_min_size_spin = None
        self._analysis_frame_label = None
        self._analysis_tab = None
        self._analysis_rows: list[dict[str, float | int]] = []
        self._analysis_all_rows: list[dict[str, float | int]] = []
        self._analysis_all_plot_info: dict[str, object] | None = None
        self._analysis_labels = None
        self._analysis_frame_state = None
        self._analysis_topology_cache: dict[int, dict[str, object]] = {}
        self._analysis_labels_layer = None
        self._analysis_highlight_layer = None
        self._analysis_branch_highlight_layer = None
        self._analysis_topology_skeleton_layer = None
        self._analysis_topology_markers_layer = None
        self._analysis_selected_ids: set[int] = set()
        self._analysis_cleanup_in_progress = False
        self._layer_combo_state = None
        self._analysis_thread = None
        self._analysis_worker = None
        self._analysis_export_thread = None
        self._analysis_export_worker = None
        self._analysis_request_id = 0
        self._analysis_pending_frame_key = None
        self._analysis_refresh_timer = None
        self._tracking_jump_timer = None
        self._tracking_pending_jump = None
        self._tracking_interaction_timer = None
        self._tracking_pending_interaction = None
        self._analysis_plot_canvas = None
        self._analysis_plot_scroll = None
        self._analysis_plot_figure = None
        self._analysis_plot_axes = None
        self._size_info_label = None
        self._time_info_label = None
        self.lbl_time = None
        self._proximity_tab = None
        self._proximity_source_raw_combo = None
        self._proximity_source_seg_combo = None
        self._proximity_target_raw_combo = None
        self._proximity_target_seg_combo = None
        self._proximity_summary_table = None
        self._proximity_components_table = None
        self._proximity_export_btn = None
        self._proximity_result: ProximityResult | None = None
        self._proximity_layer = None
        self._proximity_overlap_layer = None
        self._proximity_roi_layer = None
        self._proximity_roi_text_layer = None
        self._proximity_roi_status_label = None
        self._proximity_roi_3d_view_state = None
        self._proximity_rebuilding_layer_combos = False
        self._proximity_layer_selection_names: dict[str, str] = {}
        self._tracking_layer_combo = None
        self._tracking_raw_layer_combo = None
        self._tracking_linear_table = None
        self._tracking_fission_table = None
        self._tracking_fusion_table = None
        self._tracking_split_merge_table = None
        self._tracking_unlinked_box = None
        self._tracking_unlinked_table = None
        self._tracking_unlinked_display_frame: int | None = None
        self._tracking_tab = None
        self._tracking_layer = None
        self._tracking_layer_pending_data: np.ndarray | None = None
        self._tracking_lineages_dirty = False
        self._tracking_base_color_map = None
        self._tracking_event_layers: dict[str, Any] = {}
        self._tracking_event_layer_color_maps: dict[str, dict[int, np.ndarray]] = {}
        self._tracking_active_event_kind: str | None = None
        self._tracking_transition_combo = None
        self._tracking_event_kind_combo = None
        self._tracking_active_transition: tuple[int, int] | None = None
        self._tracking_export_btn = None
        self._tracking_statistics_export_btn = None
        self._tracking_statistics_import_btn = None
        self._tracking_event_base_layer = None
        self._tracking_event_base_source_layer = None
        self._tracking_event_base_source_was_visible: bool | None = None
        self._tracking_highlight_layer = None
        self._tracking_highlight_det_ids: set[int] = set()
        self._tracking_highlight_value_det_ids: dict[int, set[int]] = {}
        self._tracking_detection_lookup: dict[tuple[int, int], int] = {}
        self._tracking_result: TrackingResult | None = None
        self._tracking_result_source_layer = None
        self._tracking_config: TrackingConfig | None = None
        self._tracking_pending_event_import: (
            tuple[str, list[Any], list[list[Any]]] | None
        ) = None
        self._tracking_refine_object_combo = None
        self._tracking_refine_selected_table = None
        self._tracking_refine_available_table = None
        self._tracking_refine_apply_btn = None
        self._tracking_refine_remove_btn = None
        self._tracking_refine_undo_btn = None
        self._tracking_refine_undo_shortcut = None
        self._tracking_refine_pick_in_btn = None
        self._tracking_refine_pick_out_btn = None
        self._tracking_refine_pick_confirm_btn = None
        self._tracking_refine_pick_cancel_btn = None
        self._tracking_refine_image_pick_direction: str | None = None
        self._tracking_refine_image_pick_anchor_id: int | None = None
        self._tracking_refine_image_pick_det_id: int | None = None
        self._tracking_refine_undo_history: list[
            tuple[tuple[int, ...], int | None, dict[str, Any] | None]
        ] = []
        self._tracking_refine_status_label = None
        self._tracking_refine_det_id: int | None = None
        self._tracking_refine_event_scope: dict[str, Any] | None = None
        self._tracking_refine_syncing_tables = False
        self._tracking_match_summary: MatchSummary | None = None
        self._tracking_match_table = None
        self._tracking_match_src_overlay_layer = None
        self._tracking_match_dst_overlay_layer = None
        self._tracking_match_link_layer = None
        self._tracking_match_src_layer = None
        self._tracking_match_dst_layer = None
        self._tracking_match_hidden_layers = None
        self._tracking_display_events = []
        self._tracking_linear_events = []
        self._tracking_fission_events = []
        self._tracking_fusion_events = []
        self._tracking_split_merge_events = []
        self._tracking_thread = None
        self._tracking_worker = None
        self._frangi_thread = None
        self._frangi_worker = None
        self._upsample_thread = None
        self._upsample_worker = None
        self._denoise_thread = None
        self._denoise_worker = None
        self._gaussian_thread = None
        self._gaussian_worker = None
        self._seg_thread = None
        self._seg_worker = None
        self._proximity_thread = None
        self._proximity_worker = None
        self._panel_tabs = None
        self._initial_panel_width_applied = False

        # ---- Top-level: scroll area with a single content widget -------
        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        self._outer_layout = outer

        scroll = QScrollArea(self)
        scroll.setWidgetResizable(True)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        scroll.setAlignment(Qt.AlignTop | Qt.AlignLeft)
        scroll.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self._scroll = scroll
        content = QWidget()
        content.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Preferred)
        scroll.setWidget(content)

        layout = QVBoxLayout(content)
        self._content_layout = layout

        tabs = QTabWidget(content)
        tabs.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        tabs.addTab(self._build_segmentation_page(), "Segmentation")
        self._analysis_tab = self._build_analysis_page()
        tabs.addTab(self._analysis_tab, "Morphology Analysis")
        self._tracking_tab = self._build_tracking_page()
        tabs.addTab(self._tracking_tab, "Tracking")
        self._proximity_tab = self._build_proximity_page()
        tabs.addTab(self._proximity_tab, "Proximity Analysis")
        tabs.currentChanged.connect(self._on_panel_tab_changed)
        self._panel_tabs = tabs
        layout.addWidget(tabs, 0, Qt.AlignTop)

        outer.addWidget(scroll)

        # ---- listeners -------------------------------------------------
        self._install_viewer_hooks()

        # initial state
        self._configure_responsive_ui()
        self._update_info()
        self._fit_view_and_scalebar()
        self._disable_wheel_adjustment_on_inputs()
        self._analysis_refresh_timer = QTimer(self)
        self._analysis_refresh_timer.setSingleShot(True)
        self._analysis_refresh_timer.setInterval(120)
        self._analysis_refresh_timer.timeout.connect(self._run_scheduled_analysis_refresh)
        self._tracking_jump_timer = QTimer(self)
        self._tracking_jump_timer.setSingleShot(True)
        self._tracking_jump_timer.setInterval(180)
        self._tracking_jump_timer.timeout.connect(self._apply_pending_tracking_time_jump)
        self._tracking_interaction_timer = QTimer(self)
        self._tracking_interaction_timer.setSingleShot(True)
        self._tracking_interaction_timer.setInterval(90)
        self._tracking_interaction_timer.timeout.connect(self._run_pending_tracking_interaction)
        QTimer.singleShot(0, self._sync_standard_control_heights)
        QTimer.singleShot(0, self._sync_panel_tab_height)
        QTimer.singleShot(0, self._normalize_existing_layers_on_startup)
        QTimer.singleShot(0, self._apply_initial_panel_width)
        self._ui_ready = True

    def _install_viewer_hooks(self) -> None:
        if self._hooks_installed:
            return
        self._disposed = False
        self._viewer_connections = [
            (self.viewer.layers.events.inserted, self._on_layer_inserted),
            (self.viewer.layers.events.removed, self._update_info),
            (self.viewer.layers.selection.events.active, self._on_active_changed),
            (self.viewer.dims.events.current_step, self._on_dims_step_changed),
        ]
        for emitter, callback in self._viewer_connections:
            emitter.connect(callback)
        self._disable_double_click_zoom()
        if self._on_analysis_viewer_double_click not in self.viewer.mouse_double_click_callbacks:
            self.viewer.mouse_double_click_callbacks.append(self._on_analysis_viewer_double_click)
        self._patch_layer_list_context_menu()
        self._hooks_installed = True

    def dispose(self) -> None:
        """Stop work and release only hooks owned by this panel; safe to repeat."""
        if self._disposed:
            return
        self._disposed = True
        for timer in self.findChildren(QTimer):
            timer.stop()
        for name in ("analysis", "analysis_export", "tracking", "frangi", "upsample",
                     "denoise", "gaussian", "seg", "proximity"):
            thread_attr, worker_attr = f"_{name}_thread", f"_{name}_worker"
            self._request_worker_cancel(thread_attr, worker_attr)
            self._cleanup_worker_thread(thread_attr, worker_attr)
        for emitter, callback in self._viewer_connections:
            with suppress(Exception):
                emitter.disconnect(callback)
        self._viewer_connections.clear()
        if self._scale_conn is not None:
            with suppress(Exception):
                self._scale_conn[0].disconnect(self._scale_conn[1])
            self._scale_conn = None
        with suppress(ValueError):
            self.viewer.mouse_double_click_callbacks.remove(self._on_analysis_viewer_double_click)
        if self._zoom_callback is not None and self._zoom_callback not in self.viewer.mouse_double_click_callbacks:
            self.viewer.mouse_double_click_callbacks.append(self._zoom_callback)
        self._zoom_callback = None
        if self._menu_patch is not None:
            delegate, original, installed = self._menu_patch
            with suppress(RuntimeError, AttributeError):
                if delegate.show_context_menu == installed:
                    delegate.show_context_menu = original
                    delegate._segment_any_save_menu_patched = False
                    menu = getattr(delegate, "_context_menu", None)
                    if menu is not None:
                        for action in list(menu.actions()):
                            if action.property("segment_any_save_action") or action.property("segment_any_save_separator"):
                                menu.removeAction(action)
                                action.deleteLater()
            self._menu_patch = None
        self._hooks_installed = False

    def event(self, event):
        if event.type() == QEvent.ParentChange and getattr(self, "_ui_ready", False):
            parent = self.parentWidget()
            if isinstance(parent, QDockWidget):
                self._dock_widget = parent
                parent.installEventFilter(self)
                self._install_viewer_hooks()
            elif parent is None and self._dock_widget is not None:
                self._dock_widget = None
                self.dispose()
        return super().event(event)

    def closeEvent(self, event):
        self.dispose()
        super().closeEvent(event)

    def showEvent(self, event):
        if self._ui_ready and self._disposed:
            self._install_viewer_hooks()
        super().showEvent(event)

    def _disable_double_click_zoom(self) -> None:
        with suppress(ValueError, AttributeError, ImportError, RuntimeError):
            from napari.components._viewer_mouse_bindings import (
                double_click_to_zoom,
            )
            self.viewer.mouse_double_click_callbacks.remove(double_click_to_zoom)
            self._zoom_callback = double_click_to_zoom

    def _thread_is_running(self, attr_name: str) -> bool:
        thread = getattr(self, attr_name, None)
        return bool(thread is not None and thread.isRunning())

    def _execute_main_thread_callback(self, callback, args) -> None:
        if not self._disposed:
            callback(*tuple(args))

    def _connect_worker_callback(self, signal, callback) -> None:
        signal.connect(
            lambda *args: self._main_thread_dispatch.emit(callback, args)
        )

    def _request_worker_cancel(self, thread_attr: str, worker_attr: str) -> None:
        thread = getattr(self, thread_attr, None)
        worker = getattr(self, worker_attr, None)
        if thread is None and worker is None:
            return
        if worker is not None and hasattr(worker, "cancel"):
            with suppress(Exception):
                worker.cancel()
        if thread is not None:
            with suppress(Exception):
                thread.requestInterruption()

    def _cleanup_worker_thread(self, thread_attr: str, worker_attr: str) -> None:
        thread = getattr(self, thread_attr, None)
        worker = getattr(self, worker_attr, None)
        if worker is not None:
            with suppress(Exception):
                worker.deleteLater()
        if thread is not None:
            with suppress(Exception):
                thread.quit()
            with suppress(Exception):
                thread.wait()
        if thread is not None:
            with suppress(Exception):
                thread.deleteLater()
        setattr(self, worker_attr, None)
        setattr(self, thread_attr, None)

    def _start_analysis_refresh(self, layer, *, frame_key: tuple[int, int | None]) -> None:
        frame_index = self._analysis_current_frame_index(layer)
        binary = _component_mask_from_layer(layer, frame_index=frame_index)
        unit = _get_units_from_layer(layer)
        voxel_size = _get_pixel_size_tuple(layer, binary.ndim)

        self._analysis_request_id += 1
        request_id = int(self._analysis_request_id)
        self._analysis_pending_frame_key = frame_key

        current_thread = self._analysis_thread
        current_worker = self._analysis_worker
        if current_worker is not None and hasattr(current_worker, "cancel"):
            with suppress(Exception):
                current_worker.cancel()
        if current_thread is not None:
            with suppress(Exception):
                current_thread.requestInterruption()

        thread = QThread(self)
        worker = AnalysisWorker(
            binary,
            unit=unit,
            voxel_size=voxel_size,
            frame_index=frame_index,
            frame_key=frame_key,
            request_id=request_id,
        )
        worker.moveToThread(thread)
        thread.started.connect(worker.run)

        def _finish(payload, error, *, this_thread=thread, this_worker=worker, expected_request_id=request_id):
            is_latest = self._analysis_thread is this_thread and self._analysis_worker is this_worker
            if is_latest:
                self._analysis_thread = None
                self._analysis_worker = None
            with suppress(Exception):
                this_thread.quit()
            with suppress(Exception):
                this_thread.wait()
            with suppress(Exception):
                this_worker.deleteLater()
            with suppress(Exception):
                this_thread.deleteLater()

            if payload is None:
                if isinstance(error, AnalysisCancelledError):
                    return
                if expected_request_id == self._analysis_request_id:
                    self._analysis_pending_frame_key = None
                return

            if int(payload.get("request_id", -1)) != self._analysis_request_id:
                return
            current_layer = self._selected_analysis_layer()
            if current_layer is None:
                self._analysis_pending_frame_key = None
                return
            current_frame_key = self._analysis_frame_key(current_layer)
            if tuple(payload.get("frame_key")) != tuple(current_frame_key):
                return
            self._analysis_pending_frame_key = None
            self._clear_analysis_results(reset_outputs=False)
            self._apply_analysis_result(
                current_layer,
                payload["labels"],
                payload["rows"],
                payload["plot_info"],
                topology_cache=payload.get("topology_cache"),
            )

        self._connect_worker_callback(worker.finished, _finish)
        self._analysis_thread = thread
        self._analysis_worker = worker
        thread.start()

    def _start_analysis_export_all_frames(self, layer, path: str, *, export_mode: str) -> None:
        if self._analysis_export_thread is not None and self._analysis_export_worker is not None:
            self._request_worker_cancel("_analysis_export_thread", "_analysis_export_worker")
            self._set_analysis_export_running_ui(True)
            return

        dims_tag = _layer_dims_tag(layer)
        data = layer.data
        unit = _get_units_from_layer(layer)
        voxel_size_2d: tuple[float, float] | None = None
        voxel_size_3d: tuple[float, float, float] | None = None
        with suppress(Exception):
            voxel_size_2d = _get_pixel_size_tuple(layer, 2)
        with suppress(Exception):
            voxel_size_3d = _get_pixel_size_tuple(layer, 3)

        thread = QThread(self)
        worker = AnalysisExportWorker(
            data,
            dims_tag=dims_tag,
            unit=unit,
            voxel_size_2d=voxel_size_2d,
            voxel_size_3d=voxel_size_3d,
            path=path,
            export_mode=export_mode,
            min_object_size=max(0, int(getattr(self, "_analysis_min_size_spin", None).value())) if getattr(self, "_analysis_min_size_spin", None) is not None else 0,
        )
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        self._analysis_export_thread = thread
        self._analysis_export_worker = worker
        self._set_analysis_export_running_ui(True)

        def _update_progress(done: int, total: int) -> None:
            if self._analysis_export_progress is None:
                return
            self._analysis_export_progress.setMaximum(max(int(total), 1))
            self._analysis_export_progress.setValue(max(0, int(done)))
            self._analysis_export_progress.setFormat(f"Exporting {int(done)}/{int(total)} frames")

        def _finish(payload, error, *, this_thread=thread, this_worker=worker):
            is_latest = self._analysis_export_thread is this_thread and self._analysis_export_worker is this_worker
            if is_latest:
                self._analysis_export_thread = None
                self._analysis_export_worker = None
            with suppress(Exception):
                this_thread.quit()
            with suppress(Exception):
                this_thread.wait()
            with suppress(Exception):
                this_worker.deleteLater()
            with suppress(Exception):
                this_thread.deleteLater()
            self._set_analysis_export_idle_ui()
            if payload is None:
                if isinstance(error, AnalysisCancelledError):
                    return
                if isinstance(error, RuntimeError) and "openpyxl" in str(error).lower():
                    QMessageBox.warning(self, "XLSX unavailable", "openpyxl is not available in this environment. Please export as CSV/TXT or install openpyxl.")
                    return
                QMessageBox.critical(self, "Export failed", f"Failed to export all frames: {error!r}")
                return
            noun = "branch lengths" if payload.get("export_mode") == "branches" else "measurements"
            QMessageBox.information(self, "Export complete", f"Saved {noun} to:\n{payload['path']}")

        self._connect_worker_callback(worker.progress, _update_progress)
        self._connect_worker_callback(worker.finished, _finish)
        thread.start()

    def _set_tracking_idle_ui(self) -> None:
        if self._tracking_run_btn is not None:
            self._tracking_run_btn.setEnabled(True)
            self._tracking_run_btn.setText("Run Tracking")
        if self._tracking_match_btn is not None:
            self._tracking_match_btn.setEnabled(True)
            self._tracking_match_btn.setText("Match")
        for combo in (
            self._tracking_layer_combo,
            self._tracking_raw_layer_combo,
        ):
            if combo is not None:
                combo.setEnabled(combo.count() > 0)
        has_result = bool(self._tracking_result and self._tracking_result.detections)
        if self._tracking_refine_object_combo is not None:
            self._tracking_refine_object_combo.setEnabled(has_result)
        if self._tracking_refine_selected_table is not None:
            self._tracking_refine_selected_table.setEnabled(
                has_result and self._tracking_refine_selected_table.rowCount() > 0
            )
        if self._tracking_refine_available_table is not None:
            self._tracking_refine_available_table.setEnabled(
                has_result and self._tracking_refine_available_table.rowCount() > 0
            )
        if self._tracking_statistics_export_btn is not None:
            self._tracking_statistics_export_btn.setEnabled(
                bool(self._tracking_display_events)
            )
        if self._tracking_statistics_import_btn is not None:
            self._tracking_statistics_import_btn.setEnabled(
                self._selected_tracking_layer() is not None
            )
        self._update_tracking_refine_remove_enabled()
        self._update_tracking_refine_undo_enabled()

    def _set_analysis_export_idle_ui(self) -> None:
        if self._analysis_export_btn is not None:
            self._analysis_export_btn.setEnabled(True)
            self._analysis_export_btn.setText("Export")
        if self._analysis_branch_export_btn is not None:
            self._analysis_branch_export_btn.setEnabled(True)
            self._analysis_branch_export_btn.setText("Export")
        if self._analysis_export_progress is not None:
            self._analysis_export_progress.setValue(0)
            self._analysis_export_progress.setMaximum(1)
            self._analysis_export_progress.setFormat("")
            self._analysis_export_progress.hide()

    def _set_tracking_running_ui(self, mode: str) -> None:
        if self._tracking_run_btn is not None:
            self._tracking_run_btn.setEnabled(mode == "tracking")
            self._tracking_run_btn.setText("Stop Tracking" if mode == "tracking" else "Run Tracking")
        if self._tracking_match_btn is not None:
            self._tracking_match_btn.setEnabled(mode == "match")
            self._tracking_match_btn.setText("Stop Match" if mode == "match" else "Match")
        for control in (
            self._tracking_layer_combo,
            self._tracking_raw_layer_combo,
            self._tracking_refine_object_combo,
            self._tracking_refine_selected_table,
            self._tracking_refine_available_table,
            self._tracking_refine_apply_btn,
            self._tracking_refine_remove_btn,
            self._tracking_refine_undo_btn,
            self._tracking_statistics_import_btn,
        ):
            if control is not None:
                control.setEnabled(False)

    def _set_analysis_export_running_ui(self, running: bool) -> None:
        if self._analysis_export_btn is not None:
            self._analysis_export_btn.setEnabled(True)
            self._analysis_export_btn.setText("Stop Export" if running else "Export")
        if self._analysis_branch_export_btn is not None:
            self._analysis_branch_export_btn.setEnabled(True)
            self._analysis_branch_export_btn.setText("Stop Export" if running else "Export")
        if self._analysis_export_progress is not None:
            if running:
                self._analysis_export_progress.setMaximum(1)
                self._analysis_export_progress.setValue(0)
                self._analysis_export_progress.setFormat("Preparing export...")
                self._analysis_export_progress.show()
            else:
                self._analysis_export_progress.setValue(0)
                self._analysis_export_progress.setMaximum(1)
                self._analysis_export_progress.setFormat("")
                self._analysis_export_progress.hide()

    def eventFilter(self, obj: QObject, event: QEvent) -> bool:
        if obj is self._dock_widget:
            if event.type() == QEvent.Close:
                self.dispose()
            elif event.type() == QEvent.Show and self._ui_ready and self._disposed:
                self._install_viewer_hooks()
        if event.type() == QEvent.Wheel:
            plot_viewport = (
                self._analysis_plot_scroll.viewport()
                if self._analysis_plot_scroll is not None
                else None
            )
            if obj is self._analysis_plot_canvas or obj is plot_viewport:
                if self._scroll_area_from_wheel(self._analysis_plot_scroll, event):
                    return True
                if self._scroll_area_from_wheel(self._scroll, event):
                    return True
            if isinstance(obj, QSpinBox | QDoubleSpinBox | QComboBox):
                if self._scroll_area_from_wheel(self._scroll, event):
                    return True
                event.ignore()
                return True
        return super().eventFilter(obj, event)

    def _scroll_area_from_wheel(
        self,
        scroll: QScrollArea | None,
        event: QEvent,
    ) -> bool:
        if scroll is None:
            return False
        bar = scroll.verticalScrollBar()
        if bar is None or bar.maximum() <= bar.minimum():
            return False
        pixel_delta = int(event.pixelDelta().y())
        if pixel_delta:
            amount = pixel_delta
        else:
            wheel_steps = float(event.angleDelta().y()) / 120.0
            if wheel_steps == 0:
                return False
            amount = int(
                round(
                    wheel_steps
                    * max(bar.singleStep() * 3, self._scaled_px(48))
                )
            )
        previous = bar.value()
        bar.setValue(previous - amount)
        if bar.value() == previous:
            return False
        event.accept()
        return True

    def _disable_wheel_adjustment_on_inputs(self) -> None:
        for widget_type in (QSpinBox, QDoubleSpinBox, QComboBox):
            for widget in self.findChildren(widget_type):
                widget.installEventFilter(self)
                if isinstance(widget, QComboBox):
                    self._configure_combo_popup(widget)

    def _configure_combo_popup(self, combo: QComboBox) -> None:
        if combo is None:
            return
        combo.setMaxVisibleItems(5)
        view = combo.view()
        if view is None:
            return
        row_h = max(self._scaled_px(26), view.sizeHintForRow(0) if view.model() and view.model().rowCount() else 0)
        view.setMinimumHeight(row_h * 8 + self._scaled_px(8))

    def _patch_layer_list_context_menu(self) -> None:
        if not self._try_patch_layer_list_context_menu():
            QTimer.singleShot(0, self._try_patch_layer_list_context_menu)

    def _try_patch_layer_list_context_menu(self) -> bool:
        if self._disposed:
            return False
        with suppress(AttributeError, RuntimeError):
            from napari._app_model.constants import MenuId
            from napari._app_model.context import get_context
            from napari._qt._qapp_model import build_qmodel_menu

            layer_list = self.viewer.window._qt_viewer.layers
            delegate = layer_list.itemDelegate()
            if delegate is None:
                return False
            if getattr(delegate, "_segment_any_save_menu_patched", False):
                return True

            def _patched_show_context_menu(this_delegate, index, model, pos, parent):
                if not hasattr(this_delegate, "_context_menu"):
                    this_delegate._context_menu = build_qmodel_menu(MenuId.LAYERLIST_CONTEXT, parent=parent)
                layer_list_root = model.sourceModel()._root
                ctx = get_context(layer_list_root)
                this_delegate._context_menu.update_from_context(ctx)
                for action in list(this_delegate._context_menu.actions()):
                    if action.property("segment_any_save_action") or action.isSeparator() and action.property("segment_any_save_separator"):
                        this_delegate._context_menu.removeAction(action)
                separator = this_delegate._context_menu.addSeparator()
                separator.setProperty("segment_any_save_separator", True)
                save_action = this_delegate._context_menu.addAction("Save Layer with Metadata...")
                save_action.setProperty("segment_any_save_action", True)
                global_pos = pos if isinstance(pos, QPoint) else parent.mapToGlobal(pos)
                chosen = this_delegate._context_menu.exec_(global_pos)
                if chosen is save_action:
                    self._save_active_layer_with_metadata()

            original = delegate.show_context_menu
            installed = MethodType(_patched_show_context_menu, delegate)
            delegate.show_context_menu = installed
            self._menu_patch = (delegate, original, installed)
            delegate._segment_any_save_menu_patched = True
            return True
        return False

    def _scaled_px(self, px: int) -> int:
        screen = self.screen() or QApplication.primaryScreen()
        dpi = float(screen.logicalDotsPerInch()) if screen is not None else 96.0
        return max(1, int(round(px * max(dpi / 96.0, 0.9))))

    def sizeHint(self):
        hint = super().sizeHint()
        hint.setWidth(max(int(hint.width()), self._scaled_px(780)))
        return hint

    def _apply_initial_panel_width(self) -> None:
        if self._initial_panel_width_applied:
            return
        target_width = self._scaled_px(780)
        self.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Expanding)
        if self.width() < target_width:
            self.resize(target_width, self.height())

        parent = self.parentWidget()
        while parent is not None and not isinstance(parent, QDockWidget):
            parent = parent.parentWidget()
        if parent is not None:
            # Napari applies a vertical Maximum policy after constructing
            # right-side plugin widgets, so restore the scrollable dock here.
            parent.setMaximumHeight(16777215)
            dock_policy = parent.sizePolicy()
            dock_policy.setVerticalPolicy(QSizePolicy.Expanding)
            parent.setSizePolicy(dock_policy)
            main_window = parent.parentWidget()
            if main_window is not None and hasattr(main_window, "resizeDocks"):
                with suppress(TypeError, RuntimeError):
                    main_window.resizeDocks([parent], [target_width], Qt.Horizontal)
                with suppress(TypeError, RuntimeError):
                    main_window.resizeDocks(
                        [parent],
                        [main_window.height()],
                        Qt.Vertical,
                    )
            elif parent.width() < target_width:
                parent.resize(target_width, parent.height())
            parent.updateGeometry()
        self._initial_panel_width_applied = True

    def _configure_grid(self, grid: QGridLayout, h_spacing: int, v_spacing: int) -> None:
        grid.setHorizontalSpacing(self._scaled_px(h_spacing))
        grid.setVerticalSpacing(self._scaled_px(v_spacing))

    def _sync_panel_tab_height(self, *_args) -> None:
        tabs = self._panel_tabs
        if tabs is None:
            return
        current_page = tabs.currentWidget()
        if current_page is None:
            return
        tab_bar = tabs.tabBar()
        tab_bar_height = tab_bar.sizeHint().height() if tab_bar is not None else 0
        page_height = max(current_page.sizeHint().height(), current_page.minimumSizeHint().height())
        target_height = page_height + tab_bar_height + self._scaled_px(8)
        tabs.setMinimumHeight(target_height)
        tabs.setMaximumHeight(target_height)

    def _on_panel_tab_changed(self, *_args) -> None:
        self._sync_panel_tab_height()
        if self._panel_tabs is not None and self._panel_tabs.currentWidget() is self._analysis_tab:
            if self._ensure_analysis_plot():
                self._plot_distributions(getattr(self, "_analysis_pending_plot_info", None))
            with suppress(Exception):
                self._schedule_analysis_refresh_for_current_frame(force=False)
        scroll = getattr(self, "_scroll", None)
        if scroll is None:
            return

        def _scroll_to_top() -> None:
            with suppress(Exception):
                bar = scroll.verticalScrollBar()
                if bar is not None:
                    bar.setValue(bar.minimum())

        QTimer.singleShot(0, _scroll_to_top)

    def _build_panel_page(self, *boxes: QGroupBox) -> QWidget:
        page = QWidget()
        page.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        layout = QVBoxLayout(page)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(self._scaled_px(8))
        container = QWidget(page)
        container.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        container_layout = QVBoxLayout(container)
        container_layout.setContentsMargins(0, 0, 0, 0)
        container_layout.setSpacing(self._scaled_px(8))
        for box in boxes:
            self._style_group_box(box)
            container_layout.addWidget(box)
        layout.addWidget(container, 0, Qt.AlignTop)
        return page

    def _schedule_analysis_refresh_for_current_frame(self, *, force: bool = False) -> None:
        if self._panel_tabs is None or self._analysis_tab is None:
            return
        if self._panel_tabs.currentWidget() is not self._analysis_tab:
            return
        layer = getattr(self, "_selected_analysis_layer", lambda: None)()
        if layer is None:
            return
        if not force and self._analysis_labels_layer is None and not self._analysis_rows:
            return
        timer = self._analysis_refresh_timer
        if timer is None:
            return
        timer.setProperty("force_refresh", bool(force))
        timer.start()

    def _run_scheduled_analysis_refresh(self) -> None:
        timer = self._analysis_refresh_timer
        force = bool(timer.property("force_refresh")) if timer is not None else False
        with suppress(Exception):
            self._refresh_analysis_for_current_frame_if_needed(force=force)

    def _panel_font(self, point_size: float, *, bold: bool = False) -> QFont:
        font = QFont(self.font())
        font.setPointSizeF(point_size)
        font.setBold(bold)
        return font

    def _set_panel_control_font(
        self,
        widget,
        point_size: float,
        *,
        padding: str,
    ) -> None:
        widget.setFont(self._panel_font(point_size))
        base_style = widget.property("sigma_base_stylesheet")
        if base_style is None:
            base_style = widget.styleSheet()
            widget.setProperty("sigma_base_stylesheet", base_style)
        selector = type(widget).__name__
        override = (
            f"{selector} {{ font-size: {point_size:g}pt; "
            f"padding: {padding}; min-height: 0px; }}"
        )
        widget.setStyleSheet(f"{base_style}\n{override}" if base_style else override)

    def _apply_panel_typography(self, compact: bool) -> None:
        base = 10.0 if compact else 11.0
        control = base
        small = max(9.0, base - 0.5)
        strong = base + 0.5
        title = base + 1.0

        tabs = self.findChild(QTabWidget)
        if tabs is not None:
            tabs.setFont(self._panel_font(strong, bold=True))

        for box in self.findChildren(QGroupBox):
            box.setFont(self._panel_font(title, bold=True))

        for label in self.findChildren(QLabel):
            label.setFont(self._panel_font(base))

        for btn in self.findChildren(QPushButton):
            self._set_panel_control_font(
                btn,
                control,
                padding="0px 4px",
            )

        for combo in self.findChildren(QComboBox):
            self._set_panel_control_font(
                combo,
                control,
                padding="0px 10px 0px 8px",
            )
            self._configure_combo_popup(combo)

        for spin_type in (QSpinBox, QDoubleSpinBox):
            for spin in self.findChildren(spin_type):
                self._set_panel_control_font(
                    spin,
                    control,
                    padding="0px 10px",
                )

        for edit in self.findChildren(QLineEdit):
            self._set_panel_control_font(
                edit,
                control,
                padding="0px 2px",
            )
            if getattr(edit, "isReadOnly", lambda: False)():
                edit.setAttribute(Qt.WA_InputMethodEnabled, False)

        for edit in self.findChildren(QTextEdit):
            edit.setFont(self._panel_font(base))
            if getattr(edit, "isReadOnly", lambda: False)():
                edit.setAttribute(Qt.WA_InputMethodEnabled, False)

        for table in self.findChildren(QTableWidget):
            table.setFont(self._panel_font(base))
            header = table.horizontalHeader()
            if header is not None:
                header.setFont(self._panel_font(small, bold=True))

    def _sync_standard_control_heights(self) -> None:
        target_height = 20
        for widget_type in (QPushButton, QComboBox, QSpinBox, QDoubleSpinBox):
            for widget in self.findChildren(widget_type):
                widget.setFixedHeight(target_height)
        if self.path_edit is not None:
            self.path_edit.setFixedHeight(target_height)

    def _style_group_box(self, box: QGroupBox) -> None:
        if box is None:
            return
        box.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Maximum)
        box.setStyleSheet(
            "QGroupBox {"
            " margin-top: 8px;"
            " padding-top: 4px;"
            "}"
            "QGroupBox::title {"
            " subcontrol-origin: margin;"
            " left: 8px;"
            " padding: 0 3px;"
            "}"
        )

    def _style_toolbar_layout(self, layout: QHBoxLayout) -> None:
        if layout is None:
            return
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(self._scaled_px(4))

    def _style_table_widget(
        self,
        table: QTableWidget,
        *,
        min_height: int,
        max_height: int,
        expanding: bool = False,
        ) -> None:
        if table is None:
            return
        scrollbar_allowance = self._scaled_px(36)
        table.setAlternatingRowColors(True)
        table.setShowGrid(False)
        table.setWordWrap(False)
        table.setStyleSheet(
            "QTableWidget {"
            " background-color: #232831;"
            " alternate-background-color: #2a303a;"
            " selection-background-color: #4a5568;"
            " selection-color: #f3f4f6;"
            "}"
            "QHeaderView::section {"
            " background-color: #2d333d;"
            " border: 0;"
            " padding: 4px 6px;"
            "}"
        )
        table.setMinimumHeight(self._scaled_px(min_height) + scrollbar_allowance)
        table.setMaximumHeight(self._scaled_px(max_height) + scrollbar_allowance)
        table.setSizePolicy(
            QSizePolicy.Expanding,
            QSizePolicy.Expanding if expanding else QSizePolicy.Fixed,
        )
        table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        table.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel)
        header = table.horizontalHeader()
        if header is not None:
            header.setStretchLastSection(False)
            header.setMinimumSectionSize(self._scaled_px(64))
            header.setDefaultSectionSize(self._scaled_px(96))
            header.setFixedHeight(self._scaled_px(24))
            header.setDefaultAlignment(Qt.AlignCenter)
        vheader = table.verticalHeader()
        if vheader is not None:
            vheader.setVisible(False)
            vheader.setDefaultSectionSize(self._scaled_px(24))

    def _set_strict_fixed_width(self, widget, width: int) -> None:
        if widget is None:
            return
        px = self._scaled_px(width)
        widget.setMinimumWidth(px)
        widget.setMaximumWidth(px)
        widget.setSizePolicy(QSizePolicy.Fixed, QSizePolicy.Fixed)

    def _set_adaptive_width(self, widget, min_width: int) -> None:
        if widget is None:
            return
        px = self._scaled_px(min_width)
        widget.setMinimumWidth(px)
        widget.setMaximumWidth(16777215)
        widget.setSizePolicy(QSizePolicy.MinimumExpanding, QSizePolicy.Fixed)

    def _configure_responsive_ui(self) -> None:
        current_width = int(self.width()) if self.width() > 0 else self._scaled_px(360)
        compact = current_width <= self._scaled_px(520)

        self.setMinimumWidth(self._scaled_px(680))
        self.setMaximumWidth(16777215)
        self.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Expanding)
        if self._content_layout is not None:
            margin = self._scaled_px(6 if compact else 8)
            self._content_layout.setContentsMargins(margin, margin, margin, margin)
            self._content_layout.setSpacing(self._scaled_px(8 if compact else 10))

        for grid, h_spacing, v_spacing in (
            (self._top_grid, 8, 4),
            (self._info_grid, 8, 4),
            (self._denoise_grid, 8, 4),
            (self._frangi_grid, 8, 4),
            (self._seg_grid, 10, 6),
        ):
            if grid is not None:
                self._configure_grid(grid, h_spacing if not compact else max(h_spacing - 2, 6), v_spacing)

        for row_layout in (
            self._row_upsample_layout,
            self._row_med_layout,
            self._row_gaussian_layout,
        ):
            if row_layout is not None:
                row_layout.setSpacing(self._scaled_px(4 if compact else 6))

        for btn in (
            self.open_btn,
            self.btn_apply_voxel,
            self.btn_apply_upsample,
            self.btn_apply_denoise,
            self.btn_apply_gaussian,
            self.apply_btn,
            self.btn_view_vessel_frangi,
            self.btn_view_sheet_frangi,
            self.btn_view_max_frangi,
            self.seg_btn,
            self.btn_apply_time_range,
            self.btn_apply_slice_range,
        ):
            if btn is not None:
                btn.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Fixed)

        for widget in (
            self.device_combo,
            self.cmb_kernel,
            self.c_spin,
            self.t_spin,
            self.t_start_spin,
            self.t_end_spin,
            self.slice_start_spin,
            self.slice_end_spin,
            self.edit_vz,
            self.edit_vyx,
            self.upsample_factor_spin,
            self.kernel_spin,
            self.sigma_count,
            self.sigma_min,
            self.sigma_max,
            self.vessel_rescale_spin,
            self.sheet_sigma_count,
            self.sheet_sigma_min,
            self.sheet_sigma_max,
            self.sheet_rescale_spin,
            self.psf_ratio_spin,
            self.frangi_response_mode_combo,
            self.beta1_spin,
            self.beta2_spin,
            self.nfore_spin,
            self.nback_spin,
            self.em_fg_sample_combo,
            self.em_bg_sample_combo,
            self.maxiter_spin,
            self.init_combo,
            self._tracking_layer_combo,
            self._tracking_raw_layer_combo,
            self._proximity_source_raw_combo,
            self._proximity_source_seg_combo,
            self._proximity_target_raw_combo,
            self._proximity_target_seg_combo,
            self.track_max_dist,
            self.track_distance_weight,
            self.track_overlap_weight,
            self.track_point_support_weight,
            self.track_cost_cutoff,
            self.track_max_neighbors,
            self.track_min_link_size,
            self.track_sample_points,
            self._seg_raw_layer_combo,
            self._seg_frangi_layer_combo,
        ):
            if widget is not None:
                widget.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Fixed)

        if self.path_edit is not None:
            self.path_edit.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Fixed)
        if self._analysis_layer_combo is not None:
            self._analysis_layer_combo.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        for combo in (
            self._proximity_source_raw_combo,
            self._proximity_source_seg_combo,
            self._proximity_target_raw_combo,
            self._proximity_target_seg_combo,
        ):
            if combo is not None:
                combo.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)

        self.path_edit.setMinimumWidth(0)
        self.lbl_phys.setWordWrap(True)
        self.lbl_voxpix.setWordWrap(True)
        self.kw_preview.setMinimumHeight(self._scaled_px(120 if compact else 160))
        self.kw_preview.setMaximumHeight(self._scaled_px(200 if compact else 260))
        self.kw_preview.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        if self._analysis_table is not None:
            self._style_table_widget(
                self._analysis_table,
                min_height=140 if compact else 180,
                max_height=220 if compact else 280,
                expanding=False,
            )
        if self._analysis_selected_table is not None:
            self._style_table_widget(
                self._analysis_selected_table,
                min_height=72 if compact else 96,
                max_height=110 if compact else 140,
                expanding=False,
            )
        if self._analysis_plot_canvas is not None:
            self._analysis_plot_canvas.setMinimumHeight(self._scaled_px(640))
            self._analysis_plot_canvas.setMaximumHeight(16777215)
            self._analysis_plot_canvas.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Expanding)
        self._apply_panel_typography(compact)
        self._sync_standard_control_heights()

        label_width = 108 if compact else 122
        processing_action_width = 88 if compact else 104
        short_width = 72 if compact else 84
        medium_width = 108 if compact else 124
        long_width = 150 if compact else 180

        for label in (
            self.device_label,
            self.channel_label,
            self.time_label,
            getattr(self, "_time_range_label", None),
            getattr(self, "_slice_range_label", None),
            self._size_info_label,
            self._time_info_label,
        ):
            if label is not None:
                self._set_strict_fixed_width(label, label_width)

        for label in (
            self.upsample_action_label,
            self.denoise_action_label,
            self.gaussian_action_label,
        ):
            self._set_strict_fixed_width(label, processing_action_width)
        for label in (
            self.upsample_factor_label,
            self.filter_size_label,
            self.gaussian_sigma_label,
        ):
            self._set_strict_fixed_width(label, label_width)
            label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)

        for widget in (
            self.c_spin,
            self.t_spin,
            self.t_start_spin,
            self.t_end_spin,
            self.slice_start_spin,
            self.slice_end_spin,
            self.upsample_factor_spin,
            self.cmb_kernel,
            self.gaussian_sigma_spin,
            self.gaussian_alpha_spin,
            self.nfore_spin,
            self.nback_spin,
            self.maxiter_spin,
        ):
            self._set_adaptive_width(widget, short_width)

        for widget in (
            self.edit_vz,
            self.edit_vyx,
            self.kernel_spin,
            self.sigma_count,
            self.sigma_min,
            self.sigma_max,
            self.vessel_rescale_spin,
            self.sheet_sigma_count,
            self.sheet_sigma_min,
            self.sheet_sigma_max,
            self.sheet_rescale_spin,
            self.psf_ratio_spin,
            self.frangi_response_mode_combo,
            self.beta1_spin,
            self.beta2_spin,
            self.init_combo,
            self._seg_raw_layer_combo,
            self._seg_frangi_layer_combo,
        ):
            self._set_adaptive_width(widget, medium_width)
        self._set_adaptive_width(self.device_combo, 180 if compact else 220)

        for widget in (
            self.track_max_dist,
            self.track_distance_weight,
            self.track_overlap_weight,
            self.track_point_support_weight,
            self.track_cost_cutoff,
            self.track_max_neighbors,
            self.track_min_link_size,
            self.track_sample_points,
        ):
            if widget is not None:
                self._set_adaptive_width(widget, short_width)

        for combo in (
            self._tracking_layer_combo,
            self._tracking_raw_layer_combo,
        ):
            if combo is not None:
                self._set_adaptive_width(combo, 240 if compact else 280)
        for combo in (
            self._proximity_source_raw_combo,
            self._proximity_source_seg_combo,
            self._proximity_target_raw_combo,
            self._proximity_target_seg_combo,
        ):
            if combo is not None:
                self._set_adaptive_width(combo, 240 if compact else 280)

        for widget in (
            self.open_btn,
            self.btn_apply_time_range,
            self.btn_apply_slice_range,
            self.btn_apply_voxel,
            self.apply_btn,
            self.seg_btn,
        ):
            self._set_adaptive_width(widget, long_width)
        for widget in (
            self.btn_apply_upsample,
            self.btn_apply_denoise,
            self.btn_apply_gaussian,
        ):
            self._set_strict_fixed_width(widget, long_width)

        self._set_adaptive_width(self.path_edit, 260 if compact else 320)

        for widget in (
            self.c_spin,
            self.t_spin,
            self.t_start_spin,
            self.t_end_spin,
            self.slice_start_spin,
            self.slice_end_spin,
            self.edit_vz,
            self.edit_vyx,
            self.upsample_factor_spin,
            self.kernel_spin,
            self.sigma_count,
            self.sigma_min,
            self.sigma_max,
            self.vessel_rescale_spin,
            self.sheet_sigma_count,
            self.sheet_sigma_min,
            self.sheet_sigma_max,
            self.sheet_rescale_spin,
            self.psf_ratio_spin,
            self.frangi_response_mode_combo,
            self.beta1_spin,
            self.beta2_spin,
            self.nfore_spin,
            self.nback_spin,
            self.em_fg_sample_combo,
            self.em_bg_sample_combo,
            self.maxiter_spin,
            self.init_combo,
            self.cmb_kernel,
            self.device_combo,
            self._seg_raw_layer_combo,
            self._seg_frangi_layer_combo,
        ):
            self._set_adaptive_width(widget, short_width if widget in (
                self.c_spin,
                self.t_spin,
                self.t_start_spin,
                self.t_end_spin,
                self.slice_start_spin,
                self.slice_end_spin,
                self.upsample_factor_spin,
                self.cmb_kernel,
                self.nfore_spin,
                self.nback_spin,
                self.maxiter_spin,
            ) else medium_width)
        self._set_adaptive_width(self.path_edit, 260 if compact else 320)
        self._set_adaptive_width(self.em_fg_sample_combo, 180 if compact else 220)
        self._set_adaptive_width(self.em_bg_sample_combo, 180 if compact else 220)

        self._configure_scale_bar(_get_units_from_layer(self.viewer.layers.selection.active) if self.viewer.layers.selection.active else "um")
        QTimer.singleShot(0, self._sync_panel_tab_height)

    def _reset_processing_state(
        self,
        *,
        clear_upsample: bool = True,
        clear_denoise: bool = True,
        clear_frangi: bool = True,
    ):
        if clear_upsample:
            self._upsample_layer = None
        if clear_denoise:
            self._denoise_ctx = None
            self._denoise_layer = None
            self._median_layer = None
            self._gaussian_layer = None
            button = getattr(self, "btn_apply_gaussian", None)
            if button is not None and not self._thread_is_running("_gaussian_thread"):
                button.setEnabled(False)
        if clear_frangi:
            self._frangi_ctx = None
            self._frangi_layer = None
            self._frangi_component_raw = {}
            self._frangi_component_layers = {}
            self._frangi_component_specs = {}
            self._frangi_component_temporal_dims = None
            self._segmentation_layer = None
            self._refresh_frangi_component_view_buttons()
            self._update_segmentation_layer_labels()

    def _configure_scale_bar(self, unit: str = "um"):
        scale_bar = None
        with suppress(Exception):
            scale_bar = self.viewer.canvas.overlays.scale_bar
        if scale_bar is None:
            with suppress(Exception):
                scale_bar = self.viewer.scale_bar
        if scale_bar is None:
            return

        with suppress(Exception):
            scale_bar.visible = True
            active_layer = self.viewer.layers.selection.active
            if active_layer is not None:
                if hasattr(active_layer, "units"):
                    active_layer.units = (unit,) * _layer_display_ndim(active_layer)
                else:
                    # napari < 0.9 stored the display unit on the scale bar.
                    scale_bar.unit = unit
            scale_bar.position = "bottom_right"
            scale_bar.font_size = self._scaled_px(8)

    def _set_text_overlay(self, text: str) -> None:
        with suppress(Exception):
            self.viewer.text_overlay.visible = True
            self.viewer.text_overlay.position = "bottom_left"
            self.viewer.text_overlay.text = text

    def _replace_layer_data(self, layer, data: np.ndarray):
        try:
            layer.data = np.asarray(data)
            return layer
        except (TypeError, ValueError, AttributeError, RuntimeError):
            md = dict(getattr(layer, "metadata", {}) or {})
            name = layer.name
            sc = tuple(getattr(layer, "scale", (1,) * np.asarray(data).ndim))
            units = _spatial_units_for_ndim(layer, np.asarray(data).ndim)
            with suppress(Exception):
                self.viewer.layers.remove(layer)
            kwargs = {"name": name, "scale": sc, "metadata": md}
            if units is not None:
                kwargs["units"] = units
            return self.viewer.add_image(np.asarray(data), **kwargs)

    def _sync_group_layer_scales(self, group_id: str, new_triplet: tuple[float | None, float, float], source_layer) -> None:
        for layer_item in list(self.viewer.layers):
            if not hasattr(layer_item, "data") or layer_item is source_layer:
                continue
            md = getattr(layer_item, "metadata", {}) or {}
            if md.get("group_id") != group_id:
                continue
            with suppress(Exception):
                sc = list(getattr(layer_item, "scale", (1,) * layer_item.data.ndim))
                if layer_item.data.ndim >= 3 and new_triplet[0] is not None:
                    sc[-3:] = list(new_triplet)
                else:
                    sc[-2:] = [new_triplet[1], new_triplet[2]]
                layer_item.scale = tuple(sc)

    def _preferred_2d_scale(self, layer, sc) -> tuple[float, float]:
        slice_start, slice_end = self._selected_slice_range()
        if slice_start == slice_end and layer.data.ndim >= 3:
            with suppress(Exception):
                _z, _y, _x = _get_pixel_size_tuple(layer, 3)
                return (float(_y), float(_x))
        return (float(sc[-2]), float(sc[-1]))

    def _compute_updated_scale_triplet(self, layer) -> tuple[tuple[float, ...], tuple[float | None, float, float]]:
        nd = layer.data.ndim
        sc = list(getattr(layer, "scale", (1,) * nd))
        if nd >= 3:
            vz = float(self.edit_vz.value()) if self.edit_vz.isEnabled() and self.edit_vz.isVisible() else float(sc[-3])
            vxy = float(self.edit_vyx.value())
            sc[-3:] = [vz, vxy, vxy]
            new_triplet = (float(sc[-3]), float(sc[-2]), float(sc[-1]))
        else:
            vxy = float(self.edit_vyx.value())
            sc[-2:] = [vxy, vxy]
            new_triplet = (None, float(sc[-2]), float(sc[-1]))
        return tuple(sc), new_triplet

    def _apply_voxel_update_to_layer(self, layer, new_scale: tuple[float, ...]) -> str:
        layer.scale = new_scale
        md = dict(getattr(layer, "metadata", {}) or {})
        _ensure_unit_in_metadata(md)
        gid = _ensure_group_id(md)
        layer.metadata = md
        return gid

    def _build_segmentation_layer_metadata(self, layer, mask8: np.ndarray, info=None) -> dict[str, Any]:
        ctx = self._frangi_ctx
        if ctx is None:
            raise RuntimeError("Structural response context missing for segmentation result")
        md_base = dict(getattr(layer, "metadata", {}) or {}) if layer else {}
        if ctx.temporal_dims == "TYX":
            output_dims = "TYX"
        elif ctx.temporal_dims == "TZYX":
            output_dims = "TZYX"
        else:
            output_dims = "ZYX" if mask8.ndim == 3 else "YX"
        md_base["dims"] = output_dims
        md_base["dims_out"] = output_dims
        new_md = dict(md_base)
        new_md["unit"] = ctx.unit or (_get_units_from_layer(layer) if layer else "um")
        new_md["source"] = ctx.source or (_get_source_from_layer(layer) if layer else None)
        new_md["is_segmentation"] = True
        new_md["group_id"] = ctx.group_id or _get_group_id(self.viewer.layers.selection.active) or _ensure_group_id(new_md)
        if info is not None:
            with suppress(Exception):
                new_md["seg_iterations_run"] = int(getattr(info, "iterations_run", 0))
                new_md["seg_delta_loglh_trace"] = list(getattr(info, "deltas", []))
                new_md["seg_converged"] = bool(getattr(info, "converged", False))
        return new_md

    def _focus_result_layer(self, result_layer) -> None:
        with suppress(Exception):
            self.viewer.layers.selection.active = result_layer
        self._view_initialized = False
        self._fit_view_and_scalebar()
        self._update_info()

    def _default_layer_export_path(self, layer) -> str:
        md = getattr(layer, "metadata", {}) or {}
        source = md.get("source") or _get_source_from_layer(layer) or ""
        base, ext = os.path.splitext(os.path.basename(source))
        if not ext or ext.lower() not in {".tif", ".tiff", ".png", ".jpg", ".jpeg"}:
            ext = ".tif"
        if not base:
            base = layer.name or "layer"
        if md.get("is_segmentation"):
            base = f"{base}_segmentation"
        return os.path.join(os.path.dirname(source) if source else "", f"{base}{ext}")

    def _save_active_layer_with_metadata(self) -> None:
        layer = self.viewer.layers.selection.active
        if layer is None or not hasattr(layer, "data"):
            QMessageBox.information(self, "No layer", "Please select an image or segmentation layer first.")
            return
        layer_metadata = dict(getattr(layer, "metadata", {}) or {})
        if layer_metadata.get("is_proximity_roi"):
            self._export_proximity_roi(layer)
            return
        default_path = self._default_layer_export_path(layer)
        path, _selected_filter = QFileDialog.getSaveFileName(
            self,
            "Save layer with metadata",
            default_path,
            "TIFF (*.tif *.tiff);;PNG (*.png);;JPEG (*.jpg *.jpeg)",
        )
        if not path:
            return
        if (
            layer_metadata.get("is_frangi")
            or layer_metadata.get("sigma_layer_role") == _STRUCTURAL_RESPONSE_LAYER_ROLE
        ):
            layer_metadata["is_frangi"] = True
            layer_metadata["sigma_layer_role"] = _STRUCTURAL_RESPONSE_LAYER_ROLE
            layer_metadata["sigma_saved_structure"] = True
            layer.metadata = layer_metadata
        scale = getattr(layer, "scale", None)
        meta = {
            "scale": tuple(float(v) for v in scale) if scale is not None else (),
            "metadata": layer_metadata,
            "name": layer.name,
            "layer_type": getattr(layer, "_type_string", "image"),
        }
        try:
            write_single_image(path, np.asarray(layer.data), meta)
        except (OSError, ValueError, RuntimeError, TypeError) as e:
            QMessageBox.critical(self, "Save failed", f"Failed to save layer: {e!r}")
            return
        suffix = os.path.splitext(path)[1].lower()
        if suffix in {".tif", ".tiff"}:
            QMessageBox.information(self, "Save complete", f"Saved layer to:\n{path}")
        else:
            QMessageBox.information(
                self,
                "Save complete",
                "Layer saved.\n\n"
                "Note: PNG/JPG/JPEG do not reliably preserve pixel/voxel size metadata; "
                "use TIFF/TIF if you need physical size information in the file.",
            )

    def _set_z_range_for_data(self, data) -> None:
        with suppress(Exception):
            if hasattr(data, "data"):
                layer = data
                arr = layer.data
                dims_tag = _layer_dims_tag(layer)
                if dims_tag in {"YX", "CYX", "TYX"}:
                    z_size = 1
                else:
                    z_size = int(arr.shape[-3]) if arr.ndim >= 3 and int(arr.shape[-3]) > 1 else 1
            else:
                arr = data
                z_size = int(arr.shape[-3]) if arr.ndim >= 3 and int(arr.shape[-3]) > 1 else 1
            new_end = max(z_size - 1, 0)
            self.slice_start_spin.setRange(0, max(z_size - 1, 0))
            self.slice_end_spin.setRange(0, max(z_size - 1, 0))
            if not self.slice_start_spin.hasFocus():
                self.slice_start_spin.setValue(0)
            if not self.slice_end_spin.hasFocus():
                self.slice_end_spin.setValue(new_end)
            enabled = z_size > 1
            self.slice_start_spin.setEnabled(enabled)
            self.slice_end_spin.setEnabled(enabled)
            if hasattr(self, "btn_apply_slice_range") and self.btn_apply_slice_range is not None:
                self.btn_apply_slice_range.setEnabled(enabled)

    def _selected_slice_range(self) -> tuple[int, int]:
        start = int(self.slice_start_spin.value()) if hasattr(self, "slice_start_spin") else 0
        end = int(self.slice_end_spin.value()) if hasattr(self, "slice_end_spin") else 0
        if end < start:
            start, end = end, start
        return start, end

    def _on_slice_range_changed(self, *_args) -> None:
        start, end = self._selected_slice_range()
        if hasattr(self, "slice_start_spin") and self.slice_start_spin.value() != start:
            self.slice_start_spin.blockSignals(True)
            self.slice_start_spin.setValue(start)
            self.slice_start_spin.blockSignals(False)
        if hasattr(self, "slice_end_spin") and self.slice_end_spin.value() != end:
            self.slice_end_spin.blockSignals(True)
            self.slice_end_spin.setValue(end)
            self.slice_end_spin.blockSignals(False)

    def _preferred_denoised_input_layer(self):
        """Return the current denoised layer when it is still in the viewer."""
        layer = getattr(self, "_denoise_layer", None)
        if layer is not None:
            try:
                if layer in self.viewer.layers:
                    return layer
            except (TypeError, ValueError):
                pass
        for candidate in reversed(list(self.viewer.layers)):
            metadata = getattr(candidate, "metadata", {}) or {}
            if metadata.get("is_denoise") and _is_frangi_raw_candidate_layer(candidate):
                self._denoise_layer = candidate
                return candidate
        return None

    def _preferred_median_input_layer(self):
        """Return the latest median-filter output while it remains in the viewer."""
        layer = getattr(self, "_median_layer", None)
        if layer is None:
            return None
        try:
            if layer not in self.viewer.layers:
                return None
        except (TypeError, ValueError):
            return None
        return layer

    def _refresh_gaussian_button_state(self) -> None:
        checkbox = getattr(self, "gaussian_enabled_checkbox", None)
        button = getattr(self, "btn_apply_gaussian", None)
        if button is None:
            return
        enabled = bool(checkbox is not None and checkbox.isChecked())
        running = self._thread_is_running("_gaussian_thread")
        if checkbox is not None:
            checkbox.setEnabled(not running)
        for control in (
            getattr(self, "gaussian_action_label", None),
            getattr(self, "gaussian_sigma_label", None),
            getattr(self, "gaussian_sigma_spin", None),
            getattr(self, "gaussian_alpha_label", None),
            getattr(self, "gaussian_alpha_spin", None),
        ):
            if control is not None:
                control.setEnabled(enabled and not running)
        progress = getattr(self, "gaussian_progress", None)
        if progress is not None:
            progress.setEnabled(enabled or running)
        if not running:
            button.setEnabled(
                enabled and self._preferred_median_input_layer() is not None
            )

    def _refresh_upsample_controls_state(self) -> None:
        checkbox = getattr(self, "upsample_enabled_checkbox", None)
        button = getattr(self, "btn_apply_upsample", None)
        enabled = bool(checkbox is not None and checkbox.isChecked())
        running = self._thread_is_running("_upsample_thread")
        if checkbox is not None:
            checkbox.setEnabled(not running)
        for control in (
            getattr(self, "upsample_action_label", None),
            getattr(self, "upsample_factor_label", None),
            getattr(self, "upsample_factor_spin", None),
        ):
            if control is not None:
                control.setEnabled(enabled and not running)
        progress = getattr(self, "upsample_progress", None)
        if progress is not None:
            progress.setEnabled(enabled or running)
        if button is not None and not running:
            button.setEnabled(enabled)

    def _segmentation_raw_candidates(self) -> list:
        group_id = getattr(self._frangi_ctx, "group_id", None)
        candidates = []
        for layer_item in self.viewer.layers:
            if not hasattr(layer_item, "data"):
                continue
            layer_type = str(getattr(layer_item, "_type_string", "") or "").lower()
            layer_class = layer_item.__class__.__name__.lower()
            if layer_type not in {"image", "labels"} and layer_class not in {"image", "labels"}:
                continue
            md = getattr(layer_item, "metadata", {}) or {}
            if md.get("is_frangi") or md.get("is_segmentation") or md.get("is_tracking"):
                continue
            if group_id is not None and md.get("group_id") != group_id:
                continue
            candidates.append(layer_item)
        return candidates

    def _frangi_raw_candidates(self) -> list:
        return [layer_item for layer_item in self.viewer.layers if _is_frangi_raw_candidate_layer(layer_item)]

    def _segmentation_frangi_candidates(self) -> list:
        group_id = getattr(self._frangi_ctx, "group_id", None)
        candidates = []
        for layer_item in self.viewer.layers:
            if not hasattr(layer_item, "data"):
                continue
            md = getattr(layer_item, "metadata", {}) or {}
            if not (
                md.get("is_frangi")
                or md.get("sigma_layer_role") == _STRUCTURAL_RESPONSE_LAYER_ROLE
            ):
                continue
            if (
                group_id is not None
                and md.get("group_id") != group_id
                and not md.get("sigma_saved_structure")
            ):
                continue
            candidates.append(layer_item)
        return candidates

    def _preferred_segmentation_frangi_layer(self, candidates: list | None = None):
        candidates = self._segmentation_frangi_candidates() if candidates is None else list(candidates)
        mode = str(self.frangi_response_mode_combo.currentData() or "vesselness")
        result_name = _frangi_result_name_for_mode(mode)
        component_layer = (getattr(self, "_frangi_component_layers", {}) or {}).get(result_name)
        for layer_item in candidates:
            if layer_item is component_layer:
                return layer_item
        for layer_item in reversed(candidates):
            md = getattr(layer_item, "metadata", {}) or {}
            if str(md.get("frangi_result_name") or "").lower() == result_name:
                return layer_item
        return None

    def _rebuild_segmentation_layer_combos(self) -> None:
        raw_combo = getattr(self, "_seg_raw_layer_combo", None)
        frangi_combo = getattr(self, "_seg_frangi_layer_combo", None)
        if raw_combo is None or frangi_combo is None:
            return
        current_raw_name = raw_combo.currentText()
        current_frangi_name = frangi_combo.currentText()

        raw_combo.blockSignals(True)
        raw_combo.clear()
        for layer_item in self._segmentation_raw_candidates():
            raw_combo.addItem(layer_item.name, layer_item)
        if raw_combo.count() > 0:
            idx = raw_combo.findText(current_raw_name)
            if idx < 0:
                preferred = self._preferred_denoised_input_layer()
                if preferred is not None:
                    idx = raw_combo.findText(preferred.name)
            if idx < 0 and self.viewer.layers.selection.active is not None:
                idx = raw_combo.findText(self.viewer.layers.selection.active.name)
            raw_combo.setCurrentIndex(idx if idx >= 0 else 0)
        raw_combo.blockSignals(False)

        frangi_candidates = self._segmentation_frangi_candidates()
        frangi_combo.blockSignals(True)
        frangi_combo.clear()
        for layer_item in frangi_candidates:
            frangi_combo.addItem(layer_item.name, layer_item)
        if frangi_combo.count() > 0:
            preferred = self._preferred_segmentation_frangi_layer(frangi_candidates)
            idx = frangi_combo.findText(preferred.name) if preferred is not None else -1
            if idx < 0:
                idx = frangi_combo.findText(current_frangi_name)
            if idx < 0 and self._frangi_layer is not None:
                idx = frangi_combo.findText(self._frangi_layer.name)
            frangi_combo.setCurrentIndex(idx if idx >= 0 else 0)
        frangi_combo.blockSignals(False)
        self._restore_selected_structure_rescale_values()
        self._refresh_rescale_target_label()

    def _rebuild_frangi_layer_combo(self, *, prefer_denoised: bool = False) -> None:
        combo = getattr(self, "_frangi_raw_layer_combo", None)
        if combo is None:
            return
        current_name = combo.currentText()
        combo.blockSignals(True)
        combo.clear()
        for layer_item in self._frangi_raw_candidates():
            combo.addItem(layer_item.name, layer_item)
        if combo.count() > 0:
            preferred = self._preferred_denoised_input_layer()
            idx = combo.findText(preferred.name) if prefer_denoised and preferred is not None else -1
            if idx < 0:
                idx = combo.findText(current_name)
            if idx < 0 and preferred is not None:
                idx = combo.findText(preferred.name)
            if idx < 0 and self.viewer.layers.selection.active is not None:
                idx = combo.findText(self.viewer.layers.selection.active.name)
            combo.setCurrentIndex(idx if idx >= 0 else 0)
        combo.blockSignals(False)

    def _rebuild_info_raw_layer_combo(self, *, prefer_denoised: bool = False) -> None:
        combo = getattr(self, "_info_raw_layer_combo", None)
        if combo is None:
            return
        current_name = combo.currentText()
        combo.blockSignals(True)
        combo.clear()
        for layer_item in self._frangi_raw_candidates():
            combo.addItem(layer_item.name, layer_item)
        if combo.count() > 0:
            preferred = self._preferred_denoised_input_layer()
            idx = combo.findText(preferred.name) if prefer_denoised and preferred is not None else -1
            if idx < 0:
                idx = combo.findText(current_name)
            if idx < 0 and preferred is not None:
                idx = combo.findText(preferred.name)
            if idx < 0 and self.viewer.layers.selection.active is not None:
                idx = combo.findText(self.viewer.layers.selection.active.name)
            combo.setCurrentIndex(idx if idx >= 0 else 0)
        combo.blockSignals(False)

    def _on_info_raw_layer_changed(self, _index: int = 0) -> None:
        self._update_info()

    def _selected_info_raw_layer(self):
        combo = getattr(self, "_info_raw_layer_combo", None)
        if combo is None or combo.count() == 0:
            return None
        return combo.currentData()

    def _selected_segmentation_raw_layer(self):
        combo = getattr(self, "_seg_raw_layer_combo", None)
        if combo is None or combo.count() == 0:
            return None
        return combo.currentData()

    def _selected_frangi_raw_layer(self):
        combo = getattr(self, "_frangi_raw_layer_combo", None)
        if combo is None or combo.count() == 0:
            return None
        return combo.currentData()

    def _selected_segmentation_frangi_layer(self):
        combo = getattr(self, "_seg_frangi_layer_combo", None)
        if combo is None or combo.count() == 0:
            return None
        return combo.currentData()

    def _infer_segmentation_shape_context(self, raw_layer, frangi_layer):
        if raw_layer is None or frangi_layer is None:
            return None
        if not hasattr(raw_layer, "data") or not hasattr(frangi_layer, "data"):
            return None

        raw_data = raw_layer.data
        frangi_data = frangi_layer.data
        raw_shape = tuple(int(value) for value in raw_data.shape)
        frangi_shape = tuple(int(value) for value in frangi_data.shape)
        if raw_shape != frangi_shape:
            with suppress(Exception):
                raw_current = np.asarray(self._current_volume_from_layer(raw_layer))
                if raw_current.shape == frangi_shape:
                    raw_data = raw_current
                    raw_shape = tuple(int(value) for value in raw_current.shape)
            if raw_shape != frangi_shape:
                with suppress(Exception):
                    frangi_current = np.asarray(self._current_volume_from_layer(frangi_layer))
                    if raw_shape == frangi_current.shape:
                        frangi_data = frangi_current
                        frangi_shape = tuple(int(value) for value in frangi_current.shape)
        if raw_shape != frangi_shape:
            return None

        dims_tag = (_layer_dims_tag(raw_layer) or _layer_dims_tag(frangi_layer) or "").upper()
        temporal_dims = None
        if dims_tag in {"TYX", "TZYX"}:
            temporal_dims = dims_tag
            spatial_dim = 3 if dims_tag == "TZYX" else 2
        elif len(raw_shape) == 4:
            temporal_dims = "TZYX"
            spatial_dim = 3
        elif len(raw_shape) == 3:
            spatial_dim = 3
        elif len(raw_shape) == 2:
            spatial_dim = 2
        else:
            return None

        return raw_data, frangi_data, int(spatial_dim), temporal_dims

    def _ensure_frangi_context_from_selected_layers(self) -> bool:
        raw_layer = self._selected_segmentation_raw_layer()
        frangi_layer = self._selected_segmentation_frangi_layer()
        inferred = self._infer_segmentation_shape_context(raw_layer, frangi_layer)
        if inferred is None:
            return False
        raw_data, frangi_data, spatial_dim, temporal_dims = inferred
        raw_md = dict(getattr(raw_layer, "metadata", {}) or {})
        frangi_md = dict(getattr(frangi_layer, "metadata", {}) or {})
        group_id = _get_group_id(raw_layer) or _get_group_id(frangi_layer) or _ensure_group_id(raw_md)
        pixel_size = _get_pixel_size_tuple(raw_layer, spatial_dim)
        self._frangi_layer = frangi_layer
        self._frangi_ctx = FrangiContext(
            image=raw_data,
            frangi=frangi_data,
            frangi_raw=frangi_data,
            dim=spatial_dim,
            pixel_size=pixel_size,
            unit=_get_units_from_layer(raw_layer),
            source=raw_md.get("source") or _get_source_from_layer(raw_layer),
            group_id=group_id,
            channel_index=raw_md.get("channel_index"),
            device=self._resolve_device(),
            temporal_dims=temporal_dims,
            time_range=self._selected_time_range() if temporal_dims is not None else None,
        )
        frangi_md["is_frangi"] = True
        frangi_md["group_id"] = group_id
        frangi_layer.metadata = frangi_md
        raw_md["group_id"] = group_id
        raw_layer.metadata = raw_md
        return True

    def _next_layer_name(self, base_name: str) -> str:
        existing = {layer_item.name for layer_item in self.viewer.layers}
        if base_name not in existing:
            return base_name
        index = 2
        while f"{base_name} ({index})" in existing:
            index += 1
        return f"{base_name} ({index})"

    def _derived_result_layer_name(self, source_layer, suffix: str) -> str:
        base_name = getattr(source_layer, "name", "") or str(suffix)
        return self._next_layer_name(f"{base_name} {suffix}")

    def _context_from_layer(
        self,
        layer,
        image: np.ndarray,
        dim: int,
        pixel_size: tuple[float, ...],
        *,
        temporal_dims: str | None = None,
        time_range: tuple[int, int] | None = None,
    ) -> ProcessingContext:
        md = getattr(layer, "metadata", {}) or {}
        return ProcessingContext(
            image=np.asarray(image).astype(np.float32, copy=False),
            dim=dim,
            pixel_size=tuple(pixel_size),
            unit=_get_units_from_layer(layer),
            source=_get_source_from_layer(layer),
            group_id=_get_group_id(layer),
            channel_index=md.get("channel_index"),
            temporal_dims=temporal_dims,
            time_range=time_range,
        )

    def _apply_loaded_layers(self, data_tc_zyx: np.ndarray, meta: dict[str, Any], *, name: str, remove_layer=None):
        meta = _mark_plugin_managed(dict(meta))
        self.path_edit.setText(meta.get("source", ""))
        self._reset_processing_state()
        layers = self._add_tczyx_image_layers(data_tc_zyx, meta, name)
        if remove_layer is not None:
            with suppress(Exception):
                self.viewer.layers.remove(remove_layer)
        self._set_ct_controls(T=int(data_tc_zyx.shape[0]), C=int(data_tc_zyx.shape[1]))
        if layers:
            self.viewer.layers.selection.active = layers[0]
            _safe_axis_labels(self.viewer, layers[0])
            self._set_default_display_mode_for_imported_layer(layers[0])
        self._configure_scale_bar(meta.get("unit", "um"))
        self._set_z_range_for_data(data_tc_zyx)
        self._view_initialized = False
        self._fit_view_and_scalebar()
        self._update_info()
        return layers


    def _sync_single_layer(self, layer):
        with suppress(Exception):
            md = layer.metadata if layer.metadata else {}
            md["unit"] = _ensure_unit_in_metadata(md)
            _ensure_group_id(md)
            layer.metadata = md
        self._configure_scale_bar(_get_units_from_layer(layer))
        _safe_axis_labels(self.viewer, layer)
        t_size = _time_size_from_layer(layer)
        self._group_layers = [layer]
        self._set_ct_controls(T=t_size, C=1)
        self._set_z_range_for_data(layer)
        self._set_default_display_mode_for_imported_layer(layer)
        self._update_info()

    def _set_default_display_mode_for_imported_layer(self, layer) -> bool:
        if layer is None or not hasattr(layer, "data"):
            return False
        data = getattr(layer, "data", None)
        shape = tuple(int(value) for value in getattr(data, "shape", ()))
        if len(shape) < 3 or shape[-3] <= 1:
            return False

        dims_tag = _layer_dims_tag(layer)
        if dims_tag:
            is_spatial_3d = "Z" in dims_tag
        else:
            # A bare 3D image is interpreted by napari as ZYX. A bare 4D
            # image follows SIGMA's usual TZYX convention.
            is_spatial_3d = len(shape) in {3, 4}
        if not is_spatial_3d:
            return False

        try:
            self.viewer.dims.ndisplay = 3
            QApplication.processEvents()
        except (AttributeError, RuntimeError, TypeError, ValueError):
            return False
        return True

    def _current_volume_from_layer(self, layer) -> np.ndarray:
        arr = layer.data
        dims_tag = _layer_dims_tag(layer)
        t_idx0 = 0
        if dims_tag in {"TYX", "TZYX", "TCZYX"} or arr.ndim >= 4:
            time_axis = _time_axis_index(self.viewer, layer)
            if time_axis is None:
                time_axis = _viewer_axis_for_layer_axis(self.viewer, layer, 0)
            with suppress(Exception):
                if time_axis is not None:
                    t_idx0 = int(self.viewer.dims.current_step[time_axis])
        if dims_tag == "TYX":
            return arr[t_idx0]
        if dims_tag == "TZYX" or arr.ndim >= 4:
            return np.squeeze(arr[t_idx0])
        return np.squeeze(arr)

    def _add_tczyx_image_layers(self, data_tc_zyx: np.ndarray, meta: dict[str, Any], name: str):
        """Use the same axis/type conversion as the registered lightweight reader."""
        gid = meta.get("group_id") or uuid4().hex
        layers = []
        for data, kwargs, kind in tczyx_to_layer_data(data_tc_zyx, meta, name):
            md = kwargs["metadata"]
            md.update(group_id=gid, managed_by_plugin=True, layer_type=kind)
            units = _metadata_units_for_ndim(md, data.ndim)
            if units is not None:
                kwargs["units"] = units
            layer = getattr(self.viewer, f"add_{kind}")(data, visible=True, **kwargs)
            layers.append(layer)
        self._group_layers = layers
        return layers

    # -----------------------------------------------------------------
    # Pages
    # -----------------------------------------------------------------

    def _build_segmentation_page(self) -> QWidget:
        top_box = self._build_top_box()
        info_box = self._build_info_box()
        denoise_box = self._build_denoise_box()
        frangi_box = self._build_frangi_box()
        seg_box = self._build_seg_box()
        return self._build_panel_page(
            top_box,
            info_box,
            denoise_box,
            frangi_box,
            seg_box,
        )

    def _build_proximity_page(self) -> QWidget:
        ctrl_box = QGroupBox("Proximity Analysis")
        ctrl_grid = QGridLayout()
        ctrl_grid.setHorizontalSpacing(self._scaled_px(10))
        ctrl_grid.setVerticalSpacing(self._scaled_px(6))

        self._proximity_source_raw_combo = QComboBox()
        self._proximity_source_seg_combo = QComboBox()
        self._proximity_target_raw_combo = QComboBox()
        self._proximity_target_seg_combo = QComboBox()
        for role, combo in (
            ("source_raw", self._proximity_source_raw_combo),
            ("source_seg", self._proximity_source_seg_combo),
            ("target_raw", self._proximity_target_raw_combo),
            ("target_seg", self._proximity_target_seg_combo),
        ):
            combo.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            combo.currentIndexChanged.connect(
                lambda _idx, role=role: self._remember_proximity_layer_combo(role)
            )

        ctrl_grid.addWidget(QLabel("Source raw"), 0, 0, alignment=Qt.AlignRight)
        ctrl_grid.addWidget(self._proximity_source_raw_combo, 0, 1)
        ctrl_grid.addWidget(QLabel("Source segmentation"), 1, 0, alignment=Qt.AlignRight)
        ctrl_grid.addWidget(self._proximity_source_seg_combo, 1, 1)
        ctrl_grid.addWidget(QLabel("Target raw"), 2, 0, alignment=Qt.AlignRight)
        ctrl_grid.addWidget(self._proximity_target_raw_combo, 2, 1)
        ctrl_grid.addWidget(QLabel("Target segmentation"), 3, 0, alignment=Qt.AlignRight)
        ctrl_grid.addWidget(self._proximity_target_seg_combo, 3, 1)

        roi_row = QHBoxLayout()
        roi_row.setContentsMargins(0, 0, 0, 0)
        roi_row.setSpacing(self._scaled_px(8))
        create_roi_btn = QPushButton("Draw ROI")
        create_roi_btn.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        create_roi_btn.clicked.connect(self._ensure_proximity_roi_layer)
        clear_roi_btn = QPushButton("Clear ROI")
        clear_roi_btn.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        clear_roi_btn.clicked.connect(self._clear_proximity_roi_layer)
        roi_row.addWidget(create_roi_btn, 1)
        roi_row.addWidget(clear_roi_btn, 1)
        self._proximity_roi_status_label = QLabel("Full image")
        self._proximity_roi_status_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self._proximity_roi_status_label.setStyleSheet(
            "QLabel { color: #AEB6C1; font-size: 11px; }"
        )
        roi_row.addWidget(self._proximity_roi_status_label)
        ctrl_grid.addWidget(QLabel("ROI:"), 4, 0, alignment=Qt.AlignRight)
        ctrl_grid.addLayout(roi_row, 4, 1)

        compute_row = QHBoxLayout()
        compute_row.setContentsMargins(0, 0, 0, 0)
        compute_row.setSpacing(self._scaled_px(8))
        compute_full_btn = QPushButton("Compute Full Image")
        compute_full_btn.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        compute_full_btn.clicked.connect(self._on_compute_proximity_clicked)
        compute_roi_btn = QPushButton("Compute ROI(s)")
        compute_roi_btn.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        compute_roi_btn.clicked.connect(self._on_compute_proximity_roi_clicked)
        self._proximity_compute_buttons = (compute_full_btn, compute_roi_btn)
        compute_row.addWidget(compute_full_btn, 1)
        compute_row.addWidget(compute_roi_btn, 1)
        ctrl_grid.addLayout(compute_row, 5, 0, 1, 2)
        ctrl_grid.setColumnStretch(0, 0)
        ctrl_grid.setColumnStretch(1, 1)
        ctrl_box.setLayout(ctrl_grid)

        summary_box = QGroupBox("Summary")
        summary_layout = QVBoxLayout(summary_box)
        summary_toolbar = QHBoxLayout()
        summary_toolbar.addStretch(1)
        self._proximity_export_btn = QPushButton("Export")
        self._proximity_export_btn.clicked.connect(self._export_proximity_summary)
        summary_toolbar.addWidget(self._proximity_export_btn)
        self._style_toolbar_layout(summary_toolbar)
        summary_layout.addLayout(summary_toolbar)
        self._proximity_summary_table = QTableWidget()
        self._proximity_summary_table.setColumnCount(0)
        self._proximity_summary_table.setRowCount(0)
        self._proximity_summary_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self._proximity_summary_table.setSelectionBehavior(QTableWidget.SelectRows)
        self._proximity_summary_table.setSelectionMode(QAbstractItemView.SingleSelection)
        self._proximity_summary_table.itemSelectionChanged.connect(
            self._on_proximity_summary_selection_changed
        )
        self._style_table_widget(self._proximity_summary_table, min_height=96, max_height=160)
        summary_layout.addWidget(self._proximity_summary_table)

        return self._build_panel_page(ctrl_box, summary_box)

    def _proximity_shape_ndim(self, layer) -> int:
        data = getattr(layer, "data", None)
        dims_tag = str((getattr(layer, "metadata", {}) or {}).get("dims") or "").upper()
        if dims_tag == "TYX":
            return 3
        if dims_tag == "TZYX":
            return 4
        return int(data.ndim)

    def _proximity_spatial_ndim(self, layer) -> int:
        data = getattr(layer, "data", None)
        dims_tag = str((getattr(layer, "metadata", {}) or {}).get("dims") or "").upper()
        if dims_tag == "TYX":
            return 2
        if dims_tag == "TZYX":
            return 3
        shape = list(getattr(data, "shape", ()))
        while len(shape) > 4 and shape[0] == 1:
            shape.pop(0)
        if len(shape) == 4:
            return 3
        if len(shape) == 3:
            return 3
        return 2

    def _capture_proximity_view_state(self) -> dict[str, object]:
        state: dict[str, object] = {}
        with suppress(Exception):
            state["ndisplay"] = int(self.viewer.dims.ndisplay)
            state["current_step"] = tuple(int(v) for v in self.viewer.dims.current_step)
        with suppress(Exception):
            state["camera_angles"] = tuple(self.viewer.camera.angles)
            state["camera_center"] = tuple(self.viewer.camera.center)
            state["camera_zoom"] = float(self.viewer.camera.zoom)
            state["camera_perspective"] = float(self.viewer.camera.perspective)
        return state

    def _restore_proximity_view_state(self, state: dict[str, object]) -> None:
        if not state:
            return
        with suppress(Exception):
            self.viewer.dims.ndisplay = int(state["ndisplay"])
            QApplication.processEvents()
        for axis, step in enumerate(state.get("current_step", ())):
            with suppress(Exception):
                self.viewer.dims.set_current_step(axis, int(step))
        with suppress(Exception):
            self.viewer.camera.angles = state["camera_angles"]
            self.viewer.camera.center = state["camera_center"]
            self.viewer.camera.zoom = float(state["camera_zoom"])
            self.viewer.camera.perspective = float(state["camera_perspective"])

    def _ensure_proximity_roi_layer(self) -> None:
        layer = self._selected_proximity_source_seg_layer() or self._selected_proximity_target_seg_layer()
        if layer is None:
            QMessageBox.information(self, "No layer", "Choose a segmentation layer first.")
            return
        try:
            if (
                self._proximity_spatial_ndim(layer) == 3
                and int(self.viewer.dims.ndisplay) == 3
            ):
                self._proximity_roi_3d_view_state = {
                    **self._capture_proximity_view_state(),
                    "source_layer": layer,
                }
            if int(self.viewer.dims.ndisplay) != 2:
                self.viewer.dims.ndisplay = 2
                QApplication.processEvents()
        except (AttributeError, RuntimeError, TypeError, ValueError) as e:
            QMessageBox.warning(
                self,
                "ROI unavailable",
                f"Could not switch to the 2D ROI drawing view: {e!r}",
            )
            return
        shape_ndim = self._proximity_shape_ndim(layer)
        try:
            source_shape = tuple(
                int(value)
                for value in _normalized_layer_array(layer, allow_float=False).shape
            )
        except (TypeError, ValueError):
            source_shape = tuple(int(value) for value in np.asarray(layer.data).shape)
        scale = tuple(float(v) for v in getattr(layer, "scale", (1.0,) * shape_ndim))
        md = {
            "is_proximity_roi": True,
            "proximity_source_layer": getattr(layer, "name", ""),
            "proximity_source_shape": source_shape,
        }
        roi_layer = self._proximity_roi_layer
        if roi_layer is not None and roi_layer in self.viewer.layers:
            roi_md = dict(getattr(roi_layer, "metadata", {}) or {})
            roi_shape = tuple(int(value) for value in roi_md.get("proximity_source_shape", ()))
            roi_ndim = int(getattr(roi_layer, "ndim", shape_ndim))
            if roi_ndim != shape_ndim or (roi_shape and roi_shape != source_shape):
                if bool(getattr(roi_layer, "_is_creating", False)):
                    with suppress(Exception):
                        roi_layer._finish_drawing()
                with suppress(Exception):
                    self.viewer.layers.remove(roi_layer)
                self._proximity_roi_layer = None
                text_layer = self._proximity_roi_text_layer
                if text_layer is not None and text_layer in self.viewer.layers:
                    with suppress(Exception):
                        self.viewer.layers.remove(text_layer)
                self._proximity_roi_text_layer = None
        if self._proximity_roi_layer is None or self._proximity_roi_layer not in self.viewer.layers:
            self._proximity_roi_layer = self.viewer.add_shapes(
                data=[],
                shape_type=[],
                ndim=shape_ndim,
                name=self._next_layer_name("proximity roi"),
                edge_color="#F2F2F2",
                edge_width=2.0,
                face_color=[1.0, 0.82, 0.22, 0.08],
                opacity=1.0,
                blending="translucent_no_depth",
                metadata=md,
                scale=scale,
                **_spatial_unit_kwargs(layer, shape_ndim),
            )
            with suppress(Exception):
                self._proximity_roi_layer.events.data.connect(self._on_proximity_roi_changed)
            with suppress(Exception):
                self._proximity_roi_layer.mouse_move_callbacks.append(self._prime_proximity_roi_cursor_state)
        else:
            with suppress(Exception):
                self._proximity_roi_layer.metadata = md
                self._proximity_roi_layer.scale = scale
                self._proximity_roi_layer.visible = True
            if bool(getattr(self._proximity_roi_layer, "_is_creating", False)):
                with suppress(Exception):
                    self.viewer.layers.selection.active = self._proximity_roi_layer
                self._update_proximity_roi_status_label()
                return
        self._sync_proximity_roi_annotation_layers(layer)
        try:
            self.viewer.layers.selection.active = self._proximity_roi_layer
            self._proximity_roi_layer.visible = True
            self._proximity_roi_layer.mode = "add_polygon"
        except (AttributeError, RuntimeError, TypeError, ValueError) as e:
            QMessageBox.warning(
                self,
                "ROI unavailable",
                f"Could not enter ROI drawing mode: {e!r}",
            )
            return
        self._update_proximity_roi_status_label()

    def _clear_proximity_roi_layer(self) -> None:
        layer = self._proximity_roi_layer
        if layer is None or layer not in self.viewer.layers:
            return
        if bool(getattr(layer, "_is_creating", False)):
            with suppress(Exception):
                layer._finish_drawing()
            with suppress(Exception):
                layer._is_creating = False
        with suppress(Exception):
            layer.data = []
            layer.shape_type = []
        with suppress(Exception):
            layer.selected_data = set()
        text_layer = self._proximity_roi_text_layer
        if text_layer is not None and text_layer in self.viewer.layers:
            with suppress(Exception):
                self.viewer.layers.remove(text_layer)
        self._proximity_roi_text_layer = None
        self._update_proximity_roi_status_label()

    def _export_proximity_roi(self, roi_layer=None) -> None:
        if roi_layer is None:
            roi_layer = self._proximity_roi_layer
        if roi_layer is None or roi_layer not in self.viewer.layers:
            QMessageBox.information(self, "No ROI", "Draw at least one ROI first.")
            return
        if bool(getattr(roi_layer, "_is_creating", False)):
            QMessageBox.information(
                self,
                "ROI in progress",
                "Finish the current ROI polygon before exporting.",
            )
            return
        shape_data = list(getattr(roi_layer, "data", []) or [])
        if not shape_data:
            QMessageBox.information(self, "No ROI", "Draw at least one ROI first.")
            return

        roi_metadata = dict(getattr(roi_layer, "metadata", {}) or {})
        source_layer_name = str(roi_metadata.get("proximity_source_layer") or "")
        layer = None
        if source_layer_name:
            with suppress(Exception):
                layer = next(
                    candidate
                    for candidate in self.viewer.layers
                    if candidate is not roi_layer
                    and str(getattr(candidate, "name", "")) == source_layer_name
                )
        if layer is None:
            layer = self._selected_proximity_source_seg_layer() or self._selected_proximity_target_seg_layer()
        if layer is None:
            QMessageBox.information(self, "No layer", "Choose a segmentation layer first.")
            return
        try:
            data = _normalized_layer_array(layer, allow_float=False)
            specs = self._proximity_roi_specs_for_layer(layer, roi_layer=roi_layer)
            if not specs:
                raise ValueError("The ROI is incomplete or outside the image.")
            roi_labels = _proximity_roi_label_mask(
                specs,
                data_shape=tuple(int(value) for value in data.shape),
            )
        except (TypeError, ValueError) as e:
            QMessageBox.warning(self, "Invalid ROI", str(e))
            return

        source = str(_get_source_from_layer(layer) or "")
        source_stem = os.path.splitext(os.path.basename(source))[0]
        if not source_stem:
            source_stem = str(getattr(layer, "name", "proximity") or "proximity")
        default_path = os.path.join(
            os.path.dirname(source) if source else "",
            f"{source_stem}_proximity_roi.tif",
        )
        path, _selected_filter = QFileDialog.getSaveFileName(
            self,
            "Export proximity ROI",
            default_path,
            "TIFF label mask (*.tif *.tiff)",
        )
        if not path:
            return
        if os.path.splitext(path)[1].lower() not in {".tif", ".tiff"}:
            path += ".tif"

        layer_metadata = dict(getattr(layer, "metadata", {}) or {})
        dims_tag = str(
            layer_metadata.get("dims") or layer_metadata.get("dims_out") or ""
        ).upper()
        if len(dims_tag) != roi_labels.ndim:
            dims_tag = {2: "YX", 3: "ZYX", 4: "TZYX"}.get(
                roi_labels.ndim,
                "".join("ABCDEFGHIJKLMNOPQRSTUVWXYZ"[-roi_labels.ndim :]),
            )
        layer_metadata.update(
            {
                "dims": dims_tag,
                "dims_out": dims_tag,
                "is_proximity_roi_mask": True,
                "proximity_source_layer": str(getattr(layer, "name", "")),
                "proximity_roi_count": len(specs),
            }
        )
        _ensure_unit_in_metadata(layer_metadata)
        meta = {
            "scale": tuple(float(value) for value in getattr(layer, "scale", ())),
            "metadata": layer_metadata,
            "name": f"{source_stem} proximity ROI",
            "layer_type": "labels",
        }
        try:
            write_single_image(path, roi_labels, meta)
        except (OSError, ValueError, RuntimeError, TypeError) as e:
            QMessageBox.critical(self, "Export failed", f"Failed to export ROI: {e!r}")
            return
        QMessageBox.information(
            self,
            "Export complete",
            f"Exported {len(specs)} ROI(s) to:\n{path}",
        )

    def _on_proximity_roi_changed(self, event=None) -> None:
        self._update_proximity_roi_status_label()
        roi_layer = self._proximity_roi_layer
        if (
            roi_layer is not None
            and roi_layer in self.viewer.layers
            and bool(getattr(roi_layer, "_is_creating", False))
        ):
            return
        self._sync_proximity_roi_annotation_layers()

    def _proximity_roi_label_data(self, layer) -> tuple[np.ndarray, list[str]]:
        specs = self._proximity_roi_specs_for_layer(layer)
        if not specs:
            ndim = self._proximity_shape_ndim(layer) if layer is not None else 2
            return np.zeros((0, ndim), dtype=float), []
        points: list[np.ndarray] = []
        labels: list[str] = []
        for spec in specs:
            roi_mask_2d = self._proximity_mask_frame(spec.mask, spec.plane_index)
            if roi_mask_2d is not None and np.any(roi_mask_2d):
                distance = ndi.distance_transform_edt(roi_mask_2d)
                row, col = np.unravel_index(int(np.argmax(distance)), distance.shape)
            else:
                continue
            coords = np.asarray(tuple(spec.plane_index) + (float(row), float(col)), dtype=float)
            points.append(coords)
            labels.append(str(spec.name))
        if not points:
            ndim = self._proximity_shape_ndim(layer) if layer is not None else 2
            return np.zeros((0, ndim), dtype=float), []
        return np.asarray(points, dtype=float), labels

    def _sync_proximity_roi_annotation_layers(self, layer=None) -> None:
        roi_layer = self._proximity_roi_layer
        if layer is None:
            layer = self._selected_proximity_source_seg_layer() or self._selected_proximity_target_seg_layer()
        if roi_layer is None or roi_layer not in self.viewer.layers or layer is None:
            text_layer = self._proximity_roi_text_layer
            if text_layer is not None and text_layer in self.viewer.layers:
                with suppress(Exception):
                    self.viewer.layers.remove(text_layer)
            self._proximity_roi_text_layer = None
            return
        if bool(getattr(roi_layer, "_is_creating", False)):
            return
        roi_was_active = False
        roi_mode = str(getattr(roi_layer, "mode", "") or "")
        with suppress(Exception):
            roi_was_active = self.viewer.layers.selection.active is roi_layer
            roi_layer.visible = True
        text_points, text_labels = self._proximity_roi_label_data(layer)
        if text_points.size == 0 or not text_labels:
            text_layer = self._proximity_roi_text_layer
            if text_layer is not None and text_layer in self.viewer.layers:
                with suppress(Exception):
                    text_layer.data = np.zeros((0, max(2, self._proximity_shape_ndim(layer))), dtype=float)
                    text_layer.text.values = []
                    text_layer.visible = False
            with suppress(Exception):
                if self._proximity_layer is not None and self._proximity_layer in self.viewer.layers:
                    overlap_idx = self.viewer.layers.index(self._proximity_layer)
                    roi_idx = self.viewer.layers.index(roi_layer)
                    if roi_idx < overlap_idx:
                        self.viewer.layers.move(roi_idx, overlap_idx + 1)
            return
        text_ndim = int(text_points.shape[1])
        text_layer = self._proximity_roi_text_layer
        if (
            text_layer is not None
            and text_layer in self.viewer.layers
            and int(getattr(text_layer, "ndim", text_ndim)) != text_ndim
        ):
            with suppress(Exception):
                self.viewer.layers.remove(text_layer)
            self._proximity_roi_text_layer = None
        if self._proximity_roi_text_layer is None or self._proximity_roi_text_layer not in self.viewer.layers:
            self._proximity_roi_text_layer = self.viewer.add_points(
                text_points,
                name=self._next_layer_name("proximity roi labels"),
                size=1,
                opacity=1.0,
                face_color="transparent",
                text={
                    "string": text_labels,
                    "size": 14,
                    "color": "white",
                    "anchor": "center",
                    "translation": [0.0] * text_ndim,
                },
                scale=tuple(float(v) for v in getattr(layer, "scale", (1.0,) * text_points.shape[1])),
                metadata={"is_proximity_roi_label": True},
                **_spatial_unit_kwargs(layer, text_ndim),
            )
        else:
            with suppress(Exception):
                self._proximity_roi_text_layer.data = text_points
                self._proximity_roi_text_layer.text.values = text_labels
                self._proximity_roi_text_layer.text.translation = np.zeros(
                    text_ndim,
                    dtype=float,
                )
                self._proximity_roi_text_layer.scale = tuple(
                    float(v) for v in getattr(layer, "scale", (1.0,) * max(2, text_points.shape[1]))
                )
                self._proximity_roi_text_layer.visible = True
                self._proximity_roi_text_layer.face_color = "transparent"
        with suppress(Exception):
            if self._proximity_layer is not None and self._proximity_layer in self.viewer.layers:
                overlap_idx = self.viewer.layers.index(self._proximity_layer)
                roi_idx = self.viewer.layers.index(roi_layer)
                if roi_idx < overlap_idx:
                    self.viewer.layers.move(roi_idx, overlap_idx + 1)
            if (
                self._proximity_roi_text_layer is not None
                and self._proximity_roi_text_layer in self.viewer.layers
                and roi_layer in self.viewer.layers
            ):
                roi_idx = self.viewer.layers.index(roi_layer)
                text_idx = self.viewer.layers.index(self._proximity_roi_text_layer)
                if text_idx < roi_idx:
                    self.viewer.layers.move(text_idx, roi_idx + 1)
            if roi_was_active and roi_layer in self.viewer.layers:
                self.viewer.layers.selection.active = roi_layer
                if roi_mode:
                    roi_layer.mode = roi_mode

    def _prime_proximity_roi_cursor_state(self, layer, event) -> None:
        if layer is None or event is None:
            return
        with suppress(Exception):
            mode = str(getattr(layer, "mode", "") or "")
            if mode not in {"add_polygon_lasso", "add_path"}:
                return
            if getattr(layer, "_last_cursor_position", None) is None:
                layer._last_cursor_position = np.asarray(event.pos, dtype=float)

    def _update_proximity_roi_status_label(self) -> None:
        label = self._proximity_roi_status_label
        if label is None:
            return
        roi_layer = self._proximity_roi_layer
        if (
            roi_layer is not None
            and roi_layer in self.viewer.layers
            and bool(getattr(roi_layer, "_is_creating", False))
        ):
            label.setText("Drawing ROI...")
            return
        has_roi = False
        if roi_layer is not None and roi_layer in self.viewer.layers:
            with suppress(Exception):
                has_roi = len(list(getattr(roi_layer, "data", []) or [])) > 0
        if has_roi and roi_layer is not None:
            with suppress(Exception):
                count = len(list(getattr(roi_layer, "data", []) or []))
                label.setText(f"{count} ROI(s)")
                return
        label.setText("Full image")

    def _proximity_physical_unit_label(self, result: ProximityResult | None, *, power_delta: int = 0) -> str:
        if result is None:
            return "phys"
        unit = str(result.unit or "unit")
        power = max(1, int(getattr(result, "spatial_power", 2)) + int(power_delta))
        return unit if power == 1 else f"{unit}^{power}"

    def _proximity_threshold_unit_label(self, result: ProximityResult | None) -> str:
        if result is None:
            return "phys"
        return self._proximity_physical_unit_label(
            result,
            power_delta=-1 if bool(getattr(result, "surface_only", False)) else 0,
        )

    def _proximity_count_unit_label(self, result: ProximityResult | None) -> str:
        if result is None:
            return "px/vx"
        return "vx" if int(getattr(result, "spatial_power", 2)) >= 3 else "px"

    def _proximity_threshold_count_unit_label(self, result: ProximityResult | None) -> str:
        base = self._proximity_count_unit_label(result)
        if result is not None and bool(getattr(result, "surface_only", False)):
            return f"surface {base}"
        return base

    def _selected_proximity_roi_id(self) -> int | None:
        table = self._proximity_summary_table
        if table is None:
            return None
        items = table.selectedItems()
        if not items:
            return None
        roi_id = items[0].data(Qt.UserRole)
        return None if roi_id is None else int(roi_id)

    def _proximity_roi_specs_for_layer(self, layer, *, roi_layer=None) -> list[ProximityRoiSpec]:
        if roi_layer is None:
            roi_layer = self._proximity_roi_layer
        if roi_layer is None or roi_layer not in self.viewer.layers:
            return []
        shape_data = list(getattr(roi_layer, "data", []) or [])
        if not shape_data:
            return []
        try:
            data = _normalized_layer_array(layer, allow_float=False)
        except (TypeError, ValueError):
            return []
        dims_tag = str((getattr(layer, "metadata", {}) or {}).get("dims") or "").upper()
        current_step = tuple(int(v) for v in getattr(self.viewer.dims, "current_step", ()))
        return _proximity_roi_specs_from_shapes(
            shape_data,
            data_shape=tuple(int(value) for value in data.shape),
            dims_tag=dims_tag,
            current_step=current_step,
        )

    def _proximity_mask_frame(self, mask: np.ndarray, plane_index: tuple[int, ...]) -> np.ndarray:
        mask = np.asarray(mask, dtype=bool)
        if mask.ndim == 2:
            return mask
        if mask.ndim == 3:
            if len(plane_index) >= 1:
                return mask[int(plane_index[0])]
            return np.any(mask, axis=0)
        if mask.ndim == 4:
            if len(plane_index) >= 2:
                return mask[int(plane_index[0]), int(plane_index[1])]
            if len(plane_index) >= 1:
                return np.any(mask[int(plane_index[0])], axis=0)
            return np.any(mask, axis=(0, 1))
        raise ValueError(f"Unsupported proximity mask ndim: {mask.ndim}")

    def _build_tracking_page(self) -> QWidget:
        self._tracking_layer_combo = QComboBox()
        self._tracking_raw_layer_combo = QComboBox()
        self._tracking_layer_combo.currentIndexChanged.connect(
            self._rebuild_tracking_raw_layer_combo
        )
        match_btn = QPushButton("Match")
        match_btn.clicked.connect(self._on_match_tracking_clicked)
        self._tracking_match_btn = match_btn
        run_btn = QPushButton("Run Tracking")
        run_btn.clicked.connect(self._on_run_tracking_clicked)
        self._tracking_run_btn = run_btn

        self.track_max_dist = QDoubleSpinBox()
        self.track_max_dist.setDecimals(2)
        self.track_max_dist.setRange(0.1, 1e6)
        self.track_max_dist.setValue(5.0)
        self.track_max_dist.setLocale(QLocale.c())

        self.track_distance_weight = QDoubleSpinBox()
        self.track_distance_weight.setDecimals(2)
        self.track_distance_weight.setRange(0.0, 10.0)
        self.track_distance_weight.setSingleStep(0.05)
        self.track_distance_weight.setValue(0.1)
        self.track_distance_weight.setLocale(QLocale.c())

        self.track_overlap_weight = QDoubleSpinBox()
        self.track_overlap_weight.setDecimals(2)
        self.track_overlap_weight.setRange(0.0, 10.0)
        self.track_overlap_weight.setSingleStep(0.05)
        self.track_overlap_weight.setValue(0.4)
        self.track_overlap_weight.setLocale(QLocale.c())

        self.track_point_support_weight = QDoubleSpinBox()
        self.track_point_support_weight.setDecimals(2)
        self.track_point_support_weight.setRange(0.0, 10.0)
        self.track_point_support_weight.setSingleStep(0.05)
        self.track_point_support_weight.setValue(1.5)
        self.track_point_support_weight.setLocale(QLocale.c())

        self.track_cost_cutoff = QDoubleSpinBox()
        self.track_cost_cutoff.setDecimals(2)
        self.track_cost_cutoff.setRange(0.0, 100.0)
        self.track_cost_cutoff.setSingleStep(0.05)
        self.track_cost_cutoff.setValue(1.0)
        self.track_cost_cutoff.setLocale(QLocale.c())

        self.track_max_neighbors = QSpinBox()
        self.track_max_neighbors.setRange(1, 10**6)
        self.track_max_neighbors.setValue(16)

        self.track_min_link_size = QSpinBox()
        self.track_min_link_size.setRange(0, 10**9)
        self.track_min_link_size.setValue(20)
        self.track_min_link_size.setToolTip(
            "Objects at or below this size are terminal birth/death observations."
        )

        self.track_sample_points = QSpinBox()
        self.track_sample_points.setRange(0, 10**8)
        self.track_sample_points.setValue(3000)
        self.track_sample_points.setToolTip(
            "Total sampled points per frame. Object budgets use integrated raw "
            "intensity; "
            "points within each object use intensity-weighted CVT. Zero uses every "
            "object point."
        )

        match_ctrl_box = QGroupBox("Match")
        match_grid = QGridLayout(match_ctrl_box)
        match_grid.setHorizontalSpacing(self._scaled_px(10))
        match_grid.setVerticalSpacing(self._scaled_px(6))
        tracking_raw_label = QLabel("Raw image")
        self._tracking_raw_layer_combo.setSizePolicy(
            QSizePolicy.Expanding,
            QSizePolicy.Fixed,
        )
        self._tracking_raw_layer_combo.setToolTip(
            "Raw intensity image used for integrated-intensity sample budgets, "
            "intensity-weighted CVT, and intensity-weighted centroids."
        )
        match_grid.addWidget(
            tracking_raw_label,
            0,
            0,
            alignment=Qt.AlignRight | Qt.AlignVCenter,
        )
        match_grid.addWidget(self._tracking_raw_layer_combo, 0, 1, 1, 3)

        tracking_layer_label = QLabel("Segmentation")
        self._tracking_layer_combo.setSizePolicy(
            QSizePolicy.Expanding,
            QSizePolicy.Fixed,
        )
        match_grid.addWidget(
            tracking_layer_label,
            1,
            0,
            alignment=Qt.AlignRight | Qt.AlignVCenter,
        )
        match_grid.addWidget(self._tracking_layer_combo, 1, 1, 1, 3)
        match_grid.addWidget(QLabel("Max distance (px/vox):"), 2, 0, alignment=Qt.AlignRight)
        match_grid.addWidget(self.track_max_dist, 2, 1)
        match_grid.addWidget(QLabel("Sample points/frame:"), 2, 2, alignment=Qt.AlignRight)
        match_grid.addWidget(self.track_sample_points, 2, 3)
        match_grid.addWidget(match_btn, 3, 0, 1, 4)
        self.match_progress = QProgressBar()
        self.match_progress.setRange(0, 1)
        self.match_progress.setValue(0)
        self.match_progress.setFormat("Idle")
        match_grid.addWidget(self.match_progress, 4, 0, 1, 4)
        self._tracking_match_label = QLabel("Match: <not run>")
        self._tracking_match_label.setWordWrap(True)
        self._tracking_match_label.setTextInteractionFlags(Qt.TextSelectableByMouse)
        match_grid.addWidget(self._tracking_match_label, 5, 0, 1, 4)
        self._tracking_match_label.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        match_grid.setColumnStretch(0, 0)
        match_grid.setColumnStretch(1, 1)
        match_grid.setColumnStretch(2, 0)
        match_grid.setColumnStretch(3, 1)

        track_ctrl_box = QGroupBox("Track")
        track_grid = QGridLayout(track_ctrl_box)
        track_grid.setHorizontalSpacing(self._scaled_px(10))
        track_grid.setVerticalSpacing(self._scaled_px(6))
        track_grid.addWidget(QLabel("Cost cutoff:"), 0, 0, alignment=Qt.AlignRight)
        track_grid.addWidget(self.track_cost_cutoff, 0, 1)
        track_grid.addWidget(QLabel("Distance weight:"), 0, 2, alignment=Qt.AlignRight)
        track_grid.addWidget(self.track_distance_weight, 0, 3)
        track_grid.addWidget(QLabel("Coverage weight:"), 1, 0, alignment=Qt.AlignRight)
        track_grid.addWidget(self.track_overlap_weight, 1, 1)
        track_grid.addWidget(QLabel("Point support weight:"), 1, 2, alignment=Qt.AlignRight)
        track_grid.addWidget(self.track_point_support_weight, 1, 3)
        track_grid.addWidget(QLabel("Min link size:"), 2, 0, alignment=Qt.AlignRight)
        track_grid.addWidget(self.track_min_link_size, 2, 1)
        track_grid.addWidget(QLabel("Max neighbors:"), 2, 2, alignment=Qt.AlignRight)
        track_grid.addWidget(self.track_max_neighbors, 2, 3)
        track_grid.addWidget(run_btn, 3, 0, 1, 4)
        self.track_progress = QProgressBar()
        self.track_progress.setRange(0, 1)
        self.track_progress.setValue(0)
        self.track_progress.setFormat("Idle")
        track_grid.addWidget(self.track_progress, 4, 0, 1, 4)
        self._tracking_export_btn = QPushButton("Export")
        self._tracking_export_btn.clicked.connect(
            self._export_tracking_events
        )
        self._tracking_export_btn.setMinimumWidth(0)
        self._tracking_export_btn.setMaximumWidth(16777215)
        self._tracking_export_btn.setSizePolicy(
            QSizePolicy.Expanding,
            QSizePolicy.Fixed,
        )
        track_grid.addWidget(self._tracking_export_btn, 5, 0, 1, 4)
        track_grid.setColumnStretch(0, 0)
        track_grid.setColumnStretch(1, 1)
        track_grid.setColumnStretch(2, 0)
        track_grid.setColumnStretch(3, 1)

        refine_box = QGroupBox("Refine")
        refine_grid = QGridLayout(refine_box)
        refine_grid.setHorizontalSpacing(self._scaled_px(10))
        refine_grid.setVerticalSpacing(self._scaled_px(6))
        self._tracking_refine_object_combo = QComboBox()
        self._tracking_refine_object_combo.setEnabled(False)
        self._tracking_refine_object_combo.setSizePolicy(
            QSizePolicy.Expanding,
            QSizePolicy.Fixed,
        )
        self._tracking_refine_object_combo.currentIndexChanged.connect(
            self._on_tracking_refine_object_changed
        )
        self._tracking_refine_selected_table = self._make_tracking_refine_link_table()
        self._tracking_refine_available_table = self._make_tracking_refine_link_table()
        self._tracking_refine_selected_table.itemSelectionChanged.connect(
            lambda: self._on_tracking_refine_link_table_selection_changed(
                self._tracking_refine_selected_table
            )
        )
        self._tracking_refine_available_table.itemSelectionChanged.connect(
            lambda: self._on_tracking_refine_link_table_selection_changed(
                self._tracking_refine_available_table
            )
        )
        self._tracking_refine_apply_btn = QPushButton()
        self._tracking_refine_apply_btn.setIcon(
            self.style().standardIcon(QStyle.SP_ArrowLeft)
        )
        self._tracking_refine_apply_btn.setToolTip("Add selected candidate link")
        self._tracking_refine_apply_btn.setFixedWidth(self._scaled_px(34))
        self._tracking_refine_apply_btn.setEnabled(False)
        self._tracking_refine_apply_btn.clicked.connect(
            self._on_tracking_refine_apply_clicked
        )
        self._tracking_refine_remove_btn = QPushButton()
        self._tracking_refine_remove_btn.setIcon(
            self.style().standardIcon(QStyle.SP_ArrowRight)
        )
        self._tracking_refine_remove_btn.setToolTip("Remove selected effective link")
        self._tracking_refine_remove_btn.setFixedWidth(self._scaled_px(34))
        self._tracking_refine_remove_btn.setEnabled(False)
        self._tracking_refine_remove_btn.clicked.connect(
            self._on_tracking_refine_remove_clicked
        )
        self._tracking_refine_undo_btn = QPushButton()
        self._tracking_refine_undo_btn.setIcon(
            self.style().standardIcon(QStyle.SP_ArrowBack)
        )
        self._tracking_refine_undo_btn.setToolTip(
            "Undo refine (Cmd+Z / Ctrl+Z)"
        )
        self._tracking_refine_undo_btn.setFixedWidth(self._scaled_px(34))
        self._tracking_refine_undo_btn.setEnabled(False)
        self._tracking_refine_undo_btn.clicked.connect(
            self._on_tracking_refine_undo_clicked
        )
        self._tracking_refine_undo_shortcut = QShortcut(
            QKeySequence(QKeySequence.StandardKey.Undo), self
        )
        self._tracking_refine_undo_shortcut.setContext(
            Qt.WidgetWithChildrenShortcut
        )
        self._tracking_refine_undo_shortcut.activated.connect(
            self._on_tracking_refine_undo_clicked
        )
        self._tracking_refine_pick_in_btn = QPushButton("Pick IN")
        self._tracking_refine_pick_in_btn.setIcon(
            self.style().standardIcon(QStyle.SP_ArrowBack)
        )
        self._tracking_refine_pick_in_btn.setToolTip(
            "Choose a source object by double-clicking it in the previous frame"
        )
        self._tracking_refine_pick_in_btn.setEnabled(False)
        self._tracking_refine_pick_in_btn.clicked.connect(
            lambda: self._start_tracking_refine_image_pick("incoming")
        )
        self._tracking_refine_pick_out_btn = QPushButton("Pick OUT")
        self._tracking_refine_pick_out_btn.setIcon(
            self.style().standardIcon(QStyle.SP_ArrowForward)
        )
        self._tracking_refine_pick_out_btn.setToolTip(
            "Choose a target object by double-clicking it in the next frame"
        )
        self._tracking_refine_pick_out_btn.setEnabled(False)
        self._tracking_refine_pick_out_btn.clicked.connect(
            lambda: self._start_tracking_refine_image_pick("outgoing")
        )
        self._tracking_refine_pick_confirm_btn = QPushButton("Confirm")
        self._tracking_refine_pick_confirm_btn.setIcon(
            self.style().standardIcon(QStyle.SP_DialogApplyButton)
        )
        self._tracking_refine_pick_confirm_btn.setToolTip(
            "Add the highlighted viewer-selected link"
        )
        self._tracking_refine_pick_confirm_btn.setEnabled(False)
        self._tracking_refine_pick_confirm_btn.clicked.connect(
            self._confirm_tracking_refine_image_pick
        )
        self._tracking_refine_pick_cancel_btn = QPushButton("Cancel")
        self._tracking_refine_pick_cancel_btn.setIcon(
            self.style().standardIcon(QStyle.SP_DialogCancelButton)
        )
        self._tracking_refine_pick_cancel_btn.setToolTip(
            "Cancel viewer link selection"
        )
        self._tracking_refine_pick_cancel_btn.setEnabled(False)
        self._tracking_refine_pick_cancel_btn.clicked.connect(
            self._cancel_tracking_refine_image_pick
        )
        pick_row = QHBoxLayout()
        pick_row.setContentsMargins(0, 0, 0, 0)
        pick_row.setSpacing(self._scaled_px(8))
        pick_row.addWidget(self._tracking_refine_pick_in_btn)
        pick_row.addWidget(self._tracking_refine_pick_out_btn)
        pick_row.addStretch(1)
        pick_row.addWidget(self._tracking_refine_pick_confirm_btn)
        pick_row.addWidget(self._tracking_refine_pick_cancel_btn)
        self._tracking_refine_status_label = QLabel(
            "Refine: <not available>"
        )
        self._tracking_refine_status_label.setWordWrap(True)
        self._tracking_refine_status_label.setTextInteractionFlags(
            Qt.TextSelectableByMouse
        )

        selected_links_panel = QWidget()
        selected_links_layout = QVBoxLayout(selected_links_panel)
        selected_links_layout.setContentsMargins(0, 0, 0, 0)
        selected_links_layout.setSpacing(self._scaled_px(4))
        selected_links_layout.addWidget(QLabel("Effective links"))
        selected_links_layout.addWidget(self._tracking_refine_selected_table)
        selected_links_panel.setMinimumWidth(0)
        selected_links_panel.setSizePolicy(
            QSizePolicy.Expanding,
            QSizePolicy.Fixed,
        )

        transfer_widget = QWidget()
        transfer_layout = QVBoxLayout(transfer_widget)
        transfer_layout.setContentsMargins(0, 0, 0, 0)
        transfer_layout.setSpacing(self._scaled_px(8))
        transfer_layout.addStretch(1)
        transfer_layout.addWidget(self._tracking_refine_apply_btn)
        transfer_layout.addStretch(1)
        transfer_layout.addWidget(self._tracking_refine_remove_btn)
        transfer_layout.addStretch(1)
        transfer_widget.setFixedWidth(self._scaled_px(44))

        available_links_panel = QWidget()
        available_links_layout = QVBoxLayout(available_links_panel)
        available_links_layout.setContentsMargins(0, 0, 0, 0)
        available_links_layout.setSpacing(self._scaled_px(4))
        available_links_layout.addWidget(QLabel("Nearby candidates"))
        available_links_layout.addWidget(self._tracking_refine_available_table)
        available_links_panel.setMinimumWidth(0)
        available_links_panel.setSizePolicy(
            QSizePolicy.Expanding,
            QSizePolicy.Fixed,
        )

        refine_links_layout = QHBoxLayout()
        refine_links_layout.setContentsMargins(0, 0, 0, 0)
        refine_links_layout.setSpacing(self._scaled_px(8))
        refine_links_layout.addWidget(selected_links_panel, 1)
        refine_links_layout.addWidget(transfer_widget, 0)
        refine_links_layout.addWidget(available_links_panel, 1)

        refine_grid.addWidget(QLabel("Object:"), 0, 0, alignment=Qt.AlignRight)
        refine_grid.addWidget(self._tracking_refine_object_combo, 0, 1, 1, 3)
        refine_grid.addWidget(self._tracking_refine_undo_btn, 0, 4)
        refine_grid.addLayout(pick_row, 1, 0, 1, 5)
        refine_grid.addLayout(refine_links_layout, 2, 0, 1, 5)
        refine_grid.addWidget(self._tracking_refine_status_label, 3, 0, 1, 5)
        refine_grid.setColumnStretch(0, 1)
        refine_grid.setColumnStretch(1, 1)
        refine_grid.setColumnStretch(2, 0)
        refine_grid.setColumnStretch(3, 1)
        refine_grid.setColumnStretch(4, 1)

        match_box = QGroupBox("Match Frames")
        match_layout = QVBoxLayout(match_box)
        self._tracking_match_table = QTableWidget()
        self._tracking_match_table.setColumnCount(0)
        self._tracking_match_table.setRowCount(0)
        self._tracking_match_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self._tracking_match_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self._tracking_match_table.setSelectionMode(QAbstractItemView.SingleSelection)
        self._tracking_match_table.itemSelectionChanged.connect(
            self._on_tracking_match_selection_changed
        )
        self._style_table_widget(self._tracking_match_table, min_height=96, max_height=150)
        match_layout.addWidget(self._tracking_match_table)

        linear_box = QGroupBox("Linear Events")
        self._tracking_linear_box = linear_box
        linear_layout = QVBoxLayout(linear_box)
        self._tracking_linear_table = QTableWidget()
        self._tracking_linear_table.setColumnCount(0)
        self._tracking_linear_table.setRowCount(0)
        self._tracking_linear_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self._tracking_linear_table.setSelectionBehavior(QTableWidget.SelectRows)
        self._tracking_linear_table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self._tracking_linear_table.cellClicked.connect(self._on_tracking_event_row_clicked)
        self._tracking_linear_table.itemSelectionChanged.connect(
            self._on_tracking_event_selection_changed
        )
        self._style_table_widget(self._tracking_linear_table, min_height=120, max_height=180)
        linear_layout.addWidget(self._tracking_linear_table)

        fission_box = QGroupBox("Fission Events")
        self._tracking_fission_box = fission_box
        fission_layout = QVBoxLayout(fission_box)
        self._tracking_fission_table = QTableWidget()
        self._tracking_fission_table.setColumnCount(0)
        self._tracking_fission_table.setRowCount(0)
        self._tracking_fission_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self._tracking_fission_table.setSelectionBehavior(QTableWidget.SelectRows)
        self._tracking_fission_table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self._tracking_fission_table.cellClicked.connect(self._on_tracking_event_row_clicked)
        self._tracking_fission_table.itemSelectionChanged.connect(
            self._on_tracking_event_selection_changed
        )
        self._style_table_widget(self._tracking_fission_table, min_height=120, max_height=180)
        fission_layout.addWidget(self._tracking_fission_table)

        fusion_box = QGroupBox("Fusion Events")
        self._tracking_fusion_box = fusion_box
        fusion_layout = QVBoxLayout(fusion_box)
        self._tracking_fusion_table = QTableWidget()
        self._tracking_fusion_table.setColumnCount(0)
        self._tracking_fusion_table.setRowCount(0)
        self._tracking_fusion_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self._tracking_fusion_table.setSelectionBehavior(QTableWidget.SelectRows)
        self._tracking_fusion_table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self._tracking_fusion_table.cellClicked.connect(self._on_tracking_event_row_clicked)
        self._tracking_fusion_table.itemSelectionChanged.connect(
            self._on_tracking_event_selection_changed
        )
        self._style_table_widget(self._tracking_fusion_table, min_height=120, max_height=180)
        fusion_layout.addWidget(self._tracking_fusion_table)

        split_box = QGroupBox("Split-Merge Events")
        self._tracking_split_merge_box = split_box
        split_layout = QVBoxLayout(split_box)
        self._tracking_split_merge_table = QTableWidget()
        self._tracking_split_merge_table.setColumnCount(0)
        self._tracking_split_merge_table.setRowCount(0)
        self._tracking_split_merge_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self._tracking_split_merge_table.setSelectionBehavior(QTableWidget.SelectRows)
        self._tracking_split_merge_table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        self._tracking_split_merge_table.cellClicked.connect(
            self._on_tracking_event_row_clicked
        )
        self._tracking_split_merge_table.itemSelectionChanged.connect(
            self._on_tracking_event_selection_changed
        )
        self._style_table_widget(self._tracking_split_merge_table, min_height=120, max_height=180)
        split_layout.addWidget(self._tracking_split_merge_table)

        visualization_box = QGroupBox("Visualization")
        event_layout = QVBoxLayout(visualization_box)
        event_view_toolbar = QHBoxLayout()
        self._tracking_transition_combo = QComboBox()
        self._tracking_transition_combo.setEnabled(False)
        self._tracking_transition_combo.activated.connect(
            self._on_tracking_transition_changed
        )
        self._tracking_event_kind_combo = QComboBox()
        for label, kind in (
            ("Linear", "linear"),
            ("Fission", "fission"),
            ("Fusion", "fusion"),
            ("Split-Merge", "split-merge"),
        ):
            self._tracking_event_kind_combo.addItem(label, kind)
        self._tracking_event_kind_combo.setEnabled(False)
        self._tracking_event_kind_combo.activated.connect(
            self._on_tracking_event_kind_changed
        )
        event_view_toolbar.addWidget(QLabel("Transition:"))
        event_view_toolbar.addWidget(self._tracking_transition_combo, 1)
        event_view_toolbar.addWidget(QLabel("Event:"))
        event_view_toolbar.addWidget(self._tracking_event_kind_combo)
        self._tracking_statistics_export_btn = QPushButton("Export Statistics")
        self._tracking_statistics_export_btn.setEnabled(False)
        self._tracking_statistics_export_btn.clicked.connect(
            self._export_tracking_statistics
        )
        event_view_toolbar.addWidget(self._tracking_statistics_export_btn)
        self._tracking_statistics_import_btn = QPushButton("Import Events")
        self._tracking_statistics_import_btn.setEnabled(False)
        self._tracking_statistics_import_btn.setToolTip(
            "Restore effective links from an exported event table"
        )
        self._tracking_statistics_import_btn.clicked.connect(
            self._import_tracking_statistics
        )
        event_view_toolbar.addWidget(self._tracking_statistics_import_btn)
        self._style_toolbar_layout(event_view_toolbar)

        unlinked_box = QGroupBox("Unlinked Large Objects")
        self._tracking_unlinked_box = unlinked_box
        unlinked_layout = QVBoxLayout(unlinked_box)
        self._tracking_unlinked_table = QTableWidget(0, 4)
        self._tracking_unlinked_table.setHorizontalHeaderLabels(
            ("Frame", "Object", "Size", "Missing")
        )
        self._tracking_unlinked_table.setEditTriggers(
            QAbstractItemView.NoEditTriggers
        )
        self._tracking_unlinked_table.setSelectionBehavior(
            QAbstractItemView.SelectRows
        )
        self._tracking_unlinked_table.setSelectionMode(
            QAbstractItemView.SingleSelection
        )
        self._tracking_unlinked_table.setEnabled(False)
        self._tracking_unlinked_table.cellClicked.connect(
            self._on_tracking_unlinked_object_clicked
        )
        self._style_table_widget(
            self._tracking_unlinked_table,
            min_height=120,
            max_height=200,
        )
        unlinked_header = self._tracking_unlinked_table.horizontalHeader()
        unlinked_header.setSectionResizeMode(0, QHeaderView.ResizeToContents)
        unlinked_header.setSectionResizeMode(1, QHeaderView.ResizeToContents)
        unlinked_header.setSectionResizeMode(2, QHeaderView.ResizeToContents)
        unlinked_header.setSectionResizeMode(3, QHeaderView.Stretch)
        unlinked_layout.addWidget(self._tracking_unlinked_table)
        self.track_min_link_size.valueChanged.connect(
            lambda _value: self._populate_tracking_unlinked_table(force=True)
        )

        event_layout.addLayout(event_view_toolbar)
        event_layout.addWidget(linear_box)
        event_layout.addWidget(fission_box)
        event_layout.addWidget(fusion_box)
        event_layout.addWidget(split_box)
        event_layout.addWidget(unlinked_box)
        event_layout.setSpacing(self._scaled_px(8))
        self._style_group_box(match_ctrl_box)
        self._style_group_box(track_ctrl_box)
        self._style_group_box(refine_box)
        self._style_group_box(match_box)
        self._style_group_box(visualization_box)
        for box in (
            unlinked_box,
            linear_box,
            fission_box,
            fusion_box,
            split_box,
        ):
            self._style_group_box(box)
        return self._build_panel_page(
            match_ctrl_box,
            match_box,
            track_ctrl_box,
            visualization_box,
            refine_box,
        )

    def _rebuild_tracking_layer_combo(self) -> None:
        if self._tracking_layer_combo is None:
            return
        current_name = self._tracking_layer_combo.currentText()
        self._tracking_layer_combo.blockSignals(True)
        self._tracking_layer_combo.clear()
        for layer_item in self.viewer.layers:
            if _is_tracking_candidate_layer(layer_item):
                self._tracking_layer_combo.addItem(layer_item.name, layer_item)
        if self._tracking_layer_combo.count() > 0:
            idx = self._tracking_layer_combo.findText(current_name)
            self._tracking_layer_combo.setCurrentIndex(idx if idx >= 0 else 0)
        self._tracking_layer_combo.blockSignals(False)
        self._tracking_layer_combo.setEnabled(self._tracking_layer_combo.count() > 0)
        self._rebuild_tracking_raw_layer_combo()
        if self._tracking_statistics_import_btn is not None:
            self._tracking_statistics_import_btn.setEnabled(
                self._tracking_layer_combo.count() > 0
                and not self._thread_is_running("_tracking_thread")
            )

    def _rebuild_tracking_raw_layer_combo(self, _index: int = -1) -> None:
        combo = self._tracking_raw_layer_combo
        if combo is None:
            return
        reference = self._selected_tracking_layer()
        current_name = combo.currentText()
        candidates = [
            layer_item
            for layer_item in self.viewer.layers
            if _is_tracking_raw_candidate_layer(layer_item, reference)
        ]
        combo.blockSignals(True)
        combo.clear()
        for layer_item in candidates:
            combo.addItem(layer_item.name, layer_item)
        if combo.count() > 0:
            idx = combo.findText(current_name)
            if idx < 0 and reference is not None:
                group_id = _get_group_id(reference)
                if group_id is not None:
                    preferred = next(
                        (
                            layer_item
                            for layer_item in reversed(candidates)
                            if _get_group_id(layer_item) == group_id
                        ),
                        None,
                    )
                    if preferred is not None:
                        idx = combo.findText(preferred.name)
            combo.setCurrentIndex(idx if idx >= 0 else 0)
        combo.setEnabled(combo.count() > 0)
        combo.blockSignals(False)

    def _proximity_combo_for_role(self, role: str) -> QComboBox | None:
        return {
            "source_raw": self._proximity_source_raw_combo,
            "source_seg": self._proximity_source_seg_combo,
            "target_raw": self._proximity_target_raw_combo,
            "target_seg": self._proximity_target_seg_combo,
        }.get(str(role))

    def _proximity_role_combos(self) -> tuple[tuple[str, QComboBox | None], ...]:
        return (
            ("source_raw", self._proximity_source_raw_combo),
            ("source_seg", self._proximity_source_seg_combo),
            ("target_raw", self._proximity_target_raw_combo),
            ("target_seg", self._proximity_target_seg_combo),
        )

    def _remember_proximity_layer_combo(self, role: str) -> None:
        if bool(getattr(self, "_proximity_rebuilding_layer_combos", False)):
            return
        combo = self._proximity_combo_for_role(role)
        if combo is None or combo.count() == 0:
            return
        layer = combo.currentData()
        name = str(getattr(layer, "name", "") or combo.currentText() or "")
        if name:
            self._proximity_layer_selection_names[str(role)] = name

    def _selected_proximity_layer_for_role(self, role: str):
        combo = self._proximity_combo_for_role(role)
        if combo is None or combo.count() == 0:
            return None
        name = str(combo.currentText() or "")
        if name:
            self._proximity_layer_selection_names[str(role)] = name
        preferred_name = self._proximity_layer_selection_names.get(str(role), "")
        if preferred_name:
            for layer_item in self.viewer.layers:
                if str(getattr(layer_item, "name", "")) == preferred_name:
                    return layer_item
        return combo.currentData()

    def _rebuild_proximity_layer_combos(self) -> None:
        combos = (
            self._proximity_source_raw_combo,
            self._proximity_source_seg_combo,
            self._proximity_target_raw_combo,
            self._proximity_target_seg_combo,
        )
        if any(combo is None for combo in combos):
            return
        current_names = {
            combo: self._proximity_layer_selection_names.get(role) or combo.currentText()
            for role, combo in self._proximity_role_combos()
            if combo is not None
        }
        raw_candidates = [layer_item for layer_item in self.viewer.layers if is_proximity_raw_candidate_layer(layer_item)]
        seg_candidates = [layer_item for layer_item in self.viewer.layers if is_proximity_segmentation_candidate_layer(layer_item)]

        self._proximity_rebuilding_layer_combos = True
        try:
            for combo, candidates in (
                (self._proximity_source_raw_combo, raw_candidates),
                (self._proximity_target_raw_combo, raw_candidates),
                (self._proximity_source_seg_combo, seg_candidates),
                (self._proximity_target_seg_combo, seg_candidates),
            ):
                combo.blockSignals(True)
                combo.clear()
                for layer_item in candidates:
                    combo.addItem(layer_item.name, layer_item)
                if combo.count() > 0:
                    idx = combo.findText(current_names.get(combo, ""))
                    combo.setCurrentIndex(idx if idx >= 0 else 0)
                combo.blockSignals(False)

            if (
                self._proximity_source_raw_combo.count() > 1
                and self._proximity_target_raw_combo.count() > 1
                and self._proximity_source_raw_combo.currentIndex() == self._proximity_target_raw_combo.currentIndex()
            ):
                self._proximity_target_raw_combo.setCurrentIndex(1)
            if (
                self._proximity_source_seg_combo.count() > 1
                and self._proximity_target_seg_combo.count() > 1
                and self._proximity_source_seg_combo.currentIndex() == self._proximity_target_seg_combo.currentIndex()
            ):
                self._proximity_target_seg_combo.setCurrentIndex(1)
        finally:
            self._proximity_rebuilding_layer_combos = False
        for role, _combo in self._proximity_role_combos():
            self._remember_proximity_layer_combo(role)

    def _selected_proximity_source_raw_layer(self):
        return self._selected_proximity_layer_for_role("source_raw")

    def _selected_proximity_source_seg_layer(self):
        return self._selected_proximity_layer_for_role("source_seg")

    def _selected_proximity_target_raw_layer(self):
        return self._selected_proximity_layer_for_role("target_raw")

    def _selected_proximity_target_seg_layer(self):
        return self._selected_proximity_layer_for_role("target_seg")

    def _set_proximity_role_image_color(self, layer, colormap: str, *, blending: str = "additive") -> None:
        if layer is None:
            return
        with suppress(Exception):
            if hasattr(layer, "colormap"):
                layer.colormap = colormap
            if hasattr(layer, "blending"):
                layer.blending = blending

    def _is_labels_layer(self, layer) -> bool:
        if layer is None:
            return False
        layer_type = str(getattr(layer, "_type_string", "") or "").lower()
        layer_class = layer.__class__.__name__.lower()
        return layer_type == "labels" or layer_class == "labels"

    def _apply_proximity_label_colors(self, layer, colors: dict[int, np.ndarray], *, blending: str = "additive") -> None:
        if layer is None:
            return
        transparent = [0.0, 0.0, 0.0, 0.0]
        color_map: dict[int | None, list[float]] = {None: transparent, 0: transparent}
        for label, color in colors.items():
            rgba = np.asarray(color, dtype=float).ravel()
            if rgba.size < 4:
                rgba = np.pad(rgba, (0, 4 - rgba.size), constant_values=1.0)
            color_map[int(label)] = rgba[:4].tolist()
        with suppress(Exception):
            layer.colormap = color_map
        with suppress(Exception):
            layer.blending = blending
        with suppress(Exception):
            layer.editable = False
        with suppress(Exception):
            layer.rendering = "iso_categorical"
        with suppress(Exception):
            layer.refresh()

    def _set_proximity_role_mask_color(self, layer, rgba: np.ndarray, colormap: str) -> None:
        if layer is None:
            return
        if self._is_labels_layer(layer):
            with suppress(Exception):
                data = np.asarray(getattr(layer, "data", None))
                labels = [int(v) for v in np.unique(data) if int(v) > 0]
                self._apply_proximity_label_colors(
                    layer,
                    {int(label): np.asarray(rgba, dtype=float) for label in labels[:4096]},
                    blending="additive",
                )
            return
        self._set_proximity_role_image_color(layer, colormap, blending="additive")

    def _apply_proximity_role_colors(self, source_raw_layer, source_seg_layer, target_raw_layer, target_seg_layer) -> None:
        magenta = np.array([1.0, 0.0, 1.0, 0.72], dtype=float)
        green = np.array([0.0, 1.0, 0.2, 0.72], dtype=float)
        self._set_proximity_role_image_color(source_raw_layer, "magenta")
        self._set_proximity_role_mask_color(source_seg_layer, magenta, "magenta")
        self._set_proximity_role_image_color(target_raw_layer, "green")
        self._set_proximity_role_mask_color(target_seg_layer, green, "green")

    def _tracking_overlay_candidates(self) -> list:
        ref_layer = self._selected_tracking_layer()
        return [
            layer_item
            for layer_item in self.viewer.layers
            if _is_tracking_overlay_layer(layer_item, ref_layer)
        ]

    def _preferred_tracking_overlay_layer(self):
        candidates = self._tracking_overlay_candidates()
        current = self._tracking_event_base_source_layer
        if current in candidates:
            return current
        for layer in candidates:
            metadata = dict(getattr(layer, "metadata", {}) or {})
            layer_type = str(
                getattr(layer, "_type_string", "")
                or layer.__class__.__name__
            ).lower()
            if (
                layer_type == "image"
                and not metadata.get("is_segmentation")
                and not metadata.get("is_tracking")
                and not metadata.get("is_frangi")
                and not metadata.get("is_denoise")
            ):
                return layer
        reference = self._selected_tracking_layer()
        if reference in candidates:
            return reference
        return candidates[0] if candidates else None

    def _selected_tracking_layer(self):
        if self._tracking_layer_combo is None or self._tracking_layer_combo.count() == 0:
            return None
        return self._tracking_layer_combo.currentData()

    def _selected_tracking_raw_layer(self):
        combo = self._tracking_raw_layer_combo
        if combo is None or combo.count() == 0:
            return None
        return combo.currentData()

    def _selected_tracking_overlay_layer(self, _kind: str | None = None):
        return self._preferred_tracking_overlay_layer()

    def _prompt_tracking_export_settings(
        self,
    ) -> tuple[Any, str, float, tuple[str, ...]] | None:
        candidates = self._tracking_overlay_candidates()
        if not candidates:
            QMessageBox.information(
                self,
                "No overlay layer",
                "No compatible tracking overlay layer is available.",
            )
            return None

        dialog = QDialog(self)
        dialog.setWindowTitle("Export Tracking Events")
        dialog.setModal(True)
        layout = QVBoxLayout(dialog)
        form = QGridLayout()
        form.setHorizontalSpacing(self._scaled_px(10))
        form.setVerticalSpacing(self._scaled_px(8))

        overlay_combo = QComboBox(dialog)
        preferred = self._preferred_tracking_overlay_layer()
        preferred_index = 0
        for index, layer in enumerate(candidates):
            overlay_combo.addItem(str(getattr(layer, "name", "Layer")), layer)
            if layer is preferred:
                preferred_index = index
        overlay_combo.setCurrentIndex(preferred_index)
        overlay_combo.setSizePolicy(
            QSizePolicy.Expanding,
            QSizePolicy.Fixed,
        )

        format_combo = QComboBox(dialog)
        for format_name in ("GIF", "MP4"):
            format_combo.addItem(format_name)

        event_checks: dict[str, QCheckBox] = {}
        event_widget = QWidget(dialog)
        event_layout = QGridLayout(event_widget)
        event_layout.setContentsMargins(0, 0, 0, 0)
        event_layout.setHorizontalSpacing(self._scaled_px(12))
        event_layout.setVerticalSpacing(self._scaled_px(4))
        groups = self._tracking_event_groups()
        for index, kind in enumerate(
            ("linear", "fission", "fusion", "split-merge")
        ):
            checkbox = QCheckBox(
                self._tracking_event_kind_label(kind),
                event_widget,
            )
            available = bool(groups.get(kind))
            checkbox.setChecked(available)
            checkbox.setEnabled(available)
            event_checks[kind] = checkbox
            event_layout.addWidget(checkbox, index // 2, index % 2)

        fps_spin = QDoubleSpinBox(dialog)
        fps_spin.setRange(0.1, 120.0)
        fps_spin.setDecimals(1)
        fps_spin.setSingleStep(0.5)
        fps_spin.setValue(1.0)
        fps_spin.setEnabled(True)

        form.addWidget(QLabel("Overlay:"), 0, 0, alignment=Qt.AlignRight)
        form.addWidget(overlay_combo, 0, 1)
        form.addWidget(QLabel("Events:"), 1, 0, alignment=Qt.AlignRight)
        form.addWidget(event_widget, 1, 1)
        form.addWidget(QLabel("Format:"), 2, 0, alignment=Qt.AlignRight)
        form.addWidget(format_combo, 2, 1)
        form.addWidget(QLabel("FPS:"), 3, 0, alignment=Qt.AlignRight)
        form.addWidget(fps_spin, 3, 1)
        form.setColumnStretch(1, 1)
        layout.addLayout(form)

        buttons = QDialogButtonBox(dialog)
        export_button = buttons.addButton(
            "Export",
            QDialogButtonBox.AcceptRole,
        )
        buttons.addButton("Cancel", QDialogButtonBox.RejectRole)
        export_button.setDefault(True)
        update_export_enabled = lambda _checked=False: export_button.setEnabled(
            any(
                checkbox.isChecked()
                for checkbox in event_checks.values()
                if checkbox.isEnabled()
            )
        )
        for checkbox in event_checks.values():
            checkbox.toggled.connect(update_export_enabled)
        update_export_enabled()
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        dialog.setMinimumWidth(self._scaled_px(360))

        if dialog.exec() != QDialog.Accepted:
            return None
        overlay_layer = overlay_combo.currentData()
        if overlay_layer is None:
            return None
        selected_kinds = tuple(
            kind
            for kind, checkbox in event_checks.items()
            if checkbox.isEnabled() and checkbox.isChecked()
        )
        if not selected_kinds:
            return None
        return (
            overlay_layer,
            str(format_combo.currentText()).upper(),
            float(fps_spin.value()),
            selected_kinds,
        )

    def _prompt_tracking_statistics_export_settings(
        self,
    ) -> tuple[tuple[str, ...], int, int, str] | None:
        groups = self._tracking_event_groups()
        all_events = [
            event
            for kind in ("linear", "fission", "fusion", "split-merge")
            for event in groups.get(kind, ())
        ]
        if not all_events:
            QMessageBox.information(
                self,
                "No events",
                "Run tracking first to export event statistics.",
            )
            return None

        frame_min = min(int(event.frame_from) for event in all_events)
        frame_max = max(int(event.frame_to) for event in all_events)

        dialog = QDialog(self)
        dialog.setWindowTitle("Export Event Statistics")
        dialog.setModal(True)
        layout = QVBoxLayout(dialog)
        form = QGridLayout()
        form.setHorizontalSpacing(self._scaled_px(10))
        form.setVerticalSpacing(self._scaled_px(8))

        event_checks: dict[str, QCheckBox] = {}
        event_widget = QWidget(dialog)
        event_layout = QGridLayout(event_widget)
        event_layout.setContentsMargins(0, 0, 0, 0)
        event_layout.setHorizontalSpacing(self._scaled_px(12))
        event_layout.setVerticalSpacing(self._scaled_px(4))
        for index, kind in enumerate(
            ("linear", "fission", "fusion", "split-merge")
        ):
            checkbox = QCheckBox(
                self._tracking_event_kind_label(kind),
                event_widget,
            )
            available = bool(groups.get(kind))
            checkbox.setChecked(available)
            checkbox.setEnabled(available)
            event_checks[kind] = checkbox
            event_layout.addWidget(checkbox, index // 2, index % 2)

        frame_widget = QWidget(dialog)
        frame_layout = QHBoxLayout(frame_widget)
        frame_layout.setContentsMargins(0, 0, 0, 0)
        frame_layout.setSpacing(self._scaled_px(6))
        frame_from_spin = QSpinBox(frame_widget)
        frame_from_spin.setRange(frame_min, frame_max)
        frame_from_spin.setValue(frame_min)
        frame_to_spin = QSpinBox(frame_widget)
        frame_to_spin.setRange(frame_min, frame_max)
        frame_to_spin.setValue(frame_max)
        frame_layout.addWidget(QLabel("From"))
        frame_layout.addWidget(frame_from_spin)
        frame_layout.addWidget(QLabel("to"))
        frame_layout.addWidget(frame_to_spin)
        frame_layout.addStretch(1)

        format_combo = QComboBox(dialog)
        for format_name in ("CSV", "TXT", "XLSX"):
            format_combo.addItem(format_name)

        form.addWidget(QLabel("Events:"), 0, 0, alignment=Qt.AlignRight)
        form.addWidget(event_widget, 0, 1)
        form.addWidget(QLabel("Frames:"), 1, 0, alignment=Qt.AlignRight)
        form.addWidget(frame_widget, 1, 1)
        form.addWidget(QLabel("Format:"), 2, 0, alignment=Qt.AlignRight)
        form.addWidget(format_combo, 2, 1)
        form.setColumnStretch(1, 1)
        layout.addLayout(form)

        buttons = QDialogButtonBox(dialog)
        export_button = buttons.addButton(
            "Export",
            QDialogButtonBox.AcceptRole,
        )
        buttons.addButton("Cancel", QDialogButtonBox.RejectRole)
        export_button.setDefault(True)

        def update_export_enabled(_value=None) -> None:
            has_kind = any(
                checkbox.isChecked()
                for checkbox in event_checks.values()
                if checkbox.isEnabled()
            )
            valid_frames = frame_from_spin.value() <= frame_to_spin.value()
            export_button.setEnabled(has_kind and valid_frames)

        for checkbox in event_checks.values():
            checkbox.toggled.connect(update_export_enabled)
        frame_from_spin.valueChanged.connect(update_export_enabled)
        frame_to_spin.valueChanged.connect(update_export_enabled)
        update_export_enabled()
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        dialog.setMinimumWidth(self._scaled_px(380))

        if dialog.exec() != QDialog.Accepted:
            return None
        selected_kinds = tuple(
            kind
            for kind, checkbox in event_checks.items()
            if checkbox.isEnabled() and checkbox.isChecked()
        )
        if not selected_kinds:
            return None
        return (
            selected_kinds,
            int(frame_from_spin.value()),
            int(frame_to_spin.value()),
            str(format_combo.currentText()).upper(),
        )

    def _tracking_spacing_from_layer(self, layer) -> tuple[float, ...]:
        data = layer.data
        scale = tuple(float(v) for v in getattr(layer, "scale", (1,) * data.ndim))
        if data.ndim == 3:
            return (scale[-2], scale[-1])
        if data.ndim == 4:
            return (scale[-3], scale[-2], scale[-1])
        raise ValueError(f"Unsupported tracking ndim: {data.ndim}")

    def _build_tracking_config(self) -> TrackingConfig:
        return TrackingConfig(
            max_distance=float(self.track_max_dist.value()),
            max_neighbors=int(self.track_max_neighbors.value()),
            distance_weight=float(self.track_distance_weight.value()),
            overlap_weight=float(self.track_overlap_weight.value()),
            point_support_weight=float(self.track_point_support_weight.value()),
            cost_cutoff=float(self.track_cost_cutoff.value()),
            sample_points=int(self.track_sample_points.value()),
            point_support_saturation=10,
            coverage_capacity_points=5,
            coverage_cap_skip_threshold=0.7,
            point_match_method="hungarian",
            point_unmatched_cost=20.0,
            min_object_match_fraction=0.5,
            min_link_size=int(self.track_min_link_size.value()),
            tracking_workers=0,
        )

    @staticmethod
    def _tracking_lineage_color_map(result: TrackingResult) -> dict[int, np.ndarray]:
        color_map = {0: np.array([0.0, 0.0, 0.0, 0.0], dtype=float)}
        for lineage_id in sorted({int(value) for value in result.lineage_ids.values()}):
            if lineage_id > 0:
                color_map[lineage_id] = np.array(
                    _lineage_rgba(lineage_id),
                    dtype=float,
                )
        return color_map

    @staticmethod
    def _tracking_event_group_kind(kind: str) -> str | None:
        normalized = str(kind).lower()
        if normalized in {"continuation", "elongation", "shortening", "linear"}:
            return "linear"
        if normalized in {"fission", "fusion", "split-merge"}:
            return normalized
        return None

    @staticmethod
    def _tracking_event_kind_label(kind: str) -> str:
        return {
            "linear": "Linear",
            "fission": "Fission",
            "fusion": "Fusion",
            "split-merge": "Split-Merge",
        }.get(str(kind), str(kind))

    @staticmethod
    def _tracking_event_rgba(kind: str, event_index: int) -> np.ndarray:
        base_hues = {
            "linear": 0.72,
            "fission": 0.03,
            "fusion": 0.53,
            "split-merge": 0.14,
        }
        group_kind = SIGMAWidget._tracking_event_group_kind(kind) or "linear"
        hue = (
            base_hues[group_kind] + int(event_index) * 0.61803398875
        ) % 1.0
        rgb = colorsys.hsv_to_rgb(hue, 0.72, 0.95)
        return np.asarray([*rgb, 1.0], dtype=float)

    @staticmethod
    def _tracking_object_rgba(det_id: int) -> np.ndarray:
        hue = (0.11 + int(det_id) * 0.61803398875) % 1.0
        rgb = colorsys.hsv_to_rgb(hue, 0.82, 1.0)
        return np.asarray([*rgb, 1.0], dtype=float)

    def _tracking_event_groups(self) -> dict[str, list]:
        return {
            "linear": list(self._tracking_linear_events),
            "fission": list(self._tracking_fission_events),
            "fusion": list(self._tracking_fusion_events),
            "split-merge": list(self._tracking_split_merge_events),
        }

    @staticmethod
    def _tracking_statistics_events(
        groups: dict[str, list],
        selected_kinds: tuple[str, ...],
        frame_from: int,
        frame_to: int,
    ) -> list:
        kind_order = {
            "linear": 0,
            "fission": 1,
            "fusion": 2,
            "split-merge": 3,
        }
        selected = set(selected_kinds)
        events = [
            event
            for kind in kind_order
            if kind in selected
            for event in groups.get(kind, ())
            if int(event.frame_from) >= int(frame_from)
            and int(event.frame_to) <= int(frame_to)
        ]
        return sorted(
            events,
            key=lambda event: (
                int(event.frame_from),
                int(event.frame_to),
                kind_order.get(
                    SIGMAWidget._tracking_event_group_kind(event.kind),
                    len(kind_order),
                ),
                str(event.kind),
                tuple(int(value) for value in event.sources),
                tuple(int(value) for value in event.targets),
            ),
        )

    @staticmethod
    def _tracking_initial_event_context(
        groups: dict[str, list],
        frame: int,
    ) -> tuple[tuple[int, int] | None, str | None]:
        frame = int(frame)
        kind_order = ("linear", "fission", "fusion", "split-merge")
        pairs_by_kind = {
            kind: sorted(
                {
                    (int(event.frame_from), int(event.frame_to))
                    for event in groups.get(kind, ())
                    if int(event.frame_from) != int(event.frame_to)
                }
            )
            for kind in kind_order
        }

        def pair_rank(pair: tuple[int, int]) -> tuple[int, int, int, int]:
            frame_from, frame_to = pair
            if frame_from == frame:
                relation = 0
            elif frame_to == frame:
                relation = 1
            elif min(frame_from, frame_to) <= frame <= max(frame_from, frame_to):
                relation = 2
            else:
                relation = 3
            distance = min(abs(frame_from - frame), abs(frame_to - frame))
            return relation, distance, frame_from, frame_to

        linear_pairs = pairs_by_kind["linear"]
        linear_at_frame = [
            pair for pair in linear_pairs if pair_rank(pair)[0] < 3
        ]
        if linear_at_frame:
            return min(linear_at_frame, key=pair_rank), "linear"

        for kind in kind_order[1:]:
            pairs_at_frame = [
                pair
                for pair in pairs_by_kind[kind]
                if pair_rank(pair)[0] < 3
            ]
            if pairs_at_frame:
                return min(pairs_at_frame, key=pair_rank), kind

        if linear_pairs:
            return min(linear_pairs, key=pair_rank), "linear"
        for kind in kind_order[1:]:
            if pairs_by_kind[kind]:
                return min(pairs_by_kind[kind], key=pair_rank), kind
        return None, None

    def _tracking_current_viewer_frame(self, layer) -> int:
        frame_count = max(_time_size_from_layer(layer), 1)
        time_axis = _time_axis_index(self.viewer, layer)
        if time_axis is None:
            return 0
        try:
            current_step = tuple(int(v) for v in self.viewer.dims.current_step)
            frame = int(current_step[int(time_axis)])
        except (AttributeError, IndexError, TypeError, ValueError):
            frame = frame_count // 2
        return max(0, min(frame, frame_count - 1))

    def _set_tracking_initial_event_context(self, frame: int) -> None:
        transition, kind = self._tracking_initial_event_context(
            self._tracking_event_groups(),
            int(frame),
        )
        if transition is None or kind is None:
            return
        self._tracking_active_transition = transition
        self._tracking_active_event_kind = kind

        transition_combo = self._tracking_transition_combo
        if transition_combo is not None:
            transition_combo.blockSignals(True)
            for index in range(transition_combo.count()):
                value = transition_combo.itemData(index)
                if (
                    isinstance(value, tuple | list)
                    and tuple(int(v) for v in value) == transition
                ):
                    transition_combo.setCurrentIndex(index)
                    break
            transition_combo.blockSignals(False)
        self._apply_tracking_transition_table_filter()
        self._refresh_tracking_event_kind_combo()

    def _set_tracking_active_event_kind(self, kind: str) -> None:
        group_kind = self._tracking_event_group_kind(kind)
        if group_kind is None:
            return
        self._tracking_active_event_kind = group_kind
        combo = self._tracking_event_kind_combo
        if combo is None:
            return
        for index in range(combo.count()):
            if str(combo.itemData(index)) == group_kind:
                combo.blockSignals(True)
                combo.setCurrentIndex(index)
                combo.blockSignals(False)
                break

    def _refresh_tracking_event_kind_combo(self) -> None:
        combo = self._tracking_event_kind_combo
        if combo is None:
            return
        transition = self._tracking_active_transition
        groups = self._tracking_event_groups()
        counts = {
            kind: sum(
                1
                for event in events
                if transition is None
                or (
                    int(event.frame_from) == int(transition[0])
                    and int(event.frame_to) == int(transition[1])
                )
            )
            for kind, events in groups.items()
        }
        active_kind = self._tracking_active_event_kind
        if not active_kind or counts.get(active_kind, 0) <= 0:
            active_kind = next(
                (kind for kind in groups if counts.get(kind, 0) > 0),
                None,
            )

        combo.blockSignals(True)
        for index in range(combo.count()):
            kind = str(combo.itemData(index))
            count = int(counts.get(kind, 0))
            combo.setItemText(
                index,
                f"{self._tracking_event_kind_label(kind)} ({count})",
            )
            with suppress(Exception):
                combo.model().item(index).setEnabled(count > 0)
            if kind == active_kind:
                combo.setCurrentIndex(index)
        combo.setEnabled(any(counts.values()))
        combo.blockSignals(False)
        self._tracking_active_event_kind = active_kind

    def _rebuild_tracking_transition_combo(self) -> None:
        combo = self._tracking_transition_combo
        if combo is None:
            return
        pairs = sorted(
            {
                (int(event.frame_from), int(event.frame_to))
                for event in self._tracking_display_events
                if int(event.frame_to) != int(event.frame_from)
            }
        )
        current = self._tracking_active_transition
        if current not in pairs:
            current = pairs[0] if pairs else None

        combo.blockSignals(True)
        combo.clear()
        for frame_from, frame_to in pairs:
            combo.addItem(
                f"{frame_from} \N{RIGHTWARDS ARROW} {frame_to}",
                (frame_from, frame_to),
            )
        if current is not None:
            for index in range(combo.count()):
                if tuple(combo.itemData(index)) == tuple(current):
                    combo.setCurrentIndex(index)
                    break
        combo.setEnabled(bool(pairs))
        combo.blockSignals(False)
        self._tracking_active_transition = current
        self._apply_tracking_transition_table_filter()
        self._refresh_tracking_event_kind_combo()

    def _apply_tracking_transition_table_filter(self) -> None:
        for table, events, kind in (
            (
                self._tracking_linear_table,
                self._tracking_linear_events,
                "linear",
            ),
            (
                self._tracking_fission_table,
                self._tracking_fission_events,
                "fission",
            ),
            (
                self._tracking_fusion_table,
                self._tracking_fusion_events,
                "fusion",
            ),
            (
                self._tracking_split_merge_table,
                self._tracking_split_merge_events,
                "split-merge",
            ),
        ):
            self._populate_tracking_event_table(table, events, kind)

    def _on_tracking_transition_changed(self, _index: int = -1) -> None:
        combo = self._tracking_transition_combo
        if combo is None:
            return
        value = combo.currentData()
        if not isinstance(value, tuple | list) or len(value) != 2:
            return
        transition = (int(value[0]), int(value[1]))
        self._render_tracking_transition(transition)

    def _render_tracking_transition(
        self,
        transition: tuple[int, int],
    ) -> None:
        transition = (int(transition[0]), int(transition[1]))
        self._cancel_tracking_refine_image_pick(restore_context=False)
        self._tracking_active_transition = transition
        self._clear_tracking_highlight()
        self._apply_tracking_transition_table_filter()
        self._refresh_tracking_event_kind_combo()
        source_layer = (
            self._tracking_event_base_source_layer
            or self._selected_tracking_layer()
            or self._tracking_layer
        )
        if self._tracking_result is not None and source_layer is not None:
            active_kind = self._tracking_active_event_kind or "linear"
            self._rebuild_tracking_event_layers(
                source_layer,
                kinds={active_kind},
            )
            layer = self._active_tracking_event_layer() or self._tracking_layer
            self._apply_tracking_time_steps(layer, [int(transition[0])])

    def _on_tracking_event_kind_changed(self, _index: int = -1) -> None:
        combo = self._tracking_event_kind_combo
        if combo is None:
            return
        group_kind = self._tracking_event_group_kind(
            str(combo.currentData() or "")
        )
        if group_kind is None:
            return
        self._cancel_tracking_refine_image_pick(restore_context=False)
        self._tracking_active_event_kind = group_kind
        if self._tracking_result is None:
            return
        self._clear_tracking_highlight()
        self._activate_tracking_event_layer(group_kind)

    def _set_tracking_active_transition(
        self,
        frame_from: int,
        frame_to: int,
        *,
        render: bool = True,
    ) -> None:
        transition = (int(frame_from), int(frame_to))
        combo = self._tracking_transition_combo
        if combo is not None:
            for index in range(combo.count()):
                value = combo.itemData(index)
                if isinstance(value, tuple | list) and tuple(value) == transition:
                    if index != combo.currentIndex():
                        combo.blockSignals(True)
                        combo.setCurrentIndex(index)
                        combo.blockSignals(False)
                    if render:
                        self._render_tracking_transition(transition)
                    else:
                        self._tracking_active_transition = transition
                        self._apply_tracking_transition_table_filter()
                        self._refresh_tracking_event_kind_combo()
                    return
        if render:
            self._render_tracking_transition(transition)
        else:
            self._tracking_active_transition = transition
            self._apply_tracking_transition_table_filter()
            self._refresh_tracking_event_kind_combo()

    @classmethod
    def _tracking_event_layer_payload(
        cls,
        result: TrackingResult,
        events: list,
        kind: str,
        transition: tuple[int, int] | None,
    ) -> tuple[np.ndarray, dict[int, np.ndarray], int]:
        indexed_events = [
            (event_index, event)
            for event_index, event in enumerate(events)
            if transition is None
            or (
                int(event.frame_from) == int(transition[0])
                and int(event.frame_to) == int(transition[1])
            )
        ]
        if transition is not None:
            frame_start = min(int(transition[0]), int(transition[1]))
            frame_stop = max(int(transition[0]), int(transition[1]))
        elif indexed_events:
            frame_start = min(int(event.frame_from) for _, event in indexed_events)
            frame_stop = max(int(event.frame_to) for _, event in indexed_events)
        else:
            frame_start = 0
            frame_stop = 0

        det_by_id = {int(det.id): det for det in result.detections}
        label_count = len(events)
        if label_count <= np.iinfo(np.uint8).max:
            dtype = np.uint8
        elif label_count <= np.iinfo(np.uint16).max:
            dtype = np.uint16
        else:
            dtype = np.uint32

        data = np.zeros(
            (
                max(frame_stop - frame_start + 1, 1),
                *tuple(int(v) for v in result.frame_labels.shape[1:]),
            ),
            dtype=dtype,
        )
        transparent = np.array([0.0, 0.0, 0.0, 0.0], dtype=float)
        color_map: dict[int, np.ndarray] = {0: transparent}

        event_values = []
        for event_index, event in indexed_events:
            value = int(event_index) + 1
            color_map[value] = cls._tracking_event_rgba(kind, event_index)
            event_values.append(
                (
                    event,
                    {
                        int(det_id): value
                        for det_id in (*event.sources, *event.targets)
                    },
                )
            )

        for _event, values in event_values:
            for det_id, label_value in values.items():
                det = det_by_id.get(int(det_id))
                if det is None:
                    continue
                local_frame = int(det.frame) - frame_start
                if local_frame < 0 or local_frame >= data.shape[0]:
                    continue
                target = data[(local_frame, *tuple(det.slice_tuple))]
                target[np.asarray(det.image, dtype=bool)] = int(label_value)
        return data, color_map, frame_start

    @staticmethod
    def _tracking_event_full_time_data(
        local_data: np.ndarray,
        frame_start: int,
        total_frames: int,
    ) -> da.Array:
        local = np.asarray(local_data)
        if local.ndim < 1:
            raise ValueError("Tracking event data must include a time axis.")
        total = max(int(total_frames), 1)
        start = max(min(int(frame_start), total), 0)
        local = local[: max(total - start, 0)]
        spatial_shape = tuple(int(v) for v in local.shape[1:])
        chunks = (1, *spatial_shape)
        parts = []
        if start > 0:
            parts.append(
                da.zeros(
                    (start, *spatial_shape),
                    chunks=chunks,
                    dtype=local.dtype,
                )
            )
        if local.shape[0] > 0:
            parts.append(da.from_array(local, chunks=chunks))
        trailing = total - start - int(local.shape[0])
        if trailing > 0:
            parts.append(
                da.zeros(
                    (trailing, *spatial_shape),
                    chunks=chunks,
                    dtype=local.dtype,
                )
            )
        return da.concatenate(parts, axis=0)

    def _remove_tracking_event_base_layer(self) -> None:
        source_layer = self._tracking_event_base_source_layer
        if (
            source_layer is not None
            and self._tracking_event_base_source_was_visible is not None
        ):
            with suppress(Exception):
                source_layer.visible = bool(
                    self._tracking_event_base_source_was_visible
                )
        if self._tracking_event_base_layer is not None:
            with suppress(Exception):
                self.viewer.layers.remove(self._tracking_event_base_layer)
        self._tracking_event_base_layer = None
        self._tracking_event_base_source_layer = None
        self._tracking_event_base_source_was_visible = None

    def _tracking_event_base_source(
        self,
        fallback_layer,
        *,
        preferred_kind: str | None = None,
    ):
        result = self._tracking_result
        if result is None:
            return fallback_layer
        candidates = []
        if preferred_kind is not None:
            candidates.append(
                self._selected_tracking_overlay_layer(preferred_kind)
            )
        for kind in ("linear", "fission", "fusion", "split-merge"):
            candidates.append(self._selected_tracking_overlay_layer(kind))
        candidates.append(fallback_layer)
        expected_shape = tuple(int(v) for v in result.frame_labels.shape)
        for candidate in candidates:
            if (
                candidate is None
                or candidate is self._tracking_event_base_layer
                or _is_analysis_aux_layer(candidate)
            ):
                continue
            shape = tuple(
                int(v)
                for v in getattr(
                    getattr(candidate, "data", None),
                    "shape",
                    (),
                )
            )
            if shape == expected_shape:
                return candidate
        return fallback_layer

    def _place_tracking_event_base_below_layers(self) -> None:
        base_layer = self._tracking_event_base_layer
        if base_layer is None:
            return
        with suppress(Exception):
            base_index = self.viewer.layers.index(base_layer)
            event_indices = [
                self.viewer.layers.index(layer)
                for layer in self._tracking_event_layers.values()
                if layer in self.viewer.layers
            ]
            if event_indices and base_index > min(event_indices):
                self.viewer.layers.move(base_index, min(event_indices))

    def _rebuild_tracking_event_base_layer(
        self,
        fallback_layer,
        *,
        preferred_kind: str | None = None,
    ) -> None:
        existing = self._tracking_event_base_layer
        existing_is_present = False
        if existing is not None:
            with suppress(Exception):
                existing_is_present = existing in self.viewer.layers
        source_layer = (
            self._tracking_event_base_source_layer
            if existing_is_present
            else self._tracking_event_base_source(
                fallback_layer,
                preferred_kind=preferred_kind,
            )
        )
        if source_layer is None or not hasattr(source_layer, "data"):
            return
        if (
            existing_is_present
            and source_layer is self._tracking_event_base_source_layer
        ):
            with suppress(Exception):
                existing.visible = True
            self._place_tracking_event_base_below_layers()
            return
        if (
            not existing_is_present
            or source_layer is not self._tracking_event_base_source_layer
        ):
            self._remove_tracking_event_base_layer()
            metadata = dict(getattr(source_layer, "metadata", {}) or {})
            for key in (
                "is_segmentation",
                "is_frangi",
                "is_denoise",
                "is_tracking",
            ):
                metadata.pop(key, None)
            metadata.update(
                {
                    "managed_by_plugin": True,
                    "is_analysis_highlight": True,
                    "is_tracking_event_base": True,
                }
            )
            dims_tag = _layer_dims_tag(source_layer)
            if dims_tag:
                metadata["dims"] = dims_tag
            add_kwargs: dict[str, Any] = {"metadata": metadata}
            scale = getattr(source_layer, "scale", None)
            if scale is not None:
                add_kwargs["scale"] = scale
            add_kwargs.update(_spatial_unit_kwargs(source_layer, np.asarray(source_layer.data).ndim))
            contrast_limits = getattr(source_layer, "contrast_limits", None)
            if contrast_limits is not None:
                add_kwargs["contrast_limits"] = tuple(
                    float(v) for v in contrast_limits
                )
            existing = self.viewer.add_image(
                source_layer.data,
                name=self._next_layer_name("tracking raw base"),
                colormap="gray",
                opacity=0.7,
                blending="additive",
                **add_kwargs,
            )
            self._tracking_event_base_layer = existing
            self._tracking_event_base_source_layer = source_layer
            self._tracking_event_base_source_was_visible = bool(
                getattr(source_layer, "visible", True)
            )
            with suppress(Exception):
                source_layer.visible = False

        with suppress(Exception):
            existing.colormap = "gray"
            existing.opacity = 0.7
            existing.interpolation2d = "cubic"
            existing.interpolation3d = "cubic"
            existing.blending = "additive"
            existing.visible = True
        _apply_tracking_overlay_rendering(existing, source_layer)
        with suppress(Exception):
            existing.refresh()
        self._place_tracking_event_base_below_layers()

    def _remove_tracking_event_layers(self) -> None:
        self._remove_tracking_event_base_layer()
        for layer in self._tracking_event_layers.values():
            with suppress(Exception):
                self.viewer.layers.remove(layer)
        self._tracking_event_layers = {}
        self._tracking_event_layer_color_maps = {}
        self._tracking_active_event_kind = None

    def _rebuild_tracking_event_layers(
        self,
        source_layer,
        *,
        kinds: set[str] | None = None,
    ) -> None:
        result = self._tracking_result
        if result is None or source_layer is None:
            return
        groups = self._tracking_event_groups()
        active_kind = self._tracking_active_event_kind
        if active_kind not in groups or not groups.get(active_kind):
            active_kind = next(
                (kind for kind, events in groups.items() if events),
                "linear",
            )
        requested_kinds = (
            {active_kind}
            if kinds is None
            else {str(kind) for kind in kinds if str(kind) in groups}
        )
        if not requested_kinds:
            return
        render_kind = (
            active_kind
            if active_kind in requested_kinds
            else next(kind for kind in groups if kind in requested_kinds)
        )
        active_kind = render_kind
        self._set_tracking_active_event_kind(active_kind)

        source_md = dict(getattr(source_layer, "metadata", {}) or {})
        source_md.pop("is_segmentation", None)
        dims_tag = _layer_dims_tag(source_layer)
        layer_ndim = int(result.frame_labels.ndim)
        raw_scale = tuple(float(v) for v in getattr(source_layer, "scale", ()))
        scale = ((1.0,) * layer_ndim + raw_scale)[-layer_ndim:]
        raw_translate = tuple(
            float(v) for v in getattr(source_layer, "translate", ())
        )
        source_translate = ((0.0,) * layer_ndim + raw_translate)[-layer_ndim:]
        transition = self._tracking_active_transition

        # Napari slices 3D layers asynchronously.  If an active event layer is
        # updated in place, its status worker can briefly read the internal
        # 1x1x1 placeholder slice using coordinates from the full volume.
        # Park selection on the stable base layer until the new slices load.
        existing_layers = []
        for layer in self._tracking_event_layers.values():
            if not any(layer is item for item in existing_layers):
                existing_layers.append(layer)
        existing_event_layer = self._tracking_event_layers.get(render_kind)
        if existing_event_layer is None and existing_layers:
            existing_event_layer = existing_layers[0]
        for redundant_layer in existing_layers:
            if redundant_layer is existing_event_layer:
                continue
            with suppress(Exception):
                self.viewer.layers.remove(redundant_layer)
        active_selection = getattr(
            self.viewer.layers.selection,
            "active",
            None,
        )
        active_was_event_layer = active_selection is existing_event_layer
        if existing_event_layer is not None:
            with suppress(Exception):
                existing_event_layer.visible = False
        if active_was_event_layer:
            safe_layer = self._tracking_event_base_layer or source_layer
            with suppress(Exception):
                self.viewer.layers.selection.clear()
                self.viewer.layers.selection.active = safe_layer

        data, color_map, frame_start = self._tracking_event_layer_payload(
            result,
            groups[render_kind],
            render_kind,
            transition,
        )
        layer_data = self._tracking_event_full_time_data(
            data,
            frame_start,
            int(result.frame_labels.shape[0]),
        )
        event_layer = existing_event_layer
        layer_is_present = False
        if event_layer is not None:
            with suppress(Exception):
                layer_is_present = event_layer in self.viewer.layers
        metadata = dict(source_md)
        metadata.update(
            {
                "managed_by_plugin": True,
                "is_analysis_labels": True,
                "is_tracking": True,
                "is_tracking_event": True,
                "tracking_event_kind": render_kind,
                "tracking_transition": transition,
            }
        )
        if dims_tag:
            metadata["dims"] = dims_tag
        if not layer_is_present:
            event_layer = self.viewer.add_labels(
                layer_data,
                name=self._next_layer_name(
                    f"{source_layer.name} tracking events"
                ),
                metadata=metadata,
                scale=scale,
                translate=source_translate,
                **_spatial_unit_kwargs(source_layer, np.asarray(layer_data).ndim),
            )
        else:
            with suppress(Exception):
                event_layer.data = layer_data
                event_layer.scale = scale
                event_layer.translate = source_translate
                event_layer.metadata = metadata

        self._tracking_event_layers = {render_kind: event_layer}
        self._tracking_event_layer_color_maps = {render_kind: color_map}
        with suppress(Exception):
            event_layer.color = color_map
            event_layer.opacity = 0.85
            event_layer.visible = True
            event_layer.refresh()
        if self._tracking_layer is not None:
            with suppress(Exception):
                self._tracking_layer.visible = False
        if self._tracking_event_base_layer is not None:
            with suppress(Exception):
                self._tracking_event_base_layer.visible = True
        self._place_tracking_event_base_below_layers()

    def _active_tracking_event_layer(self):
        if self._tracking_active_event_kind is None:
            return None
        return self._tracking_event_layers.get(self._tracking_active_event_kind)

    def _activate_tracking_event_layer(self, kind: str):
        group_kind = self._tracking_event_group_kind(kind)
        if group_kind is None:
            return None
        self._set_tracking_active_event_kind(group_kind)
        target_layer = self._tracking_event_layers.get(group_kind)
        layer_transition = (
            (getattr(target_layer, "metadata", {}) or {}).get(
                "tracking_transition"
            )
            if target_layer is not None
            else None
        )
        if (
            target_layer is None
            or tuple(layer_transition or ()) != tuple(
                self._tracking_active_transition or ()
            )
        ):
            source_layer = (
                self._tracking_event_base_source_layer
                or self._selected_tracking_layer()
                or self._tracking_layer
            )
            if source_layer is not None:
                self._rebuild_tracking_event_layers(
                    source_layer,
                    kinds={group_kind},
                )
            target_layer = self._tracking_event_layers.get(group_kind)
        if target_layer is None:
            return None
        for layer_kind, event_layer in self._tracking_event_layers.items():
            with suppress(Exception):
                event_layer.visible = layer_kind == group_kind
                event_layer.opacity = 0.85
                event_layer.refresh()
        if self._tracking_layer is not None:
            with suppress(Exception):
                self._tracking_layer.visible = False
        if self._tracking_event_base_layer is not None:
            with suppress(Exception):
                self._tracking_event_base_layer.visible = True
        with suppress(Exception):
            self.viewer.layers.selection.active = target_layer
        return target_layer

    def _make_tracking_refine_link_table(self) -> QTableWidget:
        table = QTableWidget(0, 3)
        table.setHorizontalHeaderLabels(("Link", "Objects", "Cost"))
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setSelectionMode(QAbstractItemView.SingleSelection)
        table.setEnabled(False)
        self._style_table_widget(table, min_height=96, max_height=140)
        header = table.horizontalHeader()
        header.setSectionResizeMode(0, QHeaderView.ResizeToContents)
        header.setSectionResizeMode(1, QHeaderView.Stretch)
        header.setSectionResizeMode(2, QHeaderView.ResizeToContents)
        return table

    @staticmethod
    def _tracking_refine_table_link_id(table: QTableWidget | None) -> int | None:
        if table is None:
            return None
        rows = table.selectionModel().selectedRows() if table.selectionModel() else []
        if not rows:
            return None
        row = int(rows[0].row())
        item = table.item(row, 0)
        if item is None:
            return None
        link_id = item.data(Qt.UserRole)
        return int(link_id) if link_id is not None else None

    def _clear_tracking_refine_link_tables(self) -> None:
        for table in (
            self._tracking_refine_selected_table,
            self._tracking_refine_available_table,
        ):
            if table is None:
                continue
            table.blockSignals(True)
            table.clearContents()
            table.setRowCount(0)
            table.blockSignals(False)
            table.setEnabled(False)

    def _update_tracking_refine_image_pick_buttons(self) -> None:
        result = self._tracking_result
        det_by_id = (
            {
                int(detection.id): detection
                for detection in result.detections
            }
            if result is not None
            else {}
        )
        anchor = det_by_id.get(
            int(self._tracking_refine_det_id)
            if self._tracking_refine_det_id is not None
            else -1
        )
        max_frame = (
            int(result.frame_labels.shape[0]) - 1
            if result is not None
            else -1
        )
        if self._tracking_refine_pick_in_btn is not None:
            self._tracking_refine_pick_in_btn.setEnabled(
                anchor is not None and int(anchor.frame) > 0
            )
        if self._tracking_refine_pick_out_btn is not None:
            self._tracking_refine_pick_out_btn.setEnabled(
                anchor is not None and int(anchor.frame) < max_frame
            )
        picking = self._tracking_refine_image_pick_direction is not None
        if self._tracking_refine_pick_confirm_btn is not None:
            self._tracking_refine_pick_confirm_btn.setEnabled(
                picking and self._tracking_refine_image_pick_det_id is not None
            )
        if self._tracking_refine_pick_cancel_btn is not None:
            self._tracking_refine_pick_cancel_btn.setEnabled(picking)

    def _cancel_tracking_refine_image_pick(
        self,
        _checked: bool = False,
        *,
        restore_context: bool = True,
    ) -> None:
        was_picking = self._tracking_refine_image_pick_direction is not None
        anchor_id = self._tracking_refine_image_pick_anchor_id
        self._tracking_refine_image_pick_direction = None
        self._tracking_refine_image_pick_anchor_id = None
        self._tracking_refine_image_pick_det_id = None
        self._update_tracking_refine_image_pick_buttons()
        if (
            was_picking
            and restore_context
            and anchor_id is not None
            and self._tracking_result is not None
        ):
            self._set_tracking_refine_object(int(anchor_id))
            self._show_tracking_object_context(int(anchor_id))

    def _start_tracking_refine_image_pick(self, direction: str) -> None:
        result = self._tracking_result
        anchor_id = self._tracking_refine_det_id
        if result is None or anchor_id is None:
            return
        direction = str(direction).strip().lower()
        if direction not in {"incoming", "outgoing"}:
            return
        det_by_id = {
            int(detection.id): detection
            for detection in result.detections
        }
        anchor = det_by_id.get(int(anchor_id))
        if anchor is None:
            return
        target_frame = int(anchor.frame) + (
            -1 if direction == "incoming" else 1
        )
        if target_frame < 0 or target_frame >= int(result.frame_labels.shape[0]):
            return

        self._set_tracking_refine_object(int(anchor_id))
        self._tracking_refine_image_pick_direction = direction
        self._tracking_refine_image_pick_anchor_id = int(anchor_id)
        self._tracking_refine_image_pick_det_id = None
        self._clear_tracking_highlight()
        for event_layer in self._tracking_event_layers.values():
            with suppress(Exception):
                event_layer.visible = False
        self._update_tracking_refine_image_pick_buttons()
        role = "source" if direction == "incoming" else "target"
        if self._tracking_refine_status_label is not None:
            self._tracking_refine_status_label.setText(
                f"Pick {'IN' if direction == 'incoming' else 'OUT'}: "
                f"double-click the {role} object at t={target_frame}, "
                "then confirm."
            )
        visible_layer = (
            self._tracking_event_base_layer
            or self._selected_tracking_layer()
            or self._tracking_layer
        )
        if self._tracking_event_base_layer is not None:
            with suppress(Exception):
                self._tracking_event_base_layer.visible = True
                self.viewer.layers.selection.active = self._tracking_event_base_layer
        self._apply_tracking_time_steps(visible_layer, [target_frame])

    def _set_tracking_refine_image_pick_candidate(self, det_id: int) -> bool:
        result = self._tracking_result
        direction = self._tracking_refine_image_pick_direction
        anchor_id = self._tracking_refine_image_pick_anchor_id
        if result is None or direction is None or anchor_id is None:
            return False
        det_by_id = {
            int(detection.id): detection
            for detection in result.detections
        }
        anchor = det_by_id.get(int(anchor_id))
        picked = det_by_id.get(int(det_id))
        if anchor is None or picked is None:
            return False
        expected_frame = int(anchor.frame) + (
            -1 if direction == "incoming" else 1
        )
        if int(picked.frame) != expected_frame:
            if self._tracking_refine_status_label is not None:
                self._tracking_refine_status_label.setText(
                    f"Pick {'IN' if direction == 'incoming' else 'OUT'}: "
                    f"choose an object at t={expected_frame}; "
                    f"the clicked object is at t={int(picked.frame)}."
                )
            return False

        self._tracking_refine_image_pick_det_id = int(det_id)
        source = picked if direction == "incoming" else anchor
        target = anchor if direction == "incoming" else picked
        self._apply_tracking_highlight_colors(
            {
                int(source.id): 1,
                int(target.id): 2,
            },
            {
                1: np.array([1.0, 0.55, 0.08, 1.0], dtype=float),
                2: np.array([0.10, 0.86, 1.0, 1.0], dtype=float),
            },
            opacity=0.98,
            base_layer=(
                self._tracking_event_base_layer or self._tracking_layer
            ),
            base_opacity=0.7,
        )
        if self._tracking_refine_status_label is not None:
            self._tracking_refine_status_label.setText(
                f"Selected link: t={int(source.frame)} "
                f"object={int(source.local_label)} -> "
                f"t={int(target.frame)} object={int(target.local_label)}. "
                "Click Confirm to add it."
            )
        self._update_tracking_refine_image_pick_buttons()
        return True

    def _confirm_tracking_refine_image_pick(self) -> None:
        result = self._tracking_result
        direction = self._tracking_refine_image_pick_direction
        anchor_id = self._tracking_refine_image_pick_anchor_id
        picked_id = self._tracking_refine_image_pick_det_id
        if (
            result is None
            or direction is None
            or anchor_id is None
            or picked_id is None
        ):
            return
        det_by_id = {
            int(detection.id): detection
            for detection in result.detections
        }
        anchor = det_by_id.get(int(anchor_id))
        picked = det_by_id.get(int(picked_id))
        if anchor is None or picked is None:
            return
        source = picked if direction == "incoming" else anchor
        target = anchor if direction == "incoming" else picked
        source_layer = (
            self._selected_tracking_layer()
            or self._tracking_result_source_layer
        )
        try:
            spacing = (
                self._tracking_spacing_from_layer(source_layer)
                if source_layer is not None
                else None
            )
            refined, link_id = add_manual_tracking_link(
                result,
                self._tracking_config or self._build_tracking_config(),
                int(source.id),
                int(target.id),
                spacing=spacing,
                rebuild_lineages=False,
            )
        except (TypeError, ValueError, KeyError) as error:
            QMessageBox.critical(self, "Refine failed", str(error))
            return

        self._cancel_tracking_refine_image_pick(restore_context=False)
        if refined.selected_link_ids == result.selected_link_ids:
            self._set_tracking_refine_candidate(int(link_id))
            if self._tracking_refine_status_label is not None:
                self._tracking_refine_status_label.setText(
                    f"Link {int(link_id)} is already effective."
                )
            return
        self._record_tracking_refine_undo(result)
        transition = (int(source.frame), int(target.frame))
        self._set_tracking_active_transition(*transition)
        self._update_tracking_result_after_refine(
            refined,
            int(anchor_id),
            preferred_link_id=int(link_id),
            changed_transition=transition,
        )

    def _populate_tracking_refine_objects(
        self,
        preferred_det_id: int | None = None,
    ) -> None:
        combo = self._tracking_refine_object_combo
        result = self._tracking_result
        if combo is None:
            return

        combo.blockSignals(True)
        combo.clear()
        if result is not None:
            for det in sorted(
                result.detections,
                key=lambda item: (int(item.frame), int(item.local_label), int(item.id)),
            ):
                combo.addItem(
                    f"t={int(det.frame)} | object={int(det.local_label)} | "
                    f"size={int(det.area)}",
                    int(det.id),
                )
        combo.setEnabled(combo.count() > 0)

        selected_index = -1
        if combo.count() > 0:
            wanted = (
                int(preferred_det_id)
                if preferred_det_id is not None
                else self._tracking_refine_det_id
            )
            if wanted is not None:
                selected_index = combo.findData(int(wanted))
            if selected_index < 0:
                selected_index = 0
            combo.setCurrentIndex(selected_index)
        combo.blockSignals(False)

        if selected_index >= 0:
            self._set_tracking_refine_object(int(combo.currentData()))
            return

        self._tracking_refine_det_id = None
        self._tracking_refine_event_scope = None
        self._cancel_tracking_refine_image_pick(restore_context=False)
        self._clear_tracking_refine_link_tables()
        if self._tracking_refine_apply_btn is not None:
            self._tracking_refine_apply_btn.setEnabled(False)
        if self._tracking_refine_remove_btn is not None:
            self._tracking_refine_remove_btn.setEnabled(False)
        if self._tracking_refine_status_label is not None:
            self._tracking_refine_status_label.setText(
                "Refine: <not available>"
            )

    def _set_tracking_refine_object(self, det_id: int) -> None:
        result = self._tracking_result
        combo = self._tracking_refine_object_combo
        if result is None or combo is None:
            return
        det_id = int(det_id)
        if det_id not in {int(det.id) for det in result.detections}:
            return
        if self._tracking_refine_image_pick_direction is not None:
            self._cancel_tracking_refine_image_pick(restore_context=False)
        self._tracking_refine_event_scope = None
        self._tracking_refine_det_id = det_id
        combo_index = combo.findData(det_id)
        if combo_index >= 0 and combo.currentIndex() != combo_index:
            combo.blockSignals(True)
            combo.setCurrentIndex(combo_index)
            combo.blockSignals(False)
        self._populate_tracking_refine_candidates(det_id)
        self._update_tracking_refine_image_pick_buttons()

    def _set_tracking_refine_event(self, event) -> None:
        result = self._tracking_result
        combo = self._tracking_refine_object_combo
        if result is None:
            return
        det_ids = {int(det.id) for det in result.detections}
        source_ids = tuple(
            int(det_id) for det_id in event.sources if int(det_id) in det_ids
        )
        target_ids = tuple(
            int(det_id) for det_id in event.targets if int(det_id) in det_ids
        )
        if not source_ids and not target_ids:
            return
        self._cancel_tracking_refine_image_pick(restore_context=False)
        scope = {
            "kind": str(event.kind),
            "frame_from": int(event.frame_from),
            "frame_to": int(event.frame_to),
            "sources": source_ids,
            "targets": target_ids,
        }
        self._tracking_refine_event_scope = scope
        anchor_id = int(source_ids[0] if source_ids else target_ids[0])
        self._tracking_refine_det_id = anchor_id
        if combo is not None:
            combo_index = combo.findData(anchor_id)
            if combo_index >= 0 and combo.currentIndex() != combo_index:
                combo.blockSignals(True)
                combo.setCurrentIndex(combo_index)
                combo.blockSignals(False)
        self._populate_tracking_refine_event_candidates(scope)
        self._update_tracking_refine_image_pick_buttons()

    def _resolve_tracking_refine_event_scope(
        self,
        scope: dict[str, Any],
    ) -> dict[str, Any]:
        result = self._tracking_result
        if result is None:
            return dict(scope)
        frame_from = int(scope.get("frame_from", -1))
        frame_to = int(scope.get("frame_to", -1))
        seed_sources = {int(det_id) for det_id in scope.get("sources", ())}
        seed_targets = {int(det_id) for det_id in scope.get("targets", ())}
        det_by_id = {int(det.id): det for det in result.detections}
        selected_ids = {int(link_id) for link_id in result.selected_link_ids}
        selected_links = [
            link
            for link in result.links
            if int(link.id) in selected_ids
            and int(det_by_id[int(link.src)].frame) == frame_from
            and int(det_by_id[int(link.dst)].frame) == frame_to
        ]
        related_links = [
            link
            for link in selected_links
            if int(link.src) in seed_sources or int(link.dst) in seed_targets
        ]
        if not related_links:
            return dict(scope)
        source_ids = {int(link.src) for link in related_links}
        target_ids = {int(link.dst) for link in related_links}
        changed = True
        while changed:
            changed = False
            for link in selected_links:
                src_id = int(link.src)
                dst_id = int(link.dst)
                if src_id not in source_ids and dst_id not in target_ids:
                    continue
                if src_id not in source_ids or dst_id not in target_ids:
                    source_ids.add(src_id)
                    target_ids.add(dst_id)
                    changed = True
        if len(source_ids) == 1 and len(target_ids) == 1:
            exact_event = next(
                (
                    event
                    for event in result.events
                    if int(event.frame_from) == frame_from
                    and int(event.frame_to) == frame_to
                    and {int(det_id) for det_id in event.sources} == source_ids
                    and {int(det_id) for det_id in event.targets} == target_ids
                ),
                None,
            )
            kind = str(exact_event.kind) if exact_event is not None else "linear"
        elif len(source_ids) == 1:
            kind = "fission"
        elif len(target_ids) == 1:
            kind = "fusion"
        else:
            kind = "split-merge"
        return {
            "kind": kind,
            "frame_from": frame_from,
            "frame_to": frame_to,
            "sources": tuple(sorted(source_ids)),
            "targets": tuple(sorted(target_ids)),
        }

    def _tracking_refine_candidate_records(self, det_id: int) -> list[dict[str, Any]]:
        result = self._tracking_result
        if result is None:
            return []
        if int(det_id) not in {int(det.id) for det in result.detections}:
            return []
        selected_ids = {int(link_id) for link_id in result.selected_link_ids}
        candidates: list[tuple[int, float, int, str, Any]] = []
        for link in result.links:
            if int(link.dst) == int(det_id):
                candidates.append((0, float(link.cost), int(link.id), "incoming", link))
            elif int(link.src) == int(det_id):
                candidates.append((1, float(link.cost), int(link.id), "outgoing", link))
        candidates.sort(key=lambda item: item[:3])

        records: list[dict[str, Any]] = []
        for _order, _cost, _link_id, direction, link in candidates:
            records.append(
                {
                    "direction": direction,
                    "link": link,
                    "current": int(link.id) in selected_ids,
                }
            )
        return records

    def _tracking_refine_event_candidate_records(
        self,
        scope: dict[str, Any],
    ) -> list[dict[str, Any]]:
        result = self._tracking_result
        if result is None:
            return []
        det_by_id = {int(det.id): det for det in result.detections}
        source_ids = {int(det_id) for det_id in scope.get("sources", ())}
        target_ids = {int(det_id) for det_id in scope.get("targets", ())}
        frame_from = int(scope.get("frame_from", -1))
        frame_to = int(scope.get("frame_to", -1))
        selected_ids = {int(link_id) for link_id in result.selected_link_ids}
        records: list[dict[str, Any]] = []
        for link in result.links:
            src_det = det_by_id.get(int(link.src))
            dst_det = det_by_id.get(int(link.dst))
            if src_det is None or dst_det is None:
                continue
            if (
                int(src_det.frame) != frame_from
                or int(dst_det.frame) != frame_to
            ):
                continue
            touches_source = int(link.src) in source_ids
            touches_target = int(link.dst) in target_ids
            if not touches_source and not touches_target:
                continue
            if touches_source and touches_target:
                relation = "event"
            elif touches_source:
                relation = "source"
            else:
                relation = "target"
            records.append(
                {
                    "direction": relation,
                    "link": link,
                    "current": int(link.id) in selected_ids,
                }
            )
        records.sort(
            key=lambda record: (
                not bool(record["current"]),
                float(record["link"].cost),
                int(record["link"].id),
            )
        )
        return records

    def _populate_tracking_refine_link_records(
        self,
        records: list[dict[str, Any]],
        *,
        status_text: str,
    ) -> None:
        result = self._tracking_result
        selected_table = self._tracking_refine_selected_table
        available_table = self._tracking_refine_available_table
        if result is None or selected_table is None or available_table is None:
            return
        det_by_id = {int(det.id): det for det in result.detections}
        selected_records = [record for record in records if bool(record["current"])]
        available_records = [record for record in records if not bool(record["current"])]

        def populate_table(table: QTableWidget, table_records: list[dict[str, Any]]) -> None:
            table.blockSignals(True)
            table.clearContents()
            table.setRowCount(len(table_records))
            for row, record in enumerate(table_records):
                link = record["link"]
                src_det = det_by_id[int(link.src)]
                dst_det = det_by_id[int(link.dst)]
                relation = str(record["direction"])
                tag = {
                    "incoming": "IN",
                    "outgoing": "OUT",
                    "event": "EVENT",
                    "source": "SRC",
                    "target": "DST",
                }.get(relation, "LINK")
                values = (
                    f"#{int(link.id)} {tag}",
                    f"t={int(src_det.frame)} obj={int(src_det.local_label)} -> "
                    f"t={int(dst_det.frame)} obj={int(dst_det.local_label)}",
                    f"{float(link.cost):.4f}",
                )
                tooltip = (
                    f"Link {int(link.id)}: source detection {int(link.src)} -> "
                    f"target detection {int(link.dst)}\n"
                    f"Source: t={int(src_det.frame)}, object={int(src_det.local_label)}, "
                    f"size={int(src_det.area)}\n"
                    f"Target: t={int(dst_det.frame)}, object={int(dst_det.local_label)}, "
                    f"size={int(dst_det.area)}\n"
                    f"Cost={float(link.cost):.6f}, distance={float(link.distance):.4f}"
                )
                for column, value in enumerate(values):
                    item = QTableWidgetItem(value)
                    item.setData(Qt.UserRole, int(link.id))
                    item.setToolTip(tooltip)
                    table.setItem(row, column, item)
            if table_records:
                table.setCurrentCell(0, 0)
            else:
                table.clearSelection()
                table.setCurrentCell(-1, -1)
            table.blockSignals(False)
            table.setEnabled(bool(table_records))

        populate_table(selected_table, selected_records)
        populate_table(available_table, available_records)
        if selected_records:
            available_table.blockSignals(True)
            available_table.clearSelection()
            available_table.setCurrentCell(-1, -1)
            available_table.blockSignals(False)
        elif available_records:
            selected_table.blockSignals(True)
            selected_table.clearSelection()
            selected_table.setCurrentCell(-1, -1)
            selected_table.blockSignals(False)

        if self._tracking_refine_status_label is not None:
            self._tracking_refine_status_label.setText(status_text)
        self._update_tracking_refine_remove_enabled()

    def _populate_tracking_refine_candidates(self, det_id: int) -> None:
        result = self._tracking_result
        if result is None:
            return
        det_by_id = {int(det.id): det for det in result.detections}
        det = det_by_id.get(int(det_id))
        if det is None:
            return
        records = self._tracking_refine_candidate_records(int(det_id))
        available_incoming = sum(
            record["direction"] == "incoming" and not bool(record["current"])
            for record in records
        )
        available_outgoing = sum(
            record["direction"] == "outgoing" and not bool(record["current"])
            for record in records
        )
        selected_incoming = sum(
            record["direction"] == "incoming" and bool(record["current"])
            for record in records
        )
        selected_outgoing = sum(
            record["direction"] == "outgoing" and bool(record["current"])
            for record in records
        )
        self._populate_tracking_refine_link_records(
            records,
            status_text=(
                f"Selected t={int(det.frame)}, object={int(det.local_label)}, "
                f"size={int(det.area)}. Nearby: IN={available_incoming}, "
                f"OUT={available_outgoing}; effective: IN={selected_incoming}, "
                f"OUT={selected_outgoing}."
            ),
        )

    def _populate_tracking_refine_event_candidates(
        self,
        scope: dict[str, Any],
    ) -> None:
        records = self._tracking_refine_event_candidate_records(scope)
        effective_count = sum(bool(record["current"]) for record in records)
        candidate_count = len(records) - effective_count
        direct_candidates = sum(
            not bool(record["current"]) and record["direction"] == "event"
            for record in records
        )
        self._populate_tracking_refine_link_records(
            records,
            status_text=(
                f"Event {scope.get('kind', 'unknown')} "
                f"t={int(scope.get('frame_from', -1))} -> "
                f"{int(scope.get('frame_to', -1))}: "
                f"{len(scope.get('sources', ()))} source, "
                f"{len(scope.get('targets', ()))} target objects. "
                f"Effective links={effective_count}; nearby candidates="
                f"{candidate_count} ({direct_candidates} within event objects)."
            ),
        )

    def _show_tracking_refine_link_preview(self, link_id: int) -> None:
        result = self._tracking_result
        tracking_layer = self._tracking_layer
        if result is None or tracking_layer is None:
            return
        link = next(
            (item for item in result.links if int(item.id) == int(link_id)),
            None,
        )
        if link is None:
            return
        det_by_id = {int(det.id): det for det in result.detections}
        src_det = det_by_id[int(link.src)]
        dst_det = det_by_id[int(link.dst)]
        self._clear_tracking_match_preview()
        self._apply_tracking_highlight_colors(
            {int(src_det.id): 1, int(dst_det.id): 2},
            {
                1: np.array([1.0, 0.55, 0.08, 1.0], dtype=float),
                2: np.array([0.10, 0.86, 1.0, 1.0], dtype=float),
            },
            opacity=0.98,
        )
        selected = int(link.id) in {int(value) for value in result.selected_link_ids}
        if self._tracking_refine_status_label is not None:
            state = "IN POOL" if selected else "CANDIDATE"
            self._tracking_refine_status_label.setText(
                f"Link {int(link.id)} [{state}]: source t={int(src_det.frame)} "
                f"object={int(src_det.local_label)} -> target t={int(dst_det.frame)} "
                f"object={int(dst_det.local_label)}; cost={float(link.cost):.4f}."
            )
        _safe_axis_labels(self.viewer, tracking_layer)
        self._apply_tracking_time_steps(tracking_layer, [int(src_det.frame)])

    def _update_tracking_refine_remove_enabled(self) -> None:
        has_result = bool(self._tracking_result and self._tracking_result.detections)
        if self._tracking_refine_apply_btn is not None:
            self._tracking_refine_apply_btn.setEnabled(
                has_result and self._tracking_refine_table_link_id(
                    self._tracking_refine_available_table
                )
                is not None
            )
        if self._tracking_refine_remove_btn is not None:
            self._tracking_refine_remove_btn.setEnabled(
                has_result and self._tracking_refine_table_link_id(
                    self._tracking_refine_selected_table
                )
                is not None
            )

    def _update_tracking_refine_undo_enabled(self) -> None:
        if self._tracking_refine_undo_btn is not None:
            self._tracking_refine_undo_btn.setEnabled(
                bool(
                    self._tracking_result
                    and self._tracking_refine_undo_history
                )
            )

    def _clear_tracking_refine_undo_history(self) -> None:
        self._tracking_refine_undo_history.clear()
        self._update_tracking_refine_undo_enabled()

    def _record_tracking_refine_undo(self, result: TrackingResult) -> None:
        scope = (
            dict(self._tracking_refine_event_scope)
            if self._tracking_refine_event_scope is not None
            else None
        )
        self._tracking_refine_undo_history.append(
            (
                tuple(int(link_id) for link_id in result.selected_link_ids),
                self._tracking_refine_det_id,
                scope,
            )
        )
        self._update_tracking_refine_undo_enabled()

    def _on_tracking_refine_undo_clicked(self) -> None:
        result = self._tracking_result
        if (
            result is None
            or not self._tracking_refine_undo_history
            or self._thread_is_running("_tracking_thread")
        ):
            return
        self._cancel_tracking_refine_image_pick(restore_context=False)
        selected_link_ids, det_id, event_scope = self._tracking_refine_undo_history.pop()
        try:
            restored = restore_tracking_link_selection(
                result,
                self._tracking_config or self._build_tracking_config(),
                selected_link_ids,
                message="Manual tracking refine undone.",
                rebuild_lineages=False,
            )
        except (TypeError, ValueError, KeyError) as error:
            self._tracking_refine_undo_history.append(
                (selected_link_ids, det_id, event_scope)
            )
            QMessageBox.critical(self, "Undo refine failed", str(error))
            return
        if det_id is None:
            det_id = int(result.detections[0].id) if result.detections else 0
        self._update_tracking_result_after_refine(
            restored,
            int(det_id),
            preferred_event_scope=event_scope,
            changed_transition=self._tracking_selection_changed_transition(
                result,
                restored,
            ),
        )
        self.track_progress.setFormat("Undo refine")
        self._update_tracking_refine_undo_enabled()

    def _on_tracking_refine_object_changed(self, _index: int) -> None:
        combo = self._tracking_refine_object_combo
        if combo is None:
            return
        det_id = combo.currentData()
        if det_id is None:
            return
        self._set_tracking_refine_object(int(det_id))
        self._show_tracking_object_context(int(det_id))

    def _on_tracking_refine_link_table_selection_changed(
        self,
        table: QTableWidget | None,
    ) -> None:
        if self._tracking_refine_syncing_tables or table is None:
            return
        link_id = self._tracking_refine_table_link_id(table)
        if link_id is None:
            self._update_tracking_refine_remove_enabled()
            return

        other_table = (
            self._tracking_refine_available_table
            if table is self._tracking_refine_selected_table
            else self._tracking_refine_selected_table
        )
        self._tracking_refine_syncing_tables = True
        try:
            if other_table is not None:
                other_table.clearSelection()
                other_table.setCurrentCell(-1, -1)
        finally:
            self._tracking_refine_syncing_tables = False
        self._update_tracking_refine_remove_enabled()
        self._show_tracking_refine_link_preview(int(link_id))

    def _set_tracking_refine_candidate(self, link_id: int) -> None:
        for table in (
            self._tracking_refine_selected_table,
            self._tracking_refine_available_table,
        ):
            if table is None:
                continue
            for row in range(table.rowCount()):
                item = table.item(row, 0)
                if item is None or int(item.data(Qt.UserRole)) != int(link_id):
                    continue
                self._tracking_refine_syncing_tables = True
                try:
                    for candidate_table in (
                        self._tracking_refine_selected_table,
                        self._tracking_refine_available_table,
                    ):
                        if candidate_table is not None:
                            candidate_table.blockSignals(True)
                            candidate_table.clearSelection()
                            candidate_table.setCurrentCell(-1, -1)
                            candidate_table.blockSignals(False)
                    table.blockSignals(True)
                    table.setCurrentCell(row, 0)
                    table.selectRow(row)
                    table.blockSignals(False)
                finally:
                    self._tracking_refine_syncing_tables = False
                self._on_tracking_refine_link_table_selection_changed(table)
                return

    def _apply_tracking_refine_candidate(self, link_id: int) -> None:
        result = self._tracking_result
        det_id = self._tracking_refine_det_id
        if result is None or det_id is None:
            return
        try:
            refined = set_tracking_link_selected(
                result,
                self._tracking_config or self._build_tracking_config(),
                int(link_id),
                selected=True,
                rebuild_lineages=False,
            )
        except (TypeError, ValueError, KeyError) as error:
            QMessageBox.critical(self, "Refine failed", str(error))
            return
        if refined.selected_link_ids == result.selected_link_ids:
            return
        self._record_tracking_refine_undo(result)
        self._update_tracking_result_after_refine(
            refined,
            int(det_id),
            preferred_link_id=int(link_id),
            preferred_event_scope=self._tracking_refine_event_scope,
            changed_transition=self._tracking_link_transition(result, int(link_id)),
        )

    def _remove_tracking_refine_link(self, link_id: int) -> None:
        result = self._tracking_result
        det_id = self._tracking_refine_det_id
        if result is None or det_id is None:
            return
        try:
            refined = set_tracking_link_selected(
                result,
                self._tracking_config or self._build_tracking_config(),
                int(link_id),
                selected=False,
                rebuild_lineages=False,
            )
        except (TypeError, ValueError, KeyError) as error:
            QMessageBox.critical(self, "Refine failed", str(error))
            return
        if refined.selected_link_ids == result.selected_link_ids:
            return
        self._record_tracking_refine_undo(result)
        self._update_tracking_result_after_refine(
            refined,
            int(det_id),
            preferred_link_id=int(link_id),
            preferred_event_scope=self._tracking_refine_event_scope,
            changed_transition=self._tracking_link_transition(result, int(link_id)),
        )

    def _on_tracking_refine_apply_clicked(self) -> None:
        link_id = self._tracking_refine_table_link_id(
            self._tracking_refine_available_table
        )
        if link_id is not None:
            self._apply_tracking_refine_candidate(int(link_id))

    def _on_tracking_refine_remove_clicked(self) -> None:
        link_id = self._tracking_refine_table_link_id(
            self._tracking_refine_selected_table
        )
        if link_id is not None:
            self._remove_tracking_refine_link(int(link_id))

    def _refresh_tracking_detection_lookup(self) -> None:
        result = self._tracking_result
        self._tracking_detection_lookup = (
            {
                (int(det.frame), int(det.local_label)): int(det.id)
                for det in result.detections
            }
            if result is not None
            else {}
        )

    @staticmethod
    def _tracking_link_transition(
        result: TrackingResult,
        link_id: int,
    ) -> tuple[int, int] | None:
        link = next(
            (item for item in result.links if int(item.id) == int(link_id)),
            None,
        )
        if link is None:
            return None
        det_by_id = {int(det.id): det for det in result.detections}
        source = det_by_id.get(int(link.src))
        target = det_by_id.get(int(link.dst))
        if source is None or target is None:
            return None
        return int(source.frame), int(target.frame)

    @classmethod
    def _tracking_selection_changed_transition(
        cls,
        before: TrackingResult,
        after: TrackingResult,
    ) -> tuple[int, int] | None:
        changed_ids = {
            int(link_id) for link_id in before.selected_link_ids
        } ^ {
            int(link_id) for link_id in after.selected_link_ids
        }
        if len(changed_ids) != 1:
            return None
        return cls._tracking_link_transition(before, changed_ids.pop())

    def _on_tracking_layer_visibility_changed(self, _event=None) -> None:
        layer = self._tracking_layer
        if layer is None or not bool(getattr(layer, "visible", False)):
            return
        if self._tracking_lineages_dirty and self._tracking_result is not None:
            self.track_progress.setRange(0, 0)
            self.track_progress.setFormat("Updating tracking index…")
            QApplication.processEvents()
            try:
                self._tracking_result = rebuild_tracking_lineages(
                    self._tracking_result
                )
            except (MemoryError, TypeError, ValueError, KeyError) as error:
                layer.visible = False
                self.track_progress.setRange(0, 1)
                self.track_progress.setFormat("Tracking index failed")
                QMessageBox.critical(
                    self,
                    "Tracking index failed",
                    str(error),
                )
                return
            self._tracking_lineages_dirty = False
            self._tracking_layer_pending_data = (
                self._tracking_result.tracked_labels.astype(
                    np.int32,
                    copy=False,
                )
            )
            color_map = self._tracking_lineage_color_map(
                self._tracking_result
            )
            self._tracking_base_color_map = dict(color_map)
            layer.color = color_map
            self.track_progress.setRange(0, 1)
            self.track_progress.setValue(1)
            self.track_progress.setFormat("Tracking index updated")
        pending = self._tracking_layer_pending_data
        if pending is not None:
            layer.data = pending
            layer.refresh()
            self._tracking_layer_pending_data = None

    def _update_tracking_result_after_refine(
        self,
        result: TrackingResult,
        det_id: int,
        *,
        preferred_link_id: int | None = None,
        preferred_event_scope: dict[str, Any] | None = None,
        changed_transition: tuple[int, int] | None = None,
    ) -> None:
        layer = self._tracking_layer
        if layer is None:
            return
        source_layer = self._selected_tracking_layer()
        self._clear_tracking_highlight()
        self._tracking_result = result
        self._tracking_lineages_dirty = True
        self._tracking_layer_pending_data = None
        resolved_event_scope = (
            self._resolve_tracking_refine_event_scope(preferred_event_scope)
            if preferred_event_scope is not None
            else None
        )
        with suppress(Exception):
            layer.opacity = 0.85
            layer.visible = False
        active_transition_before = self._tracking_active_transition
        self._populate_tracking_table(list(result.events))
        self._rebuild_tracking_event_base_layer(
            source_layer or layer,
            preferred_kind=self._tracking_active_event_kind,
        )
        active_transition = self._tracking_active_transition
        refresh_event_layer = (
            not self._tracking_event_layers
            or changed_transition is None
            or active_transition_before is None
            or tuple(active_transition_before) == tuple(changed_transition)
            or active_transition is None
        )
        if refresh_event_layer:
            self._rebuild_tracking_event_layers(source_layer or layer)
        if resolved_event_scope is not None:
            self._tracking_refine_event_scope = resolved_event_scope
            self._tracking_refine_det_id = int(det_id)
            combo = self._tracking_refine_object_combo
            if combo is not None:
                combo_index = combo.findData(int(det_id))
                if combo_index >= 0 and combo.currentIndex() != combo_index:
                    combo.blockSignals(True)
                    combo.setCurrentIndex(combo_index)
                    combo.blockSignals(False)
            self._populate_tracking_refine_event_candidates(
                self._tracking_refine_event_scope
            )
        else:
            self._set_tracking_refine_object(int(det_id))
        if preferred_link_id is not None:
            self._set_tracking_refine_candidate(int(preferred_link_id))
        if resolved_event_scope is None:
            self._show_tracking_object_context(int(det_id))
        if self._tracking_refine_status_label is not None:
            self._tracking_refine_status_label.setText(
                f"{self._tracking_refine_status_label.text()} {result.message}"
            )
        self.track_progress.setFormat("Refined")

    def _format_match_summary(self, summary: MatchSummary) -> str:
        if not summary.frame_summaries:
            return f"Match: {summary.message}"
        worst = min(
            summary.frame_summaries,
            key=lambda item: min(float(item.src_coverage), float(item.dst_coverage)),
        )
        return (
            "Match: "
            f"overall src={summary.overall_src_coverage * 100.0:.1f}% "
            f"dst={summary.overall_dst_coverage * 100.0:.1f}% "
            f"pairs={summary.total_matched_pairs}; "
            f"worst {worst.frame_from}->{worst.frame_to}: "
            f"src={worst.src_coverage * 100.0:.1f}% "
            f"dst={worst.dst_coverage * 100.0:.1f}% "
            f"({worst.matched_pairs}/{worst.src_points},{worst.dst_points})"
        )

    def _populate_tracking_match_table(self, summary: MatchSummary | None) -> None:
        table = self._tracking_match_table
        if table is None:
            return
        table.clearContents()
        headers = ["from", "to", "src %", "dst %", "pairs"]
        table.setColumnCount(len(headers))
        table.setHorizontalHeaderLabels(headers)
        rows = [] if summary is None else summary.frame_summaries
        table.setRowCount(len(rows))
        for r, item in enumerate(rows):
            values = [
                str(item.frame_from),
                str(item.frame_to),
                f"{item.src_coverage * 100.0:.1f}",
                f"{item.dst_coverage * 100.0:.1f}",
                str(item.matched_pairs),
            ]
            for c, value in enumerate(values):
                cell = QTableWidgetItem(value)
                cell.setData(Qt.UserRole, r)
                cell.setTextAlignment(Qt.AlignCenter)
                table.setItem(r, c, cell)
        table.resizeColumnsToContents()
        with suppress(Exception):
            header = table.horizontalHeader()
            if header is not None:
                header.setSectionResizeMode(QHeaderView.Stretch)

    def _populate_proximity_tables(self, result: ProximityResult | None) -> None:
        summary_table = self._proximity_summary_table
        components_table = self._proximity_components_table
        physical_label = self._proximity_physical_unit_label(result)
        proximity_label = self._proximity_threshold_unit_label(result)
        count_label = self._proximity_count_unit_label(result)
        proximity_count_label = self._proximity_threshold_count_unit_label(result)
        distance_label = str(getattr(result, "distance_unit", getattr(result, "unit", "unit")) or "unit") if result is not None else "unit"
        selected_roi_id = self._selected_proximity_roi_id()
        if summary_table is not None:
            summary_table.blockSignals(True)
            summary_table.clearContents()
            headers = [
                "roi",
                f"source ({count_label})",
                f"target ({count_label})",
                f"overlap ({count_label})",
                f"overlap ({physical_label})",
                f"src mean d ({distance_label})",
                f"tgt mean d ({distance_label})",
                "M1",
                "M2",
                "Dice",
                "Jaccard",
            ]
            summary_table.setColumnCount(len(headers))
            summary_table.setHorizontalHeaderLabels(headers)
            summary_rows = [] if result is None else result.summary_rows
            summary_table.setRowCount(len(summary_rows))
            for row_idx, row in enumerate(summary_rows):
                values = [
                    row.roi_name,
                    str(int(row.source_size)),
                    str(int(row.target_size)),
                    str(int(row.overlap_size)),
                    f"{row.overlap_size_physical:.4g}",
                    f"{row.source_mean_distance:.4g}",
                    f"{row.target_mean_distance:.4g}",
                    f"{row.manders_m1:.4f}",
                    f"{row.manders_m2:.4f}",
                    f"{row.dice:.4f}",
                    f"{row.jaccard:.4f}",
                ]
                for col_idx, value in enumerate(values):
                    item = QTableWidgetItem(value)
                    item.setTextAlignment(Qt.AlignCenter)
                    item.setData(Qt.UserRole, int(row.roi_id))
                    summary_table.setItem(row_idx, col_idx, item)
            with suppress(Exception):
                header = summary_table.horizontalHeader()
                header.setSectionResizeMode(QHeaderView.ResizeToContents)
            summary_table.resizeColumnsToContents()
            with suppress(Exception):
                summary_table.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
            if selected_roi_id is None:
                if len(summary_rows) == 1:
                    summary_table.selectRow(0)
                    summary_table.setCurrentCell(0, 0)
                    selected_roi_id = int(summary_rows[0].roi_id)
                else:
                    summary_table.clearSelection()
            else:
                target_row = None
                for row_idx in range(summary_table.rowCount()):
                    item = summary_table.item(row_idx, 0)
                    if item is None:
                        continue
                    item_roi_id = item.data(Qt.UserRole)
                    if item_roi_id is not None and int(item_roi_id) == int(selected_roi_id):
                        target_row = row_idx
                        break
                if target_row is None:
                    summary_table.clearSelection()
                else:
                    summary_table.selectRow(target_row)
                    summary_table.setCurrentCell(target_row, 0)
            summary_table.blockSignals(False)

        if components_table is not None:
            components_table.blockSignals(True)
            components_table.clearContents()
            component_rows = (
                []
                if result is None or selected_roi_id is None
                else [
                    row
                    for row in result.component_rows
                    if int(row.roi_id) == int(selected_roi_id)
                ]
            )
            object_rows_mode = any(int(getattr(row, "source_size", 0)) > 0 for row in component_rows)
            headers = (
                [
                    "source object",
                    f"source size ({count_label})",
                    f"source ({physical_label})",
                    f"target prox ({proximity_count_label})",
                    f"target prox ({proximity_label})",
                ]
                if object_rows_mode
                else [
                    "roi",
                    "frame",
                    "component",
                    f"size ({proximity_count_label})",
                    f"size ({proximity_label})",
                ]
            )
            components_table.setColumnCount(len(headers))
            components_table.setHorizontalHeaderLabels(headers)
            components_table.setRowCount(len(component_rows))
            for row_idx, row in enumerate(component_rows):
                if object_rows_mode:
                    values = [
                        str(int(row.source_label or row.component_id)),
                        str(int(row.source_size)),
                        f"{row.source_size_physical:.4g}",
                        str(int(row.size)),
                        f"{row.size_physical:.4g}",
                    ]
                else:
                    values = [
                        f"ROI {int(row.roi_id)}",
                        "-" if int(row.frame) < 0 else str(int(row.frame)),
                        str(int(row.component_id)),
                        str(int(row.size)),
                        f"{row.size_physical:.4g}",
                    ]
                for col_idx, value in enumerate(values):
                    item = QTableWidgetItem(value)
                    item.setTextAlignment(Qt.AlignCenter)
                    item.setData(Qt.UserRole, int(row.global_component_id))
                    components_table.setItem(row_idx, col_idx, item)
            with suppress(Exception):
                header = components_table.horizontalHeader()
                header.setSectionResizeMode(QHeaderView.ResizeToContents)
            components_table.resizeColumnsToContents()
            components_table.clearSelection()
            components_table.blockSignals(False)

    def _on_proximity_summary_selection_changed(self) -> None:
        table = self._proximity_summary_table
        if table is None:
            return
        items = table.selectedItems()
        if not items:
            return
        roi_id = items[0].data(Qt.UserRole)
        if roi_id is None:
            return
        result = self._proximity_result
        self._populate_proximity_tables(result)

    def _proximity_overlap_display_data(self, result: ProximityResult) -> np.ndarray:
        source_mask = np.asarray(result.source_mask, dtype=bool)
        target_mask = np.asarray(result.target_mask, dtype=bool)
        display_data = np.zeros(source_mask.shape, dtype=np.uint8)
        display_data[np.logical_and(source_mask, target_mask)] = 1
        return display_data

    def _show_proximity_result(
        self,
        result: ProximityResult,
        source_raw_layer,
        source_seg_layer,
        target_raw_layer,
        target_seg_layer,
        *,
        view_state: dict[str, object] | None = None,
    ) -> None:
        if view_state is None:
            view_state = self._capture_proximity_view_state()
        old_proximity_layer = self._proximity_layer
        if old_proximity_layer is not None and old_proximity_layer in self.viewer.layers:
            with suppress(Exception):
                self.viewer.layers.remove(old_proximity_layer)
        self._proximity_layer = None
        overlap_metadata = {
            "is_proximity_overlap_display": True,
            "source_seg_layer": result.source_seg_name,
            "target_seg_layer": result.target_seg_name,
            "label_meanings": {
                1: "true source-target overlap",
            },
        }
        overlap_display = self._proximity_overlap_display_data(result)
        overlap_scale = _spatial_scale_for_ndim(source_seg_layer, overlap_display.ndim)
        if (
            self._proximity_overlap_layer is not None
            and self._proximity_overlap_layer in self.viewer.layers
            and not self._is_labels_layer(self._proximity_overlap_layer)
        ):
            with suppress(Exception):
                self.viewer.layers.remove(self._proximity_overlap_layer)
            self._proximity_overlap_layer = None
        if self._proximity_overlap_layer is None or self._proximity_overlap_layer not in self.viewer.layers:
            overlap_kwargs = {
                "name": self._next_layer_name("true overlap"),
                "opacity": 1.0,
                "metadata": overlap_metadata,
            }
            if overlap_scale is not None:
                overlap_kwargs["scale"] = tuple(overlap_scale)
            overlap_kwargs.update(_spatial_unit_kwargs(source_seg_layer, overlap_display.ndim))
            self._proximity_overlap_layer = self.viewer.add_labels(overlap_display, **overlap_kwargs)
        else:
            with suppress(Exception):
                self._proximity_overlap_layer.data = overlap_display
                self._proximity_overlap_layer.metadata = overlap_metadata
                if overlap_scale is not None:
                    self._proximity_overlap_layer.scale = tuple(overlap_scale)
        self._apply_proximity_label_colors(
            self._proximity_overlap_layer,
            {1: np.array([1.0, 1.0, 1.0, 1.0], dtype=float)},
            blending="translucent_no_depth",
        )
        with suppress(Exception):
            self._proximity_overlap_layer.opacity = 1.0
            self._proximity_overlap_layer.blending = "translucent_no_depth"
            self._proximity_overlap_layer.visible = True
            self.viewer.layers.selection.active = self._proximity_overlap_layer
        self._sync_proximity_roi_annotation_layers(source_seg_layer)
        self._restore_proximity_view_state(view_state)

    def _export_proximity_summary(self) -> None:
        result = self._proximity_result
        if result is None or not result.summary_rows:
            QMessageBox.information(self, "No proximity summary", "Compute proximity first.")
            return
        count_label = self._proximity_count_unit_label(result).replace(" ", "_")
        physical_label = self._proximity_physical_unit_label(result)
        headers = [
            "roi",
            f"source_size_{count_label}",
            f"target_size_{count_label}",
            f"overlap_size_{count_label}",
            f"overlap_size_{physical_label}",
            f"source_mean_distance_{getattr(result, 'distance_unit', result.unit)}",
            f"source_median_distance_{getattr(result, 'distance_unit', result.unit)}",
            f"target_mean_distance_{getattr(result, 'distance_unit', result.unit)}",
            f"target_median_distance_{getattr(result, 'distance_unit', result.unit)}",
            "manders_m1",
            "manders_m2",
            "dice",
            "jaccard",
        ]
        rows = [
            [
                row.roi_name,
                str(int(row.source_size)),
                str(int(row.target_size)),
                str(int(row.overlap_size)),
                f"{row.overlap_size_physical:.6g}",
                f"{row.source_mean_distance:.6g}",
                f"{row.source_median_distance:.6g}",
                f"{row.target_mean_distance:.6g}",
                f"{row.target_median_distance:.6g}",
                f"{row.manders_m1:.6f}",
                f"{row.manders_m2:.6f}",
                f"{row.dice:.6f}",
                f"{row.jaccard:.6f}",
            ]
            for row in result.summary_rows
        ]
        def _stem(name: str) -> str:
            base = os.path.basename(str(name or "").strip())
            stem, _ext = os.path.splitext(base)
            return stem or "layer"

        source_stem = _stem(result.source_raw_name)
        target_stem = _stem(result.target_raw_name)
        if source_stem == target_stem:
            default_name = f"{source_stem}_proximity_summary.csv"
        else:
            default_name = f"{source_stem}__{target_stem}_proximity_summary.csv"
        path, selected_filter = QFileDialog.getSaveFileName(
            self,
            "Export proximity summary",
            default_name,
            "CSV (*.csv);;Text (*.txt);;Excel Workbook (*.xlsx)",
        )
        if not path:
            return
        if not path.lower().endswith((".csv", ".txt", ".xlsx")):
            if "xlsx" in selected_filter.lower():
                path += ".xlsx"
            elif "txt" in selected_filter.lower():
                path += ".txt"
            else:
                path += ".csv"
        write_analysis_export_file(path, headers, rows)

    def _on_compute_proximity_clicked(self) -> None:
        self._compute_proximity(use_roi=False)

    def _on_compute_proximity_roi_clicked(self) -> None:
        self._compute_proximity(use_roi=True)

    def _compute_proximity(self, *, use_roi: bool) -> None:
        if self._thread_is_running("_proximity_thread"):
            self._request_worker_cancel("_proximity_thread", "_proximity_worker")
            return
        roi_layer = self._proximity_roi_layer
        if (
            roi_layer is not None
            and roi_layer in self.viewer.layers
            and bool(getattr(roi_layer, "_is_creating", False))
        ):
            QMessageBox.information(
                self,
                "ROI in progress",
                "Finish the current ROI polygon before computing proximity.",
            )
            return
        source_raw_layer = self._selected_proximity_source_raw_layer()
        source_seg_layer = self._selected_proximity_source_seg_layer()
        target_raw_layer = self._selected_proximity_target_raw_layer()
        target_seg_layer = self._selected_proximity_target_seg_layer()
        if any(layer is None for layer in (source_raw_layer, source_seg_layer, target_raw_layer, target_seg_layer)):
            QMessageBox.information(
                self,
                "No layers",
                "Please choose source/target raw layers and source/target segmentation layers.",
            )
            return
        if source_seg_layer is target_seg_layer:
            QMessageBox.information(
                self,
                "Choose two layers",
                "Please select two different target/source segmentation layers.",
            )
            return
        roi_specs: list[ProximityRoiSpec] = []
        if use_roi:
            if roi_layer is None or roi_layer not in self.viewer.layers:
                QMessageBox.information(
                    self,
                    "No ROI",
                    "Draw at least one ROI before computing ROI proximity.",
                )
                return
            shape_data = list(getattr(roi_layer, "data", []) or [])
            if not shape_data:
                QMessageBox.information(
                    self,
                    "No ROI",
                    "Draw at least one ROI before computing ROI proximity.",
                )
                return
            roi_specs = self._proximity_roi_specs_for_layer(source_seg_layer)
            if not roi_specs:
                QMessageBox.information(
                    self,
                    "Invalid ROI",
                    "The ROI is incomplete or outside the image. Clear it and draw a new polygon.",
                )
                return
        spatial_ndim = self._proximity_spatial_ndim(source_seg_layer)
        layers = (source_raw_layer, source_seg_layer, target_raw_layer, target_seg_layer)
        kwargs = dict(
            roi_specs=roi_specs,
            voxel_size=_get_pixel_size_tuple(source_seg_layer, spatial_ndim),
            unit=_get_units_from_layer(source_seg_layer),
            distance_threshold=0.0,
            surface_only=False,
        )
        thread = QThread(self)
        worker = ProximityWorker(layers, kwargs)
        worker.moveToThread(thread)
        thread.started.connect(worker.run)
        self._proximity_thread, self._proximity_worker = thread, worker
        for button in self._proximity_compute_buttons:
            button.setText("Stop Proximity")

        def finished(result, error):
            self._cleanup_worker_thread("_proximity_thread", "_proximity_worker")
            for button, title in zip(self._proximity_compute_buttons, ("Compute Full Image", "Compute ROI(s)")):
                button.setText(title)
            if isinstance(error, ProximityCancelledError):
                return
            if error is not None:
                QMessageBox.critical(self, "Proximity failed", f"Failed to compute proximity: {error!r}")
                return
            if not all(layer in self.viewer.layers for layer in layers):
                return
            self._finish_proximity_result(result, layers, use_roi, spatial_ndim)

        self._connect_worker_callback(worker.finished, finished)
        thread.start()

    def _finish_proximity_result(self, result, layers, use_roi, spatial_ndim):
        source_raw_layer, source_seg_layer, target_raw_layer, target_seg_layer = layers
        self._apply_proximity_role_colors(source_raw_layer, source_seg_layer, target_raw_layer, target_seg_layer)
        self._proximity_result = result
        self._populate_proximity_tables(result)
        view_state = None
        if use_roi and spatial_ndim == 3:
            saved_state = self._proximity_roi_3d_view_state
            if (
                saved_state is not None
                and saved_state.get("source_layer") is source_seg_layer
            ):
                view_state = dict(saved_state)
                view_state.pop("source_layer", None)
            else:
                view_state = {"ndisplay": 3}
        self._show_proximity_result(
            result,
            source_raw_layer,
            source_seg_layer,
            target_raw_layer,
            target_seg_layer,
            view_state=view_state,
        )
        if use_roi:
            self._proximity_roi_3d_view_state = None
        for mask_layer in (source_seg_layer, target_seg_layer):
            with suppress(Exception):
                mask_layer.visible = False

    def _clear_tracking_match_preview(self) -> None:
        hidden_layers = getattr(self, "_tracking_match_hidden_layers", None) or []
        for layer, was_visible in hidden_layers:
            with suppress(Exception):
                layer.visible = bool(was_visible)
        self._tracking_match_hidden_layers = None
        for attr in (
            "_tracking_match_src_overlay_layer",
            "_tracking_match_dst_overlay_layer",
            "_tracking_match_link_layer",
            "_tracking_match_src_layer",
            "_tracking_match_dst_layer",
        ):
            layer = getattr(self, attr, None)
            if layer is not None:
                with suppress(Exception):
                    self.viewer.layers.remove(layer)
                setattr(self, attr, None)
        stale_names = {
            "match src overlay",
            "match dst overlay",
            "match src",
            "match dst",
            "match links",
        }
        for layer in list(self.viewer.layers):
            if str(getattr(layer, "name", "")) in stale_names:
                with suppress(Exception):
                    self.viewer.layers.remove(layer)

    @staticmethod
    def _hide_tracking_match_source_layers(
        *layers,
    ) -> list[tuple[object, bool]]:
        hidden_layers: list[tuple[object, bool]] = []
        for layer in layers:
            if layer is None or any(
                layer is hidden_layer
                for hidden_layer, _was_visible in hidden_layers
            ):
                continue
            try:
                was_visible = bool(layer.visible)
                layer.visible = False
            except (AttributeError, RuntimeError, TypeError):
                continue
            hidden_layers.append((layer, was_visible))
        return hidden_layers

    def _show_tracking_match_frame(self, detail) -> None:
        """Show frame-to-frame match preview as projected overlays, points, and links."""
        summary = self._tracking_match_summary
        if summary is None:
            return
        ref_layer = self._selected_tracking_layer()
        if ref_layer is None:
            return
        self._clear_tracking_match_preview()
        hidden_layers = self._hide_tracking_match_source_layers(
            ref_layer,
            self._tracking_layer,
        )
        self._tracking_match_hidden_layers = hidden_layers
        preview_metadata = {
            "managed_by_plugin": True,
            "is_tracking": True,
            "is_tracking_match_preview": True,
        }

        def _frame_binary_volume(layer, frame: int) -> np.ndarray | None:
            if layer is None or not hasattr(layer, "data"):
                return None
            data = layer.data
            dims_tag = _layer_dims_tag(layer)
            frame = int(frame)
            if dims_tag == "TZYX" and data.ndim >= 4:
                return (np.asarray(data[frame]) > 0).astype(float, copy=False)
            return None

        def _frame_binary_projection(layer, frame: int) -> np.ndarray | None:
            if layer is None or not hasattr(layer, "data"):
                return None
            data = layer.data
            dims_tag = _layer_dims_tag(layer)
            frame = int(frame)
            if dims_tag == "TYX" and data.ndim >= 3:
                frame_data = np.asarray(data[frame]) > 0
            elif dims_tag == "TZYX" and data.ndim >= 4:
                frame_data = np.max(np.asarray(data[frame]) > 0, axis=0)
            elif data.ndim == 3:
                frame_data = np.asarray(data[frame]) > 0
            elif data.ndim == 4:
                frame_data = np.max(np.asarray(data[frame]) > 0, axis=0)
            else:
                return None
            return np.asarray(frame_data, dtype=float)

        def _tinted_rgb(mask_2d: np.ndarray | None, rgb: tuple[float, float, float]) -> np.ndarray | None:
            if mask_2d is None:
                return None
            base = np.clip(np.asarray(mask_2d, dtype=float), 0.0, 1.0)
            return np.stack([base * float(rgb[0]), base * float(rgb[1]), base * float(rgb[2])], axis=-1)

        def _update_or_add_image(attr, data, name: str, opacity: float):
            if data is None:
                return
            existing = getattr(self, attr, None)
            scale = _spatial_scale_for_ndim(ref_layer, 2)
            translate = _spatial_translate_for_ndim(ref_layer, 2)
            if existing is not None and str(getattr(existing, "_type_string", "")) != "image":
                with suppress(Exception):
                    self.viewer.layers.remove(existing)
                existing = None
                setattr(self, attr, None)
            if existing is None:
                new_layer = self.viewer.add_image(
                    data,
                    name=self._next_layer_name(name),
                    rgb=True,
                    opacity=float(opacity),
                    blending="additive",
                    metadata=dict(preview_metadata),
                    **_spatial_unit_kwargs(ref_layer, 2),
                )
                if scale is not None:
                    with suppress(Exception):
                        new_layer.scale = scale
                if translate is not None:
                    with suppress(Exception):
                        new_layer.translate = translate
                setattr(self, attr, new_layer)
            else:
                with suppress(Exception):
                    existing.data = data
                    existing.opacity = float(opacity)
                    existing.blending = "additive"
                    if scale is not None:
                        existing.scale = scale
                    if translate is not None:
                        existing.translate = translate
                    existing.visible = True

        def _update_or_add_colored_volume(
            attr: str,
            data: np.ndarray | None,
            *,
            name: str,
            colormap: str,
            opacity: float,
        ) -> None:
            if data is None:
                return
            existing = getattr(self, attr, None)
            if existing is not None and str(getattr(existing, "_type_string", "")) != "image":
                with suppress(Exception):
                    self.viewer.layers.remove(existing)
                existing = None
                setattr(self, attr, None)
            if existing is not None and bool(getattr(existing, "rgb", False)):
                with suppress(Exception):
                    self.viewer.layers.remove(existing)
                existing = None
                setattr(self, attr, None)
            scalar_data = np.asarray(data, dtype=float)
            image_kwargs = {
                "name": self._next_layer_name(name) if existing is None else name,
                "opacity": float(opacity),
                "blending": "translucent_no_depth",
                "colormap": colormap,
                "contrast_limits": (0.0, 1.0),
                "gamma": 1.0,
                "metadata": dict(preview_metadata),
            }
            scale = _spatial_scale_for_ndim(ref_layer, data.ndim)
            translate = _spatial_translate_for_ndim(ref_layer, data.ndim)
            if scale is not None:
                image_kwargs["scale"] = scale
            if translate is not None:
                image_kwargs["translate"] = translate
            image_kwargs.update(_spatial_unit_kwargs(ref_layer, data.ndim))
            if existing is None:
                new_layer = self.viewer.add_image(scalar_data, **image_kwargs)
                setattr(self, attr, new_layer)
            else:
                with suppress(Exception):
                    existing.data = scalar_data
                    existing.name = name
                    existing.opacity = float(opacity)
                    existing.blending = "translucent_no_depth"
                    existing.colormap = colormap
                    existing.contrast_limits = (0.0, 1.0)
                    existing.gamma = 1.0
                    if scale is not None:
                        existing.scale = scale
                    if translate is not None:
                        existing.translate = translate
                    existing.visible = True

        src_owner = np.asarray(detail.src_owner, dtype=np.int32)
        dst_owner = np.asarray(detail.dst_owner, dtype=np.int32)
        src_idx = np.asarray(detail.match_src_indices, dtype=np.int32)
        dst_idx = np.asarray(detail.match_dst_indices, dtype=np.int32)

        src_ids = set(src_owner[src_idx].tolist()) if src_owner.size and src_idx.size else set()
        dst_ids = set(dst_owner[dst_idx].tolist()) if dst_owner.size and dst_idx.size else set()

        det_by_id = {int(det.id): det for det in summary.detections}

        is_true_3d = (_layer_dims_tag(ref_layer) == "TZYX")
        if is_true_3d:
            with suppress(Exception):
                self.viewer.dims.ndisplay = 3
            src_volume = _frame_binary_volume(ref_layer, int(detail.frame_from))
            dst_volume = _frame_binary_volume(ref_layer, int(detail.frame_to))
            _update_or_add_colored_volume(
                "_tracking_match_dst_overlay_layer",
                dst_volume,
                name="match dst overlay",
                colormap="cyan",
                opacity=0.18,
            )
            _update_or_add_colored_volume(
                "_tracking_match_src_overlay_layer",
                src_volume,
                name="match src overlay",
                colormap="orange",
                opacity=0.14,
            )
        else:
            src_overlay = _tinted_rgb(
                _frame_binary_projection(ref_layer, int(detail.frame_from)),
                (1.0, 0.56, 0.18),
            )
            dst_overlay = _tinted_rgb(
                _frame_binary_projection(ref_layer, int(detail.frame_to)),
                (0.10, 0.84, 1.0),
            )
            _update_or_add_image(
                "_tracking_match_dst_overlay_layer",
                dst_overlay,
                "match dst overlay",
                0.82,
            )
            _update_or_add_image(
                "_tracking_match_src_overlay_layer",
                src_overlay,
                "match src overlay",
                0.42,
            )

        def _project_points_to_view(points: np.ndarray) -> np.ndarray:
            pts = np.asarray(points, dtype=float)
            if pts.ndim != 2 or pts.size == 0:
                return np.zeros((0, 2), dtype=float)
            if pts.shape[1] >= 2:
                return pts[:, -2:].astype(float, copy=False)
            return np.zeros((0, 2), dtype=float)

        def _det_coords_spatial(det_ids: set, frame: int) -> np.ndarray:
            pts: list[np.ndarray] = []
            for det_id in det_ids:
                det = det_by_id.get(det_id)
                if det is None or int(det.frame) != frame:
                    continue
                raw = getattr(det, "sampled_coords", None)
                if raw is None:
                    raw = getattr(det, "coords", None)
                if raw is None:
                    continue
                coords = np.asarray(raw, dtype=float)
                if coords.ndim != 2 or coords.size == 0:
                    continue
                if coords.shape[0] > 96:
                    coords = _limit_display_points(coords, 96)
                pts.append(coords)
            if not pts:
                spatial_ndim = 3 if is_true_3d else 2
                return np.zeros((0, spatial_ndim), dtype=float)
            return np.vstack(pts)

        def _det_coords_projected(det_ids: set, frame: int) -> np.ndarray:
            pts: list[np.ndarray] = []
            for det_id in det_ids:
                det = det_by_id.get(det_id)
                if det is None or int(det.frame) != frame:
                    continue
                raw = getattr(det, "sampled_coords", None)
                if raw is None:
                    raw = getattr(det, "coords", None)
                if raw is None:
                    continue
                coords = np.asarray(raw, dtype=float)
                if coords.ndim != 2 or coords.size == 0:
                    continue
                if coords.shape[0] > 96:
                    coords = _limit_display_points(coords, 96)
                projected = _project_points_to_view(coords)
                if projected.size > 0:
                    pts.append(projected)
            return np.vstack(pts) if pts else np.zeros((0, 2), dtype=float)

        if is_true_3d:
            src_pts = _det_coords_spatial(src_ids, int(detail.frame_from))
            dst_pts = _det_coords_spatial(dst_ids, int(detail.frame_to))
            matched_src_pts = np.asarray(detail.src_points[src_idx], dtype=float)
            matched_dst_pts = np.asarray(detail.dst_points[dst_idx], dtype=float)
        else:
            src_pts = _det_coords_projected(src_ids, int(detail.frame_from))
            dst_pts = _det_coords_projected(dst_ids, int(detail.frame_to))
            matched_src_pts = _project_points_to_view(np.asarray(detail.src_points[src_idx], dtype=float))
            matched_dst_pts = _project_points_to_view(np.asarray(detail.dst_points[dst_idx], dtype=float))

        src_color = [1.00, 0.46, 0.16, 0.90]  # orange
        dst_color = [0.08, 0.78, 1.00, 0.90]  # cyan

        def _update_or_add(attr, pts, color, name):
            existing = getattr(self, attr, None)
            empty = np.zeros((0, pts.shape[1] if pts.ndim == 2 and pts.size > 0 else (3 if is_true_3d else 2)), dtype=float)
            data = pts if pts.shape[0] > 0 else empty
            if existing is None:
                kwargs = {
                    "name": self._next_layer_name(name),
                    "size": 4.0 if is_true_3d else 5.0,
                    "face_color": color,
                    "metadata": dict(preview_metadata),
                }
                scale = _spatial_scale_for_ndim(ref_layer, data.shape[1]) if data.ndim == 2 and data.shape[1] > 0 else None
                translate = _spatial_translate_for_ndim(ref_layer, data.shape[1]) if data.ndim == 2 and data.shape[1] > 0 else None
                if scale is not None:
                    kwargs["scale"] = scale
                if translate is not None:
                    kwargs["translate"] = translate
                kwargs.update(_spatial_unit_kwargs(ref_layer, data.shape[1]))
                new_layer = self.viewer.add_points(data, **kwargs)
                with suppress(Exception):
                    new_layer.visible = True
                setattr(self, attr, new_layer)
            else:
                with suppress(Exception):
                    existing.data = data
                    existing.face_color = color
                    scale = _spatial_scale_for_ndim(ref_layer, data.shape[1]) if data.ndim == 2 and data.shape[1] > 0 else None
                    translate = _spatial_translate_for_ndim(ref_layer, data.shape[1]) if data.ndim == 2 and data.shape[1] > 0 else None
                    if scale is not None:
                        existing.scale = scale
                    if translate is not None:
                        existing.translate = translate
                    existing.visible = True

        _update_or_add("_tracking_match_src_layer", src_pts, src_color, "match src")
        _update_or_add("_tracking_match_dst_layer", dst_pts, dst_color, "match dst")

        existing_links = getattr(self, "_tracking_match_link_layer", None)
        if is_true_3d:
            vectors = np.zeros((len(matched_src_pts), 2, 3), dtype=float)
            if len(matched_src_pts) > 0:
                vectors[:, 0, :] = matched_src_pts
                vectors[:, 1, :] = matched_dst_pts - matched_src_pts
            if vectors.size == 0:
                if existing_links is not None:
                    with suppress(Exception):
                        self.viewer.layers.remove(existing_links)
                    self._tracking_match_link_layer = None
            elif existing_links is None:
                kwargs = {
                    "name": self._next_layer_name("match links"),
                    "edge_color": "#f4f7fb",
                    "edge_width": 0.8,
                    "opacity": 0.6,
                    "metadata": dict(preview_metadata),
                }
                scale = _spatial_scale_for_ndim(ref_layer, 3)
                translate = _spatial_translate_for_ndim(ref_layer, 3)
                if scale is not None:
                    kwargs["scale"] = scale
                if translate is not None:
                    kwargs["translate"] = translate
                kwargs.update(_spatial_unit_kwargs(ref_layer, 3))
                self._tracking_match_link_layer = self.viewer.add_vectors(vectors, **kwargs)
            else:
                with suppress(Exception):
                    existing_links.data = vectors
                    scale = _spatial_scale_for_ndim(ref_layer, 3)
                    translate = _spatial_translate_for_ndim(ref_layer, 3)
                    if scale is not None:
                        existing_links.scale = scale
                    if translate is not None:
                        existing_links.translate = translate
                    existing_links.edge_color = "#f4f7fb"
                    existing_links.edge_width = 0.8
                    existing_links.opacity = 0.6
                    existing_links.visible = True
        else:
            line_segments = [
                np.asarray([src_pt, dst_pt], dtype=float)
                for src_pt, dst_pt in zip(matched_src_pts, matched_dst_pts, strict=False)
            ]
            if not line_segments:
                if existing_links is not None:
                    with suppress(Exception):
                        self.viewer.layers.remove(existing_links)
                    self._tracking_match_link_layer = None
            elif existing_links is None:
                self._tracking_match_link_layer = self.viewer.add_shapes(
                    line_segments,
                    shape_type="line",
                    name=self._next_layer_name("match links"),
                    edge_color="#f4f7fb",
                    edge_width=1.2,
                    opacity=0.55,
                    face_color="transparent",
                    metadata=dict(preview_metadata),
                    **_spatial_unit_kwargs(ref_layer, 2),
                )
                scale = _spatial_scale_for_ndim(ref_layer, 2)
                translate = _spatial_translate_for_ndim(ref_layer, 2)
                if scale is not None:
                    with suppress(Exception):
                        self._tracking_match_link_layer.scale = scale
                if translate is not None:
                    with suppress(Exception):
                        self._tracking_match_link_layer.translate = translate
                with suppress(Exception):
                    self._tracking_match_link_layer.visible = True
            else:
                with suppress(Exception):
                    existing_links.data = line_segments
                    existing_links.shape_type = ["line"] * len(line_segments)
                    existing_links.edge_color = "#f4f7fb"
                    existing_links.edge_width = 1.2
                    existing_links.opacity = 0.55
                    scale = _spatial_scale_for_ndim(ref_layer, 2)
                    translate = _spatial_translate_for_ndim(ref_layer, 2)
                    if scale is not None:
                        existing_links.scale = scale
                    if translate is not None:
                        existing_links.translate = translate
                    existing_links.visible = True

        with suppress(Exception):
            self.viewer.layers.selection.active = getattr(
                self,
                "_tracking_match_dst_overlay_layer",
                None,
            ) or getattr(self, "_tracking_match_src_overlay_layer", None)

        # Navigate to frame_to so the destination frame is the reference canvas.
        with suppress(Exception):
            _apply_viewer_time_step(self.viewer, ref_layer, int(detail.frame_to))

    def _select_tracking_event_in_tables(self, event) -> None:
        table_map = {
            "continuation": (self._tracking_linear_table, self._tracking_linear_events),
            "elongation": (self._tracking_linear_table, self._tracking_linear_events),
            "shortening": (self._tracking_linear_table, self._tracking_linear_events),
            "fusion": (self._tracking_fusion_table, self._tracking_fusion_events),
            "fission": (self._tracking_fission_table, self._tracking_fission_events),
            "split-merge": (self._tracking_split_merge_table, self._tracking_split_merge_events),
        }
        target = table_map.get(str(event.kind))
        if target is None:
            return
        table, events = target
        if table is None:
            return
        with suppress(ValueError):
            event_index = events.index(event)
        if "event_index" not in locals():
            return
        row = self._tracking_table_row_for_event_index(
            table,
            int(event_index),
        )
        if row is None:
            return
        for other in (
            self._tracking_linear_table,
            self._tracking_fission_table,
            self._tracking_split_merge_table,
            self._tracking_fusion_table,
        ):
            if other is None:
                continue
            other.blockSignals(True)
            other.clearSelection()
            other.blockSignals(False)
        table.blockSignals(True)
        table.selectRow(row)
        table.blockSignals(False)
        with suppress(Exception):
            table.scrollToItem(table.item(row, 0), QAbstractItemView.PositionAtTop)

    @staticmethod
    def _tracking_table_row_for_event_index(
        table: QTableWidget,
        event_index: int,
    ) -> int | None:
        for row in range(table.rowCount()):
            item = table.item(row, 0)
            if item is not None and int(item.data(Qt.UserRole)) == int(event_index):
                return row
        return None

    def _clear_tracking_event_table_selections(self) -> None:
        for other in (
            self._tracking_linear_table,
            self._tracking_fission_table,
            self._tracking_split_merge_table,
            self._tracking_fusion_table,
        ):
            if other is None:
                continue
            other.blockSignals(True)
            other.clearSelection()
            other.blockSignals(False)

    def _select_tracking_events_in_tables(self, events: list) -> None:
        table_map = {
            "continuation": (self._tracking_linear_table, self._tracking_linear_events),
            "elongation": (self._tracking_linear_table, self._tracking_linear_events),
            "shortening": (self._tracking_linear_table, self._tracking_linear_events),
            "fusion": (self._tracking_fusion_table, self._tracking_fusion_events),
            "fission": (self._tracking_fission_table, self._tracking_fission_events),
            "split-merge": (self._tracking_split_merge_table, self._tracking_split_merge_events),
        }
        box_by_table = {
            self._tracking_linear_table: getattr(self, "_tracking_linear_box", None),
            self._tracking_fission_table: getattr(self, "_tracking_fission_box", None),
            self._tracking_fusion_table: getattr(self, "_tracking_fusion_box", None),
            self._tracking_split_merge_table: getattr(self, "_tracking_split_merge_box", None),
        }
        self._clear_tracking_event_table_selections()
        selections: list[tuple[QTableWidget, int]] = []
        for event in events:
            target = table_map.get(str(event.kind))
            if target is None:
                continue
            table, event_list = target
            if table is None:
                continue
            with suppress(ValueError):
                event_index = event_list.index(event)
                row = self._tracking_table_row_for_event_index(
                    table,
                    int(event_index),
                )
                if row is not None and (table, row) not in selections:
                    selections.append((table, row))
        rows_by_table: dict[QTableWidget, list[int]] = {}
        target_box_y: int | None = None
        for table, row in selections:
            table.blockSignals(True)
            if table.item(row, 0) is not None:
                rows_by_table.setdefault(table, []).append(row)
                model = table.selectionModel()
                index = table.model().index(row, 0)
                if model is not None and index.isValid():
                    model.select(
                        index,
                        QItemSelectionModel.Select | QItemSelectionModel.Rows,
                    )
                box = box_by_table.get(table)
                if box is not None:
                    box_y = int(box.mapTo(self._scroll.widget(), QPoint(0, 0)).y())
                    target_box_y = box_y if target_box_y is None else min(target_box_y, box_y)
            table.blockSignals(False)
        first_by_table = {table: min(rows) for table, rows in rows_by_table.items() if rows}
        for table, row in first_by_table.items():
            model = table.selectionModel()
            index = table.model().index(row, 0)
            if model is not None and index.isValid():
                model.setCurrentIndex(index, QItemSelectionModel.NoUpdate)
            item = table.item(row, 0)
            if item is not None:
                table.scrollToItem(item, QAbstractItemView.PositionAtTop)
        if target_box_y is not None and getattr(self, "_scroll", None) is not None:
            scroll_bar = self._scroll.verticalScrollBar()
            if scroll_bar is not None:
                top_y = max(0, int(target_box_y) - self._scaled_px(8))
                scroll_bar.setValue(top_y)

    def _apply_pending_tracking_time_jump(self) -> None:
        pending = self._tracking_pending_jump
        self._tracking_pending_jump = None
        if pending is None:
            return
        layer, frame = pending
        with suppress(Exception):
            _apply_viewer_time_step(self.viewer, layer, int(frame))

    def _apply_tracking_time_steps(self, layer, frames: list[int]) -> None:
        if layer is None or not frames:
            return
        if self._tracking_jump_timer is not None:
            self._tracking_jump_timer.stop()
        self._tracking_pending_jump = None
        _apply_viewer_time_step(self.viewer, layer, int(frames[0]))
        if len(frames) > 1 and self._tracking_jump_timer is not None:
            self._tracking_pending_jump = (layer, int(frames[-1]))
            self._tracking_jump_timer.start()

    def _schedule_tracking_interaction(self, kind: str, payload) -> None:
        self._tracking_pending_interaction = (str(kind), payload)
        if self._tracking_interaction_timer is not None:
            self._tracking_interaction_timer.start()
        else:
            self._run_pending_tracking_interaction()

    def _run_pending_tracking_interaction(self) -> None:
        pending = self._tracking_pending_interaction
        self._tracking_pending_interaction = None
        if pending is None:
            return
        kind, payload = pending
        if kind == "event":
            event, jump_target = payload
            self._set_tracking_refine_event(event)
            self._show_tracking_event_highlight(event, jump_target=jump_target)
        elif kind == "object":
            self._show_tracking_object_context(int(payload))
        elif kind == "match":
            self._show_tracking_match_frame(payload)

    def _tracking_events_for_kind(self, kind: str) -> list:
        if kind in {"continuation", "elongation", "shortening"}:
            return self._tracking_linear_events
        if kind == "fusion":
            return self._tracking_fusion_events
        if kind == "split-merge":
            return self._tracking_split_merge_events
        return self._tracking_fission_events

    def _sync_tracking_event_tables_for_det(self, det_id: int) -> None:
        result = self._tracking_result
        if result is None:
            return
        det_by_id = {det.id: det for det in result.detections}
        det = det_by_id.get(int(det_id))
        if det is None:
            return
        supported_kinds = {
            "continuation",
            "elongation",
            "shortening",
            "fusion",
            "fission",
            "split-merge",
        }
        related = [
            event
            for event in result.events
            if event.kind in supported_kinds
            and (
                (
                    int(det_id) in event.sources
                    and int(event.frame_from) == int(det.frame)
                )
                or (
                    int(det_id) in event.targets
                    and int(event.frame_to) == int(det.frame)
                )
            )
        ]
        if not related:
            self._clear_tracking_event_table_selections()
            return

        active_transition = self._tracking_active_transition
        selected_events = [
            event
            for event in related
            if active_transition is not None
            and (
                int(event.frame_from),
                int(event.frame_to),
            ) == tuple(active_transition)
        ]
        if not selected_events:
            active_kind = self._tracking_active_event_kind

            def event_rank(event) -> tuple[int, int, int, str]:
                is_outgoing = (
                    int(event.frame_from) == int(det.frame)
                    and int(det_id) in event.sources
                )
                group_kind = self._tracking_event_group_kind(str(event.kind))
                return (
                    0 if is_outgoing else 1,
                    0 if group_kind == active_kind else 1,
                    int(event.frame_from),
                    str(event.kind),
                )

            primary = min(related, key=event_rank)
            active_transition = (
                int(primary.frame_from),
                int(primary.frame_to),
            )
            self._set_tracking_active_transition(
                *active_transition,
                render=False,
            )
            selected_events = [
                event
                for event in related
                if (
                    int(event.frame_from),
                    int(event.frame_to),
                ) == active_transition
            ]
        self._select_tracking_events_in_tables(selected_events)

    @staticmethod
    def _tracking_unlinked_large_object_records(
        result: TrackingResult,
        min_link_size: int,
    ) -> list[tuple[Any, str]]:
        selected_ids = {
            int(link_id) for link_id in result.selected_link_ids
        }
        selected_links = [
            link
            for link in result.links
            if int(link.id) in selected_ids
        ]
        detections_with_in = {
            int(link.dst) for link in selected_links
        }
        detections_with_out = {
            int(link.src) for link in selected_links
        }
        frame_count = int(result.frame_labels.shape[0])
        records: list[tuple[Any, str]] = []
        for detection in result.detections:
            if int(detection.area) <= int(min_link_size):
                continue
            missing: list[str] = []
            if (
                int(detection.frame) > 0
                and int(detection.id) not in detections_with_in
            ):
                missing.append("IN")
            if (
                int(detection.frame) < frame_count - 1
                and int(detection.id) not in detections_with_out
            ):
                missing.append("OUT")
            if missing:
                records.append((detection, " + ".join(missing)))
        return sorted(
            records,
            key=lambda record: (
                int(record[0].frame),
                -int(record[0].area),
                int(record[0].local_label),
                int(record[0].id),
            ),
        )

    def _tracking_visualization_frame(self) -> int:
        result = self._tracking_result
        if result is None:
            return 0
        reference_layer = (
            self._active_tracking_event_layer()
            or self._tracking_event_base_layer
            or self._tracking_layer
            or self._tracking_result_source_layer
        )
        if reference_layer is not None:
            return self._tracking_current_viewer_frame(reference_layer)
        if self._tracking_active_transition is not None:
            return int(self._tracking_active_transition[0])
        return 0

    def _populate_tracking_unlinked_table(
        self,
        *,
        force: bool = False,
    ) -> None:
        table = self._tracking_unlinked_table
        if table is None:
            return
        frame = self._tracking_visualization_frame()
        if not force and frame == self._tracking_unlinked_display_frame:
            return
        self._tracking_unlinked_display_frame = frame
        threshold = int(self.track_min_link_size.value())
        result = self._tracking_result
        all_records = (
            self._tracking_unlinked_large_object_records(result, threshold)
            if result is not None
            else []
        )
        records = [
            record
            for record in all_records
            if int(record[0].frame) == int(frame)
        ]
        if self._tracking_unlinked_box is not None:
            self._tracking_unlinked_box.setTitle(
                f"Unlinked Large Objects "
                f"(t={frame}; size > {threshold}; {len(records)})"
            )
        signals_were_blocked = table.blockSignals(True)
        table.setUpdatesEnabled(False)
        try:
            table.clearContents()
            table.setRowCount(len(records))
            for row, (detection, missing) in enumerate(records):
                values = (
                    int(detection.frame),
                    int(detection.local_label),
                    int(detection.area),
                    str(missing),
                )
                for column, value in enumerate(values):
                    item = QTableWidgetItem(str(value))
                    item.setData(Qt.UserRole, int(detection.id))
                    item.setTextAlignment(Qt.AlignCenter)
                    table.setItem(row, column, item)
            table.setEnabled(bool(records))
            table.horizontalScrollBar().setValue(0)
        finally:
            table.setUpdatesEnabled(True)
            table.blockSignals(signals_were_blocked)

    def _on_tracking_unlinked_object_clicked(
        self,
        row: int,
        _column: int,
    ) -> None:
        table = self._tracking_unlinked_table
        if table is None:
            return
        item = table.item(int(row), 0)
        if item is None:
            return
        det_id = item.data(Qt.UserRole)
        if det_id is not None:
            self._show_tracking_object_context(
                int(det_id),
                sync_events=False,
            )

    def _select_tracking_unlinked_object_in_table(self, det_id: int) -> None:
        table = self._tracking_unlinked_table
        if table is None:
            return
        target_row = None
        for row in range(table.rowCount()):
            item = table.item(row, 0)
            if (
                item is not None
                and item.data(Qt.UserRole) is not None
                and int(item.data(Qt.UserRole)) == int(det_id)
            ):
                target_row = row
                break
        signals_were_blocked = table.blockSignals(True)
        try:
            table.clearSelection()
            if target_row is not None:
                table.selectRow(int(target_row))
                item = table.item(int(target_row), 0)
                if item is not None:
                    table.scrollToItem(
                        item,
                        QAbstractItemView.PositionAtCenter,
                    )
        finally:
            table.blockSignals(signals_were_blocked)

    def _populate_tracking_event_table(self, table, events, kind: str) -> None:
        if table is None:
            return
        signals_were_blocked = table.blockSignals(True)
        table.setUpdatesEnabled(False)
        try:
            table.clearContents()
            step_columns = self._tracking_event_step_columns((kind,))
            headers = self._tracking_event_export_headers(
                step_columns=step_columns,
            )
            table.setColumnCount(len(headers))
            table.setHorizontalHeaderLabels(headers)
            transition = self._tracking_active_transition
            indexed_events = [
                (event_index, event)
                for event_index, event in enumerate(events)
                if transition is None
                or (
                    int(event.frame_from) == int(transition[0])
                    and int(event.frame_to) == int(transition[1])
                )
            ]
            table.setRowCount(len(indexed_events))
            for r, (event_index, event) in enumerate(indexed_events):
                row_values = self._tracking_event_export_row(
                    event,
                    step_columns=step_columns,
                )
                for c, value in enumerate(row_values):
                    item = QTableWidgetItem(value)
                    item.setData(Qt.UserRole, int(event_index))
                    item.setData(
                        Qt.UserRole + 1,
                        event.kind if kind == "linear" else kind,
                    )
                    item.setTextAlignment(Qt.AlignCenter)
                    table.setItem(r, c, item)
            widths = (92, 56, 56, 92, 60, 60, 128, 128)
            for column, width in enumerate(widths):
                if column < len(headers):
                    table.setColumnWidth(column, self._scaled_px(width))
            for column in range(len(widths), len(headers)):
                table.setColumnWidth(column, self._scaled_px(92))
            table.horizontalScrollBar().setValue(0)
        finally:
            table.setUpdatesEnabled(True)
            table.blockSignals(signals_were_blocked)

    def _populate_tracking_table(self, events) -> None:
        linear_events = sorted(
            [event for event in events if event.kind in {"continuation", "elongation", "shortening"}],
            key=lambda event: (int(event.frame_from), int(event.frame_to), str(event.kind)),
        )
        fission_events = sorted(
            [event for event in events if event.kind == "fission"],
            key=lambda event: (int(event.frame_from), int(event.frame_to)),
        )
        split_merge_events = sorted(
            [event for event in events if event.kind == "split-merge"],
            key=lambda event: (int(event.frame_from), int(event.frame_to)),
        )
        fusion_events = sorted(
            [event for event in events if event.kind == "fusion"],
            key=lambda event: (int(event.frame_from), int(event.frame_to)),
        )
        self._tracking_linear_events = linear_events
        self._tracking_fission_events = fission_events
        self._tracking_split_merge_events = split_merge_events
        self._tracking_fusion_events = fusion_events
        self._tracking_display_events = (
            linear_events + fission_events + split_merge_events + fusion_events
        )
        if self._tracking_statistics_export_btn is not None:
            self._tracking_statistics_export_btn.setEnabled(
                bool(self._tracking_display_events)
            )
        if self._tracking_statistics_import_btn is not None:
            self._tracking_statistics_import_btn.setEnabled(
                self._tracking_result is not None
            )
        self._populate_tracking_unlinked_table(force=True)
        self._rebuild_tracking_transition_combo()

    def _tracking_event_export_headers(
        self,
        *,
        step_columns: tuple[str, ...] = (),
        include_links: bool = False,
    ) -> list[str]:
        headers = [
            "kind",
            "from",
            "to",
            "size",
            "#src",
            "#dst",
            "sources",
            "targets",
        ]
        if include_links:
            headers.extend(("link ids", "link pairs"))
        headers.extend(step_columns)
        return headers

    @staticmethod
    def _tracking_event_step_columns(
        event_kinds,
    ) -> tuple[str, ...]:
        group_kinds = {
            SIGMAWidget._tracking_event_group_kind(kind)
            for kind in event_kinds
        }
        columns = []
        if group_kinds & {"fission", "split-merge"}:
            columns.append("fission steps")
        if group_kinds & {"fusion", "split-merge"}:
            columns.append("fusion steps")
        return tuple(columns)

    @staticmethod
    def _tracking_event_step_counts_from_links(
        event,
        links,
        selected_link_ids,
    ) -> tuple[int, int]:
        source_ids = {int(value) for value in event.sources}
        target_ids = {int(value) for value in event.targets}
        selected_ids = {int(value) for value in selected_link_ids}
        outgoing = {det_id: 0 for det_id in source_ids}
        incoming = {det_id: 0 for det_id in target_ids}
        for link in links:
            if int(link.id) not in selected_ids:
                continue
            src = int(link.src)
            dst = int(link.dst)
            if src not in source_ids or dst not in target_ids:
                continue
            outgoing[src] += 1
            incoming[dst] += 1
        fission_steps = sum(max(count - 1, 0) for count in outgoing.values())
        fusion_steps = sum(max(count - 1, 0) for count in incoming.values())
        return int(fission_steps), int(fusion_steps)

    def _tracking_event_step_counts(self, event) -> tuple[int, int]:
        result = self._tracking_result
        if result is None:
            return 0, 0
        return SIGMAWidget._tracking_event_step_counts_from_links(
            event,
            getattr(result, "links", ()),
            getattr(result, "selected_link_ids", ()),
        )

    def _tracking_event_export_row(
        self,
        event,
        *,
        step_columns: tuple[str, ...] = (),
        include_links: bool = False,
    ) -> list[str]:
        det_by_id = {}
        if self._tracking_result is not None:
            det_by_id = {int(det.id): det for det in self._tracking_result.detections}
        src_area = sum(int(getattr(det_by_id.get(int(v)), "area", 0)) for v in event.sources)
        dst_area = sum(int(getattr(det_by_id.get(int(v)), "area", 0)) for v in event.targets)
        row = [
            str(event.kind),
            str(int(event.frame_from)),
            str(int(event.frame_to)),
            f"{src_area}->{dst_area}",
            str(len(event.sources)),
            str(len(event.targets)),
            ",".join(str(int(v)) for v in event.sources),
            ",".join(str(int(v)) for v in event.targets),
        ]
        if include_links:
            source_ids = {int(value) for value in event.sources}
            target_ids = {int(value) for value in event.targets}
            selected_ids = {
                int(value)
                for value in getattr(
                    self._tracking_result,
                    "selected_link_ids",
                    (),
                )
            }
            event_links = sorted(
                (
                    link
                    for link in getattr(self._tracking_result, "links", ())
                    if int(link.id) in selected_ids
                    and int(link.src) in source_ids
                    and int(link.dst) in target_ids
                ),
                key=lambda link: int(link.id),
            )
            row.extend(
                (
                    ",".join(str(int(link.id)) for link in event_links),
                    ",".join(
                        f"{int(link.src)}->{int(link.dst)}"
                        for link in event_links
                    ),
                )
            )
        if step_columns:
            fission_steps, fusion_steps = self._tracking_event_step_counts(event)
            group_kind = SIGMAWidget._tracking_event_group_kind(event.kind)
            for column in step_columns:
                if column == "fission steps":
                    row.append(
                        str(fission_steps)
                        if group_kind in {"fission", "split-merge"}
                        else ""
                    )
                elif column == "fusion steps":
                    row.append(
                        str(fusion_steps)
                        if group_kind in {"fusion", "split-merge"}
                        else ""
                    )
        return row

    def _clear_tracking_highlight(self, *, remove_layer: bool = False) -> None:
        self._tracking_highlight_det_ids = set()
        self._tracking_highlight_value_det_ids = {}
        if self._tracking_event_layers:
            active_kind = self._tracking_active_event_kind
            if self._tracking_event_base_layer is not None:
                with suppress(Exception):
                    self._tracking_event_base_layer.visible = True
            for kind, event_layer in self._tracking_event_layers.items():
                with suppress(Exception):
                    event_layer.color = self._tracking_event_layer_color_maps.get(
                        kind,
                        event_layer.color,
                    )
                    event_layer.opacity = 0.85
                    event_layer.visible = kind == active_kind
                    event_layer.refresh()
            if self._tracking_layer is not None:
                with suppress(Exception):
                    self._tracking_layer.visible = False
        elif self._tracking_layer is not None and self._tracking_base_color_map is not None:
            with suppress(Exception):
                self._tracking_layer.color = self._tracking_base_color_map
                self._tracking_layer.opacity = 0.85
                self._tracking_layer.visible = True
                self._tracking_layer.refresh()
        if self._tracking_highlight_layer is not None:
            if remove_layer:
                with suppress(Exception):
                    self.viewer.layers.remove(self._tracking_highlight_layer)
                self._tracking_highlight_layer = None
            else:
                with suppress(Exception):
                    self._tracking_highlight_layer.visible = False
        self._clear_tracking_match_preview()

    def _apply_tracking_highlight_colors(
        self,
        value_by_det: dict[int, int],
        value_colors: dict[int, np.ndarray],
        *,
        opacity: float = 0.95,
        base_layer=None,
        base_opacity: float = 0.12,
    ) -> None:
        layer = base_layer or self._active_tracking_event_layer() or self._tracking_layer
        result = self._tracking_result
        if layer is None or result is None:
            return
        self._tracking_highlight_det_ids = {int(v) for v in value_by_det}
        value_det_ids: dict[int, set[int]] = {}
        for det_id, value in value_by_det.items():
            value_det_ids.setdefault(int(value), set()).add(int(det_id))
        self._tracking_highlight_value_det_ids = value_det_ids
        highlight = self._tracking_id_value_mask_data(
            value_by_det,
            lazy_time=True,
        )
        if highlight is None:
            return
        transparent = np.array([0.0, 0.0, 0.0, 0.0], dtype=float)
        color_map: dict[int | None, np.ndarray] = {None: transparent, 0: transparent}
        for value, color in value_colors.items():
            color_map[int(value)] = np.asarray(color, dtype=float)
        scale = getattr(layer, "scale", None)
        add_kwargs = {}
        if scale is not None:
            add_kwargs["scale"] = scale
        add_kwargs.update(_spatial_unit_kwargs(layer, np.asarray(highlight).ndim))
        if self._tracking_highlight_layer is None:
            self._tracking_highlight_layer = self.viewer.add_labels(
                highlight,
                name=self._next_layer_name("tracking event highlight"),
                metadata={
                    "managed_by_plugin": True,
                    "is_analysis_highlight": True,
                },
                **add_kwargs,
            )
        else:
            with suppress(Exception):
                self._tracking_highlight_layer.data = highlight
                if scale is not None:
                    self._tracking_highlight_layer.scale = scale
        with suppress(Exception):
            self._tracking_highlight_layer.color = color_map
            self._tracking_highlight_layer.opacity = float(opacity)
            self._tracking_highlight_layer.visible = True
            self._tracking_highlight_layer.refresh()
            layer.visible = True
            layer.opacity = float(base_opacity)
            layer.refresh()
            self.viewer.layers.selection.active = layer

    def _tracking_id_value_mask_data(
        self,
        value_by_det: dict[int, int],
        *,
        lazy_time: bool = False,
    ) -> np.ndarray | da.Array | None:
        result = self._tracking_result
        if result is None:
            return None
        max_value = max((int(value) for value in value_by_det.values()), default=0)
        if max_value <= np.iinfo(np.uint8).max:
            dtype = np.uint8
        elif max_value <= np.iinfo(np.uint16).max:
            dtype = np.uint16
        else:
            dtype = np.uint32
        det_by_id = {det.id: det for det in result.detections}
        selected = [
            (det_by_id[int(det_id)], int(label_value))
            for det_id, label_value in value_by_det.items()
            if int(det_id) in det_by_id
        ]
        if lazy_time:
            if not selected:
                return self._tracking_event_full_time_data(
                    np.zeros(
                        (1, *result.frame_labels.shape[1:]),
                        dtype=dtype,
                    ),
                    0,
                    int(result.frame_labels.shape[0]),
                )
            frame_start = min(int(det.frame) for det, _value in selected)
            frame_stop = max(int(det.frame) for det, _value in selected)
            data = np.zeros(
                (
                    frame_stop - frame_start + 1,
                    *result.frame_labels.shape[1:],
                ),
                dtype=dtype,
            )
        else:
            frame_start = 0
            data = np.zeros_like(result.frame_labels, dtype=dtype)
        for det, label_value in selected:
            spatial_slices = tuple(det.slice_tuple)
            target = data[
                (int(det.frame) - int(frame_start), *spatial_slices)
            ]
            mask = np.asarray(det.image, dtype=bool)
            target[mask] = int(label_value)
        if lazy_time:
            return self._tracking_event_full_time_data(
                data,
                frame_start,
                int(result.frame_labels.shape[0]),
            )
        return data

    def _tracking_kind_overlay_data(self, kind: str, base_layer) -> np.ndarray | None:
        result = self._tracking_result
        if result is None or base_layer is None or not hasattr(base_layer, "data"):
            return None
        base = np.asarray(base_layer.data, dtype=np.float32)
        if base.shape != result.frame_labels.shape:
            return None
        arr = np.asarray(base, dtype=np.float32)
        if arr.size == 0:
            return None
        lo = float(np.percentile(arr, 1.0))
        hi = float(np.percentile(arr, 99.5))
        if not np.isfinite(lo) or not np.isfinite(hi) or hi <= lo:
            lo = float(np.min(arr))
            hi = float(np.max(arr))
        scale = max(hi - lo, 1e-6)
        norm = np.clip((arr - lo) / scale, 0.0, 1.0)
        gray = (norm * 255.0).astype(np.uint8)
        if gray.ndim == 3 or gray.ndim == 4:
            base_rgb = np.stack([gray, gray, gray], axis=-1)
            source_overlay = base_rgb.copy()
            target_overlay = base_rgb.copy()
        else:
            return None

        det_by_id = {int(det.id): det for det in result.detections}
        alpha = 0.42

        def _event_color(event_index: int) -> np.ndarray:
            base_hues = {
                "linear": 0.72,
                "fission": 0.03,
                "fusion": 0.53,
                "split-merge": 0.14,
            }
            hue = (base_hues.get(str(kind), 0.0) + event_index * 0.61803398875) % 1.0
            sat = 0.72
            val = 0.95
            rgb = colorsys.hsv_to_rgb(hue, sat, val)
            return np.asarray([int(255 * c) for c in rgb], dtype=np.uint8)

        def _paint_detection(canvas: np.ndarray, det_id: int, color: np.ndarray) -> None:
            det = det_by_id.get(int(det_id))
            if det is None:
                return
            coords = np.asarray(det.coords, dtype=int)
            if coords.size == 0:
                return
            color = color.astype(np.float32)
            frame = int(det.frame)
            if canvas.ndim == 4:
                yy = coords[:, 0]
                xx = coords[:, 1]
                current = canvas[frame, yy, xx].astype(np.float32)
                blended = (1.0 - alpha) * current + alpha * color
                canvas[frame, yy, xx] = np.clip(blended, 0, 255).astype(np.uint8)
            else:
                zz = coords[:, 0]
                yy = coords[:, 1]
                xx = coords[:, 2]
                current = canvas[frame, zz, yy, xx].astype(np.float32)
                blended = (1.0 - alpha) * current + alpha * color
                canvas[frame, zz, yy, xx] = np.clip(blended, 0, 255).astype(np.uint8)

        for event_index, event in enumerate(result.events):
            is_linear = str(kind) == "linear" and str(event.kind) in {"continuation", "elongation", "shortening"}
            if not (str(event.kind) == str(kind) or is_linear):
                continue
            color = _event_color(event_index)
            for det_id in event.sources:
                _paint_detection(source_overlay, int(det_id), color)
            for det_id in event.targets:
                _paint_detection(target_overlay, int(det_id), color)

        frame_count = base_rgb.shape[0]
        if frame_count <= 1:
            return target_overlay
        expanded_shape = (1 + (frame_count - 1) * 3, *base_rgb.shape[1:])
        expanded = np.empty(expanded_shape, dtype=np.uint8)
        expanded[0] = base_rgb[0]
        out_idx = 1
        for frame_idx in range(1, frame_count):
            expanded[out_idx] = source_overlay[frame_idx - 1]
            expanded[out_idx + 1] = target_overlay[frame_idx]
            expanded[out_idx + 2] = base_rgb[frame_idx]
            out_idx += 3
        return expanded

    def _tracking_kind_overlay_projected_frames(self, kind: str, base_layer) -> np.ndarray | None:
        result = self._tracking_result
        if result is None or base_layer is None or not hasattr(base_layer, "data"):
            return None
        base_data = base_layer.data
        base_shape = tuple(int(v) for v in getattr(base_data, "shape", ()))
        if base_shape != tuple(result.frame_labels.shape) or len(base_shape) != 4:
            return None

        det_by_id = {int(det.id): det for det in result.detections}
        source_events_by_frame: dict[int, list[tuple[int, int]]] = {}
        target_events_by_frame: dict[int, list[tuple[int, int]]] = {}

        def _event_color(event_index: int) -> np.ndarray:
            base_hues = {
                "linear": 0.72,
                "fission": 0.03,
                "fusion": 0.53,
                "split-merge": 0.14,
            }
            hue = (base_hues.get(str(kind), 0.0) + event_index * 0.61803398875) % 1.0
            sat = 0.72
            val = 0.95
            rgb = colorsys.hsv_to_rgb(hue, sat, val)
            return np.asarray([float(c) for c in rgb], dtype=float)

        event_colors: dict[int, np.ndarray] = {}
        selected_events = []
        for event in result.events:
            is_linear = str(kind) == "linear" and str(event.kind) in {"continuation", "elongation", "shortening"}
            if str(event.kind) == str(kind) or is_linear:
                selected_events.append(event)

        for event_index, event in enumerate(selected_events):
            value = event_index + 1
            rgb = _event_color(event_index)
            event_colors[value] = np.rint(rgb * 255.0).astype(np.uint8)
            for det_id in event.sources:
                det = det_by_id.get(int(det_id))
                if det is not None:
                    source_events_by_frame.setdefault(int(det.frame), []).append(
                        (int(det_id), value)
                    )
            for det_id in event.targets:
                det = det_by_id.get(int(det_id))
                if det is not None:
                    target_events_by_frame.setdefault(int(det.frame), []).append(
                        (int(det_id), value)
                    )

        def _event_channels_for_frame(
            mode: str,
            frame: int,
        ) -> np.ndarray | None:
            events_by_frame = (
                source_events_by_frame if mode == "source" else target_events_by_frame
            )
            entries = events_by_frame.get(int(frame), [])
            if not entries:
                return None
            channels = np.zeros((3, *base_shape[1:]), dtype=np.uint8)
            for det_id, value in entries:
                det = det_by_id.get(int(det_id))
                if det is None or int(det.frame) != int(frame):
                    continue
                coords = np.asarray(det.coords, dtype=int)
                if coords.ndim != 2 or coords.size == 0:
                    continue
                zz = coords[:, 0]
                yy = coords[:, 1]
                xx = coords[:, 2]
                channels[:, zz, yy, xx] = event_colors[int(value)][:, None]
            return channels

        frame_count = int(base_shape[0])
        if frame_count <= 0:
            return None

        def _volume_zyx(data: np.ndarray, *, dtype) -> np.ndarray:
            volume = np.asarray(data, dtype=dtype)
            volume = np.squeeze(volume)
            if volume.ndim == 2:
                volume = volume[np.newaxis, :, :]
            if volume.ndim != 3:
                raise ValueError(
                    f"Tracking export needs a ZYX volume, got shape {volume.shape}."
                )
            return np.ascontiguousarray(volume)

        def _contrast_limits(data: np.ndarray) -> tuple[float, float]:
            values = np.asarray(data, dtype=np.float32)
            values = values[np.isfinite(values)]
            if values.size == 0:
                return (0.0, 1.0)
            low = float(np.min(values))
            high = float(np.max(values))
            return (low, high if high > low else low + 1.0)

        def _event_channel_colormap(channel: int):
            from napari.utils import Colormap

            color = np.zeros(4, dtype=float)
            color[int(channel)] = 1.0
            color[3] = 1.0
            return Colormap(
                colors=np.asarray(
                    [
                        [0.0, 0.0, 0.0, 0.0],
                        color,
                    ],
                    dtype=float,
                ),
                name=f"tracking_event_channel_{int(channel)}",
            )

        def _refresh_viewer(rounds: int = 3) -> None:
            canvas = getattr(
                getattr(self.viewer.window, "_qt_viewer", None),
                "canvas",
                None,
            )
            if canvas is not None:
                update = getattr(canvas, "update", None)
                if callable(update):
                    update()
                native = getattr(canvas, "native", None)
                if native is not None:
                    native_update = getattr(native, "update", None)
                    if callable(native_update):
                        native_update()
                    repaint = getattr(native, "repaint", None)
                    if callable(repaint):
                        repaint()
            for _ in range(max(int(rounds), 1)):
                QApplication.processEvents()

        sequence: list[tuple[str, int]] = [("base", 0)]
        for frame_idx in range(1, frame_count):
            sequence.append(("source", frame_idx - 1))
            sequence.append(("target", frame_idx))
            sequence.append(("base", frame_idx))

        saved_ndisplay = int(self.viewer.dims.ndisplay)
        saved_step = tuple(self.viewer.dims.current_step)
        saved_camera = {
            "angles": tuple(getattr(self.viewer.camera, "angles", (0.0, 0.0, 90.0))),
            "center": tuple(getattr(self.viewer.camera, "center", (0.0, 0.0, 0.0))),
            "zoom": float(getattr(self.viewer.camera, "zoom", 1.0)),
            "perspective": float(getattr(self.viewer.camera, "perspective", 0.0)),
        }
        saved_scale_bar_visible = bool(getattr(self.viewer.scale_bar, "visible", False))
        layer_states = [(layer, bool(getattr(layer, "visible", True))) for layer in list(self.viewer.layers)]
        temp_layers = []
        event_layers = []
        captured: list[np.ndarray] = []

        try:
            for layer, _visible in layer_states:
                with suppress(Exception):
                    layer.visible = False

            self.viewer.dims.ndisplay = 3
            with suppress(Exception):
                self.viewer.scale_bar.visible = False
            _refresh_viewer(4)

            render_scale = (1.0, 1.0, 1.0)
            first_base = _volume_zyx(base_data[0], dtype=np.float32)
            temp_base = self.viewer.add_image(
                first_base,
                name=self._next_layer_name("tracking export base"),
                colormap="gray",
                contrast_limits=_contrast_limits(first_base),
                opacity=1.0,
                gamma=0.7,
                interpolation3d="cubic",
                rendering="attenuated_mip",
                blending="additive",
                scale=render_scale,
                **_spatial_unit_kwargs(base_layer, 3),
            )
            if hasattr(temp_base, "attenuation"):
                temp_base.attenuation = 0.25
            _apply_tracking_overlay_rendering(temp_base, base_layer)
            temp_layers.append(temp_base)

            def _show_event_volume(channels_czyx: np.ndarray | None) -> None:
                if channels_czyx is None:
                    for event_layer in event_layers:
                        event_layer.visible = False
                    return
                if not event_layers:
                    for channel in range(3):
                        event_layer = self.viewer.add_image(
                            channels_czyx[channel],
                            name=self._next_layer_name(
                                f"tracking export event channel {channel}"
                            ),
                            colormap=_event_channel_colormap(channel),
                            contrast_limits=(0, 255),
                            opacity=0.72,
                            gamma=1.0,
                            interpolation3d="nearest",
                            rendering="attenuated_mip",
                            blending="additive",
                            scale=render_scale,
                            **_spatial_unit_kwargs(base_layer, 3),
                        )
                        if hasattr(event_layer, "attenuation"):
                            event_layer.attenuation = 0.25
                        event_layers.append(event_layer)
                        temp_layers.append(event_layer)
                else:
                    for channel, event_layer in enumerate(event_layers):
                        event_layer.data = channels_czyx[channel]
                for event_layer in event_layers:
                    event_layer.opacity = 0.72
                    if hasattr(event_layer, "attenuation"):
                        event_layer.attenuation = 0.25
                    event_layer.visible = True

            def _force_hide_scale_bar() -> None:
                with suppress(Exception):
                    self.viewer.scale_bar.visible = False
                canvas = getattr(
                    getattr(self.viewer.window, "_qt_viewer", None),
                    "canvas",
                    None,
                )
                overlay_visuals = getattr(canvas, "_overlay_to_visual", {})
                with suppress(Exception):
                    for visual in overlay_visuals.get(
                        self.viewer.scale_bar,
                        (),
                    ):
                        visual.node.visible = False

            def _capture_fitted_view() -> np.ndarray:
                target_size = (int(base_shape[-2]), int(base_shape[-1]))
                qt_viewer = getattr(self.viewer.window, "_qt_viewer", None)
                _force_hide_scale_bar()
                self.viewer.camera.angles = (0.0, 0.0, 0.0)
                self.viewer.camera.perspective = 0.0
                spatial_shape = np.asarray(base_shape[-3:], dtype=float)
                self.viewer.camera.center = tuple(
                    float(value) for value in (spatial_shape - 1.0) / 2.0
                )
                device_pixel_ratio = 1.0
                if qt_viewer is not None:
                    with suppress(Exception):
                        device_pixel_ratio = max(
                            float(qt_viewer.devicePixelRatioF()),
                            1.0,
                        )
                logical_size = (
                    np.asarray(target_size, dtype=float) / device_pixel_ratio
                )
                display_size = np.maximum(spatial_shape[-2:], 1.0)
                self.viewer.camera.zoom = float(
                    np.min(logical_size / display_size)
                )
                _refresh_viewer()
                if qt_viewer is not None:
                    image = np.asarray(
                        qt_viewer.screenshot(
                            flash=False,
                            size=target_size,
                            scale=1.0,
                            fit_to_data_extent=False,
                        ),
                        dtype=np.uint8,
                    )
                else:
                    image = np.asarray(
                        self.viewer.screenshot(
                            canvas_only=True,
                            size=target_size,
                            scale=1.0,
                            flash=False,
                        ),
                        dtype=np.uint8,
                    )
                if tuple(int(value) for value in image.shape[:2]) != target_size:
                    raise RuntimeError(
                        "Tracking export rendered an unexpected frame size: "
                        f"{image.shape[:2]}, expected {target_size}."
                    )
                return image

            seed_channels = None
            if source_events_by_frame:
                seed_frame = min(source_events_by_frame)
                seed_channels = _event_channels_for_frame(
                    "source",
                    seed_frame,
                )
            elif target_events_by_frame:
                seed_frame = min(target_events_by_frame)
                seed_channels = _event_channels_for_frame(
                    "target",
                    seed_frame,
                )
            if seed_channels is not None:
                _show_event_volume(seed_channels)
                _capture_fitted_view()
            _show_event_volume(None)
            _capture_fitted_view()

            current_base_frame = 0
            for mode, frame_idx in sequence:
                if int(frame_idx) != current_base_frame:
                    base_volume = _volume_zyx(
                        base_data[frame_idx],
                        dtype=np.float32,
                    )
                    temp_base.data = base_volume
                    temp_base.contrast_limits = _contrast_limits(base_volume)
                    current_base_frame = int(frame_idx)
                if mode == "source":
                    _show_event_volume(
                        _event_channels_for_frame("source", frame_idx)
                    )
                elif mode == "target":
                    _show_event_volume(
                        _event_channels_for_frame("target", frame_idx)
                    )
                else:
                    _show_event_volume(None)
                image = _capture_fitted_view()
                captured.append(np.asarray(image, dtype=np.uint8)[..., :3])
        finally:
            for layer in temp_layers:
                with suppress(Exception):
                    self.viewer.layers.remove(layer)
            for layer, visible in layer_states:
                with suppress(Exception):
                    layer.visible = visible
            with suppress(Exception):
                self.viewer.dims.ndisplay = saved_ndisplay
            for axis, step in enumerate(saved_step):
                with suppress(Exception):
                    self.viewer.dims.set_current_step(axis, step)
            with suppress(Exception):
                self.viewer.camera.angles = saved_camera["angles"]
                self.viewer.camera.center = saved_camera["center"]
                self.viewer.camera.zoom = saved_camera["zoom"]
                self.viewer.camera.perspective = saved_camera["perspective"]
            with suppress(Exception):
                self.viewer.scale_bar.visible = saved_scale_bar_visible
            _refresh_viewer(4)

        if not captured:
            return None
        return np.stack(captured, axis=0)

    def _write_tracking_event_kind(
        self,
        kind: str,
        overlay_layer,
        format_name: str,
        fps: float,
        target_dir: str,
    ) -> str | None:
        suffix = {
            "GIF": ".gif",
            "MP4": ".mp4",
        }.get(str(format_name).upper())
        if suffix is None:
            raise ValueError(
                f"Unsupported tracking export format: {format_name!r}"
            )
        overlay_ndim = int(getattr(getattr(overlay_layer, "data", None), "ndim", 0))
        is_projected_3d = overlay_ndim == 4
        if is_projected_3d:
            overlay = self._tracking_kind_overlay_projected_frames(kind, overlay_layer)
        else:
            overlay = self._tracking_kind_overlay_data(kind, overlay_layer)
        if overlay is None:
            return None
        if overlay.ndim != 4:
            raise ValueError(
                f"{format_name} export currently supports 2D time series or projected 3D time series only.",
            )
        overlay_data = getattr(overlay_layer, "data", None)
        overlay_data_ndim = int(getattr(overlay_data, "ndim", 0))
        base_scale = tuple(
            float(v)
            for v in getattr(
                overlay_layer,
                "scale",
                (1,) * overlay_data_ndim,
            )
        )
        overlay_scale = (base_scale[0], base_scale[-2], base_scale[-1], 1.0)
        overlay_meta = {
            "scale": overlay_scale,
            "metadata": {
                "dims_out": "TYXC",
                "source": _get_source_from_layer(overlay_layer),
                "unit": _get_units_from_layer(overlay_layer),
                "layer_type": "image",
                "fps": fps,
                "projection_from_3d": bool(is_projected_3d),
            },
            "layer_type": "image",
            "fps": fps,
        }
        output_path = os.path.join(
            target_dir,
            f"tracking_{kind.replace('-', '_')}_overlay{suffix}",
        )
        write_single_image(
            output_path,
            overlay,
            overlay_meta,
        )
        return output_path

    def _export_tracking_events(self) -> None:
        result = self._tracking_result
        if result is None or not self._tracking_display_events:
            QMessageBox.information(
                self,
                "No events",
                "Run tracking first to export events.",
            )
            return
        settings = self._prompt_tracking_export_settings()
        if settings is None:
            return
        overlay_layer, format_name, fps, selected_kinds = settings
        target_dir = QFileDialog.getExistingDirectory(
            self,
            "Export Tracking Events",
        )
        if not target_dir:
            return

        groups = self._tracking_event_groups()
        kinds = [
            kind
            for kind in selected_kinds
            if groups.get(kind)
        ]
        exported_paths = []
        errors = []
        self.track_progress.setRange(0, max(len(kinds), 1))
        self.track_progress.setValue(0)
        for index, kind in enumerate(kinds, start=1):
            self.track_progress.setFormat(f"Exporting {kind}")
            QApplication.processEvents()
            try:
                output_path = self._write_tracking_event_kind(
                    kind,
                    overlay_layer,
                    format_name,
                    fps,
                    target_dir,
                )
                if output_path is not None:
                    exported_paths.append(output_path)
            except (OSError, RuntimeError, TypeError, ValueError) as error:
                errors.append(f"{kind}: {error}")
            self.track_progress.setValue(index)

        if errors:
            self.track_progress.setFormat("Export incomplete")
            QMessageBox.warning(
                self,
                "Export incomplete",
                "\n".join(errors),
            )
            return
        self.track_progress.setFormat("Exported")
        QMessageBox.information(
            self,
            "Export complete",
            (
                f"Exported {len(exported_paths)} {format_name} event "
                f"overlay file(s) to:\n{target_dir}"
            ),
        )

    def _export_tracking_statistics(self) -> None:
        if self._tracking_result is None or not self._tracking_display_events:
            QMessageBox.information(
                self,
                "No events",
                "Run tracking first to export event statistics.",
            )
            return
        settings = self._prompt_tracking_statistics_export_settings()
        if settings is None:
            return
        selected_kinds, frame_from, frame_to, format_name = settings
        events = self._tracking_statistics_events(
            self._tracking_event_groups(),
            selected_kinds,
            frame_from,
            frame_to,
        )
        if not events:
            QMessageBox.information(
                self,
                "No events",
                "No selected events fall completely within the chosen frame range.",
            )
            return

        format_suffix = {
            "CSV": ".csv",
            "TXT": ".txt",
            "XLSX": ".xlsx",
        }
        suffix = format_suffix[format_name]
        layer = self._selected_tracking_layer()
        layer_name = str(getattr(layer, "name", "tracking"))
        layer_stem = os.path.splitext(os.path.basename(layer_name))[0] or "tracking"
        default_name = f"{layer_stem}_event_statistics{suffix}"
        format_filter = {
            "CSV": "CSV (*.csv)",
            "TXT": "Text (*.txt)",
            "XLSX": "Excel Workbook (*.xlsx)",
        }[format_name]
        path, _selected_filter = QFileDialog.getSaveFileName(
            self,
            "Export Event Statistics",
            default_name,
            format_filter,
        )
        if not path:
            return
        base, current_suffix = os.path.splitext(path)
        if current_suffix.lower() != suffix:
            path = (
                f"{base}{suffix}"
                if current_suffix.lower() in {".csv", ".txt", ".xlsx"}
                else f"{path}{suffix}"
            )

        step_columns = self._tracking_event_step_columns(selected_kinds)
        headers = self._tracking_event_export_headers(
            step_columns=step_columns,
            include_links=True,
        )
        rows = [
            self._tracking_event_export_row(
                event,
                step_columns=step_columns,
                include_links=True,
            )
            for event in events
        ]
        try:
            write_analysis_export_file(
                path,
                headers,
                rows,
                sheet_title="tracking events",
            )
        except RuntimeError as error:
            if "openpyxl" in str(error).lower():
                QMessageBox.warning(
                    self,
                    "XLSX unavailable",
                    (
                        "openpyxl is not available in this environment. "
                        "Please export as CSV/TXT or install openpyxl."
                    ),
                )
                return
            QMessageBox.critical(self, "Export failed", str(error))
            return
        except (OSError, TypeError, ValueError) as error:
            QMessageBox.critical(self, "Export failed", str(error))
            return
        QMessageBox.information(
            self,
            "Export complete",
            f"Saved {len(rows)} event row(s) to:\n{path}",
        )

    @staticmethod
    def _tracking_table_header(value: Any) -> str:
        return " ".join(
            str(value or "").strip().lower().replace("_", " ").split()
        )

    @staticmethod
    def _tracking_table_int(value: Any, *, field: str, row_number: int) -> int:
        try:
            parsed = float(value)
        except (TypeError, ValueError) as error:
            raise ValueError(
                f"Row {row_number}: {field} must be an integer, got {value!r}."
            ) from error
        if not np.isfinite(parsed) or not parsed.is_integer():
            raise ValueError(
                f"Row {row_number}: {field} must be an integer, got {value!r}."
            )
        return int(parsed)

    @classmethod
    def _tracking_table_id_list(
        cls,
        value: Any,
        *,
        field: str,
        row_number: int,
    ) -> tuple[int, ...]:
        if value is None or str(value).strip() == "":
            return ()
        tokens = str(value).replace(";", ",").split(",")
        return tuple(
            cls._tracking_table_int(
                token.strip(),
                field=field,
                row_number=row_number,
            )
            for token in tokens
            if token.strip()
        )

    @classmethod
    def _tracking_table_link_pairs(
        cls,
        value: Any,
        *,
        row_number: int,
    ) -> tuple[tuple[int, int], ...]:
        if value is None or str(value).strip() == "":
            return ()
        pairs = []
        for token in str(value).replace(";", ",").split(","):
            token = token.strip().replace("→", "->")
            if not token:
                continue
            parts = token.split("->")
            if len(parts) != 2:
                raise ValueError(
                    f"Row {row_number}: invalid link pair {token!r}; use src->dst."
                )
            pairs.append(
                (
                    cls._tracking_table_int(
                        parts[0].strip(),
                        field="link pair source",
                        row_number=row_number,
                    ),
                    cls._tracking_table_int(
                        parts[1].strip(),
                        field="link pair target",
                        row_number=row_number,
                    ),
                )
            )
        return tuple(pairs)

    @staticmethod
    def _read_tracking_event_table(path: str) -> tuple[list[Any], list[list[Any]]]:
        suffix = os.path.splitext(path)[1].lower()
        if suffix in {".csv", ".txt"}:
            delimiter = "," if suffix == ".csv" else "\t"
            with open(path, encoding="utf-8-sig", newline="") as handle:
                raw_rows = list(csv.reader(handle, delimiter=delimiter))
        elif suffix == ".xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError as error:
                raise RuntimeError(
                    "openpyxl is required to import XLSX event tables."
                ) from error
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                raw_rows = [list(row) for row in workbook.active.iter_rows(values_only=True)]
            finally:
                workbook.close()
        else:
            raise ValueError("Event tables must be CSV, TXT, or XLSX files.")
        if not raw_rows:
            raise ValueError("The event table is empty.")
        headers = list(raw_rows[0])
        rows = [
            list(row)
            for row in raw_rows[1:]
            if any(value is not None and str(value).strip() for value in row)
        ]
        if not rows:
            raise ValueError("The event table has no event rows.")
        return headers, rows

    @classmethod
    def _tracking_selection_from_event_table(
        cls,
        result: TrackingResult,
        headers: list[Any],
        rows: list[list[Any]],
    ) -> tuple[tuple[int, ...], set[tuple[int, int]]]:
        normalized_headers = [cls._tracking_table_header(value) for value in headers]
        header_index = {
            name: index for index, name in enumerate(normalized_headers) if name
        }
        required = {"from", "to", "sources", "targets"}
        missing = sorted(required - set(header_index))
        if missing:
            raise ValueError(
                "Event table is missing required column(s): " + ", ".join(missing)
            )

        det_by_id = {int(det.id): det for det in result.detections}
        link_by_id = {int(link.id): link for link in result.links}
        link_by_pair = {
            (int(link.src), int(link.dst)): link for link in result.links
        }

        def cell(row: list[Any], name: str) -> Any:
            index = header_index.get(name)
            return row[index] if index is not None and index < len(row) else None

        imported_link_ids: set[int] = set()
        transitions: set[tuple[int, int]] = set()
        for row_number, row in enumerate(rows, start=2):
            frame_from = cls._tracking_table_int(
                cell(row, "from"), field="from", row_number=row_number
            )
            frame_to = cls._tracking_table_int(
                cell(row, "to"), field="to", row_number=row_number
            )
            sources = set(
                cls._tracking_table_id_list(
                    cell(row, "sources"),
                    field="sources",
                    row_number=row_number,
                )
            )
            targets = set(
                cls._tracking_table_id_list(
                    cell(row, "targets"),
                    field="targets",
                    row_number=row_number,
                )
            )
            if not sources or not targets:
                raise ValueError(
                    f"Row {row_number}: imported tracking events require sources and targets."
                )
            for det_id in sources:
                detection = det_by_id.get(det_id)
                if detection is None or int(detection.frame) != frame_from:
                    raise ValueError(
                        f"Row {row_number}: source detection {det_id} is not in frame {frame_from}."
                    )
            for det_id in targets:
                detection = det_by_id.get(det_id)
                if detection is None or int(detection.frame) != frame_to:
                    raise ValueError(
                        f"Row {row_number}: target detection {det_id} is not in frame {frame_to}."
                    )

            row_link_ids: set[int] = set()
            pairs = cls._tracking_table_link_pairs(
                cell(row, "link pairs"),
                row_number=row_number,
            )
            if pairs:
                for pair in pairs:
                    if pair[0] not in sources or pair[1] not in targets:
                        raise ValueError(
                            f"Row {row_number}: link pair {pair[0]}->{pair[1]} "
                            "is outside the row's sources/targets."
                        )
                    link = link_by_pair.get(pair)
                    if link is None:
                        raise ValueError(
                            f"Row {row_number}: candidate link {pair[0]}->{pair[1]} "
                            "is unavailable in the current tracking result."
                        )
                    row_link_ids.add(int(link.id))
            else:
                listed_ids = cls._tracking_table_id_list(
                    cell(row, "link ids"),
                    field="link ids",
                    row_number=row_number,
                )
                for link_id in listed_ids:
                    link = link_by_id.get(link_id)
                    if link is None:
                        raise ValueError(
                            f"Row {row_number}: candidate link id {link_id} is unavailable."
                        )
                    if int(link.src) not in sources or int(link.dst) not in targets:
                        raise ValueError(
                            f"Row {row_number}: link id {link_id} is outside the row's sources/targets."
                        )
                    row_link_ids.add(link_id)

            if not row_link_ids:
                candidates = [
                    link
                    for link in result.links
                    if int(link.src) in sources and int(link.dst) in targets
                ]
                kind = str(cell(row, "kind") or "").strip().lower()
                expected_count = None
                if len(sources) == 1 and len(targets) == 1:
                    expected_count = 1
                elif kind == "fission":
                    expected_count = len(targets)
                elif kind == "fusion":
                    expected_count = len(sources)
                elif kind == "split-merge":
                    fission_steps = cell(row, "fission steps")
                    fusion_steps = cell(row, "fusion steps")
                    if str(fission_steps or "").strip():
                        expected_count = len(sources) + cls._tracking_table_int(
                            fission_steps,
                            field="fission steps",
                            row_number=row_number,
                        )
                    if str(fusion_steps or "").strip():
                        target_count = len(targets) + cls._tracking_table_int(
                            fusion_steps,
                            field="fusion steps",
                            row_number=row_number,
                        )
                        if expected_count is not None and target_count != expected_count:
                            raise ValueError(
                                f"Row {row_number}: fission/fusion step counts disagree."
                            )
                        expected_count = target_count
                if expected_count is None or len(candidates) != expected_count:
                    raise ValueError(
                        f"Row {row_number}: links cannot be inferred uniquely. "
                        "Re-export the table with link pairs."
                    )
                row_link_ids.update(int(link.id) for link in candidates)

            imported_link_ids.update(row_link_ids)
            transitions.add((frame_from, frame_to))

        current_ids = {int(value) for value in result.selected_link_ids}
        for link_id in tuple(current_ids):
            link = link_by_id.get(link_id)
            if link is None:
                continue
            source = det_by_id[int(link.src)]
            target = det_by_id[int(link.dst)]
            if (int(source.frame), int(target.frame)) in transitions:
                current_ids.remove(link_id)
        current_ids.update(imported_link_ids)
        return tuple(sorted(current_ids)), transitions

    @classmethod
    def _tracking_link_pairs_from_event_table(
        cls,
        headers: list[Any],
        rows: list[list[Any]],
    ) -> tuple[tuple[int, int], ...]:
        normalized_headers = [
            cls._tracking_table_header(value)
            for value in headers
        ]
        try:
            pair_column = normalized_headers.index("link pairs")
        except ValueError:
            return ()
        pairs: set[tuple[int, int]] = set()
        for row_number, row in enumerate(rows, start=2):
            value = row[pair_column] if pair_column < len(row) else None
            pairs.update(
                cls._tracking_table_link_pairs(
                    value,
                    row_number=row_number,
                )
            )
        return tuple(sorted(pairs))

    def _ensure_tracking_event_table_candidates(
        self,
        result: TrackingResult,
        headers: list[Any],
        rows: list[list[Any]],
    ) -> tuple[TrackingResult, int]:
        pairs = self._tracking_link_pairs_from_event_table(headers, rows)
        if not pairs:
            return result, 0
        source_layer = (
            self._selected_tracking_layer()
            or self._tracking_result_source_layer
        )
        spacing = (
            self._tracking_spacing_from_layer(source_layer)
            if source_layer is not None
            else None
        )
        config = self._tracking_config or self._build_tracking_config()
        prepared = result
        created_count = 0
        for source_id, target_id in pairs:
            prepared, _link_id, created = (
                ensure_manual_tracking_link_candidate(
                    prepared,
                    config,
                    source_id,
                    target_id,
                    spacing=spacing,
                )
            )
            created_count += int(created)
        return prepared, created_count

    def _import_tracking_statistics(self) -> None:
        if self._thread_is_running("_tracking_thread"):
            return
        path, _selected_filter = QFileDialog.getOpenFileName(
            self,
            "Import Event Table",
            "",
            "Event Tables (*.csv *.txt *.xlsx)",
        )
        if not path:
            return
        try:
            headers, rows = self._read_tracking_event_table(path)
        except RuntimeError as error:
            QMessageBox.warning(self, "Import unavailable", str(error))
            return
        except (OSError, TypeError, ValueError, KeyError) as error:
            QMessageBox.critical(self, "Import failed", str(error))
            return
        layer = self._selected_tracking_layer()
        if layer is None or not hasattr(layer, "data"):
            QMessageBox.information(
                self,
                "No segmentation",
                "Choose a TYX or TZYX segmentation layer before importing events.",
            )
            return
        result = self._tracking_result
        if (
            result is None
            or not result.detections
            or self._tracking_result_source_layer is not layer
        ):
            self._tracking_pending_event_import = (path, headers, rows)
            self.track_progress.setRange(0, 0)
            self.track_progress.setFormat("Preparing event import…")
            self._on_run_tracking_clicked()
            if not self._thread_is_running("_tracking_thread"):
                self._tracking_pending_event_import = None
            return
        self._apply_tracking_statistics_import(path, headers, rows)

    def _apply_tracking_statistics_import(
        self,
        path: str,
        headers: list[Any],
        rows: list[list[Any]],
    ) -> None:
        result = self._tracking_result
        if result is None or not result.detections:
            QMessageBox.critical(
                self,
                "Import failed",
                "No detections were available after preparing the tracking result.",
            )
            return
        try:
            prepared_result, restored_candidate_count = (
                self._ensure_tracking_event_table_candidates(
                    result,
                    headers,
                    rows,
                )
            )
            selected_link_ids, transitions = self._tracking_selection_from_event_table(
                prepared_result,
                headers,
                rows,
            )
        except (TypeError, ValueError, KeyError) as error:
            QMessageBox.critical(self, "Import failed", str(error))
            return
        if selected_link_ids == tuple(result.selected_link_ids):
            QMessageBox.information(
                self,
                "No changes",
                "The imported event table matches the current effective links.",
            )
            return
        answer = QMessageBox.question(
            self,
            "Import Event Table",
            (
                f"Replace effective links in {len(transitions)} transition(s) "
                f"using {len(rows)} event row(s)?"
                + (
                    f"\nRecreate {restored_candidate_count} manual candidate "
                    "link(s) that were absent from automatic tracking."
                    if restored_candidate_count
                    else ""
                )
                + "\n\n"
                "Other transitions will remain unchanged. This import can be undone."
            ),
            QMessageBox.Yes | QMessageBox.Cancel,
            QMessageBox.Cancel,
        )
        if answer != QMessageBox.Yes:
            return
        try:
            restored = restore_tracking_link_selection(
                prepared_result,
                self._tracking_config or self._build_tracking_config(),
                selected_link_ids,
                message=f"Imported tracking event table: {os.path.basename(path)}.",
                rebuild_lineages=False,
            )
        except (TypeError, ValueError, KeyError) as error:
            QMessageBox.critical(self, "Import failed", str(error))
            return
        self._record_tracking_refine_undo(result)
        det_id = self._tracking_refine_det_id
        if det_id is None:
            det_id = int(result.detections[0].id)
        changed_transition = next(iter(transitions)) if len(transitions) == 1 else None
        self._update_tracking_result_after_refine(
            restored,
            int(det_id),
            preferred_event_scope=self._tracking_refine_event_scope,
            changed_transition=changed_transition,
        )
        self.track_progress.setFormat("Imported events")
        QMessageBox.information(
            self,
            "Import complete",
            (
                f"Imported {len(transitions)} transition(s). Tracking now "
                f"contains {len(selected_link_ids)} effective link(s)."
            ),
        )

    def _show_tracking_event_highlight(self, event, *, jump_target: str = "both") -> None:
        result = self._tracking_result
        if result is None:
            return
        self._clear_tracking_match_preview()
        group_kind = self._tracking_event_group_kind(str(event.kind))
        if group_kind is not None:
            self._set_tracking_active_event_kind(group_kind)
        self._set_tracking_active_transition(
            int(event.frame_from),
            int(event.frame_to),
        )
        self._select_tracking_event_in_tables(event)
        event_layer = (
            self._activate_tracking_event_layer(group_kind)
            if group_kind is not None
            else None
        )
        layer = event_layer or self._selected_tracking_layer()
        if layer is None:
            return
        det_ids = list(
            dict.fromkeys(
                int(det_id)
                for det_id in (*event.sources, *event.targets)
            )
        )
        value_by_det = {
            det_id: value
            for value, det_id in enumerate(det_ids, start=1)
        }
        value_colors = {
            value: self._tracking_object_rgba(det_id)
            for det_id, value in value_by_det.items()
        }
        self._apply_tracking_highlight_colors(
            value_by_det,
            value_colors,
            opacity=0.95,
            base_layer=event_layer,
        )
        _safe_axis_labels(self.viewer, layer)
        frame_from = int(event.frame_from)
        frame_to = int(event.frame_to)
        frame = frame_to if jump_target == "to" else frame_from
        with suppress(Exception):
            if self.t_spin.isEnabled():
                self.t_spin.setValue(frame)
        if jump_target == "both" and frame_to != frame_from:
            self._apply_tracking_time_steps(layer, [frame_from, frame_to])
        else:
            self._apply_tracking_time_steps(layer, [frame])

    def _tracking_effective_object_neighbors(
        self,
        det_id: int,
    ) -> tuple[tuple[int, ...], tuple[int, ...]]:
        result = self._tracking_result
        if result is None:
            return (), ()
        selected_ids = {int(link_id) for link_id in result.selected_link_ids}
        selected_links = sorted(
            (
                link
                for link in result.links
                if int(link.id) in selected_ids
                and (
                    int(link.src) == int(det_id)
                    or int(link.dst) == int(det_id)
                )
            ),
            key=lambda link: int(link.id),
        )
        incoming = tuple(
            dict.fromkeys(
                int(link.src)
                for link in selected_links
                if int(link.dst) == int(det_id)
            )
        )
        outgoing = tuple(
            dict.fromkeys(
                int(link.dst)
                for link in selected_links
                if int(link.src) == int(det_id)
            )
        )
        return incoming, outgoing

    def _show_tracking_object_context(
        self,
        det_id: int,
        *,
        sync_events: bool = True,
    ) -> None:
        layer = self._tracking_layer
        result = self._tracking_result
        if layer is None or result is None:
            return
        self._clear_tracking_match_preview()
        det_by_id = {det.id: det for det in result.detections}
        det = det_by_id.get(int(det_id))
        if det is None:
            return
        if (
            self._tracking_refine_det_id != int(det_id)
            or self._tracking_refine_event_scope is not None
        ):
            self._set_tracking_refine_object(int(det_id))

        incoming, outgoing = self._tracking_effective_object_neighbors(
            int(det_id)
        )
        for event_layer in self._tracking_event_layers.values():
            with suppress(Exception):
                event_layer.visible = False
        value_by_det: dict[int, int] = {}
        value_colors: dict[int, np.ndarray] = {}
        for source_id in incoming:
            value = len(value_by_det) + 1
            value_by_det[int(source_id)] = value
            value_colors[value] = np.array(
                [1.0, 0.75, 0.15, 1.0],
                dtype=float,
            )
        for target_id in outgoing:
            value = len(value_by_det) + 1
            value_by_det[int(target_id)] = value
            value_colors[value] = np.array(
                [0.2, 0.95, 0.3, 1.0],
                dtype=float,
            )
        selected_value = len(value_by_det) + 1
        value_by_det[int(det_id)] = selected_value
        value_colors[selected_value] = np.array(
            [1.0, 0.15, 0.8, 1.0],
            dtype=float,
        )
        self._apply_tracking_highlight_colors(
            value_by_det,
            value_colors,
            opacity=0.95,
            base_layer=(
                self._tracking_event_base_layer or self._tracking_layer
            ),
            base_opacity=0.7,
        )
        _safe_axis_labels(self.viewer, layer)
        with suppress(Exception):
            if self.t_spin.isEnabled():
                self.t_spin.setValue(int(det.frame))
        self._apply_tracking_time_steps(layer, [int(det.frame)])
        self._select_tracking_unlinked_object_in_table(int(det_id))
        if sync_events:
            self._sync_tracking_event_tables_for_det(det_id)

    def _tracking_det_id_from_world_position(
        self,
        position,
        *,
        highlighted_only: bool = False,
    ) -> int | None:
        layer = self._tracking_layer
        result = self._tracking_result
        if layer is None or result is None:
            return None
        try:
            data_pos = np.asarray(layer.world_to_data(position), dtype=float).ravel()
        except (AttributeError, TypeError, ValueError):
            return None
        expected_ndim = int(result.frame_labels.ndim)
        if data_pos.size < expected_ndim:
            return None
        frame = int(round(float(data_pos[0])))
        spatial = tuple(
            int(round(float(v))) for v in data_pos[1 : 1 + (expected_ndim - 1)]
        )
        if frame < 0 or frame >= result.frame_labels.shape[0]:
            return None
        spatial_shape = tuple(int(v) for v in result.frame_labels.shape[1:])
        if any(
            coord < 0 or coord >= size
            for coord, size in zip(spatial, spatial_shape, strict=False)
        ):
            return None
        local_label = int(result.frame_labels[(frame, *spatial)])
        allowed_ids = (
            self._tracking_highlight_det_ids
            if highlighted_only
            else None
        )
        return self._tracking_det_id_for_frame_label(
            frame,
            local_label,
            allowed_ids=allowed_ids,
        )

    def _tracking_det_id_for_frame_label(
        self,
        frame: int,
        local_label: int,
        *,
        allowed_ids: set[int] | None = None,
    ) -> int | None:
        if int(local_label) <= 0:
            return None
        lookup = getattr(self, "_tracking_detection_lookup", None)
        if not lookup and self._tracking_result is not None:
            self._refresh_tracking_detection_lookup()
            lookup = self._tracking_detection_lookup
        det_id = (lookup or {}).get((int(frame), int(local_label)))
        if det_id is None:
            return None
        det_id = int(det_id)
        if allowed_ids is not None and det_id not in allowed_ids:
            return None
        return det_id

    @staticmethod
    def _tracking_scalar_layer_value(value) -> int | None:
        if isinstance(value, tuple):
            if not value:
                return None
            value = value[-1]
        if value is None:
            return None
        try:
            scalar = np.asarray(value)
            if scalar.size != 1:
                return None
            return int(scalar.reshape(-1)[0])
        except (TypeError, ValueError, OverflowError):
            return None

    def _tracking_allowed_det_ids_for_pick_layer(
        self,
        layer,
        event,
    ) -> set[int] | None:
        if layer is None or not bool(getattr(layer, "visible", False)):
            return set()
        metadata = dict(getattr(layer, "metadata", {}) or {})
        is_highlight = layer is self._tracking_highlight_layer
        is_event_layer = bool(metadata.get("is_tracking_event"))
        if not is_highlight and not is_event_layer:
            return None
        try:
            raw_value = layer.get_value(
                event.position,
                view_direction=event.view_direction,
                dims_displayed=event.dims_displayed,
                world=True,
            )
        except (AttributeError, IndexError, RuntimeError, TypeError, ValueError):
            return set()
        value = self._tracking_scalar_layer_value(raw_value)
        if value is None or value <= 0:
            return set()
        if is_highlight:
            return set(
                self._tracking_highlight_value_det_ids.get(int(value), set())
            )

        kind = self._tracking_event_group_kind(
            str(metadata.get("tracking_event_kind", ""))
        )
        events = (
            self._tracking_event_groups().get(kind, [])
            if kind is not None
            else []
        )
        event_index = int(value) - 1
        if event_index < 0 or event_index >= len(events):
            return set()
        selected_event = events[event_index]
        return {
            int(det_id)
            for det_id in (*selected_event.sources, *selected_event.targets)
        }

    def _tracking_det_id_from_ray(
        self,
        layer,
        event,
        *,
        allowed_ids: set[int] | None = None,
    ) -> int | None:
        result = self._tracking_result
        if layer is None or result is None:
            return None
        try:
            start, end = layer.get_ray_intersections(
                position=event.position,
                view_direction=event.view_direction,
                dims_displayed=list(event.dims_displayed),
                world=True,
            )
        except (AttributeError, IndexError, RuntimeError, TypeError, ValueError):
            return None
        if start is None or end is None:
            return None
        start = np.asarray(start, dtype=float).ravel()
        end = np.asarray(end, dtype=float).ravel()
        shape = tuple(int(v) for v in result.frame_labels.shape)
        if start.size != len(shape) or end.size != len(shape):
            return None

        delta = end - start
        sample_count = max(
            int(np.ceil(float(np.max(np.abs(delta))) * 2.0)) + 1,
            2,
        )
        points = np.linspace(start, end, sample_count)
        coords = np.rint(points).astype(np.intp)
        valid = np.ones(coords.shape[0], dtype=bool)
        for axis, size in enumerate(shape):
            valid &= (coords[:, axis] >= 0) & (coords[:, axis] < size)
        coords = coords[valid]
        if coords.size == 0:
            return None

        flat = np.ravel_multi_index(coords.T, shape)
        keep = np.concatenate(([True], flat[1:] != flat[:-1]))
        coords = coords[keep]
        labels = np.asarray(
            result.frame_labels[tuple(coords.T)],
            dtype=np.int64,
        ).reshape(-1)
        for coord, local_label in zip(coords, labels, strict=False):
            det_id = self._tracking_det_id_for_frame_label(
                int(coord[0]),
                int(local_label),
                allowed_ids=allowed_ids,
            )
            if det_id is not None:
                return int(det_id)
        return None

    def _tracking_det_id_near_world_position(
        self,
        layer,
        event,
        *,
        allowed_ids: set[int] | None = None,
        radius: int = 2,
    ) -> int | None:
        result = self._tracking_result
        if layer is None or result is None:
            return None
        try:
            center = np.rint(
                np.asarray(layer.world_to_data(event.position), dtype=float)
            ).astype(np.intp)
        except (AttributeError, TypeError, ValueError):
            return None
        shape = tuple(int(v) for v in result.frame_labels.shape)
        if center.size != len(shape):
            return None
        displayed = [
            int(axis)
            for axis in event.dims_displayed
            if 0 <= int(axis) < len(shape) and int(axis) != 0
        ]
        if not displayed:
            return None
        offsets = [
            offset
            for offset in np.ndindex(
                *((2 * int(radius) + 1,) * len(displayed))
            )
        ]
        offsets.sort(
            key=lambda offset: sum(
                (int(value) - int(radius)) ** 2 for value in offset
            )
        )
        for offset in offsets:
            coord = center.copy()
            for axis, value in zip(displayed, offset, strict=False):
                coord[axis] += int(value) - int(radius)
            if any(
                int(value) < 0 or int(value) >= size
                for value, size in zip(coord, shape, strict=False)
            ):
                continue
            local_label = int(result.frame_labels[tuple(coord)])
            det_id = self._tracking_det_id_for_frame_label(
                int(coord[0]),
                local_label,
                allowed_ids=allowed_ids,
            )
            if det_id is not None:
                return int(det_id)
        return None

    def _tracking_det_id_from_mouse_event(self, event) -> int | None:
        result = self._tracking_result
        if result is None:
            return None
        is_3d = len(tuple(event.dims_displayed)) == 3
        if not is_3d:
            exact = self._tracking_det_id_from_world_position(event.position)
            if exact is not None:
                return exact

        active_layer = getattr(self.viewer.layers.selection, "active", None)
        layers = [
            self._tracking_highlight_layer,
            self._active_tracking_event_layer(),
            active_layer,
            self._tracking_event_base_layer,
            self._tracking_layer,
        ]
        pick_layers = []
        expected_shape = tuple(int(v) for v in result.frame_labels.shape)
        for layer in layers:
            if (
                layer is None
                or any(layer is existing for existing in pick_layers)
                or not bool(getattr(layer, "visible", False))
            ):
                continue
            layer_shape = tuple(
                int(v)
                for v in getattr(getattr(layer, "data", None), "shape", ())
            )
            if layer_shape != expected_shape:
                continue
            pick_layers.append(layer)

        for layer in pick_layers:
            allowed_ids = self._tracking_allowed_det_ids_for_pick_layer(
                layer,
                event,
            )
            if allowed_ids == set():
                continue
            if (
                layer is self._tracking_highlight_layer
                and allowed_ids is not None
                and len(allowed_ids) == 1
            ):
                return int(next(iter(allowed_ids)))
            if is_3d:
                resolved = self._tracking_det_id_from_ray(
                    layer,
                    event,
                    allowed_ids=allowed_ids,
                )
            else:
                resolved = self._tracking_det_id_near_world_position(
                    layer,
                    event,
                    allowed_ids=allowed_ids,
                )
            if resolved is not None:
                return int(resolved)
        return None

    def _on_tracking_layer_double_click(self, _viewer, event) -> None:
        panel_tabs = getattr(self, "_panel_tabs", None)
        tracking_tab = getattr(self, "_tracking_tab", None)
        if (
            panel_tabs is not None
            and tracking_tab is not None
            and panel_tabs.currentWidget() is not tracking_tab
        ):
            return
        resolved = self._tracking_det_id_from_mouse_event(event)
        if resolved is None:
            return
        if self._tracking_refine_image_pick_direction is not None:
            self._set_tracking_refine_image_pick_candidate(int(resolved))
        else:
            self._show_tracking_object_context(int(resolved))
        highlight_layer = getattr(self, "_tracking_highlight_layer", None)
        active_event_layer = getattr(
            self,
            "_active_tracking_event_layer",
            lambda: None,
        )()
        interaction_layer = (
            highlight_layer
            or active_event_layer
            or getattr(self, "_tracking_layer", None)
        )
        if interaction_layer is not None:
            with suppress(Exception):
                interaction_layer.visible = True
                self.viewer.layers.selection.active = interaction_layer
        with suppress(Exception):
            event.handled = True

    def _on_tracking_event_selection_changed(self) -> None:
        table = self.sender() if isinstance(self.sender(), QTableWidget) else None
        if table is None:
            return
        item = table.currentItem()
        if item is None:
            return
        idx = item.data(Qt.UserRole)
        kind = item.data(Qt.UserRole + 1)
        if idx is None:
            return
        idx = int(idx)
        events = self._tracking_events_for_kind(str(kind))
        if idx < 0 or idx >= len(events):
            return
        self._schedule_tracking_interaction("event", (events[idx], "both"))

    def _on_tracking_event_row_clicked(self, row: int, _col: int) -> None:
        table = self.sender() if isinstance(self.sender(), QTableWidget) else None
        if table is None:
            return
        item = table.item(row, 0)
        if item is None:
            return
        idx = item.data(Qt.UserRole)
        kind = item.data(Qt.UserRole + 1)
        if idx is None:
            return
        idx = int(idx)
        events = self._tracking_events_for_kind(str(kind))
        if idx < 0 or idx >= len(events):
            return
        self._schedule_tracking_interaction("event", (events[idx], "both"))

    def _on_tracking_match_selection_changed(self) -> None:
        table = self._tracking_match_table
        summary = self._tracking_match_summary
        if table is None or summary is None:
            return
        items = table.selectedItems()
        if not items:
            return
        idx = items[0].data(Qt.UserRole)
        if idx is None:
            return
        idx = int(idx)
        if idx < 0 or idx >= len(summary.frame_details):
            return
        self._schedule_tracking_interaction("match", summary.frame_details[idx])

    def _add_tracking_result(
        self,
        layer,
        result: TrackingResult,
        *,
        config: TrackingConfig | None = None,
    ) -> None:
        self._cancel_tracking_refine_image_pick(restore_context=False)
        self._clear_tracking_highlight(remove_layer=True)
        self._remove_tracking_event_layers()
        self._clear_tracking_refine_undo_history()
        self._tracking_result = result
        self._tracking_result_source_layer = layer
        self._tracking_layer_pending_data = None
        self._tracking_lineages_dirty = False
        self._refresh_tracking_detection_lookup()
        self._tracking_config = config or self._build_tracking_config()
        md = dict(getattr(layer, "metadata", {}) or {})
        md["managed_by_plugin"] = True
        md["is_analysis_labels"] = True
        md["is_tracking"] = True
        md["is_tracking_index"] = True
        md["source"] = md.get("source", _get_source_from_layer(layer))
        md["group_id"] = _get_group_id(layer) or _ensure_group_id(md)
        dims_tag = _layer_dims_tag(layer)
        if dims_tag:
            md["dims"] = dims_tag
        tracked_name = self._next_layer_name(f"{layer.name} tracking index")
        scale = getattr(layer, "scale", None)
        add_kwargs = {"metadata": md}
        if scale is not None:
            add_kwargs["scale"] = scale
        self._tracking_layer = NapariLabels(
            result.tracked_labels.astype(np.int32, copy=False),
            name=tracked_name,
            **add_kwargs,
        )
        with suppress(Exception):
            color_map = self._tracking_lineage_color_map(result)
            self._tracking_base_color_map = dict(color_map)
            self._tracking_layer.color = color_map
            self._tracking_layer.opacity = 0.85
            self._tracking_layer.visible = False
            self._tracking_layer.events.visible.connect(
                self._on_tracking_layer_visibility_changed
            )
        if self._on_tracking_layer_double_click not in self.viewer.mouse_double_click_callbacks:
            self.viewer.mouse_double_click_callbacks.append(
                self._on_tracking_layer_double_click
            )
        self._populate_tracking_table(
            list(result.events)
        )
        self._rebuild_tracking_event_base_layer(
            layer,
            preferred_kind=self._tracking_active_event_kind,
        )
        initial_frame = self._tracking_current_viewer_frame(layer)
        self._set_tracking_initial_event_context(initial_frame)
        initial_kind = self._tracking_active_event_kind or "linear"
        self._rebuild_tracking_event_layers(
            layer,
            kinds={initial_kind},
        )
        self._populate_tracking_refine_objects()
        _safe_axis_labels(self.viewer, layer)
        visible_layer = (
            self._active_tracking_event_layer()
            or self._tracking_event_base_layer
            or layer
        )
        _apply_viewer_time_step(
            self.viewer,
            visible_layer,
            initial_frame,
        )
        self._update_info()

    def _prepare_tracking_inputs(self, *, title: str, action: str):
        layer = self._selected_tracking_layer()
        if layer is None or not hasattr(layer, "data"):
            QMessageBox.information(
                self,
                "No segmentation",
                "Please choose a TYX or TZYX segmentation layer first.",
            )
            return None
        raw_layer = self._selected_tracking_raw_layer()
        if raw_layer is None or not hasattr(raw_layer, "data"):
            QMessageBox.information(
                self,
                "No raw image",
                f"Please choose a raw image for intensity-weighted {action}.",
            )
            return None
        try:
            data = layer.data
            intensity_data = raw_layer.data
            if tuple(int(v) for v in intensity_data.shape) != tuple(
                int(v) for v in data.shape
            ):
                raise ValueError(
                    "Raw image and segmentation must have identical TYX/TZYX shapes."
                )
            spacing = self._tracking_spacing_from_layer(layer)
            config = self._build_tracking_config()
        except (TypeError, ValueError, AttributeError, RuntimeError) as e:
            QMessageBox.critical(
                self,
                f"{title} failed",
                f"Failed to prepare {action} input: {e!r}",
            )
            return None
        return layer, data, intensity_data, spacing, config

    def _on_run_tracking_clicked(self) -> None:
        if self._thread_is_running("_tracking_thread"):
            self._request_worker_cancel("_tracking_thread", "_tracking_worker")
            self.track_progress.setRange(0, 1)
            self.track_progress.setValue(0)
            self.track_progress.setFormat("Stopping tracking…")
            self._tracking_run_btn.setEnabled(False)
            self._tracking_match_btn.setEnabled(False)
            return
        prepared = self._prepare_tracking_inputs(title="Tracking", action="tracking")
        if prepared is None:
            return
        layer, data, intensity_data, spacing, config = prepared
        self._cancel_tracking_refine_image_pick(restore_context=False)
        self._clear_tracking_highlight(remove_layer=True)

        preparing_import = self._tracking_pending_event_import is not None
        self.track_progress.setRange(0, 0)
        self.track_progress.setFormat(
            "Preparing event import…" if preparing_import else "Running tracking…"
        )
        self._set_tracking_running_ui("tracking")

        self._tracking_thread = QThread()
        self._tracking_worker = TrackingWorker(
            data,
            intensity_data,
            spacing=spacing,
            config=config,
        )
        self._tracking_worker.moveToThread(self._tracking_thread)
        self._tracking_thread.started.connect(self._tracking_worker.run)

        def _finished(result, error):
            pending_import = self._tracking_pending_event_import
            self._tracking_pending_event_import = None
            apply_pending_import = False
            self.track_progress.setRange(0, 1)
            self._set_tracking_idle_ui()
            if isinstance(error, InterruptedError):
                self.track_progress.setValue(0)
                self.track_progress.setFormat(
                    "Import cancelled" if pending_import is not None else "Stopped"
                )
            elif error is not None:
                self.track_progress.setValue(0)
                self.track_progress.setFormat("Failed")
                QMessageBox.critical(
                    self,
                    "Import failed" if pending_import is not None else "Tracking failed",
                    (
                        f"Failed to prepare tracking candidates: {error!r}"
                        if pending_import is not None
                        else f"track_segmentations_ilp() error: {error!r}"
                    ),
                )
            else:
                self.track_progress.setValue(1)
                self.track_progress.setFormat("Finished")
                self._add_tracking_result(layer, result, config=config)
                apply_pending_import = pending_import is not None
            self._cleanup_worker_thread("_tracking_thread", "_tracking_worker")
            if apply_pending_import and pending_import is not None:
                self._apply_tracking_statistics_import(*pending_import)

        self._connect_worker_callback(self._tracking_worker.finished, _finished)
        self._tracking_thread.start()

    def _on_match_tracking_clicked(self) -> None:
        if self._thread_is_running("_tracking_thread"):
            self._request_worker_cancel("_tracking_thread", "_tracking_worker")
            self.match_progress.setRange(0, 1)
            self.match_progress.setValue(0)
            self.match_progress.setFormat("Stopping match…")
            self._tracking_match_label.setText("Match: stopping…")
            self._tracking_run_btn.setEnabled(False)
            self._tracking_match_btn.setEnabled(False)
            return
        prepared = self._prepare_tracking_inputs(title="Match", action="matching")
        if prepared is None:
            return
        _layer, data, intensity_data, spacing, config = prepared
        raw_layer = self._selected_tracking_raw_layer()

        self.match_progress.setRange(0, 0)
        self.match_progress.setFormat("Running match…")
        self._tracking_match_label.setText("Match: running…")
        self._set_tracking_running_ui("match")

        self._tracking_thread = QThread()
        self._tracking_worker = MatchWorker(
            data,
            intensity_data,
            spacing=spacing,
            config=config,
        )
        self._tracking_worker.moveToThread(self._tracking_thread)
        self._tracking_thread.started.connect(self._tracking_worker.run)

        def _finished(summary, error):
            self.match_progress.setRange(0, 1)
            self._set_tracking_idle_ui()
            if isinstance(error, InterruptedError):
                self.match_progress.setValue(0)
                self.match_progress.setFormat("Stopped")
                self._tracking_match_label.setText("Match: stopped by user")
            elif error is not None:
                self.match_progress.setValue(0)
                self.match_progress.setFormat("Failed")
                self._tracking_match_label.setText(f"Match: error: {error!r}")
                QMessageBox.critical(self, "Match failed", f"compute_match_summary() error: {error!r}")
            else:
                self.match_progress.setValue(1)
                self.match_progress.setFormat("Matched")
                self._tracking_match_summary = summary
                self._tracking_match_label.setText(self._format_match_summary(summary))
                self._populate_tracking_match_table(summary)
                if raw_layer is not None:
                    with suppress(Exception):
                        raw_layer.visible = False
            self._cleanup_worker_thread("_tracking_thread", "_tracking_worker")

        self._connect_worker_callback(self._tracking_worker.finished, _finished)
        self._tracking_thread.start()

    # -----------------------------------------------------------------
    # UI builders
    # -----------------------------------------------------------------

    def _build_top_box(self) -> QGroupBox:
        box = QGroupBox("Data / Axes / Device")
        g = QGridLayout()
        self._top_grid = g
        g.setHorizontalSpacing(self._scaled_px(10))
        g.setVerticalSpacing(self._scaled_px(6))

        self.open_btn = QPushButton("Open File")
        self.open_btn.clicked.connect(self._open_image_tc_zyx)
        self.device_combo = QComboBox()
        self._populate_devices()
        self.device_combo.currentTextChanged.connect(self._on_device_changed)
        self.device_label = QLabel("Device:")
        top_row = QWidget()
        top_row_layout = QHBoxLayout()
        top_row_layout.setContentsMargins(0, 0, 0, 0)
        top_row_layout.setSpacing(self._scaled_px(8))
        top_row_layout.addWidget(self.open_btn)
        top_row_layout.addStretch(1)
        top_row_layout.addWidget(self.device_label)
        top_row_layout.addWidget(self.device_combo, 1)
        top_row.setLayout(top_row_layout)
        g.addWidget(top_row, 0, 0, 1, 4)

        self.path_edit = QLineEdit()
        self.path_edit.setReadOnly(True)
        file_row = QWidget()
        file_row_layout = QHBoxLayout()
        file_row_layout.setContentsMargins(0, 0, 0, 0)
        file_row_layout.setSpacing(self._scaled_px(8))
        file_row_layout.addWidget(QLabel("File:"))
        file_row_layout.addWidget(self.path_edit, 1)
        file_row.setLayout(file_row_layout)
        g.addWidget(file_row, 1, 0, 1, 4)

        self.c_spin = QSpinBox()
        self.c_spin.setRange(0, 0)  # 0-based
        self.c_spin.setEnabled(False)
        self.c_spin.setValue(0)
        self.c_spin.valueChanged.connect(self._on_c_changed)

        self.t_spin = QSpinBox()
        self.t_spin.setRange(0, 0)  # 0-based
        self.t_spin.setEnabled(False)
        self.t_spin.setValue(0)
        self.t_spin.valueChanged.connect(self._on_t_changed)

        self.t_start_spin = QSpinBox()
        self.t_start_spin.setRange(0, 0)
        self.t_start_spin.setEnabled(False)
        self.t_start_spin.setValue(0)
        self.t_start_spin.valueChanged.connect(self._on_time_range_changed)

        self.t_end_spin = QSpinBox()
        self.t_end_spin.setRange(0, 0)
        self.t_end_spin.setEnabled(False)
        self.t_end_spin.setValue(0)
        self.t_end_spin.valueChanged.connect(self._on_time_range_changed)

        self.btn_apply_time_range = QPushButton("Apply Time")
        self.btn_apply_time_range.setEnabled(False)
        self.btn_apply_time_range.clicked.connect(self._apply_time_range_to_group)

        self.btn_apply_slice_range = QPushButton("Apply Slice")
        self.btn_apply_slice_range.setEnabled(False)
        self.btn_apply_slice_range.clicked.connect(self._apply_slice_range_to_group)

        self.channel_label = QLabel("Channel:")
        self.time_label = QLabel("Time:")
        self._top_to_label_1 = QLabel("to")
        self._top_to_label_2 = QLabel("to")
        self._top_to_label_1.setAlignment(Qt.AlignCenter)
        self._top_to_label_2.setAlignment(Qt.AlignCenter)
        self.slice_start_spin = QSpinBox()
        self.slice_start_spin.setRange(0, 0)
        self.slice_start_spin.setValue(0)
        self.slice_start_spin.setEnabled(False)
        self.slice_end_spin = QSpinBox()
        self.slice_end_spin.setRange(0, 0)
        self.slice_end_spin.setValue(0)
        self.slice_end_spin.setEnabled(False)
        self.slice_start_spin.valueChanged.connect(self._on_slice_range_changed)
        self.slice_end_spin.valueChanged.connect(self._on_slice_range_changed)
        row_label_width = self._scaled_px(132)
        mid_label_width = self._scaled_px(44)
        button_width = self._scaled_px(148)

        channel_row = QWidget()
        channel_layout = QHBoxLayout(channel_row)
        channel_layout.setContentsMargins(0, 0, 0, 0)
        channel_layout.setSpacing(self._scaled_px(10))
        self.channel_label.setFixedWidth(row_label_width)
        self.channel_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.time_label.setFixedWidth(row_label_width)
        self.time_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        channel_layout.addWidget(self.channel_label)
        channel_layout.addWidget(self.c_spin, 1)
        channel_layout.addWidget(self.time_label)
        channel_layout.addWidget(self.t_spin, 1)
        g.addWidget(channel_row, 2, 0, 1, 4)

        self._time_range_label = QLabel("Time range:")
        self._slice_range_label = QLabel("Slice range:")
        for range_label in (self._time_range_label, self._slice_range_label):
            range_label.setFixedWidth(row_label_width)
            range_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        for to_label in (self._top_to_label_1, self._top_to_label_2):
            to_label.setFixedWidth(mid_label_width)
            to_label.setAlignment(Qt.AlignCenter)
        for button in (self.btn_apply_time_range, self.btn_apply_slice_range):
            button.setMinimumWidth(button_width)

        time_range_row = QWidget()
        time_range_layout = QHBoxLayout(time_range_row)
        time_range_layout.setContentsMargins(0, 0, 0, 0)
        time_range_layout.setSpacing(self._scaled_px(10))
        time_range_layout.addWidget(self._time_range_label)
        time_range_layout.addWidget(self.t_start_spin, 1)
        time_range_layout.addWidget(self._top_to_label_1)
        time_range_layout.addWidget(self.t_end_spin, 1)
        time_range_layout.addWidget(self.btn_apply_time_range)
        g.addWidget(time_range_row, 3, 0, 1, 4)

        slice_range_row = QWidget()
        slice_range_layout = QHBoxLayout(slice_range_row)
        slice_range_layout.setContentsMargins(0, 0, 0, 0)
        slice_range_layout.setSpacing(self._scaled_px(10))
        slice_range_layout.addWidget(self._slice_range_label)
        slice_range_layout.addWidget(self.slice_start_spin, 1)
        slice_range_layout.addWidget(self._top_to_label_2)
        slice_range_layout.addWidget(self.slice_end_spin, 1)
        slice_range_layout.addWidget(self.btn_apply_slice_range)
        g.addWidget(slice_range_row, 4, 0, 1, 4)

        g.setColumnStretch(0, 1)
        g.setColumnStretch(1, 1)
        g.setColumnStretch(2, 1)
        g.setColumnStretch(3, 1)

        box.setLayout(g)
        return box

    def _build_info_box(self) -> QGroupBox:
        box = QGroupBox("Image Info")
        g = QGridLayout()
        self._info_grid = g
        g.setHorizontalSpacing(self._scaled_px(10))
        g.setVerticalSpacing(self._scaled_px(6))
        g.setColumnStretch(0, 0)
        g.setColumnStretch(1, 1)

        self.lbl_shape = QLabel("–")
        self.lbl_phys = QLabel("–")
        self.lbl_voxpix = QLabel("–")
        self.lbl_time = QLabel("–")
        for lbl in (self.lbl_shape, self.lbl_phys, self.lbl_voxpix, self.lbl_time):
            lbl.setWordWrap(True)
            lbl.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
            lbl.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)

        g.addWidget(QLabel("Shape:"), 0, 0, alignment=Qt.AlignLeft | Qt.AlignVCenter)
        g.addWidget(self.lbl_shape, 0, 1)
        g.addWidget(QLabel("Physical size:"), 1, 0, alignment=Qt.AlignLeft | Qt.AlignVCenter)
        g.addWidget(self.lbl_phys, 1, 1)
        self._size_info_label = QLabel("Voxel size:")
        g.addWidget(self._size_info_label, 2, 0, alignment=Qt.AlignLeft | Qt.AlignVCenter)
        g.addWidget(self.lbl_voxpix, 2, 1)
        self._time_info_label = QLabel("Time series:")
        g.addWidget(self._time_info_label, 3, 0, alignment=Qt.AlignLeft | Qt.AlignVCenter)
        g.addWidget(self.lbl_time, 3, 1)
        self._time_info_label.hide()
        self.lbl_time.hide()

        # Editable voxel/pixel sizes
        self.edit_vz = QDoubleSpinBox()
        self.edit_vz.setDecimals(6)
        self.edit_vz.setRange(0, 1e3)
        self.edit_vz.setLocale(QLocale.c())

        self.edit_vyx = QDoubleSpinBox()
        self.edit_vyx.setDecimals(6)
        self.edit_vyx.setRange(0, 1e3)
        self.edit_vyx.setLocale(QLocale.c())

        self.btn_apply_voxel = QPushButton("Apply Size")
        self.btn_apply_voxel.clicked.connect(self._on_apply_voxel_clicked)

        vox_row = 4
        g.addWidget(QLabel("Edit size:"), vox_row, 0, alignment=Qt.AlignLeft | Qt.AlignVCenter)

        col_widget = QWidget()
        col_layout = QHBoxLayout()
        col_layout.setContentsMargins(0, 0, 0, 0)
        col_layout.setSpacing(self._scaled_px(8))
        self.lblZ = QLabel("Z:")
        self.lblYX = QLabel("Y/X:")
        col_layout.addWidget(self.lblZ)
        col_layout.addWidget(self.edit_vz, 1)
        col_layout.addWidget(self.lblYX)
        col_layout.addWidget(self.edit_vyx, 1)
        col_layout.addWidget(self.btn_apply_voxel)
        col_widget.setLayout(col_layout)

        g.addWidget(col_widget, vox_row, 1)

        # default scale bar
        self._configure_scale_bar("um")

        box.setLayout(g)
        return box

    def _build_denoise_box(self) -> QGroupBox:
        box = QGroupBox("Preprocessing")
        g = QGridLayout()
        self._denoise_grid = g
        g.setHorizontalSpacing(self._scaled_px(10))
        g.setVerticalSpacing(self._scaled_px(6))
        g.setColumnStretch(0, 0)
        g.setColumnStretch(1, 1)

        self._info_raw_layer_combo = QComboBox()
        self._info_raw_layer_combo.currentIndexChanged.connect(
            self._on_info_raw_layer_changed
        )
        g.addWidget(
            QLabel("Raw image layer:"),
            0,
            0,
            alignment=Qt.AlignLeft | Qt.AlignVCenter,
        )
        g.addWidget(self._info_raw_layer_combo, 0, 1)

        self.upsample_factor_spin = QSpinBox()
        self.upsample_factor_spin.setRange(2, 4)
        self.upsample_factor_spin.setValue(2)
        self.upsample_factor_spin.setSuffix("×")
        self.upsample_enabled_checkbox = QCheckBox()
        self.upsample_enabled_checkbox.setChecked(False)
        self.upsample_enabled_checkbox.setToolTip("Enable optional XY upsampling")
        self.btn_apply_upsample = QPushButton("Run Bilinear")
        self.btn_apply_upsample.clicked.connect(self._on_apply_upsample_clicked)

        row_upsample = QWidget()
        row_upsample_l = QGridLayout()
        row_upsample_l.setContentsMargins(0, 0, 0, 0)
        self._row_upsample_layout = row_upsample_l
        self.upsample_action_label = QLabel("Upsample XY")
        self.upsample_factor_label = QLabel("Factor:")
        self.upsample_factor_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        row_upsample_l.addWidget(self.upsample_enabled_checkbox, 0, 0)
        row_upsample_l.addWidget(self.upsample_action_label, 0, 1)
        row_upsample_l.addWidget(self.upsample_factor_label, 0, 2)
        row_upsample_l.addWidget(self.upsample_factor_spin, 0, 3)
        row_upsample_l.addWidget(self.btn_apply_upsample, 0, 4)
        row_upsample_l.setColumnStretch(3, 1)
        row_upsample.setLayout(row_upsample_l)
        g.addWidget(row_upsample, 5, 0, 1, 2)

        self.upsample_progress = QProgressBar()
        self.upsample_progress.setRange(0, 1)
        self.upsample_progress.setValue(0)
        self.upsample_progress.setFormat("Ready")
        self.upsample_progress.setMaximumHeight(self._scaled_px(18))
        g.addWidget(self.upsample_progress, 6, 0, 1, 2)
        self.upsample_enabled_checkbox.toggled.connect(
            self._refresh_upsample_controls_state
        )
        self._refresh_upsample_controls_state()

        # --- Denoise controls -------------------------------------------
        self.cmb_kernel = QComboBox()
        for k in (3, 5, 7, 11):
            self.cmb_kernel.addItem(str(k))
        self.cmb_kernel.setCurrentText("3")

        self.btn_apply_denoise = QPushButton("Run Median")
        self.btn_apply_denoise.clicked.connect(self._on_apply_denoise_clicked)

        row_med = QWidget()
        row_med_l = QGridLayout()
        row_med_l.setContentsMargins(0, 0, 0, 0)
        self._row_med_layout = row_med_l
        self.denoise_action_label = QLabel("Median filter")
        self.filter_size_label = QLabel("Filter size:")
        self.filter_size_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        row_med_l.addWidget(self.denoise_action_label, 0, 0)
        row_med_l.addWidget(self.filter_size_label, 0, 1)
        row_med_l.addWidget(self.cmb_kernel, 0, 2)
        row_med_l.addWidget(self.btn_apply_denoise, 0, 3)
        row_med_l.setColumnStretch(2, 1)
        row_med.setLayout(row_med_l)
        g.addWidget(row_med, 1, 0, 1, 2)

        self.denoise_progress = QProgressBar()
        self.denoise_progress.setRange(0, 1)
        self.denoise_progress.setValue(0)
        self.denoise_progress.setFormat("Ready")
        self.denoise_progress.setMaximumHeight(self._scaled_px(18))
        g.addWidget(self.denoise_progress, 2, 0, 1, 2)

        self.gaussian_sigma_spin = QDoubleSpinBox()
        self.gaussian_sigma_spin.setRange(0.5, 100.0)
        self.gaussian_sigma_spin.setDecimals(1)
        self.gaussian_sigma_spin.setSingleStep(0.5)
        self.gaussian_sigma_spin.setValue(_GAUSSIAN_BACKGROUND_SIGMA_XY)
        self.gaussian_sigma_spin.setSuffix(" px")
        self.gaussian_alpha_spin = QDoubleSpinBox()
        self.gaussian_alpha_spin.setRange(0.0, 1.0)
        self.gaussian_alpha_spin.setDecimals(2)
        self.gaussian_alpha_spin.setSingleStep(0.05)
        self.gaussian_alpha_spin.setValue(_GAUSSIAN_BACKGROUND_ALPHA)
        self.gaussian_enabled_checkbox = QCheckBox()
        self.gaussian_enabled_checkbox.setChecked(False)
        self.gaussian_enabled_checkbox.setToolTip(
            "Enable optional Gaussian background subtraction"
        )
        self.btn_apply_gaussian = QPushButton("Run Gaussian")
        self.btn_apply_gaussian.setEnabled(False)
        self.btn_apply_gaussian.clicked.connect(self._on_apply_gaussian_background_clicked)

        row_gaussian = QWidget()
        row_gaussian_l = QGridLayout()
        row_gaussian_l.setContentsMargins(0, 0, 0, 0)
        self._row_gaussian_layout = row_gaussian_l
        self.gaussian_action_label = QLabel("Gaussian background")
        self.gaussian_sigma_label = QLabel("Sigma:")
        self.gaussian_alpha_label = QLabel("Alpha:")
        self.gaussian_sigma_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.gaussian_alpha_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        row_gaussian_l.addWidget(self.gaussian_enabled_checkbox, 0, 0)
        row_gaussian_l.addWidget(self.gaussian_action_label, 0, 1)
        row_gaussian_l.addWidget(self.gaussian_sigma_label, 0, 2)
        row_gaussian_l.addWidget(self.gaussian_sigma_spin, 0, 3)
        row_gaussian_l.addWidget(self.gaussian_alpha_label, 0, 4)
        row_gaussian_l.addWidget(self.gaussian_alpha_spin, 0, 5)
        row_gaussian_l.addWidget(self.btn_apply_gaussian, 0, 6)
        row_gaussian_l.setColumnStretch(3, 1)
        row_gaussian_l.setColumnStretch(5, 1)
        row_gaussian.setLayout(row_gaussian_l)
        g.addWidget(row_gaussian, 3, 0, 1, 2)

        self.gaussian_progress = QProgressBar()
        self.gaussian_progress.setRange(0, 1)
        self.gaussian_progress.setValue(0)
        self.gaussian_progress.setFormat("Ready")
        self.gaussian_progress.setMaximumHeight(self._scaled_px(18))
        g.addWidget(self.gaussian_progress, 4, 0, 1, 2)
        self.gaussian_enabled_checkbox.toggled.connect(
            self._refresh_gaussian_button_state
        )
        self._refresh_gaussian_button_state()

        box.setLayout(g)
        return box

    def _build_frangi_box(self) -> QGroupBox:
        box = QGroupBox("Structural Awareness Extraction")
        g = QGridLayout()
        self._frangi_grid = g
        g.setHorizontalSpacing(self._scaled_px(10))
        g.setVerticalSpacing(self._scaled_px(6))

        self._frangi_raw_layer_combo = QComboBox()
        g.addWidget(QLabel("Select layer:"), 0, 0, alignment=Qt.AlignRight)
        g.addWidget(self._frangi_raw_layer_combo, 0, 1, 1, 3)

        self.psf_ratio_spin = QDoubleSpinBox()
        self.psf_ratio_spin.setDecimals(2)
        self.psf_ratio_spin.setRange(0.10, 10.00)
        self.psf_ratio_spin.setSingleStep(0.10)
        self.psf_ratio_spin.setValue(3.00)
        self.psf_ratio_spin.setLocale(QLocale.c())
        g.addWidget(QLabel("PSF z/xy:"), 1, 0, alignment=Qt.AlignRight)
        g.addWidget(self.psf_ratio_spin, 1, 1)

        self.frangi_response_mode_combo = QComboBox()
        self.frangi_response_mode_combo.addItem("vesselness (tubular)", "vesselness")
        self.frangi_response_mode_combo.addItem("sheetness (sheet-like)", "sheetness")
        self.frangi_response_mode_combo.addItem("combined", "combined")
        self.frangi_response_mode_combo.setCurrentIndex(0)
        self.frangi_response_mode_combo.currentIndexChanged.connect(self._on_frangi_response_mode_changed)

        self.kernel_spin = QSpinBox()
        self.kernel_spin.setRange(1, 20)
        self.kernel_spin.setValue(4)

        g.addWidget(QLabel("Mode:"), 2, 0, alignment=Qt.AlignRight)
        g.addWidget(self.frangi_response_mode_combo, 2, 1)
        g.addWidget(QLabel("Kernel radius:"), 2, 2, alignment=Qt.AlignRight)
        g.addWidget(self.kernel_spin, 2, 3)

        self.sigma_min = QDoubleSpinBox()
        self.sigma_min.setDecimals(2)
        self.sigma_min.setRange(0.01, 20.0)
        self.sigma_min.setSingleStep(0.01)
        self.sigma_min.setValue(0.1)
        self.sigma_min.setLocale(QLocale.c())

        self.sigma_max = QDoubleSpinBox()
        self.sigma_max.setDecimals(2)
        self.sigma_max.setRange(0.01, 20.0)
        self.sigma_max.setSingleStep(0.1)
        self.sigma_max.setValue(1.0)
        self.sigma_max.setLocale(QLocale.c())

        self.sigma_count = QSpinBox()
        self.sigma_count.setRange(1, 99)
        self.sigma_count.setValue(5)

        self.vessel_rescale_spin = QDoubleSpinBox()
        self.vessel_rescale_spin.setDecimals(2)
        self.vessel_rescale_spin.setRange(0.0, 20.0)
        self.vessel_rescale_spin.setSingleStep(0.01)
        self.vessel_rescale_spin.setSuffix(" %")
        self.vessel_rescale_spin.setValue(0.0)
        self.vessel_rescale_spin.setLocale(QLocale.c())

        self.vessel_sigma_min_label = QLabel("Vessel σ min:")
        self.vessel_sigma_max_label = QLabel("Vessel σ max:")
        self.vessel_sigma_count_label = QLabel("Vessel count:")
        self.vessel_rescale_label = QLabel("Vessel rescale:")
        g.addWidget(self.vessel_sigma_min_label, 3, 0, alignment=Qt.AlignRight)
        g.addWidget(self.sigma_min, 3, 1)
        g.addWidget(self.vessel_sigma_max_label, 3, 2, alignment=Qt.AlignRight)
        g.addWidget(self.sigma_max, 3, 3)
        g.addWidget(self.vessel_sigma_count_label, 4, 0, alignment=Qt.AlignRight)
        g.addWidget(self.sigma_count, 4, 1)
        self._frangi_vessel_mode_widgets = (
            self.vessel_sigma_min_label,
            self.vessel_sigma_max_label,
            self.vessel_sigma_count_label,
            self.vessel_rescale_label,
            self.sigma_min,
            self.sigma_max,
            self.sigma_count,
            self.vessel_rescale_spin,
        )

        self.sheet_sigma_min = QDoubleSpinBox()
        self.sheet_sigma_min.setDecimals(2)
        self.sheet_sigma_min.setRange(0.01, 20.0)
        self.sheet_sigma_min.setSingleStep(0.01)
        self.sheet_sigma_min.setValue(0.01)
        self.sheet_sigma_min.setLocale(QLocale.c())

        self.sheet_sigma_max = QDoubleSpinBox()
        self.sheet_sigma_max.setDecimals(2)
        self.sheet_sigma_max.setRange(0.01, 20.0)
        self.sheet_sigma_max.setSingleStep(0.1)
        self.sheet_sigma_max.setValue(0.4)
        self.sheet_sigma_max.setLocale(QLocale.c())

        self.sheet_sigma_count = QSpinBox()
        self.sheet_sigma_count.setRange(1, 99)
        self.sheet_sigma_count.setValue(5)

        self.sheet_rescale_spin = QDoubleSpinBox()
        self.sheet_rescale_spin.setDecimals(2)
        self.sheet_rescale_spin.setRange(0.0, 20.0)
        self.sheet_rescale_spin.setSingleStep(0.01)
        self.sheet_rescale_spin.setSuffix(" %")
        self.sheet_rescale_spin.setValue(0.0)
        self.sheet_rescale_spin.setLocale(QLocale.c())

        self.sheet_sigma_min_label = QLabel("Sheet σ min:")
        self.sheet_sigma_max_label = QLabel("Sheet σ max:")
        self.sheet_sigma_count_label = QLabel("Sheet count:")
        self.sheet_rescale_label = QLabel("Sheet rescale:")
        g.addWidget(self.sheet_sigma_min_label, 5, 0, alignment=Qt.AlignRight)
        g.addWidget(self.sheet_sigma_min, 5, 1)
        g.addWidget(self.sheet_sigma_max_label, 5, 2, alignment=Qt.AlignRight)
        g.addWidget(self.sheet_sigma_max, 5, 3)
        g.addWidget(self.sheet_sigma_count_label, 6, 0, alignment=Qt.AlignRight)
        g.addWidget(self.sheet_sigma_count, 6, 1)
        self._frangi_sheet_mode_widgets = (
            self.sheet_sigma_min_label,
            self.sheet_sigma_max_label,
            self.sheet_sigma_count_label,
            self.sheet_rescale_label,
            self.sheet_sigma_min,
            self.sheet_sigma_max,
            self.sheet_sigma_count,
            self.sheet_rescale_spin,
        )

        self.btn_view_vessel_frangi = QPushButton("View Vessel")
        self.btn_view_vessel_frangi.clicked.connect(lambda: self._show_frangi_component("vesselness"))
        self.btn_view_sheet_frangi = QPushButton("View Sheet")
        self.btn_view_sheet_frangi.clicked.connect(lambda: self._show_frangi_component("sheetness"))
        self.btn_view_max_frangi = QPushButton("View Max")
        self.btn_view_max_frangi.clicked.connect(lambda: self._show_frangi_component("combined max"))
        view_row = QWidget()
        view_row_l = QHBoxLayout()
        view_row_l.setContentsMargins(0, 0, 0, 0)
        view_row_l.addWidget(self.btn_view_vessel_frangi)
        view_row_l.addWidget(self.btn_view_sheet_frangi)
        view_row_l.addWidget(self.btn_view_max_frangi)
        view_row.setLayout(view_row_l)
        view_row.setMinimumHeight(self._scaled_px(30))
        view_row.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
        self._frangi_component_view_row = view_row
        g.addWidget(view_row, 7, 1, 1, 3)

        self.apply_btn = QPushButton("Run Extraction")
        self.apply_btn.clicked.connect(self._on_apply_frangi_clicked)
        g.addWidget(self.apply_btn, 8, 0, 1, 4)

        self.frangi_progress = QProgressBar()
        self.frangi_progress.setValue(0)
        self.frangi_progress.setTextVisible(True)
        self.frangi_progress.setMaximumHeight(self._scaled_px(18))
        g.addWidget(self.frangi_progress, 9, 0, 1, 4)

        self.frangi_rescale_checkbox = QCheckBox("Enable rescale")
        self.frangi_rescale_checkbox.setChecked(False)
        self.frangi_rescale_checkbox.setToolTip(
            "Enable optional percentile rescaling of the Frangi response"
        )
        g.addWidget(self.frangi_rescale_checkbox, 10, 0, 1, 4)

        self.rescale_target_label = QLabel("Rescale target: —")
        self.rescale_target_label.setWordWrap(True)
        g.addWidget(self.rescale_target_label, 11, 0, 1, 4)
        g.addWidget(self.vessel_rescale_label, 12, 0, alignment=Qt.AlignRight)
        g.addWidget(self.vessel_rescale_spin, 12, 1, 1, 3)
        g.addWidget(self.sheet_rescale_label, 13, 0, alignment=Qt.AlignRight)
        g.addWidget(self.sheet_rescale_spin, 13, 1, 1, 3)
        self.rescale_apply_btn = QPushButton("Apply Rescale")
        self.rescale_apply_btn.setToolTip(
            "Apply the rescale values above to the current structural response"
        )
        self.rescale_apply_btn.clicked.connect(self._on_apply_component_rescale_clicked)
        g.addWidget(self.rescale_apply_btn, 14, 0, 1, 4)
        self.frangi_rescale_checkbox.toggled.connect(
            self._refresh_frangi_rescale_controls_state
        )

        self._on_frangi_response_mode_changed()
        self._refresh_frangi_component_view_buttons()
        self._refresh_frangi_rescale_controls_state()
        g.setColumnStretch(0, 0)
        g.setColumnStretch(1, 1)
        g.setColumnStretch(2, 0)
        g.setColumnStretch(3, 1)

        box.setLayout(g)
        return box

    def _create_em_sample_combo(self, default_value: int | None) -> QComboBox:
        combo = QComboBox()
        combo.setEditable(True)
        combo.setInsertPolicy(QComboBox.NoInsert)
        for value in _EM_SAMPLE_PRESETS:
            combo.addItem(_format_em_sample_points(value), value)
        combo.addItem("All", None)
        combo.setEditText(_format_em_sample_points(default_value))
        combo._last_valid_sample_points = default_value
        combo.currentTextChanged.connect(
            lambda text, target=combo: self._on_em_sample_text_changed(target, text)
        )
        combo.lineEdit().editingFinished.connect(
            lambda target=combo: self._normalize_em_sample_combo(target)
        )
        return combo

    def _on_em_sample_text_changed(self, combo: QComboBox, text: str) -> None:
        with suppress(ValueError):
            combo._last_valid_sample_points = _parse_em_sample_points_text(text)
        if hasattr(self, "kw_preview"):
            self._update_segmentation_preview()

    def _normalize_em_sample_combo(self, combo: QComboBox) -> None:
        try:
            value = _parse_em_sample_points_text(combo.currentText())
            combo._last_valid_sample_points = value
        except ValueError:
            value = getattr(combo, "_last_valid_sample_points", None)
        normalized = _format_em_sample_points(value)
        if combo.currentText() != normalized:
            combo.setEditText(normalized)

    @staticmethod
    def _em_sample_points_from_combo(combo: QComboBox | None) -> int | None:
        if combo is None:
            return None
        try:
            return _parse_em_sample_points_text(combo.currentText())
        except ValueError:
            return getattr(combo, "_last_valid_sample_points", None)

    def _em_foreground_sample_points(self) -> int | None:
        return self._em_sample_points_from_combo(
            getattr(self, "em_fg_sample_combo", None)
        )

    def _em_background_sample_points(self) -> int | None:
        return self._em_sample_points_from_combo(
            getattr(self, "em_bg_sample_combo", None)
        )

    def _build_seg_box(self) -> QGroupBox:
        box = QGroupBox("Segmentation")
        g = QGridLayout()
        self._seg_grid = g
        g.setHorizontalSpacing(self._scaled_px(10))
        g.setVerticalSpacing(self._scaled_px(6))

        self._seg_raw_layer_combo = QComboBox()
        self._seg_raw_layer_combo.currentIndexChanged.connect(self._update_segmentation_preview)
        self._seg_frangi_layer_combo = QComboBox()
        self._seg_frangi_layer_combo.currentIndexChanged.connect(self._on_segmentation_frangi_layer_changed)
        g.addWidget(QLabel("Raw data layer:"), 0, 0, alignment=Qt.AlignRight)
        g.addWidget(self._seg_raw_layer_combo, 0, 1, 1, 3)
        g.addWidget(QLabel("Structural response layer:"), 1, 0, alignment=Qt.AlignRight)
        g.addWidget(self._seg_frangi_layer_combo, 1, 1, 1, 3)

        self.beta1_spin = QDoubleSpinBox()
        self.beta1_spin.setDecimals(2)
        self.beta1_spin.setRange(0.0, 10)
        self.beta1_spin.setSingleStep(0.01)
        self.beta1_spin.setValue(1.0)
        self.beta1_spin.setLocale(QLocale.c())
        self.beta1_spin.setToolTip("Pairwise spatial smoothness weight.")

        self.beta2_spin = QDoubleSpinBox()
        self.beta2_spin.setDecimals(2)
        self.beta2_spin.setRange(0.0, 10)
        self.beta2_spin.setSingleStep(0.01)
        self.beta2_spin.setValue(1.0)
        self.beta2_spin.setLocale(QLocale.c())
        self.beta2_spin.setToolTip("Structural-response potential weight.")

        self.nfore_spin = QSpinBox()
        self.nfore_spin.setRange(1, 16)
        self.nfore_spin.setValue(3)

        self.nback_spin = QSpinBox()
        self.nback_spin.setRange(1, 16)
        self.nback_spin.setValue(8)

        self.maxiter_spin = QSpinBox()
        self.maxiter_spin.setRange(1, 200)
        self.maxiter_spin.setValue(50)

        self.init_combo = QComboBox()
        self.init_combo.addItems(["random", "otsu"])

        self.em_fg_sample_combo = self._create_em_sample_combo(
            _DEFAULT_EM_FOREGROUND_SAMPLE_POINTS
        )
        self.em_bg_sample_combo = self._create_em_sample_combo(
            _DEFAULT_EM_BACKGROUND_SAMPLE_POINTS
        )

        g.addWidget(QLabel("beta1 (smoothness):"), 2, 0, alignment=Qt.AlignRight)
        g.addWidget(self.beta1_spin, 2, 1)
        g.addWidget(QLabel("beta2 (structure):"), 2, 2, alignment=Qt.AlignRight)
        g.addWidget(self.beta2_spin, 2, 3)
        g.addWidget(QLabel("nforeground:"), 3, 0, alignment=Qt.AlignRight)
        g.addWidget(self.nfore_spin, 3, 1)
        g.addWidget(QLabel("nbackground:"), 3, 2, alignment=Qt.AlignRight)
        g.addWidget(self.nback_spin, 3, 3)
        g.addWidget(QLabel("maxiter:"), 4, 0, alignment=Qt.AlignRight)
        g.addWidget(self.maxiter_spin, 4, 1)
        g.addWidget(QLabel("init:"), 4, 2, alignment=Qt.AlignRight)
        g.addWidget(self.init_combo, 4, 3)

        g.addWidget(QLabel("EM foreground points:"), 5, 0, alignment=Qt.AlignRight)
        g.addWidget(self.em_fg_sample_combo, 5, 1, 1, 3)
        g.addWidget(QLabel("EM background points:"), 6, 0, alignment=Qt.AlignRight)
        g.addWidget(self.em_bg_sample_combo, 6, 1, 1, 3)

        self.seg_progress = QProgressBar()
        self.seg_progress.setValue(0)
        self.seg_progress.setTextVisible(True)
        g.addWidget(self.seg_progress, 7, 0, 1, 4)
        self.seg_progress.setMaximumHeight(self._scaled_px(18))

        self.kw_preview = QTextEdit()
        self.kw_preview.setReadOnly(True)
        self.kw_preview.setAttribute(Qt.WA_InputMethodEnabled, False)
        g.addWidget(self.kw_preview, 8, 0, 1, 4)

        self.seg_btn = QPushButton("Run Segmentation")
        self.seg_btn.clicked.connect(self._on_run_segmentation_clicked)
        g.addWidget(self.seg_btn, 9, 0, 1, 4)
        g.setColumnStretch(0, 0)
        g.setColumnStretch(1, 1)
        g.setColumnStretch(2, 0)
        g.setColumnStretch(3, 1)

        box.setLayout(g)
        return box

    # -----------------------------------------------------------------
    # Open / drop interception & unified loader with 0–255 display rescale
    # -----------------------------------------------------------------


    def _normalize_existing_layers_on_startup(self) -> None:
        """Synchronize controls without reloading or replacing existing image data."""
        if not getattr(self, "_disposed", False):
            self._update_info()


    def _load_and_add_path(self, path: str):
        data_tc_zyx, meta = _load_image_tc_zyx(path)
        self._apply_loaded_layers(data_tc_zyx, meta, name=self._next_layer_name(os.path.basename(path)))

    # -----------------------------------------------------------------
    # Device helpers
    # -----------------------------------------------------------------

    def _populate_devices(self):
        """Select lazily; availability is checked only when computation needs it."""
        self.device_combo.clear()
        self.device_combo.addItems(["auto", "cpu", "cuda", "mps"])
        self.device_combo.setToolTip("Auto chooses an available GPU. Unavailable devices fall back to CPU.")

    def _resolve_device(self) -> str:
        torch = _get_torch_module()
        if torch is None:
            return "cpu"
        requested = ""
        with suppress(AttributeError, RuntimeError):
            requested = str(self.device_combo.currentText()).strip().lower()
        if requested == "cpu":
            return "cpu"
        if requested == "cuda":
            with suppress(RuntimeError, AssertionError, AttributeError):
                if torch.cuda.is_available():
                    return "cuda"
            return "cpu"
        if requested == "mps":
            with suppress(RuntimeError, AssertionError, AttributeError):
                if hasattr(torch.backends, "mps") and torch.backends.mps.is_available():
                    return "mps"
            return "cpu"
        with suppress(RuntimeError, AssertionError, AttributeError):
            if torch.cuda.is_available():
                return "cuda"
        with suppress(RuntimeError, AssertionError, AttributeError):
            if hasattr(torch.backends, "mps") and torch.backends.mps.is_available():
                return "mps"
        return "cpu"

    # -----------------------------------------------------------------
    # File dialog open & normalization
    # -----------------------------------------------------------------

    def _open_image_tc_zyx(self):
        dlg = QFileDialog(self, "Open image")
        dlg.setFileMode(QFileDialog.ExistingFile)
        dlg.setNameFilter(
            "Images (*.tif *.tiff *.png *.jpg *.jpeg);;All files (*)"
        )
        if not dlg.exec_():
            return
        path = dlg.selectedFiles()[0]
        try:
            self._load_and_add_path(path)
        except (OSError, ValueError, TypeError, RuntimeError) as e:
            QMessageBox.critical(self, "Load failed", f"Failed to load: {e!r}")

    def _on_layer_inserted(self, event=None):
        layer = getattr(event, "value", None)
        if layer is None:
            return
        if bool((getattr(layer, "metadata", {}) or {}).get("managed_by_plugin")):
            return

        def _deferred():
            with suppress(Exception):
                self._maybe_normalize_dragdrop_layer(layer)
            is_analysis_aux = _is_analysis_aux_layer(layer)
            if not is_analysis_aux:
                self._view_initialized = False
                self._fit_view_and_scalebar()
            self._update_info()

        QTimer.singleShot(0, _deferred)

    def _maybe_normalize_dragdrop_layer(self, layer):
        """Respect in-memory data, including edits made by other plugins."""
        src = _get_source_from_layer(layer)
        if src:
            self.path_edit.setText(src)
        # File loading is handled by the registered reader. An insertion event
        # is not permission to replace a layer from its original file.

    # -----------------------------------------------------------------
    # Viewer sync & UI updates
    # -----------------------------------------------------------------

    def _temporal_reference_layer(self, preferred=None):
        candidates = [preferred]
        for getter_name in (
            "_selected_info_raw_layer",
            "_selected_segmentation_raw_layer",
            "_selected_tracking_layer",
        ):
            getter = getattr(self, getter_name, None)
            if getter is None:
                continue
            with suppress(Exception):
                candidates.append(getter())
        with suppress(Exception):
            candidates.extend(list(self.viewer.layers))

        seen: set[int] = set()
        for layer in candidates:
            if layer is None or id(layer) in seen:
                continue
            seen.add(id(layer))
            if _time_axis_index(self.viewer, layer) is not None:
                return layer
        return None

    def _current_temporal_step(self, preferred=None) -> int | None:
        layer = self._temporal_reference_layer(preferred)
        if layer is None:
            return None
        time_axis = _time_axis_index(self.viewer, layer)
        if time_axis is None:
            return None
        with suppress(AttributeError, IndexError, TypeError, ValueError):
            return int(self.viewer.dims.current_step[time_axis])
        return None

    def _on_device_changed(self, *_):
        if self._frangi_ctx is not None:
            self._frangi_ctx.device = self._resolve_device()
        self._update_segmentation_preview()

    def _on_active_changed(self, event=None):
        # Keep user-selected slice/zoom
        layer = self.viewer.layers.selection.active

        with suppress(Exception):
            if self._scale_conn is not None:
                self._scale_conn[0].disconnect(self._scale_conn[1])

        if layer is not None and hasattr(layer, "events") and hasattr(layer.events, "scale"):
            layer.events.scale.connect(self._update_info)
            self._scale_conn = (layer.events.scale, self._update_info)
        else:
            self._scale_conn = None

        if layer is not None:
            md = getattr(layer, "metadata", {}) or {}
            if "channel_index" in md and self.c_spin.isEnabled():
                self.c_spin.blockSignals(True)
                self.c_spin.setValue(int(md["channel_index"]))  # 0-based
                self.c_spin.blockSignals(False)
            t_cur0 = self._current_temporal_step(layer)
            if t_cur0 is not None and self.t_spin.isEnabled():
                self.t_spin.blockSignals(True)
                self.t_spin.setValue(t_cur0)  # 0-based
                self.t_spin.blockSignals(False)

        self._update_info()

    def _on_dims_step_changed(self, event=None):
        active_layer = self.viewer.layers.selection.active
        t_cur0 = self._current_temporal_step(active_layer)
        if t_cur0 is not None and self.t_spin.isEnabled():
            self.t_spin.blockSignals(True)
            self.t_spin.setValue(t_cur0)  # 0-based
            self.t_spin.blockSignals(False)
        self._update_info(update_preview=False)
        with suppress(Exception):
            self._populate_tracking_unlinked_table()
        with suppress(Exception):
            self._schedule_analysis_refresh_for_current_frame(force=False)

    def _layer_combo_state_signature(self):
        signature = []
        for layer in self.viewer.layers:
            data = getattr(layer, "data", None)
            md = getattr(layer, "metadata", {}) or {}
            shape = getattr(data, "shape", ())
            with suppress(TypeError):
                shape = tuple(int(value) for value in shape)
            signature.append(
                (
                    id(layer),
                    id(data),
                    str(getattr(layer, "name", "")),
                    str(getattr(layer, "_type_string", "")),
                    shape,
                    str(getattr(data, "dtype", "")),
                    str(md.get("dims") or md.get("dims_out") or ""),
                    str(md.get("layer_type") or ""),
                    tuple(bool(md.get(key)) for key in _LAYER_COMBO_METADATA_KEYS),
                )
            )
        return tuple(signature)

    def _rebuild_layer_combos_if_needed(
        self,
        *,
        active_is_analysis_aux: bool = False,
        prefer_denoised: bool = False,
    ) -> None:
        signature = self._layer_combo_state_signature()
        state = (
            signature,
            bool(active_is_analysis_aux),
            bool(self._analysis_cleanup_in_progress),
            bool(prefer_denoised),
        )
        if state == self._layer_combo_state:
            return
        if not self._analysis_cleanup_in_progress and not active_is_analysis_aux:
            self._rebuild_analysis_layer_combo(clear_outputs=False)
        self._rebuild_proximity_layer_combos()
        self._rebuild_tracking_layer_combo()
        self._rebuild_frangi_layer_combo(prefer_denoised=prefer_denoised)
        self._rebuild_info_raw_layer_combo(prefer_denoised=prefer_denoised)
        self._rebuild_segmentation_layer_combos()
        self._layer_combo_state = state

    def _update_info(self, event=None, *, update_preview: bool = True):
        self._refresh_gaussian_button_state()
        removed_layer = getattr(event, "value", None)
        removed_metadata = getattr(removed_layer, "metadata", {}) or {}
        prefer_denoised = bool(
            removed_metadata.get("is_frangi")
            or removed_metadata.get("is_segmentation")
            or removed_metadata.get("sigma_layer_role") == _STRUCTURAL_RESPONSE_LAYER_ROLE
        )
        info_combo = getattr(self, "_info_raw_layer_combo", None)
        combo_layer = info_combo.currentData() if (info_combo is not None and info_combo.count() > 0) else None
        layer = combo_layer if combo_layer is not None else self.viewer.layers.selection.active
        time_text = _format_time_series_info(layer) if layer is not None else None
        if self._time_info_label is not None and self.lbl_time is not None:
            self._time_info_label.setVisible(time_text is not None)
            self.lbl_time.setVisible(time_text is not None)
            self.lbl_time.setText(time_text or "–")
        if layer is None or not hasattr(layer, "data"):
            self.lbl_shape.setText("–")
            self.lbl_phys.setText("–")
            self.lbl_voxpix.setText("–")
            self._rebuild_layer_combos_if_needed(prefer_denoised=prefer_denoised)
            return
        is_analysis_aux = _is_analysis_aux_layer(layer)
        data = getattr(layer, "data", None)
        nd = getattr(data, "ndim", None)
        if nd is None:
            self.lbl_shape.setText("–")
            self.lbl_phys.setText("–")
            self.lbl_voxpix.setText("–")
            self._rebuild_layer_combos_if_needed(
                active_is_analysis_aux=is_analysis_aux,
                prefer_denoised=prefer_denoised,
            )
            return
        try:
            info = _describe_size_and_resolution(layer)
        except (TypeError, ValueError, IndexError, AttributeError) as e:
            self.lbl_shape.setText(f"(error) {e!r}")
            return

        if info.get("ndim") == 3:
            nz, ny, nx = info["shape"]
            vz, vy, vx = info["voxel"]
            pz, py, px = info["phys"]
            self.lbl_shape.setText(f"Z={nz}  Y={ny}  X={nx}")
            self.lbl_phys.setText(f"{pz:.6g}×{py:.6g}×{px:.6g} {info['unit']}")
            self.lbl_voxpix.setText(f"vz={vz:.6g}  vy={vy:.6g}  vx={vx:.6g} {info['unit']}/px")
            if self._size_info_label is not None:
                self._size_info_label.setText("Voxel size:")
            self.lblZ.show()
            self.edit_vz.show()
            self.edit_vz.setEnabled(True)
            self.lblYX.setText("Y/X:")
            self.edit_vyx.show()
            self.edit_vyx.setEnabled(True)
            self.edit_vz.setValue(max(vz, 0))
            self.edit_vyx.setValue(max(vy, 0))
            self.btn_apply_voxel.setText("Update Voxel size")
        elif info.get("ndim") == 2:
            ny, nx = info["shape"]
            vy, vx = info["voxel"]
            py, px = info["phys"]
            self.lbl_shape.setText(f"Y={ny}  X={nx}")
            self.lbl_phys.setText(f"{py:.6g}×{px:.6g} {info['unit']}")
            self.lbl_voxpix.setText(f"py={vy:.6g}  px={vx:.6g} {info['unit']}/px")
            if self._size_info_label is not None:
                self._size_info_label.setText("Pixel size:")
            self.lblZ.hide()
            self.edit_vz.hide()
            self.edit_vz.setEnabled(False)
            self.lblYX.setText("Y/X:")
            self.edit_vyx.show()
            self.edit_vyx.setEnabled(True)
            self.edit_vyx.setValue(max(vy, 0))
            self.btn_apply_voxel.setText("Update Pixel size")
        else:
            self.lbl_shape.setText(str(info.get("shape")))
            self.lbl_phys.setText("–")
            self.lbl_voxpix.setText("–")

        _safe_axis_labels(self.viewer, layer)
        self._configure_scale_bar(_get_units_from_layer(layer))

        self._rebuild_layer_combos_if_needed(
            active_is_analysis_aux=is_analysis_aux,
            prefer_denoised=prefer_denoised,
        )
        if update_preview and not is_analysis_aux:
            self._update_segmentation_preview()

    def _fit_view_and_scalebar(self):
        layer = self.viewer.layers.selection.active
        if layer is None or not hasattr(layer, "data"):
            return
        data = getattr(layer, "data", None)
        nd = getattr(data, "ndim", None)

        if nd is not None and not self._view_initialized:
            self._view_initialized = True
            self._center_view_on_layer(layer)

            # --- Sync Time spinbox with viewer after centering ---
            cur_t = self._current_temporal_step(layer)
            if cur_t is not None and self.t_spin.isEnabled():
                self.t_spin.blockSignals(True)
                self.t_spin.setValue(cur_t)
                self.t_spin.blockSignals(False)

        unit = (getattr(layer, "metadata", {}) or {}).get("unit", "um")
        self._configure_scale_bar(unit)

        _safe_axis_labels(self.viewer, layer)

    def _center_view_on_layer(self, layer) -> None:
        if layer is None or not hasattr(layer, "data"):
            return
        data = getattr(layer, "data", None)
        nd = _layer_display_ndim(layer)
        if nd <= 0:
            return
        with suppress(Exception):
            if np.all(np.isfinite(layer.extent.world[0])) and np.all(np.isfinite(layer.extent.world[1])):
                self.viewer.reset_view()
        shape = getattr(data, "shape", ())
        for layer_axis in range(max(0, nd - 2)):
            if layer_axis < len(shape) and shape[layer_axis] > 1:
                viewer_axis = _viewer_axis_for_layer_axis(
                    self.viewer,
                    layer,
                    layer_axis,
                )
                if viewer_axis is None:
                    continue
                with suppress(Exception):
                    self.viewer.dims.set_current_step(
                        viewer_axis,
                        int(shape[layer_axis] // 2),
                    )

    # -----------------------------------------------------------------
    # Voxel/pixel size handling
    # -----------------------------------------------------------------

    def _on_apply_voxel_clicked(self):
        info_combo = getattr(self, "_info_raw_layer_combo", None)
        combo_layer = info_combo.currentData() if (info_combo is not None and info_combo.count() > 0) else None
        layer = combo_layer if combo_layer is not None else self.viewer.layers.selection.active
        if layer is None or not hasattr(layer, "data"):
            QMessageBox.information(self, "No image", "Please select an image layer first.")
            return

        try:
            new_scale, new_triplet = self._compute_updated_scale_triplet(layer)
            gid = self._apply_voxel_update_to_layer(layer, new_scale)
            self._sync_group_layer_scales(gid, new_triplet, layer)

            if self._frangi_ctx is not None:
                with suppress(Exception):
                    if _get_group_id(layer) == self._frangi_ctx.group_id:
                        self._frangi_ctx.pixel_size = _get_pixel_size_tuple(layer, self._frangi_ctx.dim)

            self._center_view_on_layer(layer)
            self._fit_view_and_scalebar()
            self._update_info()
        except (TypeError, ValueError, AttributeError, RuntimeError) as e:
            QMessageBox.critical(self, "Apply failed", f"Failed to set pixel/voxel size: {e!r}")

    # -----------------------------------------------------------------
    # C/T controls (0-based)
    # -----------------------------------------------------------------

    def _set_ct_controls(self, T: int, C: int):
        self.t_spin.blockSignals(True)
        self.t_spin.setRange(0, max(int(T) - 1, 0))
        self.t_spin.setValue(0)
        self.t_spin.setEnabled(int(T) > 1)
        self.t_spin.blockSignals(False)

        self.t_start_spin.blockSignals(True)
        self.t_start_spin.setRange(0, max(int(T) - 1, 0))
        self.t_start_spin.setValue(0)
        self.t_start_spin.setEnabled(int(T) > 1)
        self.t_start_spin.blockSignals(False)

        self.t_end_spin.blockSignals(True)
        self.t_end_spin.setRange(0, max(int(T) - 1, 0))
        self.t_end_spin.setValue(max(int(T) - 1, 0))
        self.t_end_spin.setEnabled(int(T) > 1)
        self.t_end_spin.blockSignals(False)

        self.btn_apply_time_range.setEnabled(int(T) > 1)

        self.c_spin.blockSignals(True)
        self.c_spin.setRange(0, max(int(C) - 1, 0))
        self.c_spin.setValue(0)
        self.c_spin.setEnabled(int(C) > 1)
        self.c_spin.blockSignals(False)

    def _on_t_changed(self, v0: int):
        layer = self._temporal_reference_layer(
            self.viewer.layers.selection.active,
        )
        if layer is not None:
            _apply_viewer_time_step(self.viewer, layer, max(int(v0), 0))
        self._update_info()

    def _on_time_range_changed(self, *_args) -> None:
        start = int(self.t_start_spin.value())
        end = int(self.t_end_spin.value())
        if start > end:
            sender = self.sender()
            if sender is self.t_start_spin:
                self.t_end_spin.blockSignals(True)
                self.t_end_spin.setValue(start)
                self.t_end_spin.blockSignals(False)
            else:
                self.t_start_spin.blockSignals(True)
                self.t_start_spin.setValue(end)
                self.t_start_spin.blockSignals(False)

    def _selected_time_range(self) -> tuple[int, int]:
        return int(self.t_start_spin.value()), int(self.t_end_spin.value())

    def _reload_source_group_with_subranges(self, layer) -> None:
        if layer is None:
            return
        src = _get_source_from_layer(layer)
        if not src:
            QMessageBox.information(self, "No source", "Current layer has no file path metadata to reload and slice.")
            return
        t0, t1 = self._selected_time_range()
        z0, z1 = self._selected_slice_range()
        try:
            data_tc_zyx, meta = _load_image_tc_zyx(src)
            meta = dict(meta)
            if int(data_tc_zyx.shape[0]) > 1:
                t0 = max(0, min(t0, int(data_tc_zyx.shape[0]) - 1))
                t1 = max(t0, min(t1, int(data_tc_zyx.shape[0]) - 1))
                data_tc_zyx = data_tc_zyx[t0 : t1 + 1]
                meta["time_slice"] = (t0, t1)
            z_size = int(data_tc_zyx.shape[2])
            if z_size > 1:
                z0 = max(0, min(z0, z_size - 1))
                z1 = max(z0, min(z1, z_size - 1))
                data_tc_zyx = data_tc_zyx[:, :, z0 : z1 + 1]
                meta["slice_range"] = (z0, z1)
            src_group = _get_group_id(layer)
            for layer_item in list(self.viewer.layers):
                if not hasattr(layer_item, "data"):
                    continue
                md = getattr(layer_item, "metadata", {}) or {}
                if _is_analysis_aux_layer(layer_item):
                    continue
                if src_group is not None and md.get("group_id") == src_group:
                    with suppress(Exception):
                        self.viewer.layers.remove(layer_item)
            self._clear_analysis_results(reset_outputs=True)
            self._load_and_add_path(src)
            current = self.viewer.layers.selection.active
            if current is not None and _get_source_from_layer(current) == src:
                src_group = _get_group_id(current)
                for layer_item in list(self.viewer.layers):
                    if not hasattr(layer_item, "data"):
                        continue
                    if _is_analysis_aux_layer(layer_item):
                        continue
                    if src_group is not None and _get_group_id(layer_item) == src_group:
                        with suppress(Exception):
                            self.viewer.layers.remove(layer_item)
            self._apply_loaded_layers(data_tc_zyx, meta, name=self._next_layer_name(os.path.basename(src)))
        except (OSError, ValueError, TypeError, RuntimeError) as e:
            QMessageBox.critical(self, "Range apply failed", f"Failed to apply selected ranges: {e!r}")

    def _apply_time_range_to_group(self) -> None:
        layer = self.viewer.layers.selection.active
        self._reload_source_group_with_subranges(layer)

    def _apply_slice_range_to_group(self) -> None:
        layer = self.viewer.layers.selection.active
        self._reload_source_group_with_subranges(layer)

    def _on_c_changed(self, v0: int):
        idx0 = max(int(v0), 0)  # 0-based
        for lyr in self._group_layers:
            md = getattr(lyr, "metadata", {}) or {}
            if md.get("channel_index") == idx0:
                with suppress(Exception):
                    self.viewer.layers.selection.active = lyr
                break
        self._update_info()

    # -----------------------------------------------------------------
    # XY upsampling (bilinear, after denoise)
    # -----------------------------------------------------------------

    def _switch_combos_to_upsampled(self) -> None:
        layer = getattr(self, "_upsample_layer", None)
        if layer is None:
            return
        try:
            if layer not in self.viewer.layers:
                return
        except (TypeError, ValueError):
            return

        self._rebuild_info_raw_layer_combo()
        self._rebuild_frangi_layer_combo()
        self._rebuild_segmentation_layer_combos()
        for combo_name in (
            "_info_raw_layer_combo",
            "_frangi_raw_layer_combo",
            "_seg_raw_layer_combo",
        ):
            combo = getattr(self, combo_name, None)
            if combo is None or combo.count() == 0:
                continue
            idx = combo.findText(layer.name)
            if idx >= 0 and combo.currentIndex() != idx:
                combo.setCurrentIndex(idx)

    def _on_apply_upsample_clicked(self) -> None:
        if self._thread_is_running("_upsample_thread"):
            self._request_worker_cancel("_upsample_thread", "_upsample_worker")
            self.btn_apply_upsample.setEnabled(False)
            self.btn_apply_upsample.setText("Stopping...")
            self.upsample_progress.setRange(0, 1)
            self.upsample_progress.setValue(0)
            self.upsample_progress.setFormat("Stopping...")
            return

        if not self.upsample_enabled_checkbox.isChecked():
            return

        base_layer = self._selected_info_raw_layer() or self.viewer.layers.selection.active
        if base_layer is None or not hasattr(base_layer, "data"):
            QMessageBox.information(self, "No image", "Please select a raw data layer first.")
            return

        image = base_layer.data
        shape = tuple(int(value) for value in getattr(image, "shape", ()))
        if len(shape) < 2:
            QMessageBox.warning(self, "Upsample failed", "XY upsampling requires an image with at least two axes.")
            return

        factor = int(self.upsample_factor_spin.value())
        total = int(np.prod(shape[:-2], dtype=np.int64)) if len(shape) > 2 else 1
        self.upsample_progress.setRange(0, max(total, 1))
        self.upsample_progress.setValue(0)
        self.upsample_progress.setFormat(f"Progress 0/{max(total, 1)}")
        self.upsample_factor_spin.setEnabled(False)
        self.upsample_enabled_checkbox.setEnabled(False)
        self.btn_apply_upsample.setEnabled(True)
        self.btn_apply_upsample.setText("Stop Upsample")

        self._upsample_thread = QThread()
        self._upsample_worker = UpsampleWorker(image, factor)
        self._upsample_worker.moveToThread(self._upsample_thread)
        self._upsample_thread.started.connect(self._upsample_worker.run)
        self._upsample_worker.progress.connect(self._on_upsample_progress)

        def _finished(output, error):
            self.btn_apply_upsample.setText("Run Bilinear")
            if isinstance(error, InterruptedError):
                self.upsample_progress.setRange(0, 1)
                self.upsample_progress.setValue(0)
                self.upsample_progress.setFormat("Stopped")
            elif error is not None:
                self.upsample_progress.setRange(0, 1)
                self.upsample_progress.setValue(0)
                self.upsample_progress.setFormat("Failed")
                QMessageBox.critical(self, "Upsample failed", f"Bilinear XY upsampling error: {error!r}")
            else:
                output = np.asarray(output)
                base_scale = tuple(float(value) for value in getattr(base_layer, "scale", (1.0,) * output.ndim))
                new_scale = _upsampled_xy_scale(base_scale, factor)
                metadata = dict(getattr(base_layer, "metadata", {}) or {})
                input_is_denoised = bool(metadata.get("is_denoise"))
                try:
                    previous_factor = int(metadata.get("upsample_factor_xy", metadata.get("XYUpsamplingFactor", 1)))
                except (TypeError, ValueError):
                    previous_factor = 1
                cumulative_factor = max(previous_factor, 1) * factor
                metadata_scale = list(metadata.get("scale_per_axis", base_scale))
                if len(metadata_scale) >= 2:
                    metadata_scale[-2] = float(metadata_scale[-2]) / factor
                    metadata_scale[-1] = float(metadata_scale[-1]) / factor
                metadata.update(
                    {
                        "managed_by_plugin": True,
                        "is_upsample": True,
                        "upsample_factor_xy": cumulative_factor,
                        "upsample_step_factor_xy": factor,
                        "upsample_interpolation": "bilinear",
                        "upsample_input_shape": shape,
                        "upsample_original_shape": tuple(
                            int(value)
                            for value in metadata.get("upsample_original_shape", shape)
                        ),
                        "upsample_output_shape": tuple(int(value) for value in output.shape),
                        "scale_per_axis": tuple(metadata_scale),
                        "XYUpsamplingFactor": cumulative_factor,
                        "XYInterpolation": "bilinear",
                        "group_id": _get_group_id(base_layer) or _ensure_group_id(metadata),
                    }
                )
                name = self._derived_result_layer_name(base_layer, f"{factor}x XY bilinear")
                translate = tuple(float(value) for value in getattr(base_layer, "translate", (0.0,) * output.ndim))
                self._upsample_layer = self.viewer.add_image(
                    output,
                    name=name,
                    scale=new_scale,
                    translate=translate,
                    metadata=metadata,
                    **_spatial_unit_kwargs(base_layer, output.ndim),
                )
                for attr in ("colormap", "contrast_limits", "gamma", "opacity", "blending"):
                    with suppress(Exception):
                        setattr(self._upsample_layer, attr, getattr(base_layer, attr))
                self._reset_processing_state(
                    clear_upsample=False,
                    clear_denoise=not input_is_denoised,
                    clear_frangi=True,
                )
                self._focus_result_layer(self._upsample_layer)
                self._switch_combos_to_upsampled()
                self._update_info()
                self.upsample_progress.setRange(0, 1)
                self.upsample_progress.setValue(1)
                self.upsample_progress.setFormat(f"Finished: {factor}x XY")
            self._cleanup_worker_thread("_upsample_thread", "_upsample_worker")
            self.upsample_enabled_checkbox.setEnabled(True)
            self._refresh_upsample_controls_state()

        self._connect_worker_callback(self._upsample_worker.finished, _finished)
        self._upsample_thread.start()

    def _on_upsample_progress(self, done: int, total: int) -> None:
        total = max(int(total), 1)
        done = max(0, min(int(done), total))
        if self.upsample_progress.maximum() != total:
            self.upsample_progress.setRange(0, total)
        self.upsample_progress.setValue(done)
        self.upsample_progress.setFormat(f"Progress {done}/{total}")

    # -----------------------------------------------------------------
    # Denoise (median filter via button)
    # -----------------------------------------------------------------

    def _switch_combos_to_denoised(self) -> None:
        """Select the newly created denoised layer for subsequent processing."""
        layer = self._preferred_denoised_input_layer()
        if layer is None:
            return
        # Make the denoised result the input for upsampling and later steps.
        self._rebuild_info_raw_layer_combo()
        self._rebuild_frangi_layer_combo()
        self._rebuild_segmentation_layer_combos()
        for combo_name in (
            "_info_raw_layer_combo",
            "_frangi_raw_layer_combo",
            "_seg_raw_layer_combo",
        ):
            combo = getattr(self, combo_name, None)
            if combo is None or combo.count() == 0:
                continue
            idx = combo.findText(layer.name)
            if idx >= 0 and combo.currentIndex() != idx:
                combo.setCurrentIndex(idx)

    def _on_apply_denoise_clicked(self):
        if self._thread_is_running("_denoise_thread"):
            self._request_worker_cancel("_denoise_thread", "_denoise_worker")
            self.btn_apply_denoise.setEnabled(False)
            self.btn_apply_denoise.setText("Stopping...")
            self.denoise_progress.setRange(0, 1)
            self.denoise_progress.setValue(0)
            self.denoise_progress.setFormat("Stopping...")
            return
        base_layer = self._selected_info_raw_layer() or self.viewer.layers.selection.active
        if base_layer is None or not hasattr(base_layer, "data"):
            QMessageBox.information(self, "No image", "Please select an image layer first.")
            return

        try:
            arr = base_layer.data
            dims_tag = _layer_dims_tag(base_layer)
            # Infer dims from array shape when metadata tag is absent
            if dims_tag is None:
                if arr.ndim == 4:
                    dims_tag = "TZYX"
                elif arr.ndim == 3:
                    dims_tag = "ZYX"
                elif arr.ndim == 2:
                    dims_tag = "YX"
            k = int(self.cmb_kernel.currentText())
            scale_all = tuple(float(v) for v in getattr(base_layer, "scale", (1,) * arr.ndim))
            temporal = False
            z_filter_range = None
            squeeze_time = False

            if dims_tag == "TYX":
                t0, t1 = self._selected_time_range()
                t0 = max(0, min(t0, arr.shape[0] - 1))
                t1 = max(t0, min(t1, arr.shape[0] - 1))
                denoise_input = arr[t0 : t1 + 1]
                temporal = True
                scale = scale_all[-3:]
                dims_out = "TYX" if denoise_input.shape[0] > 1 else "YX"
                squeeze_time = denoise_input.shape[0] == 1
                if squeeze_time:
                    scale = scale_all[-2:]

            elif dims_tag == "TZYX":
                t0, t1 = self._selected_time_range()
                t0 = max(0, min(t0, arr.shape[0] - 1))
                t1 = max(t0, min(t1, arr.shape[0] - 1))
                denoise_input = arr[t0 : t1 + 1]
                temporal = True
                z0, z1 = self._selected_slice_range()
                nz = denoise_input.shape[1]
                # z0==z1==0 is the spinbox default — treat as "full Z range"
                full_z = (z0 == z1 == 0 and nz > 1) or (z0 == 0 and z1 >= nz - 1)
                if not full_z:
                    zstart = int(np.clip(z0, 0, nz - 1))
                    zend = int(np.clip(max(z0, z1), zstart, nz - 1))
                    z_filter_range = (zstart, zend)
                scale = scale_all[-4:]
                dims_out = "TZYX" if denoise_input.shape[0] > 1 else "ZYX"
                squeeze_time = denoise_input.shape[0] == 1
                if squeeze_time:
                    scale = scale_all[-3:]

            elif dims_tag in {"ZYX", None}:
                vol = arr if arr.ndim == 3 else self._current_volume_from_layer(base_layer)
                z0, z1 = self._selected_slice_range()
                nz = vol.shape[0] if vol.ndim == 3 else 1
                full_z = vol.ndim < 3 or (z0 == z1 == 0 and nz > 1) or (z0 == 0 and z1 >= nz - 1)
                if vol.ndim == 3 and full_z:
                    denoise_input = vol
                    scale = tuple(float(v) for v in base_layer.scale[-3:])
                    dims_out = "ZYX"
                elif vol.ndim == 3:
                    zstart = int(np.clip(z0, 0, nz - 1))
                    zend = int(np.clip(max(z0, z1), zstart, nz - 1))
                    if zstart == zend:
                        denoise_input = vol[zstart]
                        scale = tuple(float(v) for v in base_layer.scale[-2:])
                        dims_out = "YX"
                    else:
                        denoise_input = vol[zstart : zend + 1]
                        scale = tuple(float(v) for v in base_layer.scale[-3:])
                        dims_out = "ZYX"
                elif vol.ndim == 2:
                    denoise_input = vol
                    scale = tuple(float(v) for v in base_layer.scale[-2:])
                    dims_out = "YX"
                else:
                    raise ValueError(f"Unsupported data ndim for denoise: {vol.ndim}")

            else:
                vol = self._current_volume_from_layer(base_layer)
                if vol.ndim == 3:
                    denoise_input = vol
                    scale = tuple(float(v) for v in base_layer.scale[-3:])
                    dims_out = "ZYX"
                elif vol.ndim == 2:
                    denoise_input = vol
                    scale = tuple(float(v) for v in base_layer.scale[-2:])
                    dims_out = "YX"
                else:
                    raise ValueError(f"Unsupported data ndim for denoise: {vol.ndim}")
        except (TypeError, ValueError, AttributeError, IndexError, RuntimeError) as e:
            QMessageBox.critical(self, "Denoise failed", f"Median filter error: {e!r}")
            return

        worker_count = min(os.cpu_count() or 1, 4)
        self.cmb_kernel.setEnabled(False)
        self.btn_apply_gaussian.setEnabled(False)
        self.btn_apply_denoise.setEnabled(True)
        self.btn_apply_denoise.setText("Stop Median")
        self.denoise_progress.setRange(0, 0)
        self.denoise_progress.setFormat("Starting...")
        self._denoise_thread = QThread()
        self._denoise_worker = DenoiseWorker(
            denoise_input,
            k,
            temporal=temporal,
            z_range=z_filter_range,
            n_workers=worker_count,
        )
        self._denoise_worker.moveToThread(self._denoise_thread)
        self._denoise_thread.started.connect(self._denoise_worker.run)
        self._denoise_worker.progress.connect(self._on_denoise_progress)

        def _finished(output, error):
            self.cmb_kernel.setEnabled(True)
            self.btn_apply_denoise.setEnabled(True)
            self.btn_apply_denoise.setText("Run Median")
            if isinstance(error, InterruptedError):
                self.denoise_progress.setRange(0, 1)
                self.denoise_progress.setValue(0)
                self.denoise_progress.setFormat("Stopped")
            elif error is not None:
                self.denoise_progress.setRange(0, 1)
                self.denoise_progress.setValue(0)
                self.denoise_progress.setFormat("Failed")
                QMessageBox.critical(self, "Denoise failed", f"Denoise error: {error!r}")
            else:
                out = np.asarray(output)
                if squeeze_time:
                    out = out[0]
                name = self._next_layer_name("denoised")
                md = dict(getattr(base_layer, "metadata", {}) or {})
                md.update(
                    {
                        "managed_by_plugin": True,
                        "is_denoise": True,
                        "dims_out": dims_out,
                        "denoise_method": "parallel_chunked_median",
                        "denoise_kernel_size": k,
                        "denoise_workers": worker_count,
                        "median_rescale_0_255": True,
                        "group_id": _get_group_id(base_layer) or _ensure_group_id(md),
                    }
                )
                if dims_tag in {"TYX", "TZYX"}:
                    md["time_range"] = self._selected_time_range()
                denoise_kwargs = {"name": name, "scale": scale, "metadata": md}
                units = _spatial_units_for_ndim(base_layer, out.ndim)
                if units is not None:
                    denoise_kwargs["units"] = units
                self._denoise_layer = self.viewer.add_image(out, **denoise_kwargs)
                self._median_layer = self._denoise_layer
                self._gaussian_layer = None
                with suppress(Exception):
                    self._denoise_layer.contrast_limits = (0, 255)
                self._focus_result_layer(self._denoise_layer)
                self._denoise_ctx = self._context_from_layer(
                    base_layer,
                    out,
                    3 if dims_out in {"ZYX", "TZYX"} else 2,
                    tuple(scale),
                    temporal_dims=dims_out if dims_out in {"TYX", "TZYX"} else None,
                    time_range=self._selected_time_range() if dims_tag in {"TYX", "TZYX"} else None,
                )
                self._reset_processing_state(clear_upsample=True, clear_denoise=False, clear_frangi=True)
                self.denoise_progress.setRange(0, 1)
                self.denoise_progress.setValue(1)
                self.denoise_progress.setFormat("Finished")
                self.kw_preview.append(
                    f"Median filter OK | size={k} -> rescale | workers={worker_count}"
                )
                self._switch_combos_to_denoised()
                self._update_segmentation_preview()
                self._refresh_gaussian_button_state()
            self._cleanup_worker_thread("_denoise_thread", "_denoise_worker")
            self._refresh_gaussian_button_state()

        self._connect_worker_callback(self._denoise_worker.finished, _finished)
        self._denoise_thread.start()

    def _on_denoise_progress(self, done: int, total: int) -> None:
        total = max(int(total), 1)
        done = max(0, min(int(done), total))
        if self.denoise_progress.maximum() != total:
            self.denoise_progress.setRange(0, total)
        self.denoise_progress.setValue(done)
        self.denoise_progress.setFormat(f"Progress {done}/{total}")

    def _on_apply_gaussian_background_clicked(self) -> None:
        if self._thread_is_running("_gaussian_thread"):
            self._request_worker_cancel("_gaussian_thread", "_gaussian_worker")
            self.btn_apply_gaussian.setEnabled(False)
            self.btn_apply_gaussian.setText("Stopping...")
            self.gaussian_progress.setRange(0, 1)
            self.gaussian_progress.setValue(0)
            self.gaussian_progress.setFormat("Stopping...")
            return

        if not self.gaussian_enabled_checkbox.isChecked():
            return

        base_layer = self._preferred_median_input_layer()
        if base_layer is None or not hasattr(base_layer, "data"):
            self.btn_apply_gaussian.setEnabled(False)
            QMessageBox.information(
                self,
                "No median result",
                "Run Median filter before Gaussian background subtraction.",
            )
            return

        try:
            arr = base_layer.data
            dims_tag = _layer_dims_tag(base_layer)
            if dims_tag is None:
                if arr.ndim == 4:
                    dims_tag = "TZYX"
                elif arr.ndim == 3:
                    dims_tag = "ZYX"
                elif arr.ndim == 2:
                    dims_tag = "YX"
            sigma_xy = float(self.gaussian_sigma_spin.value())
            alpha = float(self.gaussian_alpha_spin.value())
            scale_all = tuple(float(v) for v in getattr(base_layer, "scale", (1,) * arr.ndim))
            temporal = dims_tag in {"TYX", "TZYX"}
            squeeze_time = False
            if temporal:
                t0, t1 = self._selected_time_range()
                t0 = max(0, min(t0, arr.shape[0] - 1))
                t1 = max(t0, min(t1, arr.shape[0] - 1))
                gaussian_input = arr[t0 : t1 + 1]
                squeeze_time = gaussian_input.shape[0] == 1
                if dims_tag == "TZYX":
                    dims_out = "ZYX" if squeeze_time else "TZYX"
                    scale = scale_all[-3:] if squeeze_time else scale_all[-4:]
                else:
                    dims_out = "YX" if squeeze_time else "TYX"
                    scale = scale_all[-2:] if squeeze_time else scale_all[-3:]
            else:
                gaussian_input = arr
                if gaussian_input.ndim not in {2, 3}:
                    gaussian_input = self._current_volume_from_layer(base_layer)
                if gaussian_input.ndim == 3:
                    dims_out = "ZYX"
                    scale = scale_all[-3:]
                elif gaussian_input.ndim == 2:
                    dims_out = "YX"
                    scale = scale_all[-2:]
                else:
                    raise ValueError(
                        f"Gaussian background subtraction requires 2D/3D data, got {gaussian_input.ndim}D."
                    )
        except (TypeError, ValueError, AttributeError, IndexError, RuntimeError) as error:
            QMessageBox.critical(
                self,
                "Gaussian background failed",
                f"Gaussian background setup error: {error!r}",
            )
            return

        total = int(gaussian_input.shape[0]) if temporal else 1
        self.gaussian_sigma_spin.setEnabled(False)
        self.gaussian_alpha_spin.setEnabled(False)
        self.gaussian_enabled_checkbox.setEnabled(False)
        self.btn_apply_gaussian.setEnabled(True)
        self.btn_apply_gaussian.setText("Stop Gaussian")
        self.gaussian_progress.setRange(0, max(total, 1))
        self.gaussian_progress.setValue(0)
        self.gaussian_progress.setFormat(f"Progress 0/{max(total, 1)}")

        self._gaussian_thread = QThread()
        self._gaussian_worker = GaussianBackgroundWorker(
            gaussian_input,
            sigma_xy,
            alpha,
            temporal=temporal,
        )
        self._gaussian_worker.moveToThread(self._gaussian_thread)
        self._gaussian_thread.started.connect(self._gaussian_worker.run)
        self._gaussian_worker.progress.connect(self._on_gaussian_background_progress)

        def _finished(output, error):
            self.btn_apply_gaussian.setText("Run Gaussian")
            if isinstance(error, InterruptedError):
                self.gaussian_progress.setRange(0, 1)
                self.gaussian_progress.setValue(0)
                self.gaussian_progress.setFormat("Stopped")
            elif error is not None:
                self.gaussian_progress.setRange(0, 1)
                self.gaussian_progress.setValue(0)
                self.gaussian_progress.setFormat("Failed")
                QMessageBox.critical(
                    self,
                    "Gaussian background failed",
                    f"Gaussian background subtraction error: {error!r}",
                )
            else:
                out = np.asarray(output)
                if squeeze_time:
                    out = out[0]
                md = dict(getattr(base_layer, "metadata", {}) or {})
                previous_method = str(md.get("denoise_method") or "").strip()
                gaussian_method = "gaussian_background_subtraction"
                md.update(
                    {
                        "managed_by_plugin": True,
                        "is_denoise": True,
                        "dims_out": dims_out,
                        "denoise_method": (
                            f"{previous_method}+{gaussian_method}"
                            if previous_method
                            else gaussian_method
                        ),
                        "gaussian_background_enabled": True,
                        "gaussian_background_sigma_xy": sigma_xy,
                        "gaussian_background_alpha": alpha,
                        "gaussian_background_rescale_0_255": True,
                        "group_id": _get_group_id(base_layer) or _ensure_group_id(md),
                    }
                )
                if dims_tag in {"TYX", "TZYX"}:
                    md["time_range"] = self._selected_time_range()
                name = self._derived_result_layer_name(
                    base_layer,
                    f"Gaussian background sigma {sigma_xy:g} alpha {alpha:g}",
                )
                translate_all = tuple(
                    float(v) for v in getattr(base_layer, "translate", (0.0,) * out.ndim)
                )
                translate = translate_all[-out.ndim:]
                self._denoise_layer = self.viewer.add_image(
                    out,
                    name=name,
                    scale=scale,
                    translate=translate,
                    metadata=md,
                    **_spatial_unit_kwargs(base_layer, out.ndim),
                )
                self._gaussian_layer = self._denoise_layer
                for attr in ("colormap", "gamma", "opacity", "blending"):
                    with suppress(Exception):
                        setattr(self._denoise_layer, attr, getattr(base_layer, attr))
                with suppress(Exception):
                    self._denoise_layer.contrast_limits = (0, 255)
                self._denoise_ctx = self._context_from_layer(
                    base_layer,
                    out,
                    3 if dims_out in {"ZYX", "TZYX"} else 2,
                    tuple(scale),
                    temporal_dims=dims_out if dims_out in {"TYX", "TZYX"} else None,
                    time_range=self._selected_time_range() if dims_tag in {"TYX", "TZYX"} else None,
                )
                self._reset_processing_state(
                    clear_upsample=True,
                    clear_denoise=False,
                    clear_frangi=True,
                )
                self._focus_result_layer(self._denoise_layer)
                self._switch_combos_to_denoised()
                self._update_segmentation_preview()
                self.gaussian_progress.setRange(0, 1)
                self.gaussian_progress.setValue(1)
                self.gaussian_progress.setFormat("Finished")
                self.kw_preview.append(
                    f"Gaussian background OK | sigma XY={sigma_xy:g}, "
                    f"alpha={alpha:g} -> rescale"
                )
            self._cleanup_worker_thread("_gaussian_thread", "_gaussian_worker")
            self.gaussian_enabled_checkbox.setEnabled(True)
            self._refresh_gaussian_button_state()

        self._connect_worker_callback(self._gaussian_worker.finished, _finished)
        self._gaussian_thread.start()

    def _on_gaussian_background_progress(self, done: int, total: int) -> None:
        total = max(int(total), 1)
        done = max(0, min(int(done), total))
        if self.gaussian_progress.maximum() != total:
            self.gaussian_progress.setRange(0, total)
        self.gaussian_progress.setValue(done)
        self.gaussian_progress.setFormat(f"Progress {done}/{total}")

    # -----------------------------------------------------------------
    # Frangi
    # -----------------------------------------------------------------

    def _make_sigma_list(self, kind: str = "vessel") -> list[float] | None:
        kind = str(kind).lower()
        if kind == "sheet":
            min_spin = self.sheet_sigma_min
            max_spin = self.sheet_sigma_max
            count_spin = self.sheet_sigma_count
            label = "Sheet sigma"
        else:
            min_spin = self.sigma_min
            max_spin = self.sigma_max
            count_spin = self.sigma_count
            label = "Vessel sigma"
        smin = float(min_spin.value())
        smax = float(max_spin.value())
        cnt = int(count_spin.value())
        if smin < 0.01:
            min_spin.setValue(0.01)
            smin = 0.01
        if smax > 20.0:
            max_spin.setValue(20.0)
            smax = 20.0
        if smax <= smin:
            QMessageBox.warning(self, "Sigma range invalid", f"{label}: ensure max > min within [0.01, 20.0].")
            return None
        if cnt < 1:
            count_spin.setValue(1)
            cnt = 1
        return np.linspace(smin, smax, cnt, dtype=float).tolist()

    def _on_frangi_response_mode_changed(self, *_args) -> None:
        mode = str(self.frangi_response_mode_combo.currentData() or "vesselness")
        vessel_visible = mode in {"vesselness", "combined"}
        sheet_visible = mode in {"sheetness", "combined"}
        for widget in getattr(self, "_frangi_vessel_mode_widgets", ()):
            with suppress(Exception):
                widget.setVisible(vessel_visible)
        for widget in getattr(self, "_frangi_sheet_mode_widgets", ()):
            with suppress(Exception):
                widget.setVisible(sheet_visible)
        view_row = getattr(self, "_frangi_component_view_row", None)
        if view_row is not None:
            view_row.setVisible(mode == "combined")
        QTimer.singleShot(0, self._sync_panel_tab_height)
        self._rebuild_segmentation_layer_combos()
        self._refresh_rescale_target_label()
        self._refresh_frangi_rescale_controls_state()
        if getattr(self, "_seg_frangi_layer_combo", None) is not None and hasattr(self, "kw_preview"):
            self._update_segmentation_preview()

    def _frangi_rescale_is_enabled(self) -> bool:
        checkbox = getattr(self, "frangi_rescale_checkbox", None)
        return True if checkbox is None else bool(checkbox.isChecked())

    def _refresh_frangi_rescale_controls_state(self, *_args) -> None:
        enabled = self._frangi_rescale_is_enabled()
        running = self._thread_is_running("_frangi_thread")
        checkbox = getattr(self, "frangi_rescale_checkbox", None)
        if checkbox is not None:
            checkbox.setEnabled(not running)
        for control in (
            getattr(self, "rescale_target_label", None),
            getattr(self, "vessel_rescale_label", None),
            getattr(self, "vessel_rescale_spin", None),
            getattr(self, "sheet_rescale_label", None),
            getattr(self, "sheet_rescale_spin", None),
        ):
            if control is not None:
                control.setEnabled(enabled and not running)
        self._refresh_rescale_target_label()

    def _on_segmentation_frangi_layer_changed(self, *_args) -> None:
        self._restore_selected_structure_rescale_values()
        self._refresh_rescale_target_label()
        self._update_segmentation_preview()

    @staticmethod
    def _set_rescale_spin_from_metadata(spin, value) -> None:
        try:
            parsed = float(value)
        except (TypeError, ValueError):
            return
        if not np.isfinite(parsed):
            return
        parsed = min(max(parsed, float(spin.minimum())), float(spin.maximum()))
        spin.blockSignals(True)
        spin.setValue(parsed)
        spin.blockSignals(False)

    def _restore_selected_structure_rescale_values(self) -> None:
        layer = self._selected_segmentation_frangi_layer()
        if layer is None:
            return
        metadata = getattr(layer, "metadata", {}) or {}
        if not (
            metadata.get("is_frangi")
            or metadata.get("sigma_layer_role") == _STRUCTURAL_RESPONSE_LAYER_ROLE
        ):
            return
        result_name = str(metadata.get("frangi_result_name") or "").lower()
        if result_name == "vesselness":
            self._set_rescale_spin_from_metadata(
                self.vessel_rescale_spin,
                metadata.get("frangi_rescale_low_percent"),
            )
        elif result_name == "sheetness":
            self._set_rescale_spin_from_metadata(
                self.sheet_rescale_spin,
                metadata.get("frangi_rescale_low_percent"),
            )
        elif result_name == "combined max":
            self._set_rescale_spin_from_metadata(
                self.vessel_rescale_spin,
                metadata.get("frangi_vessel_rescale_low_percent"),
            )
            self._set_rescale_spin_from_metadata(
                self.sheet_rescale_spin,
                metadata.get("frangi_sheet_rescale_low_percent"),
            )

    def _refresh_rescale_target_label(self) -> None:
        label = getattr(self, "rescale_target_label", None)
        if label is None:
            return
        mode_combo = getattr(self, "frangi_response_mode_combo", None)
        mode = str(mode_combo.currentData() or "vesselness") if mode_combo is not None else "vesselness"
        component_specs = (
            (("Vessel", "vesselness"), ("Sheet", "sheetness"))
            if mode == "combined"
            else (("Sheet", "sheetness"),)
            if mode == "sheetness"
            else (("Vessel", "vesselness"),)
        )
        parts = []
        has_target = False
        for display_name, component_name in component_specs:
            layer = self._rescale_component_layer_candidate(component_name)
            if layer is None:
                parts.append(f"{display_name}: —")
                continue
            has_target = True
            metadata = getattr(layer, "metadata", {}) or {}
            applied = metadata.get("frangi_rescale_low_percent")
            try:
                applied_text = f" ({float(applied):.2f}%)"
            except (TypeError, ValueError):
                applied_text = ""
            parts.append(f"{display_name}: {layer.name}{applied_text}")
        label.setText("Rescale target — " + " | ".join(parts))
        apply_button = getattr(self, "rescale_apply_btn", None)
        if apply_button is not None:
            checkbox = getattr(self, "frangi_rescale_checkbox", None)
            rescale_enabled = (
                True if checkbox is None else bool(checkbox.isChecked())
            )
            thread = getattr(self, "_frangi_thread", None)
            running = bool(thread is not None and thread.isRunning())
            apply_button.setEnabled(
                rescale_enabled and not running and has_target
            )

    def _refresh_frangi_component_view_buttons(self) -> None:
        layers = getattr(self, "_frangi_component_layers", {}) or {}
        button_specs = (
            ("vesselness", getattr(self, "btn_view_vessel_frangi", None)),
            ("sheetness", getattr(self, "btn_view_sheet_frangi", None)),
            ("combined max", getattr(self, "btn_view_max_frangi", None)),
        )
        for name, button in button_specs:
            if button is not None:
                with suppress(Exception):
                    button.setEnabled(layers.get(name) is not None)

    def _show_frangi_component(self, component_name: str) -> None:
        component_name = str(component_name)
        layers = getattr(self, "_frangi_component_layers", {}) or {}
        layer = layers.get(component_name)
        if layer is None:
            return
        for item in layers.values():
            with suppress(Exception):
                item.visible = item is layer
        self._frangi_layer = layer
        self._focus_result_layer(layer)
        self._update_segmentation_preview()

    def _rescale_component_layer_candidate(self, component_name: str):
        component_name = str(component_name)
        component_layers = getattr(self, "_frangi_component_layers", {}) or {}
        mapped_layer = component_layers.get(component_name)
        candidates = []
        with suppress(Exception):
            candidates.append(self._selected_segmentation_frangi_layer())
        with suppress(Exception):
            candidates.append(self.viewer.layers.selection.active)
        candidates.extend((self._frangi_layer, mapped_layer))

        seen: set[int] = set()
        for layer in candidates:
            if layer is None or id(layer) in seen or not hasattr(layer, "data"):
                continue
            seen.add(id(layer))
            metadata = dict(getattr(layer, "metadata", {}) or {})
            result_name = str(metadata.get("frangi_result_name") or "").lower()
            if result_name and result_name != component_name:
                continue
            if (
                not result_name
                and not metadata.get("is_frangi")
                and metadata.get("sigma_layer_role")
                != _STRUCTURAL_RESPONSE_LAYER_ROLE
                and layer is not self._frangi_layer
                and layer is not mapped_layer
            ):
                continue
            data = getattr(layer, "data", None)
            if data is not None and int(getattr(data, "size", 0)) > 0:
                return layer
        return None

    def _ensure_rescalable_frangi_component(self, component_name: str) -> bool:
        component_name = str(component_name)
        raw_components = getattr(self, "_frangi_component_raw", {}) or {}
        component_layers = getattr(self, "_frangi_component_layers", {}) or {}
        layer = self._rescale_component_layer_candidate(component_name)
        if layer is None:
            return False
        if component_layers.get(component_name) is layer and component_name in raw_components:
            return True

        data = np.asarray(getattr(layer, "data", None))
        if data.size == 0:
            return False
        self._frangi_component_raw[component_name] = data
        self._frangi_component_layers[component_name] = layer
        self._frangi_component_specs.setdefault(component_name, {})
        dims_tag = str(_layer_dims_tag(layer) or "").upper()
        self._frangi_component_temporal_dims = (
            dims_tag if dims_tag in {"TYX", "TZYX"} else None
        )
        metadata = dict(getattr(layer, "metadata", {}) or {})
        metadata["is_frangi"] = True
        metadata["sigma_layer_role"] = _STRUCTURAL_RESPONSE_LAYER_ROLE
        metadata["frangi_result_name"] = component_name
        layer.metadata = metadata
        self._frangi_layer = layer
        self._refresh_frangi_component_view_buttons()
        return True

    def _component_rescale_low(self, component_name: str) -> float:
        if not self._frangi_rescale_is_enabled():
            return 0.0
        if component_name == "vesselness":
            return float(self.vessel_rescale_spin.value())
        if component_name == "sheetness":
            return float(self.sheet_rescale_spin.value())
        return float(self._frangi_component_specs.get(component_name, {}).get("rescale_low", 0.0))

    def _rescaled_frangi_component(self, component_name: str) -> np.ndarray | None:
        raw = (getattr(self, "_frangi_component_raw", {}) or {}).get(component_name)
        if raw is None:
            return None
        low = self._component_rescale_low(component_name)
        return np.asarray(
            _rescale_frangi_per_frame(raw, low, self._frangi_component_temporal_dims),
            dtype=np.float32,
        )

    def _update_component_layer_data(self, component_name: str, data: np.ndarray) -> None:
        layer = (getattr(self, "_frangi_component_layers", {}) or {}).get(component_name)
        if layer is None:
            return
        updated = self._replace_layer_data(layer, np.asarray(data))
        self._frangi_component_layers[component_name] = updated
        if self._frangi_layer is layer:
            self._frangi_layer = updated
        md = dict(getattr(updated, "metadata", {}) or {})
        md["is_frangi"] = True
        md["sigma_layer_role"] = _STRUCTURAL_RESPONSE_LAYER_ROLE
        md["frangi_result_name"] = component_name
        md["frangi_rescale_enabled"] = self._frangi_rescale_is_enabled()
        if component_name in {"vesselness", "sheetness"}:
            md["frangi_rescale_low_percent"] = self._component_rescale_low(component_name)
        elif component_name == "combined max":
            md["frangi_vessel_rescale_low_percent"] = self._component_rescale_low("vesselness")
            md["frangi_sheet_rescale_low_percent"] = self._component_rescale_low("sheetness")
        updated.metadata = md

    def _apply_component_rescales(
        self,
        component_names: tuple[str, ...],
    ) -> bool:
        scaled_components: dict[str, np.ndarray] = {}
        for component_name in component_names:
            if not self._ensure_rescalable_frangi_component(component_name):
                continue
            scaled = self._rescaled_frangi_component(component_name)
            if scaled is not None:
                scaled_components[component_name] = scaled
        if not scaled_components:
            return False

        raw_components = getattr(self, "_frangi_component_raw", {}) or {}
        for component_name, scaled in scaled_components.items():
            self._frangi_component_specs.setdefault(component_name, {})[
                "rescale_low"
            ] = self._component_rescale_low(component_name)
            self._frangi_component_specs[component_name]["rescale_enabled"] = True
            self._update_component_layer_data(component_name, scaled)

        max_data = None
        if "vesselness" in raw_components and "sheetness" in raw_components:
            vessel = self._rescaled_frangi_component("vesselness")
            sheet = self._rescaled_frangi_component("sheetness")
            if vessel is not None and sheet is not None:
                max_data = np.maximum(vessel, sheet).astype(np.float32, copy=False)
                self._update_component_layer_data("combined max", max_data)

        active_name = None
        if self._frangi_layer is not None:
            active_md = getattr(self._frangi_layer, "metadata", {}) or {}
            active_name = str(active_md.get("frangi_result_name") or "")
        if self._frangi_ctx is not None:
            if active_name in scaled_components:
                scaled = scaled_components[active_name]
                self._frangi_ctx.frangi = np.asarray(scaled)
                self._frangi_ctx.frangi_raw = np.asarray(scaled)
            elif active_name == "combined max" and max_data is not None or self._frangi_layer is None and max_data is not None:
                self._frangi_ctx.frangi = np.asarray(max_data)
                self._frangi_ctx.frangi_raw = np.asarray(max_data)
        self._refresh_rescale_target_label()
        self._update_segmentation_preview()
        return True

    def _on_apply_component_rescale_clicked(self) -> None:
        if not self._frangi_rescale_is_enabled():
            return
        mode = str(self.frangi_response_mode_combo.currentData() or "vesselness")
        components = (
            ("vesselness", "sheetness")
            if mode == "combined"
            else ("sheetness",)
            if mode == "sheetness"
            else ("vesselness",)
        )
        self._apply_component_rescales(components)

    def _extract_current_2d_or_3d(self, layer) -> tuple[np.ndarray, int, float]:
        """Return the selected layer as (image, spatial ndim, z/x ratio)."""
        vol = self._current_volume_from_layer(layer)
        if vol.ndim == 3:
            sc = getattr(layer, "scale", (1,) * layer.data.ndim)
            z = float(sc[-3]) if len(sc) >= 3 else 1.0
            x = float(sc[-1]) if len(sc) >= 1 else 1.0
            ratio = (z / x) if (z > 0 and x > 0) else 1.0
            z0, z1 = self._selected_slice_range()
            if z0 == z1:
                z_idx = int(np.clip(z0, 0, vol.shape[0] - 1))
                return vol[z_idx].astype(np.float32, copy=False), 2, 1.0
            zstart = int(np.clip(z0, 0, vol.shape[0] - 1))
            zend = int(np.clip(z1, zstart, vol.shape[0] - 1))
            return vol[zstart : zend + 1].astype(np.float32, copy=False), 3, float(ratio)
        if vol.ndim == 2:
            return vol.astype(np.float32, copy=False), 2, 1.0
        raise ValueError(f"Unsupported data shape after slicing: {vol.shape}")

    def _extract_processing_series(self, layer) -> tuple[np.ndarray, int, float, str | None]:
        arr = layer.data
        dims_tag = _layer_dims_tag(layer)
        if dims_tag is None and getattr(arr, "ndim", None) == 4:
            # SIGMA splits channels into separate layers, so an unlabeled 4D
            # image layer is the common napari TZYX representation.
            dims_tag = "TZYX"
        if dims_tag not in {"TYX", "TZYX"}:
            img, dim, ratio = self._extract_current_2d_or_3d(layer)
            return img, dim, ratio, None

        t0, t1 = self._selected_time_range()
        t0 = max(0, min(t0, arr.shape[0] - 1))
        t1 = max(t0, min(t1, arr.shape[0] - 1))
        subset = arr[t0 : t1 + 1]
        if subset.shape[0] == 1:
            single = subset[0]
            if dims_tag == "TZYX":
                sc = getattr(layer, "scale", (1,) * arr.ndim)
                z = float(sc[-3]) if len(sc) >= 3 else 1.0
                x = float(sc[-1]) if len(sc) >= 1 else 1.0
                ratio = (z / x) if (z > 0 and x > 0) else 1.0
                return single, 3, ratio, None
            return single, 2, 1.0, None

        if dims_tag == "TYX":
            return subset, 2, 1.0, "TYX"
        sc = getattr(layer, "scale", (1,) * arr.ndim)
        z = float(sc[-3]) if len(sc) >= 3 else 1.0
        x = float(sc[-1]) if len(sc) >= 1 else 1.0
        ratio = (z / x) if (z > 0 and x > 0) else 1.0
        return subset, 3, ratio, "TZYX"

    def _on_apply_frangi_clicked(self):
        if self._thread_is_running("_frangi_thread"):
            self._request_worker_cancel("_frangi_thread", "_frangi_worker")
            self.apply_btn.setEnabled(False)
            self.apply_btn.setText("Stopping Extraction…")
            self.frangi_progress.setRange(0, 1)
            self.frangi_progress.setValue(0)
            self.frangi_progress.setFormat("Stopping…")
            return
        frangi_response_mode = str(self.frangi_response_mode_combo.currentData() or "vesselness")
        rescale_enabled = self._frangi_rescale_is_enabled()
        response_specs: list[dict[str, Any]] = []
        if frangi_response_mode in {"vesselness", "combined"}:
            vessel_sigmas = self._make_sigma_list("vessel")
            if vessel_sigmas is None:
                return
            response_specs.append(
                {
                    "name": "vesselness",
                    "response_mode": "vesselness",
                    "sigmas": vessel_sigmas,
                    "rescale_enabled": rescale_enabled,
                    "rescale_low": (
                        float(self.vessel_rescale_spin.value())
                        if rescale_enabled
                        else 0.0
                    ),
                }
            )
        if frangi_response_mode in {"sheetness", "combined"}:
            sheet_sigmas = self._make_sigma_list("sheet")
            if sheet_sigmas is None:
                return
            response_specs.append(
                {
                    "name": "sheetness",
                    "response_mode": "sheetness",
                    "sigmas": sheet_sigmas,
                    "rescale_enabled": rescale_enabled,
                    "rescale_low": (
                        float(self.sheet_rescale_spin.value())
                        if rescale_enabled
                        else 0.0
                    ),
                }
            )
        if not response_specs:
            QMessageBox.warning(self, "Response mode invalid", "Choose vesselness, sheetness, or combined.")
            return
        layer = self._selected_frangi_raw_layer() or self.viewer.layers.selection.active
        if layer is None or not hasattr(layer, "data"):
            QMessageBox.information(self, "No image", "Please choose a raw data layer first.")
            return

        device = self._resolve_device()
        try:
            img, dim, ratio, temporal_dims = self._extract_processing_series(layer)
        except (TypeError, ValueError, AttributeError, IndexError, RuntimeError) as e:
            QMessageBox.critical(self, "Slice error", f"Failed to get 2D/3D data: {e!r}")
            return

        img_shape = tuple(int(value) for value in img.shape)
        z_depth = int(img_shape[-3]) if dim == 3 and len(img_shape) >= 3 else 1
        process_z_as_2d = dim == 3 and z_depth <= int(self.kernel_spin.value())
        logical_units = (
            int(img.shape[0]) * z_depth
            if temporal_dims and process_z_as_2d
            else int(img.shape[0])
            if temporal_dims
            else z_depth
            if process_z_as_2d
            else 1
        )
        initial_total = max(int(sum(logical_units * len(spec["sigmas"]) for spec in response_specs)), 1)
        self.frangi_progress.setRange(0, initial_total)
        self.frangi_progress.setValue(1)
        self.frangi_progress.setFormat(f"Progress 1/{initial_total}")
        self.apply_btn.setEnabled(True)
        self.apply_btn.setText("Stop Extraction")
        self.frangi_rescale_checkbox.setEnabled(False)
        for control in (
            self.rescale_target_label,
            self.vessel_rescale_label,
            self.vessel_rescale_spin,
            self.sheet_rescale_label,
            self.sheet_rescale_spin,
            self.rescale_apply_btn,
        ):
            control.setEnabled(False)

        self._frangi_thread = QThread()
        self._frangi_worker = FrangiWorker(
            img,
            spatial_dim=dim,
            sigmas=response_specs[0]["sigmas"],
            kernel_size=2 * int(self.kernel_spin.value()) + 1,
            device=device,
            zx_ratio=ratio,
            temporal_dims=temporal_dims,
            psf_ratio=float(self.psf_ratio_spin.value()),
            alpha=0.5,
            beta=0.5,
            gamma=2.0,
            response_mode=frangi_response_mode,
            response_specs=response_specs,
        )
        self._frangi_worker.moveToThread(self._frangi_thread)
        self._frangi_thread.started.connect(self._frangi_worker.run)
        self._frangi_worker.progress.connect(self._on_frangi_progress)
        self._frangi_worker.sigma_progress.connect(self._on_frangi_sigma_progress)

        def _finished(frangi_result_raw, error):
            self.apply_btn.setEnabled(True)
            self.apply_btn.setText("Run Extraction")
            if isinstance(error, InterruptedError):
                self.frangi_progress.setRange(0, 1)
                self.frangi_progress.setValue(0)
                self.frangi_progress.setFormat("Stopped")
            elif error is not None:
                self.frangi_progress.setRange(0, 1)
                self.frangi_progress.setValue(0)
                self.frangi_progress.setFormat("Failed")
                QMessageBox.critical(
                    self,
                    "Extraction failed",
                    f"Structural awareness extraction error: {error!r}",
                )
            else:
                self.frangi_progress.setRange(0, 1)
                self.frangi_progress.setValue(1)
                self.frangi_progress.setFormat("Finished")
                self._add_frangi_result(
                    layer,
                    img,
                    frangi_result_raw,
                    dim,
                    device,
                    temporal_dims,
                    response_specs,
                    frangi_response_mode,
                )
            self._cleanup_worker_thread("_frangi_thread", "_frangi_worker")
            self.frangi_rescale_checkbox.setEnabled(True)
            self._refresh_frangi_rescale_controls_state()

        self._connect_worker_callback(self._frangi_worker.finished, _finished)
        self._frangi_thread.start()

    def _on_frangi_progress(self, done: int, total: int) -> None:
        total = max(int(total), 1)
        done = max(0, min(int(done), total))
        if self.frangi_progress.maximum() != total:
            self.frangi_progress.setRange(0, total)
        self.frangi_progress.setValue(done)
        self.frangi_progress.setFormat(f"Progress {done}/{total}")

    def _on_frangi_sigma_progress(self, frame_done: int, frame_total: int, sigma_done: int, sigma_total: int) -> None:
        frame_total = max(int(frame_total), 1)
        sigma_total = max(int(sigma_total), 1)
        frame_done = max(1, min(int(frame_done), frame_total))
        sigma_done = max(0, min(int(sigma_done), sigma_total))
        total_steps = frame_total * sigma_total
        done_steps = (frame_done - 1) * sigma_total + sigma_done
        self.frangi_progress.setRange(0, total_steps)
        self.frangi_progress.setValue(done_steps)
        self.frangi_progress.setFormat(f"Progress {done_steps}/{total_steps}")

    def _add_frangi_result(
        self,
        layer,
        img: np.ndarray,
        frangi_payload,
        dim: int,
        device: str,
        temporal_dims: str | None,
        response_specs: list[dict[str, Any]],
        frangi_response_mode: str,
    ) -> None:
        img_shape = tuple(int(value) for value in img.shape)
        if isinstance(frangi_payload, dict) and isinstance(frangi_payload.get("responses"), dict):
            raw_responses = {
                str(name): np.asarray(value)
                for name, value in frangi_payload.get("responses", {}).items()
            }
        else:
            raw_responses = {str(frangi_response_mode): np.asarray(frangi_payload)}
        if not raw_responses:
            raise RuntimeError("Structural extraction returned no response data.")
        result_sample = np.asarray(next(iter(raw_responses.values())))
        self._frangi_component_raw = {name: np.asarray(raw) for name, raw in raw_responses.items()}
        self._frangi_component_layers = {}
        self._frangi_component_temporal_dims = temporal_dims

        specs_by_name = {
            str(spec.get("name") or spec.get("response_mode") or "frangi"): dict(spec)
            for spec in response_specs
        }
        self._frangi_component_specs = {name: dict(spec) for name, spec in specs_by_name.items()}

        base_md = dict(getattr(layer, "metadata", {}) or {})
        base_md["source"] = base_md.get("source", _get_source_from_layer(layer))
        base_md["unit"] = _get_units_from_layer(layer)
        base_md["is_frangi"] = True
        base_md["sigma_layer_role"] = _STRUCTURAL_RESPONSE_LAYER_ROLE
        base_md["frangi_response_mode"] = str(frangi_response_mode)
        base_md["frangi_rescale_enabled"] = any(
            bool(spec.get("rescale_enabled")) for spec in response_specs
        )
        base_md["group_id"] = _get_group_id(layer) or _ensure_group_id(base_md)

        sc = getattr(layer, "scale", (1,) * layer.data.ndim)
        if temporal_dims == "TYX":
            desired_scale = tuple(float(value) for value in sc[-3:])
            base_md["dims_out"] = "TYX"
        elif temporal_dims == "TZYX":
            desired_scale = tuple(float(value) for value in sc[-4:])
            base_md["dims_out"] = "TZYX"
        elif result_sample.ndim == 3:
            desired_scale = (float(sc[-3]), float(sc[-2]), float(sc[-1]))
            base_md["dims_out"] = "ZYX"
        else:
            desired_scale = self._preferred_2d_scale(layer, sc)
            base_md["dims_out"] = "YX"

        # Keep one canonical displayed-axis contract on every derived layer.
        # Source metadata may say TCZYX even though napari received a single
        # channel TZYX layer.
        base_md["dims"] = base_md["dims_out"]

        def _scaled_response(name: str, raw: np.ndarray) -> np.ndarray:
            spec = specs_by_name.get(name, {})
            low = float(spec.get("rescale_low", 0.0))
            return np.asarray(_rescale_frangi_per_frame(raw, low, temporal_dims), dtype=np.float32)

        def _spec_summary(name: str) -> str:
            spec = specs_by_name.get(name, {})
            sigmas_text = np.array(spec.get("sigmas", []), dtype=float)
            low = float(spec.get("rescale_low", 0.0))
            return f"{name}: sigmas={sigmas_text} rescale={low:.2f}%"

        display_responses: dict[str, np.ndarray] = {
            name: _scaled_response(name, raw)
            for name, raw in raw_responses.items()
        }

        created_layers = []

        def _add_response_layer(response_name: str, data: np.ndarray, *, suffix: str, extra_md: dict[str, Any] | None = None):
            md = dict(base_md)
            md["frangi_result_name"] = response_name
            md["frangi_response_mode"] = response_name if response_name in {"vesselness", "sheetness"} else str(frangi_response_mode)
            spec = specs_by_name.get(response_name, {})
            if spec:
                md["frangi_sigmas"] = [float(v) for v in spec.get("sigmas", [])]
                md["frangi_rescale_enabled"] = bool(spec.get("rescale_enabled"))
                md["frangi_rescale_low_percent"] = float(spec.get("rescale_low", 0.0))
            if extra_md:
                md.update(extra_md)
            out_layer = self.viewer.add_image(
                np.asarray(data),
                name=self._derived_result_layer_name(layer, suffix),
                scale=desired_scale,
                metadata=md,
                **_spatial_unit_kwargs(layer, np.asarray(data).ndim),
            )
            created_layers.append(out_layer)
            self._frangi_component_layers[response_name] = out_layer
            return out_layer

        if str(frangi_response_mode) == "combined":
            vessel = display_responses.get("vesselness")
            sheet = display_responses.get("sheetness")
            if vessel is None or sheet is None:
                raise RuntimeError("Combined mode expected both vesselness and sheetness responses.")
            _add_response_layer(
                "vesselness",
                vessel,
                suffix="vesselness response",
                extra_md={"frangi_combined_component": True},
            )
            _add_response_layer(
                "sheetness",
                sheet,
                suffix="sheetness response",
                extra_md={"frangi_combined_component": True},
            )
            frangi_result_raw = np.maximum(vessel, sheet).astype(np.float32, copy=False)
            max_md = {
                "frangi_result_name": "combined max",
                "frangi_response_mode": "combined",
                "frangi_combined_method": "max(vesselness, sheetness)",
                "frangi_vessel_sigmas": [float(v) for v in specs_by_name.get("vesselness", {}).get("sigmas", [])],
                "frangi_sheet_sigmas": [float(v) for v in specs_by_name.get("sheetness", {}).get("sigmas", [])],
                "frangi_vessel_rescale_low_percent": float(specs_by_name.get("vesselness", {}).get("rescale_low", 0.0)),
                "frangi_sheet_rescale_low_percent": float(specs_by_name.get("sheetness", {}).get("rescale_low", 0.0)),
            }
            self._frangi_layer = _add_response_layer(
                "combined max",
                frangi_result_raw,
                suffix="combined structural response",
                extra_md=max_md,
            )
            for response_name, response_layer in self._frangi_component_layers.items():
                with suppress(Exception):
                    response_layer.visible = response_name == "combined max"
        else:
            response_name = "sheetness" if str(frangi_response_mode) == "sheetness" else "vesselness"
            if response_name not in display_responses:
                response_name = next(iter(display_responses))
            frangi_result_raw = display_responses[response_name]
            self._frangi_layer = _add_response_layer(
                response_name,
                frangi_result_raw,
                suffix=f"{response_name} response",
            )

        self._focus_result_layer(self._frangi_layer)

        self._frangi_ctx = FrangiContext(
            image=img,
            frangi=np.asarray(frangi_result_raw),
            frangi_raw=np.asarray(frangi_result_raw),
            dim=dim,
            pixel_size=_get_pixel_size_tuple(layer, dim),
            device=device,
            source=base_md.get("source"),
            group_id=base_md.get("group_id"),
            channel_index=(getattr(layer, "metadata", {}) or {}).get("channel_index"),
            unit=base_md.get("unit", "um"),
            temporal_dims=temporal_dims,
            time_range=self._selected_time_range() if temporal_dims is not None else None,
        )

        self._refresh_frangi_component_view_buttons()
        self._rebuild_segmentation_layer_combos()

        self._update_segmentation_preview()

        slice_start, slice_end = self._selected_slice_range()
        dimtxt = temporal_dims or (
            f"2D (slice {slice_start})" if dim == 2 else ("3D" if dim == 3 else "2D")
        )
        if dim == 3 and img_shape[-3] <= int(self.kernel_spin.value()):
            dimtxt = f"{dimtxt} thin-Z→2D-per-slice"
        psf_txt = f" | PSF z/xy={float(self.psf_ratio_spin.value()):.2f}" if dim == 3 else ""
        spec_txt = " | ".join(_spec_summary(name) for name in display_responses)
        self.kw_preview.append(
            f"Structural extraction OK | {dimtxt} | mode={str(frangi_response_mode)}"
            f" | {spec_txt}{psf_txt} | device={device}"
        )
        if created_layers:
            self.kw_preview.append(
                "Structural response layers | "
                + ", ".join(layer_item.name for layer_item in created_layers)
                + (f" | segmentation uses: {self._frangi_layer.name}" if self._frangi_layer is not None else "")
            )
        with suppress(TypeError, ValueError):
            stats = np.asarray(frangi_result_raw, dtype=np.float32)
            p99 = float(np.percentile(stats, 99.0))
            self.kw_preview.append(
                "Structural response stats | "
                f"shape={stats.shape} min={float(np.nanmin(stats)):.6g} "
                f"max={float(np.nanmax(stats)):.6g} mean={float(np.nanmean(stats)):.6g} "
                f"p99={p99:.6g}"
            )

    def _update_segmentation_layer_labels(self) -> None:
        self._rebuild_segmentation_layer_combos()

    # -----------------------------------------------------------------
    # Segmentation
    # -----------------------------------------------------------------

    def _build_segmentation_kwargs(self) -> dict[str, Any] | None:
        self._segmentation_input_error = None
        raw_layer = self._selected_segmentation_raw_layer()
        frangi_layer = self._selected_segmentation_frangi_layer()
        if raw_layer is None or frangi_layer is None:
            self._segmentation_input_error = (
                "Select both a raw data layer and a structural response layer."
            )
            return None

        # Always rebuild the context from the layers currently selected in the
        # segmentation controls. Reusing a stale context can silently combine an
        # original-resolution image with the pixel size of an upsampled layer.
        if not self._ensure_frangi_context_from_selected_layers():
            raw_shape = tuple(int(value) for value in getattr(raw_layer.data, "shape", ()))
            frangi_shape = tuple(int(value) for value in getattr(frangi_layer.data, "shape", ()))
            self._segmentation_input_error = (
                f"Selected layers have different shapes: raw {raw_shape}, structural response "
                f"{frangi_shape}. Run Structural Awareness Extraction on "
                f"'{getattr(raw_layer, 'name', 'the selected raw layer')}' and select the new "
                "matching response before segmentation."
            )
            return None

        base_image = self._frangi_ctx.image
        frangi_image = self._frangi_ctx.frangi
        base_dim = self._frangi_ctx.dim
        fresh_px = self._frangi_ctx.pixel_size

        downsample_xy_factor = 1
        downsample_target_xy = None
        raw_metadata = dict(getattr(raw_layer, "metadata", {}) or {}) if raw_layer is not None else {}
        if bool(raw_metadata.get("is_upsample")):
            with suppress(TypeError, ValueError, IndexError):
                candidate_factor = int(
                    raw_metadata.get("upsample_factor_xy", raw_metadata.get("XYUpsamplingFactor", 1))
                )
                original_shape = tuple(
                    int(value)
                    for value in raw_metadata.get(
                        "upsample_original_shape",
                        raw_metadata.get("upsample_input_shape", ()),
                    )
                )
                candidate_xy = tuple(original_shape[-2:])
                current_xy = tuple(int(value) for value in base_image.shape[-2:])
                if (
                    candidate_factor > 1
                    and len(candidate_xy) == 2
                    and current_xy
                    == (candidate_xy[0] * candidate_factor, candidate_xy[1] * candidate_factor)
                ):
                    downsample_xy_factor = candidate_factor
                    downsample_target_xy = candidate_xy

        return {
            "image": base_image,
            "frangi": frangi_image,
            "pixel_size": fresh_px,
            "beta1": float(self.beta1_spin.value()),
            "beta2": float(self.beta2_spin.value()),
            "n_fore": int(self.nfore_spin.value()),
            "n_back": int(self.nback_spin.value()),
            "max_iter": int(self.maxiter_spin.value()),
            "init_method": str(self.init_combo.currentText()),
            "em_foreground_sample_points": self._em_foreground_sample_points(),
            "em_background_sample_points": self._em_background_sample_points(),
            "device": self._resolve_device(),
            "run_framewise": bool(self._frangi_ctx.temporal_dims),
            "spatial_dim": int(base_dim),
            "postprocess_downsample_xy_factor": downsample_xy_factor,
            "postprocess_downsample_target_xy": downsample_target_xy,
        }

    def _update_segmentation_preview(self):
        kwargs = self._build_segmentation_kwargs()
        if kwargs is None:
            self.kw_preview.setPlainText(
                getattr(self, "_segmentation_input_error", None)
                or "Run structural awareness extraction first to preview segmentation()."
            )
            return
        lines = [
            "segmentation() will be called with:",
            f"  image: float32 array, shape={tuple(kwargs['image'].shape)}",
            f"  structural response: float32 array, shape={tuple(kwargs['frangi'].shape)}",
            f"  pixel_size: {kwargs['pixel_size']} ({'z,y,x' if len(kwargs['pixel_size'])==3 else 'y,x'})",
            f"  beta1 (smoothness)={kwargs['beta1']}  "
            f"beta2 (structure)={kwargs['beta2']}",
            f"  n_fore={kwargs['n_fore']}  n_back={kwargs['n_back']}  max_iter={kwargs['max_iter']}",
            f"  init_method='{kwargs['init_method']}'",
            (
                "  EM sampling: "
                "foreground="
                f"{_format_em_sample_points(kwargs['em_foreground_sample_points'])}, "
                "background="
                f"{_format_em_sample_points(kwargs['em_background_sample_points'])}"
            ),
            f"  device='{kwargs['device']}'",
            (
                "  output: component-preserving XY downsample "
                f"{kwargs['postprocess_downsample_xy_factor']}x -> {kwargs['postprocess_downsample_target_xy']}"
                if kwargs["postprocess_downsample_xy_factor"] > 1
                else "  output: native segmentation grid"
            ),
            "Notes:",
            f"  • Slice range: {self._selected_slice_range()} (single slice => 2D, multiple slices => 3D subset).",
            (
                "  • Input denoised: "
                f"{'YES' if bool((getattr(self._selected_segmentation_raw_layer(), 'metadata', {}) or {}).get('is_denoise')) else 'NO'}"
                f"; median filter size={self.cmb_kernel.currentText()}"
            ),
        ]
        self.kw_preview.setPlainText("\n".join(lines))

    def _on_run_segmentation_clicked(self):
        if self._thread_is_running("_seg_thread"):
            self._request_worker_cancel("_seg_thread", "_seg_worker")
            self.seg_btn.setEnabled(False)
            self.seg_btn.setText("Stopping Segmentation…")
            self.seg_progress.setRange(0, 100)
            self.seg_progress.setValue(0)
            self.seg_progress.setFormat("Stopping…")
            return
        kwargs = self._build_segmentation_kwargs()
        if kwargs is None:
            QMessageBox.information(
                self,
                "Segmentation inputs do not match",
                getattr(self, "_segmentation_input_error", None)
                or "Run structural awareness extraction first, or select matching raw data and structural response layers.",
            )
            return

        self._set_text_overlay("Segmentation starting...")

        self.seg_progress.setRange(0, 0)
        self.seg_progress.setFormat("Running segmentation…")
        self.seg_btn.setEnabled(True)
        self.seg_btn.setText("Stop Segmentation")

        self._seg_thread = QThread()
        self._seg_worker = SegWorker(kwargs)
        self._seg_worker.moveToThread(self._seg_thread)
        self._seg_thread.started.connect(self._seg_worker.run)
        self._seg_worker.started.connect(lambda: None)
        self._seg_worker.progress.connect(self._on_seg_progress)
        self._seg_worker.frame_progress.connect(self._on_seg_frame_progress)

        def _finished(payload, error):
            self.seg_progress.setRange(0, 100)
            if isinstance(error, InterruptedError):
                self.seg_progress.setValue(0)
                self.seg_progress.setFormat("Stopped")
            elif error is not None:
                self.seg_progress.setValue(0)
                QMessageBox.critical(self, "Segmentation failed", f"segmentation() error: {error!r}")
            else:
                self.seg_progress.setValue(100)
                seg, info = payload
                self._add_segmentation_result(seg, info)
            self.seg_btn.setEnabled(True)
            self.seg_btn.setText("Run Segmentation")
            self._cleanup_worker_thread("_seg_thread", "_seg_worker")

        self._connect_worker_callback(self._seg_worker.finished, _finished)
        self._seg_thread.start()

    def _on_seg_progress(self, iteration: int, delta: float):
        if np.isfinite(delta):
            self._set_text_overlay(f"Segmentation iter: {iteration} | ΔlogL: {delta:.3e}")
        else:
            self._set_text_overlay(f"Segmentation iter: {iteration} | ΔlogL: --")

        with suppress(Exception):
            base = self.kw_preview.toPlainText() if hasattr(self.kw_preview, "toPlainText") else ""
            head = f"[iter={iteration}] ΔlogL={delta:.3e}" if np.isfinite(delta) else f"[iter={iteration}] ΔlogL=--"
            lines = [head] + base.splitlines()
            self.kw_preview.setPlainText("\n".join(lines[:30]))

    def _on_seg_frame_progress(self, done: int, total: int) -> None:
        total = max(int(total), 1)
        self.seg_progress.setRange(0, total)
        self.seg_progress.setValue(max(0, min(int(done), total)))
        self.seg_progress.setFormat(f"Segmenting frame {done}/{total}")

    def _add_segmentation_result(self, seg, info=None):
        try:
            layer = self.viewer.layers.selection.active
            expected_dim = self._frangi_ctx.dim if self._frangi_ctx is not None else None
            seg_np = np.asarray(seg)
            temporal_dims = self._frangi_ctx.temporal_dims if self._frangi_ctx is not None else None
            if temporal_dims is None and expected_dim in {2, 3}:
                seg_np = _squeeze_leading_singletons(seg_np, expected_dim)
            elif temporal_dims is not None and self._frangi_ctx is not None:
                expected_frames = int(self._frangi_ctx.image.shape[0])
                expected_ndim = int(self._frangi_ctx.dim) + 1
                if seg_np.ndim != expected_ndim or int(seg_np.shape[0]) != expected_frames:
                    raise ValueError(
                        "Temporal segmentation changed shape: "
                        f"expected {expected_frames} frames with ndim={expected_ndim}, "
                        f"got shape={tuple(seg_np.shape)}."
                    )
            source_layer = self._selected_segmentation_raw_layer() or layer
            downsample_xy_factor = max(int(getattr(info, "downsample_xy_factor", 1)), 1)
            sc_layer = getattr(source_layer, "scale", None) if source_layer else None
            add_kwargs = {}
            if sc_layer is not None:
                output_scale = list(float(value) for value in sc_layer)
                if downsample_xy_factor > 1 and len(output_scale) >= 2:
                    output_scale[-2] *= downsample_xy_factor
                    output_scale[-1] *= downsample_xy_factor
                if temporal_dims == "TYX":
                    add_kwargs["scale"] = tuple(output_scale[-3:])
                elif temporal_dims == "TZYX":
                    add_kwargs["scale"] = tuple(output_scale[-4:])
                elif seg_np.ndim == 3:
                    add_kwargs["scale"] = (output_scale[-3], output_scale[-2], output_scale[-1])
                elif seg_np.ndim == 2:
                    add_kwargs["scale"] = (output_scale[-2], output_scale[-1])

            name = self._derived_result_layer_name(source_layer, "segmentation")

            mask8 = ((seg_np.astype(np.int64) > 0).astype(np.uint8) * 255)
            mask8 = np.ascontiguousarray(mask8)

            # Build metadata BEFORE add_image so napari's "layer added"
            # listeners (combo refresh, scale-bar config, etc.) already see
            # is_segmentation/unit/source/group_id. Otherwise the first run
            # adds a layer without metadata and downstream UI misses it.
            new_md = self._build_segmentation_layer_metadata(source_layer, mask8, info)
            if downsample_xy_factor > 1:
                fallback_metadata_scale = sc_layer if sc_layer is not None else ()
                metadata_scale = list(new_md.get("scale_per_axis", fallback_metadata_scale))
                if len(metadata_scale) >= 2:
                    metadata_scale[-2] = float(metadata_scale[-2]) * downsample_xy_factor
                    metadata_scale[-1] = float(metadata_scale[-1]) * downsample_xy_factor
                    new_md["scale_per_axis"] = tuple(metadata_scale)
                source_interpolation = new_md.pop("upsample_interpolation", None)
                source_factor = new_md.pop("upsample_factor_xy", downsample_xy_factor)
                new_md.pop("XYUpsamplingFactor", None)
                new_md.pop("XYInterpolation", None)
                new_md["is_upsample"] = False
                new_md["downsampled_from_upsample"] = True
                new_md["downsample_xy_factor"] = downsample_xy_factor
                new_md["source_upsample_factor_xy"] = int(source_factor)
                if source_interpolation is not None:
                    new_md["source_upsample_interpolation"] = str(source_interpolation)
                new_md["downsample_output_shape"] = tuple(int(value) for value in mask8.shape)
                new_md["components_before_downsample"] = int(
                    getattr(info, "components_before_downsample", 0)
                )
                new_md["components_after_downsample"] = int(
                    getattr(info, "components_after_downsample", 0)
                )
                new_md["downsample_relocated_seeds"] = int(
                    getattr(info, "downsample_relocated_seeds", 0)
                )
                new_md["downsample_frame_stats"] = tuple(
                    getattr(info, "downsample_frame_stats", ())
                )

            new_layer = self.viewer.add_image(
                mask8,
                name=name,
                rgb=False,
                blending="translucent_no_depth",
                metadata=new_md,
                **_spatial_unit_kwargs(source_layer, mask8.ndim),
                **add_kwargs,
            )
            with suppress(Exception):
                new_layer.contrast_limits = (0, 255)

            new_layer.metadata = new_md
            self._segmentation_layer = new_layer
            self._focus_result_layer(new_layer)
            self._rebuild_analysis_layer_combo(clear_outputs=False)

            self._configure_scale_bar(new_md["unit"])

            ctx = self._frangi_ctx
            if ctx is None:
                raise RuntimeError("Structural response context missing for segmentation result")
            iter_text = ""
            downsample_text = ""
            if info is not None:
                with suppress(Exception):
                    iter_text = (
                        f" iterations_run={int(getattr(info, 'iterations_run', 0))} "
                        f"converged={bool(getattr(info, 'converged', False))} |"
                    )
                with suppress(Exception):
                    if int(getattr(info, "downsample_xy_factor", 1)) > 1:
                        downsample_text = (
                            f" downsample={int(info.downsample_xy_factor)}x XY "
                            f"components={int(info.components_before_downsample)}"
                            f"->{int(info.components_after_downsample)} |"
                        )
            self.kw_preview.append(
                "Segmentation OK | "
                f"{iter_text}{downsample_text} "
                f"dim={ctx.dim} | "
                f"β1(smoothness)={float(self.beta1_spin.value())} "
                f"β2(structure)={float(self.beta2_spin.value())} "
                f"nfore={int(self.nfore_spin.value())} "
                f"nback={int(self.nback_spin.value())} maxiter={int(self.maxiter_spin.value())} "
                f"init={self.init_combo.currentText()} | "
                f"device={ctx.device}"
            )
        except (TypeError, ValueError, AttributeError, RuntimeError) as e:
            QMessageBox.critical(self, "Add result failed", f"Failed to add segmentation result: {e!r}")


# ---------------------------------------------------------------------
# napari reader
# ---------------------------------------------------------------------



if __name__ == "__main__":  # pragma: no cover
    try:
        import napari
        v = napari.Viewer()
        w = SIGMAWidget(v)
        v.window.add_dock_widget(w, area="right")
        napari.run()
    except (ImportError, RuntimeError, AttributeError) as e:
        import traceback
        print("Failed to launch napari demo:", e)
        traceback.print_exc()
